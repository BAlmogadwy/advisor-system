"""Student lists endpoints: access, gates, CSRF, body cap, the fail-closed audit,
no student in any URL, the viewing cache and screen-versus-file parity."""

from __future__ import annotations

import json
import re
from html.parser import HTMLParser
from io import BytesIO

import pytest
from django.contrib.auth.models import Group
from django.test import Client, RequestFactory
from django.urls import URLPattern, URLResolver, get_resolver, resolve, reverse
from openpyxl import load_workbook

from core import exam_roster_views as views
from core import exam_student_export_views as export_views
from core.middleware import ExamCommitteeAccessMiddleware
from core.models import AuditLog, ExamTimetableRun, StudentTermSection
from core.services import exam_roster_view as rv
from core.services import exam_student_export as export
from core.services.audit import validate_hash_chain
from core.services.rbac import ROLE_ADVISOR, ROLE_EXAM_COMMITTEE, ROLE_STUDENT, ensure_role_groups
from tests.exam_source_factory import scraped_exam_registration
from tests.exam_student_export_fixture import (
    ALL_IDS,
    FEMALE_IS,
    MALE_AI,
    OVERLOADED,
    build_population,
    build_saved_run,
    save_payload,
    saved_payload,
    student_name,
)

pytestmark = pytest.mark.django_db

PAGE = "exam_rosters_page"
INDEX = "exam_roster_index"
DETAIL = "exam_roster_detail"
LOOKUP = "exam_roster_lookup"
ROSTER_VIEWS = (PAGE, INDEX, DETAIL, LOOKUP)
COURSE = {"scope": {"kind": "course", "exam": "MATH101"}}


@pytest.fixture(autouse=True)
def _fast_passwords(settings):
    settings.PASSWORD_HASHERS = ["django.contrib.auth.hashers.MD5PasswordHasher"]


@pytest.fixture(autouse=True)
def _fresh_cache():
    rv.ROSTER_VIEW_CACHE.clear()
    yield
    rv.ROSTER_VIEW_CACHE.clear()


@pytest.fixture
def run():
    build_population()
    return build_saved_run()


def _user(model, username, role=None, **kwargs):
    ensure_role_groups()
    user = model.objects.create_user(username=username, password="x-Pass-1234", **kwargs)
    if role:
        user.groups.add(Group.objects.get(name=role))
    return user


@pytest.fixture
def committee(client, django_user_model):
    client.force_login(_user(django_user_model, "exam.committee1", ROLE_EXAM_COMMITTEE))
    return client


def _post(client, run_id, payload, view=DETAIL, **extra):
    return client.post(
        reverse(view, args=[run_id]), json.dumps(payload), content_type="application/json", **extra
    )


def _index(client, run_id, **params):
    return client.get(reverse(INDEX, args=[run_id]), params)


def _tokens(text: str) -> set[str]:
    return set(re.findall(r"[A-Za-z0-9.\-]+", text))


def _names_a_student(text: str) -> bool:
    return bool({str(sid) for sid in ALL_IDS} & _tokens(text)) or any(
        student_name(sid) in text for sid in ALL_IDS
    )


def _audit_rows(action=None):
    rows = AuditLog.objects.order_by("id")
    return list(rows.filter(action=action) if action else rows)


# ── Access ─────────────────────────────────────────────────────


def test_every_view_is_on_the_committee_allow_list_under_its_resolved_name(run):
    assert set(ROSTER_VIEWS) <= ExamCommitteeAccessMiddleware.allowed_views
    assert resolve(reverse(PAGE)).view_name == PAGE
    for name in ROSTER_VIEWS[1:]:
        assert resolve(reverse(name, args=[run.pk])).view_name == name


def test_committee_and_super_admin_reach_every_view(run, committee, django_user_model):
    admin = Client()
    admin.force_login(django_user_model.objects.create_superuser(username="root-admin"))
    for client in (committee, admin):
        assert client.get(reverse(PAGE), {"run": run.pk}).status_code == 200
        assert _index(client, run.pk).status_code == 200
        assert _post(client, run.pk, COURSE).status_code == 200
        assert _post(client, run.pk, {"student_id": OVERLOADED}, LOOKUP).status_code == 200


@pytest.mark.parametrize("role", [ROLE_ADVISOR, ROLE_STUDENT, None])
def test_other_roles_are_refused_before_any_list_is_read(
    run, client, django_user_model, role, monkeypatch
):
    monkeypatch.setattr(views, "ROSTER_VIEW_CACHE", None)  # any use would raise
    client.force_login(_user(django_user_model, f"user-{role}", role))
    responses = [
        client.get(reverse(PAGE), {"run": run.pk}),
        _index(client, run.pk),
        _post(client, run.pk, COURSE),
        _post(client, run.pk, {"student_id": OVERLOADED}, LOOKUP),
    ]
    assert [r.status_code for r in responses] == [403] * 4
    assert not any(_names_a_student(r.content.decode()) for r in responses)
    assert not AuditLog.objects.exists()


def test_anonymous_requests_go_to_login_and_the_redirect_names_no_student(run, client):
    responses = [
        client.get(reverse(PAGE), {"run": run.pk}),
        _index(client, run.pk),
        _post(client, run.pk, COURSE),
        _post(client, run.pk, {"student_id": OVERLOADED}, LOOKUP),
        _post(client, run.pk, {"query": student_name(OVERLOADED)}, LOOKUP),
    ]
    for response in responses:
        assert response.status_code == 302 and "/login" in response["Location"]
        assert not _names_a_student(response["Location"])
        assert "testname" not in response["Location"].lower()
    assert f"run%3D{run.pk}" in responses[0]["Location"]


def test_posts_need_csrf_and_each_view_answers_only_its_method(run, django_user_model):
    strict = Client(enforce_csrf_checks=True)
    strict.force_login(_user(django_user_model, "strict-committee", ROLE_EXAM_COMMITTEE))
    assert _post(strict, run.pk, COURSE).status_code == 403
    assert _post(strict, run.pk, {"student_id": OVERLOADED}, LOOKUP).status_code == 403
    assert not AuditLog.objects.exists()
    assert strict.get(reverse(DETAIL, args=[run.pk])).status_code == 405
    assert strict.get(reverse(LOOKUP, args=[run.pk]), {"student_id": OVERLOADED}).status_code == 405
    assert strict.post(reverse(INDEX, args=[run.pk])).status_code in (403, 405)
    assert strict.post(reverse(PAGE)).status_code in (403, 405)
    token_client = Client()
    token_client.force_login(_user(django_user_model, "token-committee", ROLE_EXAM_COMMITTEE))
    assert token_client.post(reverse(INDEX, args=[run.pk])).status_code == 405
    assert token_client.post(reverse(PAGE)).status_code == 405


# ── No student in a URL ────────────────────────────────────────


def _named_patterns(resolver=None, prefix=""):
    resolver = resolver or get_resolver()
    for entry in resolver.url_patterns:
        if isinstance(entry, URLResolver):
            yield from _named_patterns(entry, prefix + str(entry.pattern))
        elif isinstance(entry, URLPattern) and entry.name in ROSTER_VIEWS:
            yield entry.name, prefix + str(entry.pattern), entry.pattern.converters


def test_routes_carry_the_run_id_and_nothing_else():
    routes = {name: (route, set(converters)) for name, route, converters in _named_patterns()}
    assert routes == {
        PAGE: ("exam-timetable/rosters/", set()),
        INDEX: ("ops/exam-timetable/<int:run_id>/rosters/index/", {"run_id"}),
        DETAIL: ("ops/exam-timetable/<int:run_id>/rosters/", {"run_id"}),
        LOOKUP: ("ops/exam-timetable/<int:run_id>/rosters/lookup/", {"run_id"}),
    }


def test_no_view_reads_a_student_from_the_query_string(run, committee):
    sid = str(OVERLOADED)
    # The lookup ignores the address: the ID must travel in the body.
    response = committee.post(
        reverse(LOOKUP, args=[run.pk]) + f"?student_id={sid}&query={sid}",
        "{}",
        content_type="application/json",
    )
    assert response.status_code == 400 and response.json()["field"] == "body"
    assert not AuditLog.objects.exists()
    # Neither the page nor the navigator acts on what an address carries. (The
    # sidebar's language form echoes the current address back as ``next``; the
    # page's own region must not.)
    page = committee.get(reverse(PAGE), {"run": run.pk, "student_id": sid, "q": "TESTNAME"})
    assert page.status_code == 200
    main = re.search(r'<main id="main-content".*?</main>', page.content.decode(), re.S)[0]
    assert f'data-run-id="{run.pk}"' in main
    assert sid not in main and "TESTNAME" not in main
    index = _index(committee, run.pk, student_id=sid)
    assert index.status_code == 200 and not _names_a_student(index.content.decode())
    for answer in (page, index):
        assert "Location" not in answer


def test_answers_with_student_data_are_never_redirects_or_cacheable(run, committee):
    for response in (
        _index(committee, run.pk),
        _post(committee, run.pk, COURSE),
        _post(committee, run.pk, {"student_id": OVERLOADED}, LOOKUP),
        _post(committee, run.pk, {"query": "4401"}, LOOKUP),
    ):
        assert response.status_code == 200 and "Location" not in response
        assert response["Cache-Control"] == "private, no-store"
        assert response["X-Content-Type-Options"] == "nosniff"


@pytest.mark.parametrize("page", [PAGE, "exam_timetable_page"])
def test_pages_that_show_student_rows_are_never_stored_for_back(run, committee, page):
    # Rows the page or the drawer showed must not come back from the browser's
    # back/forward cache after a sign-out on a shared exam-desk machine.
    response = committee.get(reverse(page), {"run": run.pk} if page == PAGE else {})
    assert response.status_code == 200
    directives = {part.strip() for part in response["Cache-Control"].split(",")}
    assert {"no-store", "private"} <= directives


# ── Gates and validation ───────────────────────────────────────


def test_a_missing_run_is_404_everywhere(committee):
    assert _index(committee, 987654).status_code == 404
    for view, body in ((DETAIL, COURSE), (LOOKUP, {"student_id": OVERLOADED})):
        response = _post(committee, 987654, body, view)
        assert response.status_code == 404 and response.json()["code"] == "not_found"
    assert not AuditLog.objects.exists()


def test_phase_one_gates_answer_409_and_nothing_is_audited(run, committee):
    data = saved_payload(run)
    for rows in data["section_enrollment"].values():
        for row in rows:
            row["term"] = "2"
    save_payload(run, data)
    for response in (
        _index(committee, run.pk),
        _post(committee, run.pk, COURSE),
        _post(committee, run.pk, {"student_id": OVERLOADED}, LOOKUP),
    ):
        assert response.status_code == 409
        body = response.json()
        assert body["code"] == "lists_term_mismatch"
        assert (body["live_term"], body["saved_term"]) == (["1448", "1"], ["1448", "2"])
    data["section_enrollment"] = {}
    save_payload(run, data)
    response = _post(committee, run.pk, COURSE)
    assert response.status_code == 409 and response.json()["code"] == "rebuild_required"
    assert not AuditLog.objects.exists()


def test_no_imported_lists_is_409_lists_unavailable(run, committee):
    StudentTermSection.objects.all().delete()
    for response in (_index(committee, run.pk), _post(committee, run.pk, COURSE)):
        assert response.status_code == 409 and response.json()["code"] == "lists_unavailable"


@pytest.mark.parametrize(
    "view, body, field",
    [
        (DETAIL, {"scope": {"kind": "all"}}, "scope.kind"),
        (DETAIL, {"scope": {"kind": "course", "exam": "NOPE"}}, "scope.exam"),
        (DETAIL, {**COURSE, "student_id": OVERLOADED}, "body"),
        (DETAIL, {**COURSE, "refresh": "yes"}, "refresh"),
        (DETAIL, [1, 2], "body"),
        (LOOKUP, {"student_id": "abc"}, "student_id"),
        (LOOKUP, {"query": "ab"}, "query"),
        (LOOKUP, {"query": "e e e"}, "query"),
        (LOOKUP, {"query": "---"}, "query"),
        (LOOKUP, {"query": "MATH101"}, "query"),
        (LOOKUP, {"student_id": OVERLOADED, "query": "4401"}, "body"),
    ],
)
def test_bad_requests_are_400_named_against_their_field_and_not_audited(
    run, committee, view, body, field
):
    response = _post(committee, run.pk, body, view)
    assert response.status_code == 400
    assert response.json()["field"] == field
    assert not AuditLog.objects.exists()


@pytest.mark.parametrize(
    "view", [DETAIL, LOOKUP, "exam_student_export_preflight", "exam_student_export"]
)
def test_a_form_body_is_refused_by_its_type_never_a_500(run, django_user_model, view):
    # The CSRF check parses a form body before the view runs; a multipart
    # stream cannot then be read again. Refused as not JSON, never a crash.
    strict = Client(enforce_csrf_checks=True)
    strict.force_login(_user(django_user_model, "form-committee", ROLE_EXAM_COMMITTEE))
    strict.get(reverse(PAGE), {"run": run.pk})
    token = strict.cookies["csrftoken"].value
    url = reverse(view, args=[run.pk])
    multipart = strict.post(url, {"student_id": str(OVERLOADED)}, HTTP_X_CSRFTOKEN=token)
    encoded = strict.post(
        url,
        f"student_id={OVERLOADED}",
        content_type="application/x-www-form-urlencoded",
        HTTP_X_CSRFTOKEN=token,
    )
    # JSON only: even a valid JSON object is refused under another type.
    body = COURSE if view == DETAIL else {"student_id": str(OVERLOADED)} if view == LOOKUP else {}
    plain = strict.post(url, json.dumps(body), content_type="text/plain", HTTP_X_CSRFTOKEN=token)
    for response in (multipart, encoded, plain):
        assert response.status_code == 400
        assert response.json()["field"] == "body"
    assert not AuditLog.objects.exists()


@pytest.mark.parametrize(
    "reader, refusal",
    [
        (views._json_body, rv.RosterRequestError),
        (export_views._json_body, export.ExportOptionsError),
    ],
)
def test_a_body_another_reader_took_is_refused_not_a_500(reader, refusal):
    request = RequestFactory().post("/x/", json.dumps(COURSE), content_type="application/json")
    request.read()  # the stream is gone: Django refuses ``request.body`` now
    with pytest.raises(refusal) as caught:
        reader(request)
    assert caught.value.field == "body"


def test_bodies_over_the_cap_are_refused_before_they_are_read(run, committee):
    huge = {"student_id": "1" * 5000}
    for view in (DETAIL, LOOKUP):
        response = _post(committee, run.pk, huge, view)
        assert response.status_code == 400 and response.json()["field"] == "body"
    # A declared length over the cap is refused without reading the body at all.
    declared = committee.post(
        reverse(DETAIL, args=[run.pk]),
        json.dumps(COURSE),
        content_type="application/json",
        CONTENT_LENGTH=str(views.MAX_BODY_BYTES + 1),
    )
    assert declared.status_code == 400 and declared.json()["field"] == "body"
    fits = {**COURSE, "pad": "x"}
    assert _post(committee, run.pk, fits).json()["field"] == "body", "unknown key, not size"
    for malformed in ("{not json", "[" * 3000 + "]" * 3000, "\xff"):
        response = committee.post(
            reverse(LOOKUP, args=[run.pk]), malformed, content_type="application/json"
        )
        assert response.status_code == 400 and response.json()["field"] == "body"
    assert views.MAX_BODY_BYTES == 4096
    assert not AuditLog.objects.exists()


# ── The fail-closed audit ──────────────────────────────────────


def test_each_scope_load_writes_one_row_before_any_row_is_serialized(run, committee, monkeypatch):
    order: list[str] = []
    real_record, real_row = views.record_audit_event, rv.roster_row

    def recording_audit(**kwargs):
        order.append("audit")
        return real_record(**kwargs)

    def recording_row(*args, **kwargs):
        order.append("row")
        return real_row(*args, **kwargs)

    monkeypatch.setattr(views, "record_audit_event", recording_audit)
    monkeypatch.setattr(rv, "roster_row", recording_row)
    response = _post(committee, run.pk, COURSE)
    assert response.status_code == 200 and len(response.json()["rows"]) == 43
    assert order[0] == "audit" and order.count("audit") == 1 and order.count("row") == 43
    [entry] = _audit_rows()
    assert entry.action == "exam_timetable.roster_view"
    assert entry.method == "POST" and entry.endpoint == reverse(DETAIL, args=[run.pk])
    assert (entry.actor_username, entry.actor_role) == ("exam.committee1", ROLE_EXAM_COMMITTEE)
    details = json.loads(entry.details_json)
    assert details["scope"] == {"kind": "course", "value": "MATH101"}
    assert (details["rows"], details["students"], details["cached"]) == (43, 43, False)
    assert not _names_a_student(entry.details_json)
    # The navigator is not a scope load; a second scope load is its own row.
    _index(committee, run.pk)
    _post(committee, run.pk, {"scope": {"kind": "room", "slot_index": 0, "room_code": "M-A"}})
    assert [e.action for e in _audit_rows()] == ["exam_timetable.roster_view"] * 2
    assert json.loads(_audit_rows()[-1].details_json)["cached"] is True
    assert validate_hash_chain()["ok"]


def test_a_failed_audit_is_503_and_no_student_data_is_made(run, committee, monkeypatch):
    def refuse(*args, **kwargs):
        raise RuntimeError("audit table unavailable")

    monkeypatch.setattr(AuditLog.objects, "create", refuse)
    monkeypatch.setattr(rv, "roster_row", lambda *a, **k: pytest.fail("row serialized"))
    monkeypatch.setattr(views, "scope_payload", lambda *a, **k: pytest.fail("roster serialized"))
    monkeypatch.setattr(views, "lookup_payload", lambda *a, **k: pytest.fail("lookup serialized"))
    monkeypatch.setattr(views, "search_payload", lambda *a, **k: pytest.fail("search serialized"))
    for body, view in (
        (COURSE, DETAIL),
        ({"student_id": OVERLOADED}, LOOKUP),
        ({"query": "testname"}, LOOKUP),
    ):
        response = _post(committee, run.pk, body, view)
        assert response.status_code == 503
        assert response.json()["code"] == "audit_unavailable"
        assert set(response.json()) == {"ok", "code", "error"}
        assert not _names_a_student(response.content.decode())


def test_each_settled_lookup_writes_one_row_naming_its_students_by_id(run, committee, monkeypatch):
    order: list[str] = []
    real_record, real_row = views.record_audit_event, rv.roster_row
    monkeypatch.setattr(
        views, "record_audit_event", lambda **kw: order.append("audit") or real_record(**kw)
    )
    monkeypatch.setattr(rv, "roster_row", lambda *a, **k: order.append("row") or real_row(*a, **k))
    found = _post(committee, run.pk, {"student_id": str(OVERLOADED)}, LOOKUP).json()
    assert order == ["audit", "row", "row", "row", "row"]
    assert found["found"] is True and {r["student_id"] for r in found["rows"]} == {OVERLOADED}
    missing = _post(committee, run.pk, {"student_id": 4499999}, LOOKUP).json()
    assert missing["found"] is False and missing["rows"] == []
    search = _post(committee, run.pk, {"query": "4402"}, LOOKUP).json()
    assert search["mode"] == "search" and search["total"] == 12 and search["more"] is True
    assert len(search["matches"]) == rv.SEARCH_LIMIT
    named = _post(committee, run.pk, {"query": "TestName"}, LOOKUP).json()
    assert named["total"] == len(ALL_IDS)
    entries = _audit_rows("exam_timetable.roster_lookup")
    assert len(entries) == 4 and len(_audit_rows()) == 4
    details = [json.loads(e.details_json) for e in entries]
    assert all(
        set(d)
        == {
            "run_id",
            "mode",
            "search_kind",
            "search_length",
            "student_id",
            "shown_student_ids",
            "matches",
            "lists_code_now",
            "lists_checked_at",
            "cached",
        }
        for d in details
    )
    shown = [[m["student_id"] for m in answer["matches"]] for answer in (search, named)]
    # "Who viewed which student": the student asked for and each one shown, by
    # plain ID; of the search, only its kind, its length and how many matched.
    assert [
        (
            d["mode"],
            d["search_kind"],
            d["search_length"],
            d["student_id"],
            d["shown_student_ids"],
            d["matches"],
        )
        for d in details
    ] == [
        ("student", "id", 7, OVERLOADED, [OVERLOADED], 1),
        ("student", "id", 7, 4499999, [], 0),
        ("search", "id_prefix", 4, None, shown[0], 12),
        ("search", "name", 8, None, shown[1], 43),
    ]
    # As the Audit Explorer shows and exports the row: plain IDs, readable
    # without a tool; never the text typed or a name.
    for entry, ids in zip(entries, ([OVERLOADED], [4499999], shown[0], shown[1]), strict=True):
        assert {str(sid) for sid in ids} <= _tokens(entry.details_json)
        assert not any(student_name(sid) in entry.details_json for sid in ALL_IDS)
        assert "4402" not in _tokens(entry.details_json)
        assert "testname" not in entry.details_json.lower()
    assert validate_hash_chain()["ok"]


def test_lookup_answers_carry_the_student_and_facts_only(run, committee):
    body = _post(committee, run.pk, {"student_id": FEMALE_IS[0]}, LOOKUP).json()
    assert set(body) == {
        "ok",
        "run",
        "checked_at",
        "cached",
        "cache_ttl_seconds",
        "check",
        "mode",
        "found",
        "counts",
        "exams",
        "rooms",
        "departments",
        "rows",
    }
    assert all(tuple(row) == rv.ROSTER_ROW_KEYS for row in body["rows"])
    text = json.dumps(body)
    others = [sid for sid in ALL_IDS if sid != FEMALE_IS[0]]
    assert not {str(sid) for sid in others} & _tokens(text)


def test_scope_answers_have_their_documented_shape(run, committee):
    body = _post(committee, run.pk, COURSE).json()
    assert set(body) == {
        "ok",
        "run",
        "checked_at",
        "cached",
        "cache_ttl_seconds",
        "check",
        "scope",
        "room",
        "sections",
        "counts",
        "exams",
        "rooms",
        "departments",
        "rows",
    }
    assert (
        tuple(body["check"])
        == rv.CHECK_KEYS
        == (
            "status",
            "sections_changed",
            "sections_new",
            "sections_gone",
            "program_mix_changed",
            "lists_code_saved",
            "lists_code_now",
        )
    )
    index = _index(committee, run.pk).json()
    assert body["check"] == {key: index["check"][key] for key in rv.CHECK_KEYS}, "one build"
    assert body["run"]["id"] == run.pk and body["scope"]["exam"] == "MATH101"
    room = _post(
        committee, run.pk, {"scope": {"kind": "room", "slot_index": 0, "room_code": "M-A"}}
    ).json()
    assert room["room"]["room_code"] == "M-A" and room["room"]["exams"] == ["IS201"]
    index = _index(committee, run.pk).json()
    assert {"check", "days", "slots", "rooms", "exams", "not_assigned", "programs"} <= set(index)
    assert not _names_a_student(json.dumps(index))


# ── The viewing cache ──────────────────────────────────────────


class _Clock:
    now = 1000.0

    def __call__(self):
        return self.now


def test_views_share_a_60_second_cache_that_refresh_bypasses(run, committee, monkeypatch):
    clock = _Clock()
    cache = rv.RosterViewCache(clock=clock)
    monkeypatch.setattr(views, "ROSTER_VIEW_CACHE", cache)
    first = _index(committee, run.pk).json()
    assert first["cached"] is False and first["cache_ttl_seconds"] == 60
    clock.now += 30
    assert _post(committee, run.pk, COURSE).json()["cached"] is True
    lookup = _post(committee, run.pk, {"student_id": OVERLOADED}, LOOKUP).json()
    assert lookup["cached"] is True and lookup["checked_at"] == first["checked_at"]
    # The lists change: the cached view does not see it until Refresh.
    StudentTermSection.objects.filter(term_section__section="F2").delete()
    assert _index(committee, run.pk).json()["check"]["status"] == "matches"
    refreshed = _index(committee, run.pk, refresh="1").json()
    assert refreshed["cached"] is False and refreshed["check"]["status"] == "changed"
    assert refreshed["checked_at"] > first["checked_at"]
    clock.now += 61
    assert _post(committee, run.pk, COURSE).json()["cached"] is False
    clock.now += 1
    assert _post(committee, run.pk, {**COURSE, "refresh": True}).json()["cached"] is False
    assert (
        _post(committee, run.pk, {"student_id": 1234, "refresh": True}, LOOKUP).json()["cached"]
        is False
    )


def test_downloads_never_read_the_viewing_cache(run, committee, monkeypatch):
    assert _index(committee, run.pk).json()["check"]["status"] == "matches"
    StudentTermSection.objects.filter(term_section__section="F2").delete()
    monkeypatch.setattr(rv.RosterViewCache, "get", lambda *a, **k: pytest.fail("cache read"))
    response = committee.post(
        reverse("exam_student_export_preflight", args=[run.pk]),
        "{}",
        content_type="application/json",
    )
    assert response.status_code == 200 and response.json()["check"]["status"] == "changed"
    download = committee.post(
        reverse("exam_student_export", args=[run.pk]),
        json.dumps({"scope": {"kind": "course", "exam": "CS101"}, "language": "en"}),
        content_type="application/json",
    )
    assert download.status_code == 200


def test_a_build_held_by_another_request_answers_503_busy(run, committee, monkeypatch):
    cache = rv.RosterViewCache(wait=0)
    monkeypatch.setattr(views, "ROSTER_VIEW_CACHE", cache)
    assert cache.build_lock.acquire()
    try:
        responses = [
            _index(committee, run.pk),
            _post(committee, run.pk, COURSE),
            _post(committee, run.pk, {"student_id": OVERLOADED}, LOOKUP),
        ]
    finally:
        cache.build_lock.release()
    for response in responses:
        assert response.status_code == 503 and response.json()["code"] == "roster_busy"
    assert not AuditLog.objects.exists()


# ── Parity through the endpoints: the screen and the file ──────


def _file_records(content: bytes) -> list[dict]:
    book = load_workbook(BytesIO(content))
    sheet = book["Student exams"]
    ref = sheet.tables["StudentExams"].ref
    rows = [[cell.value for cell in row] for row in sheet[ref]]
    keys = {column.header("en"): column.key for column in export.STUDENT_EXAMS.columns}
    return [dict(zip([keys[h] for h in rows[0]], row, strict=True)) for row in rows[1:]]


@pytest.mark.parametrize(
    "scope",
    [
        {"kind": "course", "exam": "MATH101"},
        {"kind": "course", "exam": "CS101"},
        {"kind": "room", "slot_index": 0, "room_code": "M-A"},
        {"kind": "room", "slot_index": 0, "room_code": "F-A"},
    ],
)
def test_the_list_on_screen_is_the_downloaded_files_list(run, committee, scope):
    scraped_exam_registration(MALE_AI[1], "CS101", section_label="M9")  # a new section
    screen = _post(committee, run.pk, {"scope": scope}).json()
    download = committee.post(
        reverse("exam_student_export", args=[run.pk]),
        json.dumps({"scope": scope, "language": "en", "one_file_per_group": False}),
        content_type="application/json",
    )
    assert download.status_code == 200
    content = b"".join(download.streaming_content)
    records = _file_records(content)
    basis = {key: export._w(key, "en") for key in rv.SEATED | {"no_seat", "unassigned"}}
    assert [
        (
            r["student_id"],
            r["name"],
            r["program"],
            r["exam"],
            r["section"] or export._w(r["section_status"], "en"),
            r["room"],
            basis[r["room_basis"]],
            " · ".join(r["clash_with"]) or None,
            " · ".join(f"{screen['exams'][c]['start']} {c}" for c in r["same_day_with"]) or None,
            {"changed": "Section changed", "new": "New since save"}.get(r["change"]),
        )
        for r in screen["rows"]
    ] == [
        (
            f["student_id"],
            f["name"],
            f["program"],
            f["exam"],
            f["section"],
            f["exam_room"],
            f["room_basis"],
            f["clash_with"],
            f["same_day_with"],
            f["change"],
        )
        for f in records
    ]
    assert len(records) == screen["counts"]["rows"] > 0


# ── The page ───────────────────────────────────────────────────


def test_the_page_opens_the_asked_run_or_the_newest(run, committee):
    newer = build_saved_run(label="Newer save")
    for params, expected in (({"run": run.pk}, run), ({}, newer)):
        response = committee.get(reverse(PAGE), params)
        assert response.status_code == 200
        page = response.content.decode()
        assert f'data-run-id="{expected.pk}"' in page
        assert 'data-roster-state="ready"' in page
        for view in (INDEX, DETAIL, LOOKUP, "exam_student_export"):
            assert f'"{reverse(view, args=[expected.pk])}"' in page
        assert _names_a_student(page) is False


@pytest.mark.parametrize("raw", ["987654", "abc", "1e3", "-1", chr(0x0661) + chr(0x0662)])
def test_an_unknown_or_unreadable_run_is_a_404_page(run, committee, raw):
    response = committee.get(reverse(PAGE), {"run": raw})
    assert response.status_code == 404
    page = response.content.decode()
    assert 'data-roster-state="run_not_found"' in page and "data-run-id" not in page


def test_without_saved_runs_the_page_says_so(committee):
    assert not ExamTimetableRun.objects.exists()
    response = committee.get(reverse(PAGE))
    assert response.status_code == 200
    page = response.content.decode()
    assert 'data-roster-state="no_saved_run"' in page
    assert "No saved timetable yet. Student lists come from a saved timetable." in page


def test_the_page_is_arabic_only_when_arabic_is_asked_for(run, committee):
    english = committee.get(reverse(PAGE), {"run": run.pk}).content.decode()
    arabic = committee.get(
        reverse(PAGE), {"run": run.pk}, HTTP_ACCEPT_LANGUAGE="ar"
    ).content.decode()
    heading = ("Student lists", "قوائم الطلاب")
    audit = (
        "Views and downloads of student lists are recorded under your name.",
        "يُسجَّل عرض قوائم الطلاب وتنزيلها باسمك.",
    )
    for en, ar in (heading, audit):
        assert en in english and ar not in english
        assert ar in arabic and en not in arabic
    assert '<html lang="ar" dir="rtl"' in arabic and '<html lang="en" dir="ltr"' in english
    for page in (english, arabic):
        assert 'dir="auto"' not in page
        assert 'aria-current="page"' in page


class _PageCopy(HTMLParser):
    """The words the page hands its script (``#examRosterCopy``), as it reads them."""

    def __init__(self) -> None:
        super().__init__()
        self.words: dict[str, str] = {}

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if dict(attrs).get("id") == "examRosterCopy":
            self.words = {name[5:]: value or "" for name, value in attrs if name[:5] == "data-"}


#: The navigator's Needs review headings By room, one then many: each names
#: what it counts - sections with no room, students with no seat, online
#: rooms. Arabic is label-first ("عدد الطلاب: 3"), so one form fits any count.
REVIEW_HEADINGS = {
    "en": {
        "not-assigned-group": ("Not assigned · {n} section", "Not assigned · {n} sections"),
        "no-seat-group": ("No seat · {n} student", "No seat · {n} students"),
        "online-group": ("Online · {n} room", "Online · {n} rooms"),
    },
    "ar": {
        "not-assigned-group": ("لم تُحدَّد لها قاعة · عدد الشعب: {n}",) * 2,
        "no-seat-group": ("بلا مقعد · عدد الطلاب: {n}",) * 2,
        "online-group": ("عن بُعد · عدد القاعات: {n}",) * 2,
    },
}


@pytest.mark.parametrize("language", ["en", "ar"])
def test_needs_review_headings_name_their_unit_in_the_language_asked_for(run, committee, language):
    page = committee.get(
        reverse(PAGE), {"run": run.pk}, HTTP_ACCEPT_LANGUAGE=language
    ).content.decode()
    # Forced, not the site's default: the page and its words are in this language.
    assert f'<html lang="{language}"' in page
    reader = _PageCopy()
    reader.feed(page)
    assert reader.words["locale"] == language
    for key, (one, many) in REVIEW_HEADINGS[language].items():
        assert (reader.words[f"{key}-one"], reader.words[key]) == (one, many), key
    # Every other language's heading is absent: no English unit in Arabic.
    other = REVIEW_HEADINGS["ar" if language == "en" else "en"]
    assert not {form for pair in other.values() for form in pair} & set(reader.words.values())


def test_the_committee_page_keeps_the_export_endpoints_it_reuses(run, committee, monkeypatch):
    # The page hands the phase-1 export endpoints to its script; they stay fresh.
    spy = []
    monkeypatch.setattr(
        export_views, "build_roster_model", lambda r: spy.append(r.pk) or _real_build(r)
    )
    page = committee.get(reverse(PAGE), {"run": run.pk}).content.decode()
    url = re.search(r'data-export-preflight-url="([^"]+)"', page)[1]
    assert committee.post(url, "{}", content_type="application/json").status_code == 200
    assert spy == [run.pk]


def _real_build(run):
    from core.services.exam_rosters import build_roster_model

    return build_roster_model(run)
