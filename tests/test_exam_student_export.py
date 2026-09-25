"""The exam student workbook: schema, Excel safety, privacy and determinism.

Workbooks are made from a real built and saved run (see the fixture module)
and read back with openpyxl and as raw XML, because some promises - no formula
anywhere it does not belong, no forbidden value in any part - are only
checkable on the bytes.
"""

from __future__ import annotations

import ast
import dataclasses
import hashlib
import json
import os
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, time
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet import _writer as openpyxl_writer

from core.models import Course, ProgrammeRequirement, Student, StudentTermSection
from core.services import exam_student_export as export
from core.services.exam_rosters import CHANGED, RebuildRequired, build_roster_model
from core.services.exam_student_export import (
    EmptyScope,
    ExportOptionsError,
    parse_export_options,
    prepare_export,
    render_export,
)
from tests.exam_source_factory import scraped_exam_registration
from tests.exam_student_export_fixture import (
    ALL_IDS,
    FEMALE_CS2,
    FEMALE_IS,
    MALE_AI,
    MALE_CS,
    MALE_IS,
    NO_COHORT,
    SENTINELS,
    build_population,
    build_saved_run,
    save_payload,
    saved_payload,
    student_name,
)

pytestmark = pytest.mark.django_db

REFERENCE = "EXR-1A2B3C4D"
AUDIT_HASH = "1a2b3c4d" + "0" * 56


@pytest.fixture
def run():
    build_population()
    return build_saved_run()


def _prepare(run, **payload):
    model = build_roster_model(run)
    options = parse_export_options(
        {"scope": {"kind": "all"}, "one_file_per_group": False, **payload}, model
    )
    return prepare_export(
        model, options, generated_by="exam.committee1", generated_role="EXAM_COMMITTEE"
    )


def _export(run, reference=REFERENCE, **payload):
    prepared = _prepare(run, **payload)
    content, name, content_type = render_export(
        prepared, reference=reference, audit_hash=AUDIT_HASH
    )
    return content, name, content_type, prepared


def _book(content):
    return load_workbook(BytesIO(content))


def _parts(content) -> dict[str, str]:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        return {name: archive.read(name).decode("utf-8") for name in archive.namelist()}


def _sheet_xml(content) -> dict[str, str]:
    parts = _parts(content)
    book = _book(content)
    return {
        title: parts[f"xl/worksheets/sheet{index}.xml"]
        for index, title in enumerate(book.sheetnames, 1)
    }


def _rows(book, sheet, table):
    ws = book[sheet]
    ref = ws.tables[table].ref
    rows = [[cell.value for cell in row] for row in ws[ref]]
    return rows[0], rows[1:]


def _records(book, sheet, table):
    header, rows = _rows(book, sheet, table)
    return [dict(zip(header, row, strict=True)) for row in rows]


EN = {
    "StudentExams": [
        "Student ID",
        "Name",
        "Program",
        "Department",
        "Group",
        "Exam",
        "Course code",
        "Course name",
        "Section",
        "Day no.",
        "Day",
        "Weekday",
        "Date",
        "Period",
        "Start",
        "Exam room",
        "Building",
        "Floor",
        "Room basis",
        "Online",
        "Clash",
        "Clash with",
        "Same day",
        "Same day with",
        "Exams that day",
        "Change since save",
        "First exam that day",
    ],
    "ExamStudents": [
        "Student ID",
        "Name",
        "Program",
        "Department",
        "Group",
        "Exams (timetable)",
        "Exams (file)",
        "Exam days (file)",
        "Clashes (timetable)",
        "Days with 2+ exams (timetable)",
        "Most exams in a day (timetable)",
        "First exam (file)",
        "First exam date (file)",
        "Last exam (file)",
        "Last exam date (file)",
        "Without a seat (file)",
        "Exam schedule (file)",
    ],
    "ExamFlags": [
        "Student ID",
        "Name",
        "Program",
        "Group",
        "Day no.",
        "Day",
        "Date",
        "Flag",
        "Exams that day",
        "Exams",
    ],
    "ExamSections": [
        "Exam",
        "Course code",
        "Course name",
        "Section",
        "Group",
        "Day no.",
        "Day",
        "Date",
        "Period",
        "Online",
        "Section record",
        "Students at save",
        "Students now",
        "Difference",
        "Membership",
        "Program mix",
        "Selected programs at save",
        "Rows in this file",
        "No seat",
        "Clash students",
        "Same-day students",
        "Room count",
        "Programs in this file",
        "Rooms and ID ranges",
    ],
    "ExamRooms": [
        "Day no.",
        "Day",
        "Date",
        "Period",
        "Start",
        "Room",
        "Building",
        "Floor",
        "Group",
        "Capacity",
        "Seated at save",
        "Seated now",
        "Free seats",
        "Use",
        "Exams",
        "Online",
        "Rows in this file",
        "Sections in room",
    ],
    "DaySummary": [
        "Day no.",
        "Day",
        "Weekday",
        "Date",
        "Exams",
        "Sections",
        "Rooms used",
        "Seats",
        "Sittings",
        "Students",
        "Male",
        "Female",
        "Students with 2+ exams",
        "Clash students",
    ],
    "PeriodSummary": [
        "Day no.",
        "Day",
        "Date",
        "Period",
        "Start",
        "Exams",
        "Rooms used",
        "Seats",
        "Sittings",
        "Students",
        "Clash students",
        "Use",
    ],
    "ProgramDays": ["Program", "Department", "Students", "Sun", "Mon", "Not scheduled"],
    "ExportChecks": [
        "Check",
        "Saved (timetable screen)",
        "Now / this file",
        "Result",
        "Where on the screen",
        "Note",
    ],
    "ChangeLog": [
        "Exam",
        "Section",
        "Group",
        "Students at save",
        "Students now",
        "Difference",
        "What changed",
        "Effect in this file",
        "What to do",
    ],
    "FileInfo": ["Key", "Value", "Meaning"],
    "ColumnGuide": ["Sheet", "Column", "Meaning", "Type", "Example"],
    "SheetGuide": ["Sheet", "Rows", "One row is…"],
}

AR = {
    "StudentExams": [
        "الرقم الجامعي",
        "الاسم",
        "البرنامج",
        "القسم",
        "الفئة",
        "الاختبار",
        "رمز المقرر",
        "اسم المقرر",
        "الشعبة",
        "ترتيب اليوم",
        "اليوم",
        "يوم الأسبوع",
        "التاريخ",
        "الفترة",
        "وقت البدء",
        "قاعة الاختبار",
        "المبنى",
        "الدور",
        "أساس القاعة",
        "عن بُعد",
        "تعارض",
        "يتعارض مع",
        "اختبار آخر في اليوم نفسه",
        "في اليوم نفسه مع",
        "اختبارات ذلك اليوم",
        "تغيّر بعد الحفظ",
        "أول اختبار في اليوم",
    ],
    "ExamStudents": [
        "الرقم الجامعي",
        "الاسم",
        "البرنامج",
        "القسم",
        "الفئة",
        "الاختبارات (الجدول)",
        "الاختبارات (الملف)",
        "أيام الاختبار (الملف)",
        "التعارضات (الجدول)",
        "أيام بها اختباران أو أكثر (الجدول)",
        "أكثر اختبارات في يوم (الجدول)",
        "أول اختبار (الملف)",
        "تاريخ أول اختبار (الملف)",
        "آخر اختبار (الملف)",
        "تاريخ آخر اختبار (الملف)",
        "دون مقعد (الملف)",
        "جدول الاختبارات (الملف)",
    ],
    "ExamFlags": [
        "الرقم الجامعي",
        "الاسم",
        "البرنامج",
        "الفئة",
        "ترتيب اليوم",
        "اليوم",
        "التاريخ",
        "التنبيه",
        "اختبارات ذلك اليوم",
        "الاختبارات",
    ],
    "ExamSections": [
        "الاختبار",
        "رمز المقرر",
        "اسم المقرر",
        "الشعبة",
        "الفئة",
        "ترتيب اليوم",
        "اليوم",
        "التاريخ",
        "الفترة",
        "عن بُعد",
        "سجل الشعبة",
        "الطلاب عند الحفظ",
        "الطلاب الآن",
        "الفرق",
        "العضوية",
        "توزيع البرامج",
        "البرامج المختارة عند الحفظ",
        "الأسطر في هذا الملف",
        "بلا مقعد",
        "طلاب لديهم تعارض",
        "طلاب لديهم اختبار آخر في اليوم",
        "عدد القاعات",
        "البرامج في هذا الملف",
        "القاعات ونطاقات الأرقام",
    ],
    "ExamRooms": [
        "ترتيب اليوم",
        "اليوم",
        "التاريخ",
        "الفترة",
        "وقت البدء",
        "القاعة",
        "المبنى",
        "الدور",
        "الفئة",
        "السعة",
        "المقاعد عند الحفظ",
        "المقاعد الآن",
        "المقاعد الشاغرة",
        "الإشغال",
        "الاختبارات",
        "عن بُعد",
        "الأسطر في هذا الملف",
        "الشعب في القاعة",
    ],
    "DaySummary": [
        "ترتيب اليوم",
        "اليوم",
        "يوم الأسبوع",
        "التاريخ",
        "الاختبارات",
        "الشعب",
        "القاعات المستخدمة",
        "المقاعد",
        "الجلسات",
        "الطلاب",
        "طلاب",
        "طالبات",
        "طلاب لديهم اختباران أو أكثر",
        "طلاب لديهم تعارض",
    ],
    "PeriodSummary": [
        "ترتيب اليوم",
        "اليوم",
        "التاريخ",
        "الفترة",
        "وقت البدء",
        "الاختبارات",
        "القاعات المستخدمة",
        "المقاعد",
        "الجلسات",
        "الطلاب",
        "طلاب لديهم تعارض",
        "الإشغال",
    ],
    "ProgramDays": ["البرنامج", "القسم", "الطلاب", "Sun", "Mon", "غير مجدول"],
    "ExportChecks": [
        "التحقق",
        "المحفوظ (شاشة الجدول)",
        "الآن / هذا الملف",
        "النتيجة",
        "مكانها في الشاشة",
        "ملاحظة",
    ],
    "ChangeLog": [
        "الاختبار",
        "الشعبة",
        "الفئة",
        "الطلاب عند الحفظ",
        "الطلاب الآن",
        "الفرق",
        "ما الذي تغيّر",
        "الأثر في هذا الملف",
        "الإجراء",
    ],
    "FileInfo": ["المفتاح", "القيمة", "المعنى"],
    "ColumnGuide": ["الورقة", "العمود", "المعنى", "النوع", "مثال"],
    "SheetGuide": ["الورقة", "الأسطر", "السطر الواحد…"],
}

SHEETS = {
    "en": [
        "About",
        "Student exams",
        "Students",
        "Flags",
        "Sections",
        "Rooms",
        "By day",
        "Programs by day",
        "Checks and changes",
        "File info",
    ],
    "ar": [
        "حول الملف",
        "اختبارات الطلاب",
        "الطلاب",
        "التنبيهات",
        "الشعب",
        "القاعات",
        "حسب اليوم",
        "البرامج حسب الأيام",
        "المطابقة والتغييرات",
        "بيانات الملف",
    ],
}

FILE_INFO_KEYS = [
    "format",
    "format_version",
    "generator_commit",
    "run_id",
    "run_label",
    "run_saved_at_utc",
    "run_status",
    "run_input_fingerprint16",
    "enrollment_source",
    "academic_year",
    "term",
    "lists_code_saved",
    "lists_code_now",
    "lists_match",
    "sections_total",
    "sections_matching",
    "sections_changed",
    "sections_new",
    "sections_gone",
    "lists_checked_at_utc",
    "scope_kind",
    "scope_value",
    "programs",
    "groups",
    "file_group",
    "file_part",
    "rows_option",
    "contents",
    "language",
    "exam_dates_source",
    "prepared_for",
    "rows",
    "students",
    "exams",
    "sections_in_file",
    "data_sha256",
    "seat_rule",
    "flag_rule",
    "generated_at_utc",
    "generated_by",
    "reference",
    "audit_entry_hash16",
]


# ── Schema ─────────────────────────────────────────────────────


@pytest.mark.parametrize("language", ["en", "ar"])
def test_every_table_has_its_exact_ordered_columns_and_matching_table_columns(run, language):
    content, *_ = _export(run, language=language)
    book = _book(content)
    expected = EN if language == "en" else AR
    assert book.sheetnames == SHEETS[language]
    seen = {}
    for ws in book.worksheets:
        assert ws.auto_filter.ref is None
        for name in list(ws.tables):
            table = ws.tables[name]
            header, _rows_ = _rows(book, ws.title, name)
            assert header == expected[name], name
            assert [column.name for column in table.tableColumns] == expected[name]
            first, last = table.ref.split(":")
            assert first.rstrip("0123456789") == "A" and len(header) == len(table.tableColumns)
            seen[name] = ws.title
    assert set(seen) == set(expected)
    info = dict((row[0], row[1]) for row in _rows(book, seen["FileInfo"], "FileInfo")[1])
    assert list(info) == FILE_INFO_KEYS


def test_no_level_column_or_rule_exists_anywhere(run):
    for language in ("en", "ar"):
        content, *_ = _export(run, language=language)
        text = "".join(_parts(content).values())
        for word in ("Level", "level_rule", "المستوى", "intake"):
            assert word not in text, word


def test_data_tables_start_at_a1_and_refs_follow_explicit_row_counts(run):
    content, _name, _type, prepared = _export(run, language="en")
    book = _book(content)
    item = prepared.files[0]
    expected_rows = {
        "StudentExams": item.rows,
        "ExamStudents": item.students,
        "ExamSections": len(item.groups),
    }
    for sheet, table in [
        ("Student exams", "StudentExams"),
        ("Students", "ExamStudents"),
        ("Sections", "ExamSections"),
    ]:
        ref = book[sheet].tables[table].ref
        assert ref.startswith("A1:")
        assert int(re.sub(r"[A-Z]+", "", ref.split(":")[1])) == expected_rows[table] + 1


def test_summaries_only_keeps_the_order_and_drops_the_three_personal_sheets(run):
    content, name, *_ = _export(run, language="en", contents="summary")
    assert _book(content).sheetnames == [
        "About",
        "Sections",
        "Rooms",
        "By day",
        "Programs by day",
        "Checks and changes",
        "File info",
    ]
    assert "_summary_" in name


# ── Types and Excel safety ─────────────────────────────────────


def test_ids_are_integers_formatted_0_and_dates_times_are_typed(run):
    content, *_ = _export(run, language="en", dates={"Sun": "2026-12-13", "Mon": "2026-12-14"})
    book = _book(content)
    for sheet, table in [
        ("Student exams", "StudentExams"),
        ("Students", "ExamStudents"),
        ("Flags", "ExamFlags"),
    ]:
        ws = book[sheet]
        ref = ws.tables[table].ref
        cells = [row[0] for row in ws[ref]][1:]
        assert cells and all(isinstance(c.value, int) and c.number_format == "0" for c in cells)
    exams = book["Student exams"]
    header = [c.value for c in exams[1]]
    row = [c for c in exams[2]]
    date_cell = row[header.index("Date")]
    start_cell = row[header.index("Start")]
    assert date_cell.value.date() == date(2026, 12, 13) and date_cell.number_format == "yyyy-mm-dd"
    assert start_cell.value == time(8, 0) and start_cell.number_format == "hh:mm"
    rooms = book["Rooms"]
    use = rooms.cell(2, [c.value for c in rooms[1]].index("Use") + 1)
    assert isinstance(use.value, float) and use.number_format == "0%"
    # No placeholder dash in any cell. Scoped to <sheetData>: the print header
    # legitimately reads "Confidential — exam administration only", and whether
    # that dash is escaped depends on which XML writer openpyxl picked.
    sheet_data = re.search(r"<sheetData>.*</sheetData>", _sheet_xml(content)["Student exams"], re.S)
    assert sheet_data and "—" not in sheet_data.group(0)


def test_difference_is_a_signed_delta(run):
    scraped_exam_registration(MALE_AI[0], "IS201", section_label="M3")
    content, *_ = _export(run, language="en")
    sections = _book(content)["Sections"]
    header = [c.value for c in sections[1]]
    column = header.index("Difference") + 1
    cells = [sections.cell(r, column) for r in range(2, sections.max_row + 1)]
    changed = [c for c in cells if c.value == 1]
    assert changed and changed[0].number_format == "+0;-0;0"


def test_strings_starting_with_equals_stay_text_and_only_two_formulas_exist(run):
    Student.objects.filter(student_id=MALE_CS[0]).update(name='=HYPERLINK("x")')
    run.label = "=1+1"
    run.save(update_fields=["label"])
    content, *_ = _export(run, language="en", prepared_for="=cmd")
    book = _book(content)
    values = {}
    for ws in book.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if (
                    isinstance(cell.value, str)
                    and cell.value.startswith("=")
                    and cell.data_type == "s"
                ):
                    values[cell.value] = cell.data_type
    assert {'=HYPERLINK("x")', "=cmd"} <= set(values)
    formulas = {
        title: re.findall(r"<f>(.*?)</f>", xml) for title, xml in _sheet_xml(content).items()
    }
    assert set(t for t, f in formulas.items() if f) == {"Programs by day", "Checks and changes"}
    assert all(f.startswith("SUBTOTAL(109,ProgramDays[[") for f in formulas["Programs by day"])
    assert formulas["Checks and changes"] == ["COUNTA('Student exams'!A:A)-1"]


def test_sheet_names_are_fixed_legal_and_only_about_is_selected(run):
    for language in ("en", "ar"):
        content, *_ = _export(run, language=language)
        book = _book(content)
        assert len(set(book.sheetnames)) == len(book.sheetnames)
        for title in book.sheetnames:
            assert len(title) <= 31 and not re.search(r"[\[\]:*?/\\]", title)
            assert not title.startswith("'") and not title.endswith("'")
        selected = [ws.title for ws in book.worksheets if ws.sheet_view.tabSelected]
        assert selected == [book.sheetnames[0]] and book.active.title == book.sheetnames[0]


def test_page_headers_hold_no_free_text_and_escape_ampersands(run):
    run.label = "Final & secret LABEL-FREE-TEXT"
    run.save(update_fields=["label"])
    content, *_ = _export(run, language="en", prepared_for="Head & PREPARED-FREE-TEXT")
    for title, xml in _sheet_xml(content).items():
        header = re.search(r"<headerFooter>(.*?)</headerFooter>", xml, re.S).group(1)
        assert "FREE-TEXT" not in header, title
        assert len(re.sub(r"<[^>]+>", "", header)) < 255
        assert "&amp;P / &amp;N" in header
        assert "&amp;amp;" not in header or "&amp;&amp;" in header


def test_arabic_files_are_right_to_left_with_no_forced_reading_order(run):
    content, *_ = _export(run, language="ar")
    book = _book(content)
    assert all(ws.sheet_view.rightToLeft for ws in book.worksheets)
    assert 'readingOrder="2"' not in _parts(content)["xl/styles.xml"]
    about = book[SHEETS["ar"][0]]
    reference_row = next(row for row in about.iter_rows(values_only=True) if row[0] == "المرجع")
    assert "\u200e" + REFERENCE + "\u200f" in reference_row[1]
    assert "\u2066" not in "".join(_parts(content).values())


def test_internal_links_use_locations_never_external_targets(run):
    content, *_ = _export(run, language="en")
    parts = _parts(content)
    assert not any("hyperlink" in xml for name, xml in parts.items() if name.endswith(".rels"))
    about = parts["xl/worksheets/sheet1.xml"]
    assert "location=\"'Student exams'!A1\"" in about


def test_workbook_default_font_is_arial_11(run):
    content, *_ = _export(run, language="en")
    # Parsed, not string-matched: lxml writes <name val="Arial"/> and the stdlib
    # writer production uses writes <name val="Arial" />.
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    font0 = ET.fromstring(_parts(content)["xl/styles.xml"].encode("utf-8")).find(
        "m:fonts/m:font", ns
    )
    assert font0.find("m:name", ns).get("val") == "Arial"
    assert font0.find("m:sz", ns).get("val") == "11"


def test_the_writer_never_uses_max_row_and_is_write_only():
    tree = ast.parse(Path(export.__file__).read_text(encoding="utf-8"))
    names = {node.attr for node in ast.walk(tree) if isinstance(node, ast.Attribute)}
    assert "max_row" not in names
    assert export._new_workbook().write_only is True


def test_a_failure_mid_write_leaves_no_spool_file_with_student_rows(run, monkeypatch):
    prepared = _prepare(run, language="en")
    before = list(openpyxl_writer.ALL_TEMP_FILES)
    created = []
    original = openpyxl_writer.create_temporary_file

    def tracking(*args, **kwargs):
        path = original(*args, **kwargs)
        created.append(path)
        return path

    monkeypatch.setattr(openpyxl_writer, "create_temporary_file", tracking)

    def explode(*args, **kwargs):
        raise RuntimeError("forced failure while writing the Rooms sheet")

    monkeypatch.setattr(
        export,
        "_apply_formatting",
        lambda ws, spec, *a: explode() if spec.name == "ExamRooms" else None,
    )
    with pytest.raises(RuntimeError):
        render_export(prepared, reference=REFERENCE, audit_hash=AUDIT_HASH)
    assert created, "the writer never spooled, so the test proves nothing"
    assert openpyxl_writer.ALL_TEMP_FILES == before
    assert not any(Path(path).exists() for path in created)


# ── Rows: seat rule, flags, change markers ─────────────────────


def test_rows_follow_seat_order_and_carry_the_room_basis(run):
    content, *_ = _export(run, language="en")
    rows = _records(_book(content), "Student exams", "StudentExams")
    m1 = [r for r in rows if (r["Exam"], r["Section"]) == ("MATH101", "M1")]
    ids = [r["Student ID"] for r in m1]
    assert ids == sorted(ids)
    rooms = [r["Exam room"] for r in m1]
    boundary = rooms.index(rooms[-1])
    assert set(rooms[:boundary]) == {rooms[0]} and set(rooms[boundary:]) == {rooms[-1]}
    assert {r["Room basis"] for r in m1} == {"Split by student ID"}
    unrecorded = next(r for r in rows if r["Student ID"] == NO_COHORT)
    assert unrecorded["Room basis"] == "Room not assigned" and unrecorded["Exam room"] is None
    assert unrecorded["Section"] == "Not recorded" and unrecorded["Group"] == "Not recorded"


def test_clash_and_same_day_columns_are_independent(run):
    content, *_ = _export(run, language="en")
    rows = _records(_book(content), "Student exams", "StudentExams")
    by = {(r["Student ID"], r["Exam"]): r for r in rows}
    is_student = by[(MALE_IS[0], "MATH101")]
    assert (is_student["Clash"], is_student["Clash with"]) == ("Yes", "IS201")
    assert is_student["Same day"] == "No" and is_student["Exams that day"] == 2
    cs_student = by[(MALE_CS[1], "MATH101")]
    assert cs_student["Clash"] == "No" and cs_student["Same day"] == "Yes"
    assert cs_student["Same day with"] == "13:00 CS101"
    overloaded = by[(MALE_CS[0], "CS101")]
    assert overloaded["Same day with"] == "08:00 IS201 · 08:00 MATH101"
    assert overloaded["Exams that day"] == 3
    for row in rows:
        for column in ("Clash with", "Same day with"):
            value = row[column] or ""
            assert re.fullmatch(r"((\d\d:\d\d )?[A-Z]+\d+( \(\d\))?( · )?)*", value), value


def test_changed_and_new_sections_are_marked_on_every_affected_row(run):
    StudentTermSection.objects.filter(
        student_id=MALE_IS[-1], term_section__course_key="IS201"
    ).delete()
    scraped_exam_registration(MALE_AI[0], "IS201", section_label="M3")
    scraped_exam_registration(MALE_AI[1], "CS101", section_label="M9")
    content, *_ = _export(run, language="en")
    book = _book(content)
    rows = _records(book, "Student exams", "StudentExams")
    m3 = [r for r in rows if (r["Exam"], r["Section"]) == ("IS201", "M3")]
    assert m3 and {r["Change since save"] for r in m3} == {"Section changed"}
    new = next(r for r in rows if (r["Exam"], r["Section"]) == ("CS101", "M9"))
    assert new["Change since save"] == "New since save" and new["Room basis"] == "No seat"
    untouched = [r for r in rows if r["Exam"] == "PHYS103 (1)"]
    assert {r["Change since save"] for r in untouched} == {None}
    changes = _records(book, "Checks and changes", "ChangeLog")
    assert {(c["Exam"], c["Section"], c["What changed"]) for c in changes} == {
        ("IS201", "M3", "Changed"),
        ("CS101", "M9", "New since save"),
    }
    about = [r for r in book["About"].iter_rows(values_only=True) if r[0] == "Status"][0]
    assert about[1].startswith("≠ Changed since #")


def test_unchanged_lists_reconcile_every_scope_programme_and_group(run):
    data = saved_payload(run)
    for programs, groups in [([], ["M", "F", "U"]), (["IS"], ["M", "F"]), (["CS", "CS2"], ["F"])]:
        content, *_ = _export(
            run, language="en", programs=programs, groups=groups, one_file_per_group=False
        )
        checks = _records(_book(content), "Checks and changes", "ExportChecks")
        assert {c["Result"] for c in checks} <= {"OK", "Info"}, checks
        rows_check = checks[10]
        expected = sum(
            count
            for course in data["operations_snapshot"]["courses"].values()
            for section in course["sections"]
            if section["gender"] in groups
            for program, count in section["program_counts"].items()
            if not programs or program in programs
        )
        assert rows_check["Saved (timetable screen)"] == expected == rows_check["Now / this file"]
        assert rows_check["Result"] == "OK"
        changes = _records(_book(content), "Checks and changes", "ChangeLog")
        assert len(changes) == 1 and changes[0]["Exam"].startswith("No changes")


def test_first_exam_that_day_counts_students_within_the_file(run):
    content, *_ = _export(run, language="en", scope={"kind": "period", "slot_index": 1})
    rows = _records(_book(content), "Student exams", "StudentExams")
    assert {r["Exam"] for r in rows} == {"CS101"}
    assert sum(r["First exam that day"] for r in rows) == len({r["Student ID"] for r in rows})


def test_overflow_rows_are_not_scheduled_and_carry_no_day(run):
    data = saved_payload(run)
    entry = next(e for e in data["schedule"] if e["course_code"] == "CS101")
    entry.update(day="OVERFLOW", period="Extra-9", slot_index=9, rooms=[])
    save_payload(run, data)
    content, *_ = _export(run, language="en")
    rows = _records(_book(content), "Student exams", "StudentExams")
    cs101 = [r for r in rows if r["Exam"] == "CS101"]
    assert cs101 and {r["Room basis"] for r in cs101} == {"Not scheduled"}
    assert {(r["Day no."], r["Day"], r["Clash"], r["First exam that day"]) for r in cs101} == {
        (None, "Not scheduled", "No", None)
    }
    assert rows[-1]["Exam"] == "CS101"


def _status(book):
    return next(r[1] for r in book["About"].iter_rows(values_only=True) if r[0] == "Status")


def _first_cells(book, sheet, table):
    _header, rows = _rows(book, sheet, table)
    return [row[0] for row in rows]


def test_a_changed_section_with_no_row_in_the_file_is_listed_explained_and_reconciled(run):
    # Every IS student leaves MATH101 M1; its CS and AI students stay, so an
    # IS file of MATH101 has no M1 row, yet M1 is what changed for it.
    StudentTermSection.objects.filter(
        student_id__in=MALE_IS, term_section__course_key="MATH101"
    ).delete()
    for rows in ("all", "flagged"):
        content, *_ = _export(
            run,
            language="en",
            scope={"kind": "course", "exam": "MATH101"},
            programs=["IS"],
            rows=rows,
        )
        book = _book(content)
        sections = {
            (s["Section"], s["Group"]): s for s in _records(book, "Sections", "ExamSections")
        }
        m1 = sections[("M1", "Male")]
        assert (m1["Membership"], m1["Rows in this file"], m1["Selected programs at save"]) == (
            "Changed",
            0,
            10,
        ), rows
        changes = _records(book, "Checks and changes", "ChangeLog")
        assert [(c["Exam"], c["Section"], c["What changed"]) for c in changes] == [
            ("MATH101", "M1", "Changed")
        ], rows
        if rows == "all":
            # Check 11's difference is exactly what the Sections sheet shows.
            check = _records(book, "Checks and changes", "ExportChecks")[10]
            assert check["Result"] == "Differs"
            assert check["Saved (timetable screen)"] == sum(
                s["Selected programs at save"] for s in sections.values()
            )
            assert check["Now / this file"] == sum(
                s["Rows in this file"] for s in sections.values()
            )
    # A changed section with no student of the file's programmes, saved or
    # live, is another programme's change: not this file's.
    StudentTermSection.objects.filter(
        student_id=MALE_CS[1], term_section__course_key="CS101"
    ).delete()
    content, *_ = _export(run, language="en", programs=["IS"])
    changed = {
        (c["Exam"], c["Section"])
        for c in _records(_book(content), "Checks and changes", "ChangeLog")
    }
    assert ("MATH101", "M1") in changed and ("CS101", "M2") not in changed


def test_a_file_whose_sections_did_not_change_never_says_every_section_matches(run):
    StudentTermSection.objects.filter(term_section__section="F2").delete()
    # A programme change elsewhere is a difference elsewhere too.
    Student.objects.filter(student_id=MALE_AI[0]).update(program="CS")
    content, *_ = _export(run, language="en", scope={"kind": "course", "exam": "PHYS103 (1)"})
    book = _book(content)
    assert _first_cells(book, "Checks and changes", "ChangeLog") == [
        f"No changes in this file's sections; 2 sections elsewhere in the timetable "
        f"differ from saved timetable #{run.pk}."
    ]
    assert _status(book).startswith("≠ Changed since #")
    content, *_ = _export(run, language="ar", scope={"kind": "course", "exam": "PHYS103 (1)"})
    assert _first_cells(_book(content), SHEETS["ar"][8], "ChangeLog") == [
        "لا تغييرات في شعب هذا الملف. الشعب المختلفة في بقية الجدول عن الجدول المحفوظ رقم "
        f"\u200e{run.pk}\u200f: 2."
    ]


@pytest.mark.parametrize(
    "student, program, elsewhere, differ_en",
    [
        (MALE_AI[0], "CS", 1, "1 section elsewhere in the timetable differs"),
        (MALE_CS[5], "CS2", 2, "2 sections elsewhere in the timetable differ"),
    ],
)
def test_a_programme_change_only_outside_the_file_is_no_change_here_and_never_a_match(
    run, student, program, elsewhere, differ_en
):
    # The same students everywhere, so the lists codes agree (lists_match);
    # the programme counts of sections outside PHYS103 (1) do not (unchanged
    # is False). Only the second may decide "every section matches".
    Student.objects.filter(student_id=student).update(program=program)
    model = build_roster_model(run)
    assert model.lists_match and not model.unchanged
    assert len(model.differing_groups) == elsewhere
    assert all(g.exam != "PHYS103 (1)" for g in model.differing_groups)
    scope = {"kind": "course", "exam": "PHYS103 (1)"}
    content, *_ = _export(run, language="en", scope=scope)
    assert _first_cells(_book(content), "Checks and changes", "ChangeLog") == [
        f"No changes in this file's sections; {differ_en} from saved timetable #{run.pk}."
    ]
    content, *_ = _export(run, language="ar", scope=scope)
    assert _first_cells(_book(content), SHEETS["ar"][8], "ChangeLog") == [
        "لا تغييرات في شعب هذا الملف. الشعب المختلفة في بقية الجدول عن الجدول المحفوظ رقم "
        f"\u200e{run.pk}\u200f: {elsewhere}."
    ]
    # The whole timetable holds the changed sections: no part of it says they match.
    said_to_match = {
        "en": ("every section matches", "Student lists match"),
        "ar": ("كل الشعب مطابقة", "قوائم الطلاب مطابقة"),
    }
    for language, phrases in said_to_match.items():
        content, *_ = _export(run, language=language)
        for name, text in _parts(content).items():
            assert not [phrase for phrase in phrases if phrase in text], (language, name)


def test_a_changed_section_belongs_to_a_programme_file_by_its_saved_or_live_students(run):
    model = build_roster_model(run)
    m1 = next(g for g in model.groups if (g.exam, g.section) == ("MATH101", "M1"))
    options = parse_export_options({"scope": {"kind": "all"}, "programs": ["IS"]}, model)
    everyone = ("M", "F", "U")

    def belongs(**changes):
        group = dataclasses.replace(m1, **changes)
        return export._changed_in_file(group, options, model, everyone)

    assert not belongs(), "an unchanged section is not a change"
    changed = {"membership": CHANGED}
    assert belongs(**changed, saved_program_counts={"IS": 3}, live_program_counts={})
    assert belongs(**changed, saved_program_counts={}, live_program_counts={"IS": 1})
    assert not belongs(**changed, saved_program_counts={"CS": 3}, live_program_counts={"CS": 4})
    assert belongs(live_program_counts={"IS": 11}), "a programme-mix change is a change"
    assert not export._changed_in_file(
        dataclasses.replace(m1, **changed), options, model, ("F",)
    ), "another group's section"
    other = parse_export_options({"scope": {"kind": "course", "exam": "CS101"}}, model)
    assert not export._changed_in_file(
        dataclasses.replace(m1, **changed), other, model, everyone
    ), "another exam's section"


def test_a_programme_change_alone_is_a_change_wherever_the_status_is_said(run):
    # Same students, so the lists codes still agree; the saved programme
    # counts do not, and neither About, the dialog nor the audit says "matches".
    Student.objects.filter(student_id=MALE_CS[5]).update(program="CS2")
    content, _name, _type, prepared = _export(
        run, language="en", scope={"kind": "course", "exam": "MATH101"}, programs=["CS2"]
    )
    book = _book(content)
    assert _status(book).startswith("≠ Changed since #")
    assert "2 of 9 sections differ" in _status(book)
    assert prepared.audit_details()["check"] == {
        "status": "changed",
        "sections_changed": 0,
        "program_mix_changed": 2,
        "exams_missing": 0,
        "no_seat": 0,
    }
    summary = export.preflight_summary(prepared.model, None)
    assert (summary["check"]["status"], summary["check"]["program_mix_changed"]) == ("changed", 2)
    info = {r["Key"]: r["Value"] for r in _records(book, "File info", "FileInfo")}
    assert info["lists_match"] == "true", "lists_match is the lists codes: students only"
    changes = _records(book, "Checks and changes", "ChangeLog")
    assert [(c["Exam"], c["What changed"]) for c in changes] == [("MATH101", "Program mix changed")]


def test_the_change_log_says_why_an_unscheduled_or_unroomed_exam_seats_nobody(run):
    scraped_exam_registration(FEMALE_IS[0], "CS101", section_label="F2")
    scraped_exam_registration(FEMALE_IS[1], "CS101", section_label="F9")

    def effects():
        content, *_ = _export(run, language="en", scope={"kind": "course", "exam": "CS101"})
        changes = _records(_book(content), "Checks and changes", "ChangeLog")
        return {c["Section"]: (c["Effect in this file"], c["What to do"]) for c in changes}

    # Scheduled and roomed: the extra student and the new section have no seat.
    assert effects() == {
        "F2": (
            "1 students have no seat; rooms were sized for 6.",
            "Timetable › Check changes › Save to resize rooms.",
        ),
        "F9": (
            "1 students have no seat: this section did not exist when the timetable was saved.",
            "Timetable › Check changes › Save to resize rooms.",
        ),
    }
    data = saved_payload(run)
    entry = next(e for e in data["schedule"] if e["course_code"] == "CS101")
    entry.update(day="OVERFLOW", period="Extra-9", slot_index=9, rooms=[])
    save_payload(run, data)
    assert effects() == {
        "F2": (
            "1 more students than saved; the exam is not scheduled.",
            "Timetable › Check changes › Save to update the saved counts.",
        ),
        "F9": (
            "1 students: this section did not exist when the timetable was saved, "
            "and the exam is not scheduled.",
            "Timetable › Check changes › Save to update the saved counts.",
        ),
    }
    data = saved_payload(run)
    for item in data["schedule"]:
        item["rooms"] = []
    entry = next(e for e in data["schedule"] if e["course_code"] == "CS101")
    entry.update(day="Sun", period="13:00-15:00", slot_index=1)
    data["assign_rooms"] = False
    save_payload(run, data)
    assert effects() == {
        "F2": (
            "1 more students than saved; rooms were not assigned in this timetable.",
            "Timetable › Check changes › Save to update the saved counts.",
        ),
        "F9": (
            "1 students: this section did not exist when the timetable was saved, "
            "and rooms were not assigned in this timetable.",
            "Timetable › Check changes › Save to update the saved counts.",
        ),
    }


@pytest.fixture
def shared_label_run():
    """Adds SHR101, whose one section label names no cohort, taken by both."""
    build_population()
    course = Course.objects.create(course_code="SHR101", description="SHARED", credit_hours=3)
    for program in ("AI", "IS"):
        ProgrammeRequirement.objects.create(
            program=program,
            course_code="SHR101",
            course_name="SHARED",
            programme_term=1,
            credit_hours=3,
        )
    for sid in [*MALE_AI[:4], *FEMALE_IS[:4]]:
        scraped_exam_registration(sid, course, section_label="01")
    return build_saved_run()


def test_a_label_shared_by_both_cohorts_is_counted_and_numbered_like_the_build(shared_label_run):
    data = saved_payload(shared_label_run)
    shared = [d for d in data["qa"]["multi_sitting_details"] if d["course_code"] == "SHR101"]
    assert [(d["sittings"], d["gender"]) for d in shared] == [(2, "M/F")], "the build's own view"
    content, *_ = _export(shared_label_run, language="en")
    book = _book(content)
    check = _records(book, "Checks and changes", "ExportChecks")[9]
    assert check["Check"] == "Sections split across rooms"
    assert (check["Saved (timetable screen)"], check["Now / this file"], check["Result"]) == (
        2,
        2,
        "OK",
    )
    sections = {
        (s["Exam"], s["Section"], s["Group"]): s for s in _records(book, "Sections", "ExamSections")
    }
    # Each cohort sits whole in one room: its part is 1 of 1, as its rows say.
    for group, ids in (("Male", MALE_AI[:4]), ("Female", FEMALE_IS[:4])):
        ranges = sections[("SHR101", "01", group)]["Rooms and ID ranges"]
        assert re.fullmatch(rf"[MF]-[ABC] \(1 of 1\): {ids[0]}–{ids[-1]} \(4\)", ranges), ranges
    rows = _records(book, "Student exams", "StudentExams")
    assert {r["Room basis"] for r in rows if r["Exam"] == "SHR101"} == {"Whole section"}
    m1 = sections[("MATH101", "M1", "Male")]["Rooms and ID ranges"]
    assert re.fullmatch(
        r"M-[ABC] \(1 of 2\): \d+–\d+ \(\d+\) · M-[ABC] \(2 of 2\): \d+–\d+ \(\d+\)", m1
    ), m1


# ── Privacy ────────────────────────────────────────────────────


def _tokens(text: str) -> set[str]:
    return set(re.findall(r"[A-Za-z0-9.\-]+", text))


SCOPES = [
    {"kind": "all"},
    {"kind": "course", "exam": "MATH101"},
    {"kind": "day", "day": "Sun"},
    {"kind": "period", "slot_index": 0},
]


@pytest.mark.parametrize("contents", ["full", "summary"])
@pytest.mark.parametrize("language", ["en", "ar"])
def test_no_forbidden_student_value_reaches_any_part_of_any_file(run, contents, language):
    sentinels = {str(value) for value in SENTINELS.values()}
    for scope in SCOPES:
        content, *_ = _export(run, language=language, contents=contents, scope=scope)
        text = "".join(_parts(content).values())
        assert not sentinels & _tokens(text), scope


def test_summaries_only_contain_no_student_id_or_name(run):
    content, *_ = _export(run, language="en", contents="summary")
    text = "".join(_parts(content).values())
    assert not {str(sid) for sid in ALL_IDS} & _tokens(text)
    assert not any(student_name(sid) in text for sid in ALL_IDS)


def test_a_department_file_never_shows_another_programmes_ids_even_in_ranges(run):
    content, _name, _type, prepared = _export(
        run, language="en", programs=["CS", "CS2"], one_file_per_group=False
    )
    allowed = {str(sid) for sid in [*MALE_CS, *FEMALE_CS2, NO_COHORT]}
    others = {str(sid) for sid in [*MALE_IS, *MALE_AI, *FEMALE_IS]}
    text = "".join(_parts(content).values())
    assert not others & _tokens(text)
    sections = _records(_book(content), "Sections", "ExamSections")
    m1 = next(s for s in sections if (s["Exam"], s["Section"]) == ("MATH101", "M1"))
    endpoints = set(re.findall(r"\d{7}", m1["Rooms and ID ranges"]))
    assert endpoints and endpoints <= allowed
    assert m1["Programs in this file"] == "CS 10"
    assert m1["Selected programs at save"] == 10 and m1["Students at save"] == 30


def test_scoped_files_list_only_their_own_exams_rooms_and_sections(run):
    content, *_ = _export(run, language="en", scope={"kind": "course", "exam": "CS101"})
    book = _book(content)
    assert {r["Exam"] for r in _records(book, "Student exams", "StudentExams")} == {"CS101"}
    assert {r["Exam"] for r in _records(book, "Sections", "ExamSections")} == {"CS101"}
    rooms = _records(book, "Rooms", "ExamRooms")
    assert rooms and all(r["Sections in room"].startswith("CS101 ") for r in rooms)
    content, *_ = _export(
        run, language="en", scope={"kind": "room", "slot_index": 0, "room_code": "M-B"}
    )
    rows = _records(_book(content), "Student exams", "StudentExams")
    assert rows and {r["Exam room"] for r in rows} == {"M-B"}


# ── Characters XML cannot hold ─────────────────────────────────

XML_ILLEGAL = "\ufffe\uffff\x0b\x1f"


def test_characters_xml_cannot_hold_never_reach_a_cell(run):
    run.label = f"Final{XML_ILLEGAL} v3"
    run.save(update_fields=["label"])
    Student.objects.filter(student_id=MALE_CS[0]).update(name=f"BAD{XML_ILLEGAL}NAME")
    data = saved_payload(run)
    entry = next(e for e in data["schedule"] if e["course_code"] == "CS101")
    entry["course_name"] = "PROGRAMMING\uffff I"
    # A day label heads a Programs-by-day column, and names its Table column.
    for item in [*data["slots"], *data["schedule"]]:
        if item["day"] == "Mon":
            item["day"] = "Mon\uffff"
    save_payload(run, data)
    content, *_ = _export(run, language="en")
    for text in _parts(content).values():
        assert not export.XML_ILLEGAL_RE.search(text)
    book = _book(content)
    info = {r["Key"]: r["Value"] for r in _records(book, "File info", "FileInfo")}
    assert info["run_label"] == "Final v3"
    rows = _records(book, "Student exams", "StudentExams")
    # The roster folds the whitespace controls into one space; the rest is dropped.
    assert {r["Name"] for r in rows if r["Student ID"] == MALE_CS[0]} == {"BAD NAME"}
    assert {r["Course name"] for r in rows if r["Exam"] == "CS101"} == {"PROGRAMMING I"}
    header, _rows_ = _rows(book, "Programs by day", "ProgramDays")
    assert "Mon" in header
    assert [c.name for c in book["Programs by day"].tables["ProgramDays"].tableColumns] == header


def test_a_lone_surrogate_in_the_saved_run_is_stripped_and_the_workbook_opens(run):
    # A JSON text may carry an unpaired \uXXXX escape: the database holds the
    # ASCII escape, and loading it gives a character UTF-8 cannot encode and
    # XML cannot hold. A course name carrying one is refused before any cell;
    # the saved status reaches File info as it is, so only the cell cleaning
    # stands between it and a workbook that cannot be written or opened.
    data = saved_payload(run)
    status = data["primary_status"]
    assert status
    entry = next(e for e in data["schedule"] if e["course_code"] == "CS101")
    entry["course_name"] = "PROGRAMMING\ud800 I"
    run.result_json = json.dumps(data)
    run.save(update_fields=["result_json"])
    with pytest.raises(RebuildRequired):
        build_roster_model(run)
    entry["course_name"] = "PROGRAMMING I"
    data["primary_status"] = f"{status[:2]}\ud800{status[2:]}\udfff"
    run.result_json = json.dumps(data)
    run.save(update_fields=["result_json"])
    for language in ("en", "ar"):
        content, *_ = _export(run, language=language)
        for text in _parts(content).values():
            assert not re.search("[\ud800-\udfff]", text)
        info_sheet = "File info" if language == "en" else SHEETS["ar"][9]
        _header, rows = _rows(_book(content), info_sheet, "FileInfo")
        info = {row[0]: row[1] for row in rows}
        assert info["run_status"] == status, language


def test_without_lxml_as_deployed_only_cleaned_text_saves_a_readable_workbook():
    """Production has no lxml: openpyxl's stdlib writer escapes nothing it cannot hold."""
    script = """
import io, re, sys
import openpyxl.xml
from openpyxl import Workbook, load_workbook
assert openpyxl.xml.LXML is False, "the stdlib writer must be under test"
pattern = re.compile(sys.argv[1])
def saved(text):
    book = Workbook(write_only=True)
    sheet = book.create_sheet("S")
    sheet.append([text])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
raw = "Dean \\uffff office \\ufffe"
try:
    load_workbook(io.BytesIO(saved(raw)))
except Exception:
    pass
else:
    raise SystemExit("an uncleaned workbook opened, so this proves nothing")
book = load_workbook(io.BytesIO(saved(pattern.sub("", raw))))
assert book["S"]["A1"].value == "Dean  office ", book["S"]["A1"].value
print("ok")
"""
    env = {**os.environ, "OPENPYXL_LXML": "False"}
    result = subprocess.run(
        [sys.executable, "-c", script, export.XML_ILLEGAL_RE.pattern],
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=120,
        check=False,
    )
    assert result.returncode == 0 and result.stdout.strip() == "ok", result.stdout + result.stderr


# ── Determinism and naming ─────────────────────────────────────


def _data_xml(content, contents="full"):
    sheets = _sheet_xml(content)
    return {
        title: re.search(r"<sheetData>.*</sheetData>", xml, re.S).group(0)
        for title, xml in sheets.items()
        if title not in {"About", "File info"}
    }


def test_same_run_lists_and_options_give_identical_data_sheets_and_hash(run):
    first, _n1, _t1, one = _export(run, reference="EXR-AAAAAAAA", language="en")
    second, _n2, _t2, two = _export(run, reference="EXR-BBBBBBBB", language="en")
    assert _data_xml(first) == _data_xml(second)
    assert one.files[0].data_sha256 == two.files[0].data_sha256


def test_data_sha256_is_recomputable_from_the_student_exam_rows(run):
    content, _name, _type, prepared = _export(run, language="en", dates={"Sun": "2026-12-13"})
    book = _book(content)
    _header, rows = _rows(book, "Student exams", "StudentExams")
    digest = hashlib.sha256()
    for row in rows:
        values = []
        for value in row:
            if value is None:
                values.append("")
            elif isinstance(value, time):
                values.append(value.strftime("%H:%M"))
            elif hasattr(value, "date") and hasattr(value, "hour"):
                values.append(value.date().isoformat())
            else:
                values.append(str(value))
        digest.update(("\x1f".join(values) + "\n").encode())
    info = {r["Key"]: r["Value"] for r in _records(book, "File info", "FileInfo")}
    assert info["data_sha256"] == digest.hexdigest() == prepared.files[0].data_sha256
    props = _parts(content)["docProps/custom.xml"]
    assert digest.hexdigest() in props and REFERENCE in props


@pytest.mark.parametrize(
    "payload, expected",
    [
        ({"scope": {"kind": "all"}}, "exam_students_all_r{run}_en_1A2B3C4D.xlsx"),
        (
            {"scope": {"kind": "course", "exam": "PHYS103 (2)"}},
            "exam_students_PHYS103-2_r{run}_en_1A2B3C4D.xlsx",
        ),
        (
            {"scope": {"kind": "room", "slot_index": 0, "room_code": "M-B"}},
            "exam_students_M-B_Sun-0800_r{run}_en_1A2B3C4D.xlsx",
        ),
        (
            {"scope": {"kind": "period", "slot_index": 0}, "dates": {"Sun": "2026-12-13"}},
            "exam_students_2026-12-13-0800_r{run}_en_1A2B3C4D.xlsx",
        ),
        (
            {"scope": {"kind": "all"}, "programs": ["CS", "CS2"], "contents": "summary"},
            "exam_students_all_dept-cs_summary_r{run}_en_1A2B3C4D.xlsx",
        ),
        (
            {"scope": {"kind": "day", "day": "Sun"}, "rows": "flagged"},
            "exam_students_Sun_flagged_r{run}_en_1A2B3C4D.xlsx",
        ),
    ],
)
def test_file_names_are_ascii_and_carry_no_student_data(run, payload, expected):
    payload = {"groups": ["M", "F", "U"], "one_file_per_group": False, **payload}
    _content, name, content_type, _prepared = _export(run, language="en", **payload)
    assert name == expected.format(run=run.pk)
    assert name.isascii() and content_type == export.XLSX_TYPE


def test_one_file_per_group_is_an_uncompressed_zip_of_single_group_files(run):
    content, name, content_type, prepared = _export(
        run, language="en", groups=["M", "F"], one_file_per_group=True
    )
    assert (
        content_type == export.ZIP_TYPE and name == f"exam_students_all_r{run.pk}_en_1A2B3C4D.zip"
    )
    with zipfile.ZipFile(BytesIO(content)) as archive:
        infos = archive.infolist()
        assert [i.filename for i in infos] == [
            f"exam_students_all_r{run.pk}_en_M_1A2B3C4D.xlsx",
            f"exam_students_all_r{run.pk}_en_F_1A2B3C4D.xlsx",
        ]
        assert {i.compress_type for i in infos} == {zipfile.ZIP_STORED}
        for info, gender in zip(infos, ["Male", "Female"], strict=True):
            rows = _records(_book(archive.read(info)), "Student exams", "StudentExams")
            assert rows and {r["Group"] for r in rows} == {gender}


def test_reference_appears_in_about_file_info_and_properties(run):
    content, *_ = _export(run, language="en")
    book = _book(content)
    about = {r[0]: r[1] for r in book["About"].iter_rows(values_only=True) if r[0]}
    assert about["Reference"].startswith(REFERENCE)
    info = {r["Key"]: r["Value"] for r in _records(book, "File info", "FileInfo")}
    assert info["reference"] == REFERENCE and info["audit_entry_hash16"] == AUDIT_HASH[:16]
    assert book.properties.subject == REFERENCE and book.properties.keywords == f"run-{run.pk}"


# ── Options ────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "payload, field",
    [
        ({"scope": {"kind": "course", "exam": "NOPE"}}, "scope.exam"),
        (
            {"scope": {"kind": "section", "exam": "MATH101", "section_key": "x", "gender": "M"}},
            "scope.section_key",
        ),
        ({"scope": {"kind": "room", "slot_index": 1, "room_code": "M-B"}}, "scope.room_code"),
        ({"scope": {"kind": "day", "day": "Fri"}}, "scope.day"),
        ({"scope": {"kind": "all"}, "programs": ["XX"]}, "programs"),
        ({"scope": {"kind": "all"}, "groups": []}, "groups"),
        ({"scope": {"kind": "all"}, "prepared_for": "x" * 81}, "prepared_for"),
        ({"scope": {"kind": "all"}, "prepared_for": "Dean \uffff office"}, "prepared_for"),
        ({"scope": {"kind": "all"}, "prepared_for": "Dean \ufffe office"}, "prepared_for"),
        # Lone surrogates: a JSON body may escape one, and no XML part can hold it.
        ({"scope": {"kind": "all"}, "prepared_for": "Dean \ud800 office"}, "prepared_for"),
        ({"scope": {"kind": "all"}, "prepared_for": "Dean \udfff office"}, "prepared_for"),
        ({"scope": {"kind": "all"}, "student_ids": [1]}, "student_ids"),
        ({"scope": {"kind": "all"}, "dates": {"Sun": "2026-12-14"}}, "dates"),
    ],
)
def test_invalid_options_name_their_field(run, payload, field):
    model = build_roster_model(run)
    with pytest.raises(ExportOptionsError) as error:
        parse_export_options(payload, model)
    assert error.value.field == field and error.value.code == "invalid_options"


@pytest.mark.parametrize(
    "dates, day",
    [
        ({"Sun": "2026-12-14"}, "Sun"),  # a Monday
        ({"Sun": "2026-12-13", "Mon": "2026-12-07"}, "Mon"),  # a Monday, out of day order
        ({"Mon": "2026-13-01"}, "Mon"),
        ({"Fri": "2026-12-18"}, ""),  # no such day: no one input to mark
    ],
)
def test_a_refused_date_names_the_day_to_fix(run, dates, day):
    model = build_roster_model(run)
    with pytest.raises(ExportOptionsError) as error:
        parse_export_options({"scope": {"kind": "all"}, "dates": dates}, model)
    assert (error.value.field, error.value.day) == ("dates", day)


def test_a_choice_with_no_students_is_an_empty_scope(run):
    model = build_roster_model(run)
    options = parse_export_options(
        {"scope": {"kind": "course", "exam": "CS101"}, "programs": ["AI"]}, model
    )
    with pytest.raises(EmptyScope) as error:
        prepare_export(model, options, generated_by="x", generated_role="EXAM_COMMITTEE")
    assert error.value.code == "empty_scope"


# ── Conditional formatting and empty tables ────────────────────


def _cf_rules(content, title):
    """Each conditional-formatting block as (sqref, [formula text, ...]), XML-decoded."""
    root = ET.fromstring(_sheet_xml(content)[title].encode("utf-8"))
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    return [
        (block.get("sqref"), [f.text for f in block.findall(".//m:formula", ns)])
        for block in root.findall("m:conditionalFormatting", ns)
    ]


@pytest.mark.parametrize("language, yes, clash", [("en", "Yes", "Clash"), ("ar", "نعم", "تعارض")])
def test_flag_colours_are_conditional_rules_on_the_right_columns(run, language, yes, clash):
    content, *_ = _export(run, language=language)
    book = _book(content)
    exams_title, flags_title = SHEETS[language][1], SHEETS[language][3]
    ((sqref, formulas),) = _cf_rules(content, exams_title)
    header, rows = _rows(book, exams_title, "StudentExams")
    assert sqref == f"A2:AA{len(rows) + 1}"
    letters = {name: get_column_letter(index + 1) for index, name in enumerate(header)}
    names = EN_OR_AR[language]
    assert formulas[0] == f'${letters[names[20]]}2="{yes}"' == f'$U2="{yes}"'
    assert formulas[1] == f'${letters[names[22]]}2="{yes}"' == f'$W2="{yes}"'
    assert formulas[2].startswith(f"OR(${letters[names[18]]}2=") and formulas[2].count("$S2=") == 3
    assert formulas[3] == f'${letters[names[25]]}2<>""' == '$Z2<>""'
    ((_ref, flag_formulas),) = _cf_rules(content, flags_title)
    assert flag_formulas == [f'$H2="{clash}"']


EN_OR_AR = {"en": EN["StudentExams"], "ar": AR["StudentExams"]}


def test_a_file_with_no_flagged_student_keeps_one_empty_flags_row(run):
    content, *_ = _export(run, language="en", scope={"kind": "day", "day": "Mon"})
    book = _book(content)
    assert book["Flags"].tables["ExamFlags"].ref == "A1:J2"
    assert all(cell.value is None for cell in book["Flags"][2])
    guide = {row[0]: (row[1], row[2]) for row in book["About"].iter_rows(values_only=True)}
    assert guide["Flags"] == (0, "0 rows — no student in this file has two exams on one day")


def test_flagged_only_keeps_every_row_of_flagged_students_and_nobody_else(run):
    content, name, *_ = _export(run, language="en", rows="flagged")
    rows = _records(_book(content), "Student exams", "StudentExams")
    students = {r["Student ID"] for r in rows}
    flagged = {r["Student ID"] for r in rows if "Yes" in (r["Clash"], r["Same day"])}
    assert students == flagged and students
    assert not students & set(MALE_AI)
    assert {r["Exam"] for r in rows if r["Student ID"] == MALE_CS[1]} == {
        "MATH101",
        "CS101",
        "PHYS103 (1)",
    }
    assert "_flagged_" in name


def test_a_gone_section_is_listed_with_no_rows_and_explained(run):
    StudentTermSection.objects.filter(term_section__section="F2").delete()
    content, *_ = _export(run, language="en", scope={"kind": "course", "exam": "CS101"})
    book = _book(content)
    sections = {s["Section"]: s for s in _records(book, "Sections", "ExamSections")}
    gone = sections["F2"]
    assert (gone["Membership"], gone["Students now"], gone["Rows in this file"]) == (
        "Gone since save",
        0,
        0,
    )
    assert gone["Difference"] == -gone["Students at save"]
    changes = _records(book, "Checks and changes", "ChangeLog")
    assert [(c["Section"], c["What changed"]) for c in changes] == [("F2", "Gone since save")]


@pytest.mark.parametrize(
    "text, token",
    [
        ("PHYS103 (2)", "PHYS103-2"),
        ("W1-Sun", "W1-Sun"),
        ("172:FA 001", "172FA-001"),
        ("a - b", "a-b"),
        ("قاعة", ""),
    ],
)
def test_file_name_tokens_are_ascii(text, token):
    assert export._token(text) == token
