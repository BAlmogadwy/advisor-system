"""A room linked courses share is one physical room: staffed, seated and counted once.

Decision 5: invigilators are counted per physical (period, room) on the room's
combined head-count, under the department rule if ANY course in the room is a
department course. Room QA counts a room's seats once per period and reports
a double booking only between different owners; the building footprint
counts the room once; the master Excel staffs it on one row; each Department
file says the room is shared, with the room's total and one set of
invigilators. The student-data export and the Student lists already hold
several exams in a room: they are checked to seat both courses' sections,
each on its own seat parts, and to count the room's students together.
"""

import io
import json
from collections import defaultdict

import pytest
from openpyxl import load_workbook

from core import models
from core.services import exam_room_allocation
from core.services.exam_department_export import (
    DepartmentExportUnavailable,
    export_department_workbooks,
)
from core.services.exam_evaluation import evaluate_exam_schedule
from core.services.exam_roster_view import build_roster_view, room_facts
from core.services.exam_rosters import build_roster_model
from core.services.exam_run_schema import derive_building_footprint
from core.services.exam_student_export import parse_export_options, prepare_export, render_export
from core.services.exam_timetable import (
    _build_room_qa,
    _invigilators_needed,
    _rebalance_invigilators_pass,
    _room_invigilators_needed,
    assign_rooms_to_schedule,
    build_exam_timetable,
    export_exam_timetable_xlsx,
    period_cohort_count,
)
from core.services.linked_exams import NO_LINKS, resolve_linked_exams
from core.services.xlsx_bidi import LRM, RLM
from tests import exam_linked_rooms_corpus as corpus


@pytest.fixture(autouse=True)
def _fresh_room_cache():
    with exam_room_allocation._CACHE_LOCK:
        exam_room_allocation._CACHE.clear()
    yield
    with exam_room_allocation._CACHE_LOCK:
        exam_room_allocation._CACHE.clear()


def _links(*groups):
    codes = sorted({code for group in groups for code in group})
    return resolve_linked_exams(
        [{"members": [{"course_identity": code} for code in group]} for group in groups],
        {code: {"course_identity": code} for code in codes},
    )


# ── the rule ─────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("codes", "students", "staff"),
    [
        (["PHYS103 (1)", "PHYS103 (2)"], 38, 1),  # external: one of us past 30
        (["PHYS103 (1)", "PHYS103 (2)"], 30, 0),
        (["MATH106", "CS101"], 25, 1),  # any department course: department rule
        (["MATH106", "CS101"], 35, 2),
        (["CS101", "MATH106"], 35, 2),  # whatever the order
        (["GS151", "IS113"], 35, 2),  # the external course first
        (["XYZ100", "STAT301"], 12, 1),  # an unknown prefix counts as department
        (["AI212", "AI225"], 29, 1),
    ],
)
def test_a_shared_room_is_staffed_on_its_whole_head_count(codes, students, staff):
    assert _room_invigilators_needed(codes, students) == staff


def test_one_course_in_a_room_is_staffed_as_before():
    for code in ("CS101", "MATH106", "XYZ100"):
        for students in (5, 29, 30, 31, 45):
            assert _room_invigilators_needed([code], students) == _invigilators_needed(
                code, students
            )


# ── room QA and the building footprint ───────────────────────────────────────


def _row(code, students, capacity, gender="M", building="North"):
    return {
        "room_code": code,
        "student_count": students,
        "room_capacity": capacity,
        "gender": gender,
        "section": f"{gender}1",
        "building": building,
    }


def _entry(code, slot, day, rooms):
    return {
        "course_code": code,
        "course_identity": code,
        "slot_index": slot,
        "day": day,
        "period": "08:00-10:00",
        "rooms": rooms,
    }


def _shared_board(partner_gender="M"):
    return [
        _entry("AI212", 0, "Sun", [_row("R-40", 10, 40)]),
        _entry("AI225", 0, "Sun", [_row("R-40", 15, 40, partner_gender)]),
        _entry("CS101", 0, "Sun", [_row("R-20", 5, 20)]),
        _entry("PHYS103 (1)", 1, "Mon", [_row("R-40", 20, 40)]),
        _entry("PHYS103 (2)", 1, "Mon", [_row("R-40", 18, 40)]),
        _entry("GS151", 2, "Tue", [_row("R-40", 25, 40, "F")]),
        _entry("IS113", 2, "Tue", [_row("R-40", 10, 40, "F")]),
    ]


LINKS = (("AI212", "AI225"), ("PHYS103 (1)", "PHYS103 (2)"), ("GS151", "IS113"))


def test_room_qa_counts_each_shared_room_once():
    qa = _build_room_qa(_shared_board(), [], links=_links(*LINKS))
    assert qa["room_double_bookings"] == []
    assert qa["rooms_used"] == 4
    assert qa["total_demand"] == 103
    assert qa["total_capacity_used"] == 40 + 20 + 40 + 40
    assert qa["avg_utilization"] == round(103 / 140, 4)
    # Sun: R-40 holds 25 department students (1), R-20 five (1).
    # Mon: 38 external students in one room need one of us - two rooms would not.
    # Tue: 35 in a room with a department course: the department rule, two.
    assert qa["invigilators_per_day"] == {
        "Sun": {"M": 2, "F": 0, "total": 2},
        "Mon": {"M": 1, "F": 0, "total": 1},
        "Tue": {"M": 0, "F": 2, "total": 2},
    }
    assert (qa["invigilators_total"], qa["invigilators_total_M"], qa["invigilators_total_F"]) == (
        5,
        3,
        2,
    )


def test_two_courses_in_one_room_are_a_double_booking_unless_one_link_holds_it():
    unlinked = _build_room_qa(_shared_board(), [], links=NO_LINKS)
    assert [
        (item["slot_index"], item["room_code"], item["courses"])
        for item in unlinked["room_double_bookings"]
    ] == [
        (0, "R-40", ["AI212", "AI225"]),
        (1, "R-40", ["PHYS103 (1)", "PHYS103 (2)"]),
        (2, "R-40", ["GS151", "IS113"]),
    ]
    other_link = _build_room_qa(
        _shared_board(), [], links=_links(("AI212", "CS101"), LINKS[1], LINKS[2])
    )
    assert [item["courses"] for item in other_link["room_double_bookings"]] == [["AI212", "AI225"]]


def test_one_link_in_one_room_code_across_two_cohorts_is_still_a_double_booking():
    qa = _build_room_qa(_shared_board(partner_gender="F"), [], links=_links(*LINKS))
    assert [item["courses"] for item in qa["room_double_bookings"]] == [["AI212", "AI225"]]
    # Never folded into one room's head-count either: each cohort is staffed.
    assert qa["invigilators_per_day"]["Sun"] == {"M": 2, "F": 1, "total": 3}


def test_the_building_footprint_counts_a_shared_room_once():
    board = [
        _entry("AI212", 0, "Sun", [_row("R-40", 10, 40), _row("R-41", 3, 40)]),
        _entry("AI225", 0, "Sun", [_row("R-40", 15, 40)]),
        _entry("CS101", 0, "Sun", [_row("R-20", 5, 20, building="South")]),
    ]
    footprint = derive_building_footprint(board)
    assert footprint["max_rooms_per_building_per_slot"] == {"Sun:08:00-10:00": 2}
    assert (
        footprint["largest_slot_footprint_summary"]
        == "Sun:08:00-10:00 uses 3 rooms across 2 buildings"
    )
    assert footprint["buildings_used_per_slot"] == {"Sun:08:00-10:00": ["North", "South"]}
    assert footprint["cross_building_clusters_per_dept"] == {}


def test_the_building_footprint_never_folds_the_other_cohorts_room():
    """A code the women's catalogue reuses is a second room, not a shared one."""
    board = [
        _entry("AI212", 0, "Sun", [_row("R-40", 10, 40)]),
        _entry("AI225", 0, "Sun", [_row("R-40", 15, 40, "F")]),
    ]
    footprint = derive_building_footprint(board)
    assert footprint["max_rooms_per_building_per_slot"] == {"Sun:08:00-10:00": 2}
    assert (
        footprint["largest_slot_footprint_summary"]
        == "Sun:08:00-10:00 uses 2 rooms across 1 building"
    )


# ── the invigilator pass ─────────────────────────────────────────────────────


def test_the_invigilator_pass_staffs_a_shared_room_once():
    """Sunday holds a four-course link in one room and CS301; Monday CS401.

    Counted once, the link's room needs one invigilator: two on Sunday, one on
    Monday, nothing worth moving. Counted per course it would look like five
    against one, and the pass would move CS301 for a balance that is not real.
    """
    slots = [
        {"index": 0, "day": "Sun", "period": "08:00-10:00"},
        {"index": 1, "day": "Mon", "period": "08:00-10:00"},
    ]
    linked = ["CS201", "CS202", "CS203", "CS204"]
    entries = [
        {**_entry(code, 0, "Sun", []), "course_name": code} for code in [*linked, "CS301"]
    ] + [{**_entry("CS401", 1, "Mon", []), "course_name": "CS401"}]
    sections = {
        entry["course_code"]: [
            {
                "section": "M1",
                "section_key": f"term-section:{entry['course_code']}",
                "mapping_status": "mapped",
                "student_count": 3,
                "gender": "M",
                "preferred_room": "",
            }
        ]
        for entry in entries
    }
    rooms = [{"room_code": f"M-20-{n}", "capacity": 20, "section": "M"} for n in range(6)]
    links = _links(tuple(linked))
    assign_rooms_to_schedule(entries, sections, rooms, links=links)
    assert len({row["room_code"] for entry in entries[:4] for row in entry["rooms"]}) == 1
    enrolled = {
        entry["course_code"]: {1000 * n + s for s in range(3)} for n, entry in enumerate(entries)
    }
    moves = _rebalance_invigilators_pass(
        entries,
        sections,
        rooms,
        slots,
        {},
        None,
        None,
        enrolled_sets=enrolled,
        links=links,
    )
    assert moves == 0
    assert {entry["course_code"]: entry["day"] for entry in entries}["CS301"] == "Sun"


# ── a saved build with shared rooms, and everything read from it ─────────────

LINKED_GROUPS = (("AI212", "AI225"), ("CYB152", "IS102"), ("PHYS103 (1)", "PHYS103 (2)"))


def _linked_build(rebalance: bool) -> dict:
    common = {
        "days": corpus.POPULATION_DAYS,
        "periods": corpus.POPULATION_PERIODS,
        "max_per_day": 2,
        "programs": ["AI", "CS", "IS"],
        "seed": 5,
        "assign_rooms": True,
        "rebalance_invigilators": rebalance,
        "persist": False,
    }
    plain = build_exam_timetable("Unlinked", **common)
    identity = {entry["course_code"]: entry["course_identity"] for entry in plain["schedule"]}
    result = build_exam_timetable(
        "Linked",
        linked_exams=[
            {"members": [{"course_identity": identity[code]} for code in group]}
            for group in LINKED_GROUPS
        ],
        **common,
    )
    assert result["status"] == "ok"
    return result


@pytest.fixture
def linked_run(db):
    corpus.create_rooms_population(models)
    result = _linked_build(rebalance=True)
    return result, corpus.save_fixed_run(models, result)


def test_a_build_without_the_invigilator_pass_shares_rooms_too(db):
    corpus.create_rooms_population(models)
    result = _linked_build(rebalance=False)
    assert _shared(result)
    assert result["qa"]["rooms"]["room_double_bookings"] == []


def _physical(result):
    """{(slot, room, cohort): {course: students}} of the saved rows."""
    rooms = defaultdict(dict)
    for entry in result["schedule"]:
        for row in entry["rooms"]:
            if row["room_code"] != "UNASSIGNED":
                rooms[(entry["slot_index"], row["room_code"], row["gender"])][
                    entry["course_code"]
                ] = row["student_count"]
    return dict(rooms)


def _shared(result):
    return {key: courses for key, courses in _physical(result).items() if len(courses) > 1}


def test_a_linked_build_shares_rooms_and_its_qa_counts_them_once(linked_run):
    result, _run = linked_run
    physical, shared = _physical(result), _shared(result)
    assert shared, "the population's links should share at least one room"
    for (_slot, room, _cohort), courses in shared.items():
        assert len({corpus_link(code) for code in courses}) == 1
        for code, students in courses.items():
            [row] = [
                row
                for entry in result["schedule"]
                if entry["course_code"] == code
                for row in entry["rooms"]
                if row["room_code"] == room
            ]
            assert row["room_shared_with"] == sorted(set(courses) - {code})
            assert row["room_student_total"] == sum(courses.values())
            assert row["student_count"] == students
    qa = result["qa"]["rooms"]
    assert qa["unassigned_room_sections"] == []
    assert qa["room_double_bookings"] == []
    assert qa["rooms_used"] == len(physical)
    capacity = {
        (entry["slot_index"], row["room_code"], row["gender"]): row["room_capacity"]
        for entry in result["schedule"]
        for row in entry["rooms"]
        if row["room_code"] != "UNASSIGNED"
    }
    assert qa["total_capacity_used"] == sum(capacity.values())
    assert all(sum(courses.values()) <= capacity[key] for key, courses in physical.items())
    assert qa["invigilators_total"] == sum(
        _room_invigilators_needed(list(courses), sum(courses.values()))
        for courses in physical.values()
    )


def corpus_link(code):
    return next((group for group in LINKED_GROUPS if code in group), (code,))


def _sheet_rows(sheet):
    return [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]


def test_the_master_excel_staffs_each_shared_room_on_one_row(linked_run):
    result, run = linked_run
    path = export_exam_timetable_xlsx(run.pk)
    try:
        book = load_workbook(path)
    finally:
        path.unlink(missing_ok=True)
    rows = _sheet_rows(book["Invigilators"])
    summary = next(
        index
        for index, row in enumerate(rows)
        if row[:4] == ["Day", "M Invigilators", "F Invigilators", "Total"]
    )
    total = next(row for row in rows[summary:] if row[0] == "TOTAL")
    detail = next(
        index for index, row in enumerate(rows) if row[:4] == ["Day", "Period", "Course", "Type"]
    )
    details = [row for row in rows[detail + 1 :] if row[0]]
    assert sum(row[8] for row in details) == total[3] == result["qa"]["rooms"]["invigilators_total"]
    assert len(details) == len(_physical(result))
    by_room = {(row[0], row[1], row[7], row[5]): row for row in details}
    for (slot, room, cohort), courses in _shared(result).items():
        entry = next(e for e in result["schedule"] if e["slot_index"] == slot)
        row = by_room[(entry["day"], entry["period"], room, cohort)]
        assert row[2] == " + ".join(sorted(courses))
        assert row[3].endswith("(shared room)")
        assert row[6] == sum(courses.values())
        assert row[8] == _room_invigilators_needed(list(courses), sum(courses.values()))
    assignments = _sheet_rows(book["Room Assignments"])
    header = assignments[0]
    shared_rows = [
        dict(zip(header, row, strict=False))
        for row in assignments[1:]
        if row[0] and "Shared with" in str(row[header.index("Room Group")] or "")
    ]
    assert len(shared_rows) == sum(len(courses) for courses in _shared(result).values())
    for row in shared_rows:
        courses = _shared(result)[
            next(
                key
                for key in _shared(result)
                if key[1] == row["Room"] and row["Course"] in _shared(result)[key]
            )
        ]
        total_in_room = sum(courses.values())
        assert f"room total {total_in_room}" in row["Room Group"]
        assert row["Utilization"] == pytest.approx(total_in_room / row["Capacity"])


def test_the_master_excel_staffs_a_mixed_shared_room_by_the_department_rule(linked_run):
    """An external course beside a department course: the department rule.

    The population links department courses to department courses, so the
    saved run is relabelled: AI212 becomes ENV212 (external, and listed first).
    """
    result, run = linked_run
    room, cohort, courses = _shared_pair(result, "AI212", "AI225")
    total = sum(courses.values())
    run.result_json = run.result_json.replace("AI212", "ENV212").replace("AI225", "IS225")
    run.save(update_fields=["result_json"])
    path = export_exam_timetable_xlsx(run.pk)
    try:
        rows = _sheet_rows(load_workbook(path)["Invigilators"])
    finally:
        path.unlink(missing_ok=True)
    [row] = [row for row in rows if row[2] == "ENV212 + IS225" and row[7] == room]
    assert (row[3], row[5], row[6]) == ("Department (shared room)", cohort, total)
    assert row[8] == _invigilators_needed("IS225", total) != _invigilators_needed("ENV212", total)


def _department(run, department, language):
    content, _name, _type = export_department_workbooks(
        run, {"departments": [department], "genders": ["M", "F"], "language": language}
    )
    return load_workbook(io.BytesIO(content))


def _cells(book):
    return [
        str(cell)
        for sheet in book
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell
    ]


def _shared_pair(result, first, second):
    for (_slot, room, cohort), courses in _shared(result).items():
        if set(courses) == {first, second}:
            return room, cohort, courses
    raise AssertionError(f"{first} and {second} share no room")


def test_department_files_say_a_room_is_shared_with_one_set_of_invigilators(linked_run):
    result, run = linked_run
    room, _cohort, courses = _shared_pair(result, "AI212", "AI225")
    total = sum(courses.values())
    english = _cells(_department(run, "ai-ds", "en"))
    note = (
        f"Room {room} shared with AI225, room total {total} — one set of invigilators for this room"
    )
    assert any(note in cell for cell in english)
    assert any(f"{room}: {courses['AI212']} (room total: {total})" in cell for cell in english)
    # The partner's own department file says the same of the same room.
    partner = _cells(_department(run, "cs", "en"))
    assert any(
        f"Room {room} shared with AI212, room total {total} — one set of invigilators" in cell
        for cell in partner
    )
    arabic = _cells(_department(run, "ai-ds", "ar"))
    arabic_note = (
        f"القاعة {LRM}{room}{RLM} مشتركة مع {LRM}AI225{RLM}، "
        f"إجمالي القاعة {total} — طاقم مراقبة واحد لهذه القاعة"
    )
    assert any(arabic_note in cell for cell in arabic)
    assert any(f"{room}: {courses['AI212']} (إجمالي القاعة: {total})" in cell for cell in arabic)


@pytest.mark.parametrize("tamper", ["partner_dropped", "wrong_total", "wrong_partner", "claimed"])
def test_department_files_refuse_shared_room_facts_the_timetable_does_not_hold(linked_run, tamper):
    result, run = linked_run
    room, _cohort, _courses = _shared_pair(result, "AI212", "AI225")
    data = json.loads(run.result_json)
    for entry in data["schedule"]:
        for row in entry["rooms"]:
            if tamper == "claimed" and entry["course_code"] == "MATH106":
                row["room_shared_with"], row["room_student_total"] = ["AI212"], 99
            if entry["course_code"] != "AI212" or row["room_code"] != room:
                continue
            if tamper == "partner_dropped":
                del row["room_shared_with"]
            elif tamper == "wrong_total":
                row["room_student_total"] += 1
            elif tamper == "wrong_partner":
                row["room_shared_with"] = ["CS101"]
    run.result_json = json.dumps(data, ensure_ascii=False)
    run.save(update_fields=["result_json"])
    with pytest.raises(DepartmentExportUnavailable):
        export_department_workbooks(
            run, {"departments": ["ai-ds"], "genders": ["M", "F"], "language": "en"}
        )


def test_the_student_export_seats_both_courses_in_the_room_and_counts_them_together(linked_run):
    result, run = linked_run
    model = build_roster_model(run, now=corpus.FIXED_AT)
    options = parse_export_options(
        {"scope": {"kind": "all"}, "one_file_per_group": False, "language": "en"}, model
    )
    prepared = prepare_export(
        model,
        options,
        generated_by="committee",
        generated_role="EXAM_COMMITTEE",
        now=corpus.FIXED_AT,
    )
    content, _name, _type = render_export(prepared, reference="EXR-SHARED01", audit_hash="0" * 64)
    book = load_workbook(io.BytesIO(content))
    sheet = book["Rooms"]
    ref = sheet.tables["ExamRooms"].ref
    header, *rows = [[cell.value for cell in row] for row in sheet[ref]]
    records = [dict(zip(header, row, strict=True)) for row in rows]
    for (slot, room, _cohort), courses in _shared(result).items():
        entry = next(e for e in result["schedule"] if e["slot_index"] == slot)
        [record] = [
            r
            for r in records
            if r["Room"] == room and r["Period"] == entry["period"] and r["Day"] == entry["day"]
        ]
        total = sum(courses.values())
        assert record["Seated at save"] == record["Seated now"] == total
        assert record["Exams"] == len(courses)
        for code in courses:
            assert f"{code} " in record["Sections in room"]
        # Each student sits in the room once, under their own course's section.
        seated = [
            row
            for row in _records(book, "Student exams", "StudentExams")
            if row["Exam room"] == room
            and row["Exam"] in courses
            and row["Period"] == entry["period"]
        ]
        assert len(seated) == total
        assert {row["Exam"] for row in seated} == set(courses)


def _records(book, sheet_name, table):
    sheet = book[sheet_name]
    header, *rows = [[cell.value for cell in row] for row in sheet[sheet.tables[table].ref]]
    return [dict(zip(header, row, strict=True)) for row in rows]


def test_student_lists_hold_a_shared_room_as_one_room_of_both_courses(linked_run):
    result, run = linked_run
    view = build_roster_view(run, built=0.0)
    for (slot, room, _cohort), courses in _shared(result).items():
        facts = room_facts(view.model, (slot, room), len(view.by_room.get((slot, room), [])))
        total = sum(courses.values())
        assert facts["exams"] == sorted(courses)
        assert facts["seated_at_save"] == facts["seated_now"] == total
        # Per-section seat parts: each course's sections keep their own part here.
        parts = [
            (group.exam, part.student_count)
            for group in view.model.groups
            for part in group.parts
            if part.slot_index == slot and part.room_code == room
        ]
        assert {exam for exam, _count in parts} == set(courses)
        assert sum(count for _exam, count in parts) == total
        seated = view.by_room[(slot, room)]
        assert {sitting.group.exam for sitting in seated} == set(courses)


def test_the_room_budget_counts_a_linked_period_twice():
    """A linked period is allocated twice, sharing and not, and is budgeted so."""
    entries = [
        _entry(code, slot, "Sun", []) for code, slot in (("AI212", 0), ("AI225", 0), ("CS101", 1))
    ]
    entries.append({**_entry("IS102", 9, "OVERFLOW", []), "period": "Extra-9"})
    sections = {"AI212": [{"gender": "M"}], "CS101": [{"gender": "F"}]}
    assert period_cohort_count(entries, sections) == 4
    assert period_cohort_count(entries, sections, _links(("AI212", "AI225"))) == 6
    assert period_cohort_count(entries, sections, _links(("CS101", "IS102"))) == 6


def test_check_keeps_a_links_shared_rooms(linked_run):
    """Check (and Save, which re-checks) rooms the saved board as the build did."""
    result, _run = linked_run
    checked = evaluate_exam_schedule(
        days=corpus.POPULATION_DAYS,
        periods=corpus.POPULATION_PERIODS,
        max_per_day=2,
        schedule_raw=result["schedule"],
        selected_courses=[entry["course_code"] for entry in result["schedule"]],
        assign_rooms=True,
        seed=5,
        thin_conflict_threshold=0,
        programs=["AI", "CS", "IS"],
        linked_exams=result["linked_exams"],
    )
    assert _shared(checked) == _shared(result)
    assert checked["qa"]["rooms"]["room_double_bookings"] == []
    unlinked = evaluate_exam_schedule(
        days=corpus.POPULATION_DAYS,
        periods=corpus.POPULATION_PERIODS,
        max_per_day=2,
        schedule_raw=result["schedule"],
        selected_courses=[entry["course_code"] for entry in result["schedule"]],
        assign_rooms=True,
        seed=5,
        thin_conflict_threshold=0,
        programs=["AI", "CS", "IS"],
        linked_exams=[],
    )
    assert _shared(unlinked) == {}
    assert unlinked["qa"]["rooms"]["rooms_used"] > checked["qa"]["rooms"]["rooms_used"]


def test_the_master_excel_counts_rooms_once_even_without_saved_staff_totals(linked_run):
    """An older room QA without daily staff totals is re-derived per room."""
    result, run = linked_run
    data = json.loads(run.result_json)
    del data["qa"]["rooms"]["invigilators_per_day"]
    run.result_json = json.dumps(data, ensure_ascii=False)
    run.save(update_fields=["result_json"])
    path = export_exam_timetable_xlsx(run.pk)
    try:
        rows = _sheet_rows(load_workbook(path)["Invigilators"])
    finally:
        path.unlink(missing_ok=True)
    total = next(row for row in rows if row[0] == "TOTAL")
    assert total[3] == result["qa"]["rooms"]["invigilators_total"]


def test_a_room_code_the_other_cohort_reuses_is_not_a_shared_room(linked_run):
    """Catalogues may reuse a code across cohorts: two rooms, not one shared."""
    result, run = linked_run
    data = json.loads(run.result_json)
    women = [
        (entry["slot_index"], entry["course_code"], row["room_code"])
        for entry in data["schedule"]
        for row in entry["rooms"]
        if row["gender"] == "F"
    ]
    # A men's room of its own takes the code of another course's women's room.
    row, code = next(
        (row, code)
        for entry in data["schedule"]
        for row in entry["rooms"]
        if row["gender"] == "M" and not row.get("room_shared_with")
        for slot, other, code in women
        if slot == entry["slot_index"] and other != entry["course_code"]
    )
    row["room_code"] = code
    run.result_json = json.dumps(data, ensure_ascii=False)
    run.save(update_fields=["result_json"])
    content, _name, _type = export_department_workbooks(
        run, {"departments": ["ai-ds", "cs", "is"], "genders": ["M", "F"], "language": "en"}
    )
    assert content
