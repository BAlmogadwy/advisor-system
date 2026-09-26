"""Department packs use saved membership and preserve operational uncertainty."""

import json
import re
import tempfile
import xml.etree.ElementTree as ET
from collections import Counter
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from core.models import ExamTimetableRun
from core.services.exam_department_export import (
    DepartmentExportUnavailable,
    department_export_options,
    export_department_workbooks,
)
from core.services.exam_run_schema import (
    EXAM_RUN_SCHEMA_VERSION,
    normalise_exam_run_payload,
    stamp_schema_version,
)
from core.services.xlsx_bidi import LRM, RLM

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _section(number, label, gender, programs, *, status="mapped", instructors=None):
    return {
        "section": label,
        "section_key": f"term-section:{number}"
        if status == "mapped"
        else f"unmapped:{gender}:{status}",
        "term_section_id": number if status == "mapped" else None,
        "mapping_status": status,
        "mapping_source": "scraper_timetable" if status == "mapped" else "",
        "academic_year": "2026",
        "term": "1",
        "gender": gender,
        "student_count": sum(programs.values()),
        "program_counts": dict(programs),
        "instructors": instructors or [],
        "instructor_source": "recorded_section_meetings" if instructors else "",
        "instructor_status": "recorded"
        if instructors
        else "missing"
        if status == "mapped"
        else "unavailable",
    }


def _room(code, *sections, count=None, group=""):
    return {
        "room_code": code,
        "room_capacity": 40 if code != "UNASSIGNED" else 0,
        "gender": sections[0]["gender"],
        "section": "+".join(section["section"] for section in sections),
        "section_parts": deepcopy(list(sections)),
        "student_count": count
        if count is not None
        else sum(section["student_count"] for section in sections),
        "room_group": group,
        "building": "Science College",
        "floor": "2",
    }


def make_department_data():
    """A deliberately shared course/section population with known totals."""
    shared = _section(1, "F01", "F", {"AI": 12, "CS": 8}, instructors=["Dr Saved"])
    irrelevant = _section(2, "F99", "F", {"CS": 9}, instructors=["Dr Irrelevant"])
    male_ai = _section(3, "M01", "M", {"AI2": 5}, instructors=["Dr Male AI"])
    male_ds = _section(4, "M02", "M", {"DS2": 7}, instructors=["Dr Male DS"])
    math = _section(5, "F02", "F", {"DS": 6})
    unassigned = _section(6, "M03", "M", {"DS2": 4})
    missing = _section(7, "", "F", {"AI": 3}, status="missing")
    ambiguous = _section(8, "", "F", {"AI2": 2}, status="ambiguous")
    other = _section(9, "M99", "M", {"CS": 3})
    unknown = _section(10, "", "U", {"DS": 2}, status="missing")
    specs = [
        ("CS111 (1)", "CS111", "Foundations", [shared, irrelevant], "Sun", False),
        ("CS111 (2)", "CS111", "Programming I", [male_ai, male_ds], "Mon", False),
        ("MATH101", "MATH101", "Mathematics", [math], "Mon", True),
        ("AI200", "AI200", "Machine Learning", [unassigned], "Tue", False),
        ("GS101", "GS101", "General Studies", [missing, ambiguous], "OVERFLOW", False),
        ("CS299", "CS299", "CS Only", [other], "Tue", False),
        ("DS300", "DS300", "Unknown Gender", [unknown], "Tue", False),
    ]
    schedule, courses, section_enrollment = [], {}, {}
    for code, source, name, sections, day, online in specs:
        counts = Counter()
        for section in sections:
            for program, count in section["program_counts"].items():
                counts[program, section["gender"]] += count
        identity = f"{source}|{name.lower()}"
        slot = {"Sun": 0, "Mon": 1, "Tue": 2, "OVERFLOW": 3}[day]
        schedule.append(
            {
                "course_code": code,
                "source_course_code": source,
                "course_identity": identity,
                "course_name": name,
                "programs": sorted({program for program, _ in counts}),
                "enrolled_count": sum(counts.values()),
                "day": day,
                "period": "" if day == "OVERFLOW" else "08:00-10:00",
                "slot_index": slot,
                "is_online": online,
                "rooms": [],
            }
        )
        courses[code] = {
            "course_identity": identity,
            "source_course_code": source,
            "course_name": name,
            "program_counts": [
                {"program": program, "gender": gender, "student_count": count}
                for (program, gender), count in sorted(counts.items())
            ],
            "sections": deepcopy(sections),
        }
        section_enrollment[code] = deepcopy(sections)
    # Each split room has ten students, but the saved evidence cannot say how
    # many of those ten are AI students. Department counts are section counts.
    split = {**shared, "student_count": 10, "room_group_count": 2}
    schedule[0]["rooms"] = [
        _room("F-101", {**split, "room_group_index": 1}, count=10, group="1/2"),
        _room("F-102", {**split, "room_group_index": 2}, count=10, group="2/2"),
        _room("F-999", irrelevant),
    ]
    schedule[1]["rooms"] = [_room("M-101", male_ai, male_ds)]
    schedule[3]["rooms"] = [_room("UNASSIGNED", unassigned)]
    return stamp_schema_version(
        {
            "status": "ok",
            "enrollment_source": "scraper_timetable",
            "schedule": schedule,
            "slots": [
                {"index": index, "day": day, "period": "08:00-10:00"}
                for index, day in enumerate(["Sun", "Mon", "Tue"])
            ],
            "courses": [entry["course_code"] for entry in schedule],
            "courses_count": len(schedule),
            "students_count": 61,
            "enrollment_scope": {
                "programs": ["AI", "AI2", "DS", "DS2", "CS"],
                "sections": ["M", "F", "U"],
            },
            "section_enrollment": section_enrollment,
            "operations_snapshot": {
                "version": 1,
                "enrollment_source": "scraper_timetable",
                "courses": courses,
            },
            "assign_rooms": True,
            "pinned": [{"course_code": "CS111 (2)", "day": "Mon", "period": "08:00-10:00"}],
            "qa": {},
        }
    )


@pytest.fixture
def department_data():
    return make_department_data()


def _export(data, **payload):
    # An unsaved model means this also fails if export attempts to reload the
    # run or reconstruct membership from any current database source.
    run = ExamTimetableRun(
        pk=417,
        label="Saved department evidence",
        created_at=timezone.now(),
        result_json=json.dumps(data),
    )
    return export_department_workbooks(
        run,
        {
            "departments": ["ai-ds"],
            "genders": ["M", "F", "U"],
            "language": "en",
            **payload,
        },
    )


def _book(data, **payload):
    content, filename, content_type = _export(data, **payload)
    assert filename.endswith(".xlsx")
    assert content_type == XLSX_TYPE
    return load_workbook(BytesIO(content))


def _text(sheet):
    return "\n".join(str(cell.value) for row in sheet for cell in row if cell.value is not None)


def _full_print_sheets(book):
    return [
        sheet
        for sheet in book
        if sheet.title.startswith("Print (") or sheet.title in {"الطلاب", "الطالبات", "غير محدد"}
    ]


def _records(sheet):
    rows = list(sheet.values)
    index = next(
        index for index, row in enumerate(rows) if "Course" in row and "Department students" in row
    )
    headers = rows[index]
    return [
        record
        for row in rows[index + 1 :]
        if (record := dict(zip(headers, row, strict=True))).get("Course")
        and record.get("Record ID")
    ]


def test_department_options_follow_enrolled_programmes_and_canonical_identities(department_data):
    options = department_export_options(department_data)
    departments = {item["id"]: item for item in options["departments"]}
    combined = departments["ai-ds"]
    assert set(combined["programs"]) == {"AI", "AI2", "DS", "DS2"}
    assert combined["course_count"] == 6
    assert combined["student_sittings"] == 41
    assert options["days"] == ["Sun", "Mon", "Tue"]
    assert set(options["genders"]) == {"M", "F", "U"}
    assert not {"ai", "ai2", "ds", "ds2"}.intersection(departments)


@pytest.mark.parametrize(
    "sentinel", ["future_schema", "future_version_unrenderable", "unrenderable"]
)
def test_unrenderable_status_rejects_otherwise_valid_saved_department_evidence(
    department_data, sentinel
):
    if sentinel == "future_schema":
        department_data["schema_version"] = EXAM_RUN_SCHEMA_VERSION + 1
    else:
        department_data["status"] = sentinel
    normalized = normalise_exam_run_payload(department_data)
    expected_status = "future_version_unrenderable" if sentinel == "future_schema" else sentinel
    assert normalized["status"] == expected_status
    # Forward-compatible payloads preserve their valid-looking course and
    # snapshot keys. The status must be honored even when evidence survives.
    assert normalized["operations_snapshot"] == department_data["operations_snapshot"]
    assert normalized["schedule"] == department_data["schedule"]
    with pytest.raises(DepartmentExportUnavailable):
        department_export_options(normalized)
    with pytest.raises(DepartmentExportUnavailable):
        _export(department_data)


def test_shared_courses_and_separate_identities_survive_and_irrelevant_sections_do_not(
    department_data,
):
    book = _book(department_data)
    details = _text(book["Details"])
    assert "CS111 (1)" in details and "CS111 (2)" in details
    assert "Foundations" in details and "Programming I" in details
    assert "MATH101" in details
    assert "Dr Saved" in details
    assert "CS299" not in details
    assert "F99" not in details and "Dr Irrelevant" not in details
    assert "F-999" not in details


def test_exact_department_counts_are_distinct_from_shared_sections_and_split_rooms(department_data):
    rows = _records(_book(department_data)["Details"])
    shared = [row for row in rows if row["Course"] == "CS111 (1)"]
    assert len(shared) == 1
    row = shared[0]
    assert row["Department students"] == 12
    assert row["Whole section students"] == 20
    assert row["Teaching section"] == "F01"
    assert row["Course instructor"] == "Dr Saved"
    assert "F-101: 10" in row["Room allocation (whole section)"]
    assert "F-102: 10" in row["Room allocation (whole section)"]
    assert "Department students per room are not determined" in row["Notes"]
    assert "shared" in row["Notes"].lower() and "split" in row["Notes"].lower()
    assert sum(row["Department students"] for row in rows) == 41
    assert all(row["Invigilators"] is None and row["Supervisor"] is None for row in rows)


def test_combined_sections_keep_their_own_counts_without_duplicating_room_total(department_data):
    rows = [
        row for row in _records(_book(department_data)["Details"]) if row["Course"] == "CS111 (2)"
    ]
    assert len(rows) == 2
    by_section = {row["Teaching section"]: row for row in rows}
    assert (
        by_section["M01"]["Department students"] == by_section["M01"]["Whole section students"] == 5
    )
    assert (
        by_section["M02"]["Department students"] == by_section["M02"]["Whole section students"] == 7
    )
    assert sum(row["Whole section students"] for row in rows) == 12
    assert "M-101: 5" in by_section["M01"]["Room allocation (whole section)"]
    assert "M-101: 7" in by_section["M02"]["Room allocation (whole section)"]
    assert all("room total: 12" in row["Room allocation (whole section)"] for row in rows)
    assert by_section["M01"]["Course instructor"] == "Dr Male AI"
    assert by_section["M02"]["Course instructor"] == "Dr Male DS"


def _building_cells(book, language, code):
    details = book["التفاصيل" if language == "ar" else "Details"]
    heading = "المبنى" if language == "ar" else "Building"
    assert details.cell(4, 10).value == heading
    result = [(row[10], row[9]) for row in details if row[0].value == code]
    printed_rows = []
    for sheet in _full_print_sheets(book):
        assert sheet.cell(4, 5).value == heading
        printed_rows.extend((row[5], row[4]) for row in sheet if code in str(row[0].value))
    assert result and len(printed_rows) == len(result)
    return result + printed_rows


@pytest.mark.parametrize("language", ["en", "ar"])
@pytest.mark.parametrize(
    "allocation", ["same", "different", "missing", "all_missing", "unassigned", "all_unassigned"]
)
def test_saved_buildings_have_a_separate_column_in_details_and_print_views(
    department_data, language, allocation
):
    rooms = department_data["schedule"][0]["rooms"]
    unknown = "غير مسجل" if language == "ar" else "Not recorded"
    unassigned = "قاعة غير مخصصة" if language == "ar" else "UNASSIGNED"
    if allocation == "same":
        expected = "Science College"
    elif allocation == "different":
        rooms[1]["building"] = "Campus B"
        expected = "F-101: Science College\nF-102: Campus B"
    elif allocation == "missing":
        rooms[1].pop("building")
        expected = f"F-101: Science College\nF-102: {unknown}"
    elif allocation == "all_missing":
        rooms[0]["building"] = ""
        rooms[1].pop("building")
        expected = None
    elif allocation == "unassigned":
        rooms[1]["room_code"] = "UNASSIGNED"
        rooms[1]["building"] = "Fake Location"
        expected = f"F-101: Science College\n{unassigned}: {unknown}"
    else:
        for room in rooms[:2]:
            room["room_code"] = "UNASSIGNED"
            room["building"] = "Fake Location"
        expected = None
    book = _book(department_data, language=language)
    for room_cell, building_cell in _building_cells(book, language, "CS111 (1)"):
        displayed = (
            "—"
            if expected is None and building_cell.parent.title not in {"Details", "التفاصيل"}
            else expected
        )
        assert building_cell.value == displayed
        assert all(
            name not in str(room_cell.value)
            for name in ("Science College", "Campus B", "Fake Location")
        )
        if expected is not None:
            assert building_cell.data_type == "s"
    # A saved building string on an unassigned placeholder is not a location,
    # and an online exam without an allocation cannot acquire one.
    for code in ("AI200", "MATH101"):
        for _, building_cell in _building_cells(book, language, code):
            assert building_cell.value == (
                None if building_cell.parent.title in {"Details", "التفاصيل"} else "—"
            )


@pytest.mark.parametrize("language", ["en", "ar"])
def test_formula_like_saved_building_names_remain_literal_in_separate_columns(
    department_data, language
):
    dangerous = '=HYPERLINK("https://invalid.example","Building")'
    for room in department_data["schedule"][0]["rooms"][:2]:
        room["building"] = dangerous
    for room_cell, building_cell in _building_cells(
        _book(department_data, language=language), language, "CS111 (1)"
    ):
        assert building_cell.value == dangerous
        assert building_cell.data_type == "s"
        assert dangerous not in room_cell.value


@pytest.mark.parametrize(
    "problem",
    [
        "wrong_total",
        "negative_fragment",
        "boolean_fragment",
        "unknown_section",
        "duplicate_room",
        "invalid_section_key",
        "invalid_room_array",
        "invalid_gender_type",
        "gender_mismatch",
    ],
)
def test_tampered_room_distributions_are_rejected_instead_of_guessing(department_data, problem):
    rooms = department_data["schedule"][0]["rooms"]
    if problem == "wrong_total":
        rooms[0]["student_count"] = 11
    elif problem == "negative_fragment":
        rooms[0]["section_parts"][0]["student_count"] = -1
    elif problem == "boolean_fragment":
        rooms[0]["section_parts"][0]["student_count"] = True
    elif problem == "unknown_section":
        rooms[0]["section_parts"][0]["section_key"] = "term-section:99999"
    elif problem == "duplicate_room":
        rooms[1]["room_code"] = rooms[0]["room_code"]
    elif problem == "invalid_section_key":
        rooms[0]["section_parts"][0]["section_key"] = ["invalid key"]
    elif problem == "invalid_room_array":
        department_data["schedule"][0]["rooms"] = {}
    elif problem == "invalid_gender_type":
        rooms[0]["gender"] = []
    elif problem == "gender_mismatch":
        rooms[0]["gender"] = "M"
    with pytest.raises(DepartmentExportUnavailable):
        _export(department_data)


def test_same_room_code_in_distinct_student_groups_preserves_separate_allocations(department_data):
    # Room catalogues can reuse a visible code in different student groups.
    male = _section(11, "M11", "M", {"AI2": 4})
    entry = department_data["schedule"][0]
    entry["enrolled_count"] += 4
    entry["programs"].append("AI2")
    entry["rooms"].append(_room("F-101", male))
    course = department_data["operations_snapshot"]["courses"]["CS111 (1)"]
    course["sections"].append(deepcopy(male))
    course["program_counts"].append({"program": "AI2", "gender": "M", "student_count": 4})
    department_data["section_enrollment"]["CS111 (1)"].append(deepcopy(male))
    rows = [
        row for row in _records(_book(department_data)["Details"]) if row["Course"] == "CS111 (1)"
    ]
    assert len(rows) == 2
    by_gender = {row["Gender"]: row for row in rows}
    assert by_gender["Female"]["Department students"] == 12
    assert by_gender["Male"]["Department students"] == 4
    assert "F-101: 10" in by_gender["Female"]["Room allocation (whole section)"]
    assert "F-101: 4" in by_gender["Male"]["Room allocation (whole section)"]


def test_instructors_require_saved_recorded_meeting_provenance(department_data):
    section = department_data["operations_snapshot"]["courses"]["CS111 (1)"]["sections"][0]
    section["instructor_source"] = "guessed_from_room"
    rows = _records(_book(department_data)["Details"])
    shared = next(row for row in rows if row["Course"] == "CS111 (1)")
    assert shared["Course instructor"] is None
    assert "Dr Saved" not in _text(_book(department_data)["Details"])


def test_unresolved_sections_rooms_and_online_exams_stay_visible(department_data):
    book = _book(department_data)
    text = _text(book["Details"])
    assert "General Studies" in text and "OVERFLOW" in text
    assert "UNASSIGNED" in text
    assert "missing" in text.lower() or "not recorded" in text.lower()
    assert "ambiguous" in text.lower()
    assert "online" in text.lower()
    assert "Unknown Gender" in text
    math_rows = [row for row in book["Details"].values if "MATH101" in row]
    assert math_rows
    assert not any(
        str(value).startswith(("F-", "M-"))
        for row in math_rows
        for value in row
        if value is not None
    )
    unresolved = [row for row in _records(book["Details"]) if row["Course"] == "GS101"]
    assert {row["Teaching section"]: row["Department students"] for row in unresolved} == {
        "Not recorded": 3,
        "Ambiguous section": 2,
    }
    assert all(row["Course instructor"] is None for row in unresolved)
    assert all(row["Day"] == "OVERFLOW" for row in unresolved)


def test_gender_filter_changes_population_without_live_membership_queries(department_data):
    book = _book(department_data, genders=["F"])
    text = _text(book["Details"])
    assert "Foundations" in text and "Mathematics" in text
    assert "Programming I" not in text and "Machine Learning" not in text
    assert "Unknown Gender" not in text


def test_arabic_workbook_has_rtl_and_printable_frozen_headers(department_data):
    book = _book(department_data, language="ar", genders=["M", "F"])
    assert "التفاصيل" in book.sheetnames
    assert len(book.sheetnames) >= 3
    assert book["التفاصيل"].sheet_view.pane.ySplit == 4
    assert book["التفاصيل"].freeze_panes == "C5"
    for sheet in book:
        assert sheet.sheet_view.rightToLeft is True
        assert sheet.print_title_rows
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.page_setup.fitToWidth == 1
        if sheet.title != "التفاصيل":
            assert sheet.freeze_panes is None


@pytest.mark.parametrize("language", ["en", "ar"])
def test_details_table_keeps_editable_duties_and_hidden_record_identity_together(
    department_data, language
):
    details = _book(department_data, language=language)[
        "التفاصيل" if language == "ar" else "Details"
    ]
    assert details.freeze_panes == "C5"
    assert details.column_dimensions["P"].hidden is True
    source_rows = [row for row in details.iter_rows(min_row=5) if row[15].value]
    ids = [row[15].value for row in source_rows]
    assert ids and all(isinstance(value, str) and value for value in ids)
    assert len(ids) == len(set(ids))
    assert len(details.tables) == 1
    table = next(iter(details.tables.values()))
    assert table.ref == f"A4:P{4 + len(source_rows)}"
    assert table.autoFilter.ref == table.ref
    assert len(table.tableColumns) == 16
    for row in source_rows:
        assert row[11].value is None and row[12].value is None
        assert row[11].fill.patternType == row[12].fill.patternType == "solid"
        assert row[11].fill.fgColor == row[12].fill.fgColor
        assert row[11].fill.fgColor != row[10].fill.fgColor


def test_record_ids_disambiguate_equal_visible_sections_and_same_code_variants(department_data):
    other = _section(11, "F01", "F", {"AI2": 4}, instructors=["Dr Alternate"])
    entry = department_data["schedule"][0]
    entry["enrolled_count"] += 4
    entry["programs"].append("AI2")
    entry["rooms"].append(_room("F-103", other))
    course = department_data["operations_snapshot"]["courses"]["CS111 (1)"]
    course["sections"].append(deepcopy(other))
    course["program_counts"].append({"program": "AI2", "gender": "F", "student_count": 4})
    department_data["section_enrollment"]["CS111 (1)"].append(deepcopy(other))
    rows = _records(_book(department_data)["Details"])
    equal_labels = [
        row for row in rows if row["Course"] == "CS111 (1)" and row["Teaching section"] == "F01"
    ]
    assert len(equal_labels) == 2
    assert {row["Department students"] for row in equal_labels} == {4, 12}
    assert len({row["Record ID"] for row in rows}) == len(rows)
    first_ids = {row["Record ID"] for row in rows if row["Course"] == "CS111 (1)"}
    second_ids = {row["Record ID"] for row in rows if row["Course"] == "CS111 (2)"}
    assert first_ids and second_ids and first_ids.isdisjoint(second_ids)


@pytest.mark.parametrize("language", ["en", "ar"])
def test_print_manual_fields_reference_the_matching_canonical_section_in_details(
    department_data, language
):
    book = _book(department_data, language=language)
    details = book["التفاصيل" if language == "ar" else "Details"]
    source_rows = {row[15].value: row for row in details.iter_rows(min_row=5) if row[15].value}
    end = 4 + len(source_rows)
    visited = set()
    formulas = []
    for sheet in _full_print_sheets(book):
        for row in sheet:
            if row[6].data_type != "f":
                continue
            record_ids = re.findall(r'MATCH\("([^\"]+)"', row[6].value)
            assert record_ids and len(set(record_ids)) == 1
            record_id = record_ids[0]
            assert record_id in source_rows and record_id not in visited
            visited.add(record_id)
            source = source_rows[record_id]
            assert row[0].data_type == "s"
            assert source[0].value in row[0].value and source[6].value in row[0].value
            if source[13].value:
                assert all(name in row[0].value for name in source[13].value.splitlines())
            assert [cell.value for cell in row[1:4]] == [
                source[1].value,
                source[7].value,
                source[8].value,
            ]
            assert row[4].value == (source[9].value or "—")
            assert row[5].value == source[10].value
            for cell, source_column in ((row[6], "L"), (row[7], "M")):
                assert cell.data_type == "f"
                assert cell.value.startswith("=IFERROR(") and "INDEX(" in cell.value
                assert '="","",' in cell.value
                assert set(re.findall(r'MATCH\("([^\"]+)"', cell.value)) == {record_id}
                references = set(
                    re.findall(r"'([^']+)'!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", cell.value)
                )
                assert references == {
                    (details.title, source_column, "5", source_column, str(end)),
                    (details.title, "P", "5", "P", str(end)),
                }
                # A broken match must leave an actionable visible value.
                assert not cell.value.endswith(',"")')
                formulas.append((record_id, source_column, cell.value))
    assert visited == set(source_rows)
    # Stable IDs are independent of the physical order of the saved exams.
    department_data["schedule"].reverse()
    reordered = _book(department_data, language=language)
    assert list(reordered[details.title].values) == list(details.values)
    # Simulate sorting the complete Excel Table after manual duties are filled.
    # Resolve each formula's MATCH key against the reordered lookup range and
    # inspect its INDEX result; never rely on its original source row number.
    expected_duties = {}
    for index, (record_id, row) in enumerate(source_rows.items(), 1):
        row[11].value = f"Invigilator {index}"
        row[12].value = f"Supervisor {index}"
        expected_duties[record_id] = {"L": row[11].value, "M": row[12].value}
    values = [list(cell.value for cell in row) for row in source_rows.values()]
    for row_number, values_row in enumerate(reversed(values), 5):
        for column, value in enumerate(values_row, 1):
            details.cell(row_number, column).value = value
    lookup_ids = [details.cell(row_number, 16).value for row_number in range(5, end + 1)]
    for record_id, column, formula in formulas:
        matched_row = 5 + lookup_ids.index(record_id)
        assert details[f"{column}{matched_row}"].value == expected_duties[record_id][column]
        assert record_id in formula


@pytest.mark.parametrize("language", ["en", "ar"])
def test_print_headers_repeat_by_page_and_breaks_preserve_session_scope(department_data, language):
    book = _book(department_data, language=language)
    details = book["التفاصيل" if language == "ar" else "Details"]
    saw_break = False
    for sheet in book:
        if sheet is details:
            continue
        assert sheet.max_column == 8
        assert sheet.print_title_rows == "$1:$4"
        header = tuple(cell.value for cell in sheet[4])
        assert all(value is not None for value in header)
        assert sum(tuple(cell.value for cell in row) == header for row in sheet) == 1
        # Session labels must span the sheet and stay immediately attached to
        # their first course row, including when a page starts mid-session.
        bands = {
            merged.min_row
            for merged in sheet.merged_cells.ranges
            if merged.min_col == 1
            and merged.max_col == 8
            and merged.min_row > 4
            and merged.min_row == merged.max_row
        }
        for page_break in sheet.row_breaks.brk:
            saw_break = True
            next_row = page_break.id + 1
            assert 4 < next_row < sheet.max_row
            assert next_row in bands
            assert sheet.cell(next_row + 1, 7).data_type == "f"
    assert saw_break


def test_a_long_print_session_repeats_its_band_without_losing_or_repeating_sections(
    department_data,
):
    sections = [_section(100 + index, f"M{index:03}", "M", {"AI2": 1}) for index in range(40)]
    entry = department_data["schedule"][1]
    entry["enrolled_count"] = 40
    entry["programs"] = ["AI2"]
    entry["rooms"] = [_room(f"M-{200 + index}", section) for index, section in enumerate(sections)]
    course = department_data["operations_snapshot"]["courses"]["CS111 (2)"]
    course["sections"] = deepcopy(sections)
    course["program_counts"] = [{"program": "AI2", "gender": "M", "student_count": 40}]
    department_data["section_enrollment"]["CS111 (2)"] = deepcopy(sections)
    sheet = _book(department_data)["Print (M)"]
    rows = [row for row in sheet if "CS111 (2)" in str(row[0].value)]
    assert len(rows) == 40
    assert {row[1].value for row in rows} == {f"M{index:03}" for index in range(40)}
    assert sum(row[2].value for row in rows) == 40
    continuations = [
        page_break
        for page_break in sheet.row_breaks.brk
        if "Mon" in str(sheet.cell(page_break.id + 1, 1).value)
    ]
    assert len(continuations) >= 2
    for page_break in continuations:
        assert "08:00-10:00" in sheet.cell(page_break.id + 1, 1).value
        assert "CS111 (2)" in sheet.cell(page_break.id + 2, 1).value
        assert sheet.cell(page_break.id + 2, 7).data_type == "f"
    header = tuple(cell.value for cell in sheet[4])
    assert sum(tuple(cell.value for cell in row) == header for row in sheet) == 1


def test_print_notes_preserve_actionable_gaps_without_repeating_routine_explanations(
    department_data,
):
    book = _book(department_data)
    detail_text = _text(book["Details"])
    print_text = "\n".join(_text(sheet) for sheet in book if sheet.title != "Details")
    assert "Section shared with other programmes" in detail_text
    assert "Section shared with other programmes" not in print_text
    assert "Department students per room are not determined" in print_text
    assert "Room allocation needs review" in print_text
    assert "Exam time not assigned" in print_text
    assert "Ambiguous" in print_text


@pytest.mark.parametrize(
    "day,valid,invalid",
    [
        ("Sunday", "2026-09-20", "2026-09-21"),
        ("monday", "2026-09-21", "2026-09-20"),
        ("الأحد", "2026-09-20", "2026-09-21"),
        ("الاثنين", "2026-09-21", "2026-09-20"),
    ],
)
def test_full_english_and_arabic_weekdays_validate_entered_dates(
    department_data, day, valid, invalid
):
    department_data["slots"] = [{"index": 0, "day": day, "period": "08:00-10:00"}]
    department_data["schedule"][0]["day"] = day
    book = _book(department_data, dates={day: valid})
    first = next(row for row in _records(book["Details"]) if row["Course"] == "CS111 (1)")
    assert first["Date"].date().isoformat() == valid
    with pytest.raises(ValueError, match="weekday"):
        _export(department_data, dates={day: invalid})


@pytest.mark.parametrize(
    "dangerous", ['=HYPERLINK("https://invalid.example","click")', "+SUM(1,2)", "-1+2", "@SUM(1,2)"]
)
def test_formula_like_recorded_text_remains_literal(department_data, dangerous):
    department_data["schedule"][0]["course_name"] = dangerous
    snapshot = department_data["operations_snapshot"]["courses"]["CS111 (1)"]
    snapshot["course_name"] = dangerous
    snapshot["sections"][0]["instructors"] = [dangerous]
    book = _book(department_data)
    relevant = [
        cell for sheet in book for row in sheet for cell in row if dangerous in str(cell.value)
    ]
    assert relevant
    assert all(cell.data_type == "s" for cell in relevant)


def test_multiple_departments_zip_separate_complete_workbooks_without_leftover_files(
    department_data, tmp_path, monkeypatch
):
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    ids = [item["id"] for item in department_export_options(department_data)["departments"]]
    assert len(ids) == 2
    content, filename, content_type = _export(department_data, departments=ids)
    assert filename.endswith(".zip") and content_type == "application/zip"
    with ZipFile(BytesIO(content)) as archive:
        assert archive.testzip() is None
        assert len(archive.namelist()) == 2
        assert len(set(archive.namelist())) == 2
        for name in archive.namelist():
            assert name.endswith(".xlsx") and "/" not in name and "\\" not in name
            book = load_workbook(BytesIO(archive.read(name)))
            assert "Details" in book.sheetnames
            text = _text(book["Details"])
            if "_ai-ds_" in name:
                assert "CS299" not in text and "MATH101" in text
            else:
                assert "CS299" in text and "MATH101" not in text
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("language", ["en", "ar"])
@pytest.mark.parametrize("all_departments", [False, True])
def test_excel_tables_own_the_only_filter_in_single_and_zipped_exports(
    department_data, language, all_departments
):
    departments = ["ai-ds"]
    if all_departments:
        departments = [
            profile["id"] for profile in department_export_options(department_data)["departments"]
        ]
    content, filename, _ = _export(department_data, language=language, departments=departments)
    if all_departments:
        with ZipFile(BytesIO(content)) as outer:
            workbooks = [(name, outer.read(name)) for name in outer.namelist()]
        assert len(workbooks) == len(departments)
    else:
        workbooks = [(filename, content)]
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    for name, raw in workbooks:
        assert name.endswith(".xlsx")
        book = load_workbook(BytesIO(raw))
        details = book["التفاصيل" if language == "ar" else "Details"]
        headers = [cell.value for cell in details[4]]
        assert len(headers) == len({value.casefold() for value in headers}) == 16
        with ZipFile(BytesIO(raw)) as archive:
            workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
            defined_names = workbook_xml.findall("s:definedNames/s:definedName", ns)
            assert all(
                item.get("name") not in {"_xlnm._FilterDatabase", "_FilterDatabase"}
                for item in defined_names
            )
            table_names = [
                item
                for item in archive.namelist()
                if item.startswith("xl/tables/") and item.endswith(".xml")
            ]
            assert len(table_names) == 1
            table = ET.fromstring(archive.read(table_names[0]))
            filters = table.findall("s:autoFilter", ns)
            assert len(filters) == 1
            assert filters[0].get("ref") == table.get("ref")
            columns = table.findall("s:tableColumns/s:tableColumn", ns)
            assert [column.get("name") for column in columns] == headers
            assert len({column.get("id") for column in columns}) == 16
            for part in archive.namelist():
                if part.startswith("xl/worksheets/sheet") and part.endswith(".xml"):
                    sheet_xml = ET.fromstring(archive.read(part))
                    # A second worksheet-level filter over the table range is
                    # accepted by openpyxl but makes Excel repair the table.
                    assert sheet_xml.find("s:autoFilter", ns) is None
        record_ids = [row[15].value for row in details.iter_rows(min_row=5) if row[15].value]
        assert len(record_ids) == len(set(record_ids))
        referenced = Counter()
        for sheet in _full_print_sheets(book):
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        keys = set(re.findall(r'MATCH\("([^\"]+)"', cell.value))
                        assert len(keys) == 1
                        referenced.update(keys)
        assert referenced == Counter({key: 2 for key in record_ids})


def test_optional_dates_are_saved_in_export_but_do_not_mutate_run(department_data):
    before = deepcopy(department_data)
    without_dates = _book(department_data)
    with_dates = _book(
        department_data, dates={"Sun": "2026-09-20", "Mon": "2026-09-21", "Tue": "2026-09-22"}
    )
    assert "2026-09-20" in _text(with_dates["Details"])
    assert "2026-09-20" not in _text(without_dates["Details"])
    assert department_data == before


def _weekly_fixture():
    data = make_department_data()
    days = ["W1 Sun", "W2 Sun", "W2 Mon", "W10 Sun", "Special session"]
    data["slots"] = [
        {"index": index, "day": day, "period": "08:00-10:00"} for index, day in enumerate(days)
    ]
    assignments = {
        "CS111 (1)": "W1 Sun",
        "CS111 (2)": "W2 Sun",
        "MATH101": "W2 Mon",
        "AI200": "W10 Sun",
        "CS299": "Special session",
        "DS300": "W2 Mon",
    }
    for entry in data["schedule"]:
        if entry["course_code"] in assignments:
            entry["day"] = assignments[entry["course_code"]]
            entry["slot_index"] = days.index(entry["day"])
    return data


def _week_sheet_name(language, week, gender):
    return (
        f"الأسبوع {week} - {'الطلاب' if gender == 'M' else 'الطالبات'}"
        if language == "ar"
        else f"Week {week} ({gender})"
    )


def _full_sheet_name(language, gender):
    return (
        {"M": "الطلاب", "F": "الطالبات", "U": "غير محدد"}[gender]
        if language == "ar"
        else f"Print ({gender})"
    )


def _print_record_rows(sheet):
    records = {}
    for row in sheet:
        if row[6].data_type != "f":
            continue
        ids = set(re.findall(r'MATCH\("([^\"]+)"', row[6].value))
        assert len(ids) == 1
        record_id = next(iter(ids))
        assert record_id not in records
        records[record_id] = tuple(cell.value for cell in row)
    return records


@pytest.mark.parametrize("language", ["en", "ar"])
def test_weekly_sheets_append_nonempty_groups_without_changing_full_views_or_source(language):
    data = _weekly_fixture()
    original = deepcopy(data)
    original_options = department_export_options(data)
    book = _book(data, language=language)
    details_name = "التفاصيل" if language == "ar" else "Details"
    original_names = [details_name] + [
        _full_sheet_name(language, gender) for gender in ("M", "F", "U")
    ]
    expected = [(1, "F", 12), (2, "M", 12), (2, "F", 6), (10, "M", 4)]
    assert book.sheetnames == original_names + [
        _week_sheet_name(language, week, gender) for week, gender, _ in expected
    ]
    details = book[details_name]
    source_ids = {row[15].value for row in details.iter_rows(min_row=5) if row[15].value}
    full_ids = set()
    for gender, total in (("M", 16), ("F", 23), ("U", 2)):
        full_rows = _print_record_rows(book[_full_sheet_name(language, gender)])
        assert sum(row[2] for row in full_rows.values()) == total
        assert full_ids.isdisjoint(full_rows)
        full_ids.update(full_rows)
    assert full_ids == source_ids
    for week, gender, total in expected:
        weekly = book[_week_sheet_name(language, week, gender)]
        full = book[_full_sheet_name(language, gender)]
        weekly_rows, full_rows = _print_record_rows(weekly), _print_record_rows(full)
        assert sum(row[2] for row in weekly_rows.values()) == total
        assert all(row == full_rows[record_id] for record_id, row in weekly_rows.items())
        assert tuple(cell.value for cell in weekly[4]) == tuple(cell.value for cell in full[4])
        assert weekly.print_title_rows == full.print_title_rows == "$1:$4"
        assert weekly.freeze_panes is None
        assert weekly.sheet_view.rightToLeft == full.sheet_view.rightToLeft
        assert weekly.page_setup.orientation == full.page_setup.orientation == "landscape"
        assert weekly.page_setup.fitToWidth == full.page_setup.fitToWidth == 1
        for column in "ABCDEFGH":
            assert weekly.column_dimensions[column].width == full.column_dimensions[column].width
        assert "General Studies" not in _text(weekly)
        assert "Unknown Gender" not in _text(weekly)
        assert "CS Only" not in _text(weekly)
    assert "General Studies" in _text(book[_full_sheet_name(language, "F")])
    assert "Unknown Gender" in _text(book[_full_sheet_name(language, "U")])
    assert department_export_options(data) == original_options
    assert data == original


@pytest.mark.parametrize("language", ["en", "ar"])
def test_explicit_week_order_is_numeric_and_omitted_weeks_are_not_filled(language):
    data = _weekly_fixture()
    data["slots"] = [data["slots"][3], *data["slots"][:3], data["slots"][4]]
    for index, slot in enumerate(data["slots"]):
        slot["index"] = index
    for entry in data["schedule"]:
        if entry["day"] != "OVERFLOW":
            entry["slot_index"] = next(
                slot["index"] for slot in data["slots"] if slot["day"] == entry["day"]
            )
    book = _book(data, language=language)
    expected = [
        _week_sheet_name(language, week, gender)
        for week, gender in [(1, "F"), (2, "M"), (2, "F"), (10, "M")]
    ]
    assert book.sheetnames[4:] == expected


@pytest.mark.parametrize("language", ["en", "ar"])
def test_bare_weekdays_use_full_slot_order_and_roll_over_from_partial_thursday(language):
    data = make_department_data()
    data["slots"] = [
        {"index": index, "day": day, "period": period}
        for index, (day, period) in enumerate(
            [
                ("Thu", "08:00-10:00"),
                ("Thu", "10:00-12:00"),
                ("Sun", "08:00-10:00"),
                ("Sun", "10:00-12:00"),
                ("Mon", "08:00-10:00"),
                ("Special session", "08:00-10:00"),
            ]
        )
    ]
    assignments = {"CS111 (1)": 3, "CS111 (2)": 2, "MATH101": 4, "AI200": 5, "CS299": 0, "DS300": 4}
    for entry in data["schedule"]:
        if entry["course_code"] in assignments:
            slot = data["slots"][assignments[entry["course_code"]]]
            entry.update(day=slot["day"], period=slot["period"], slot_index=slot["index"])
    book = _book(data, language=language)
    assert book.sheetnames[4:] == [
        _week_sheet_name(language, 2, "M"),
        _week_sheet_name(language, 2, "F"),
    ]
    assert (
        sum(row[2] for row in _print_record_rows(book[_week_sheet_name(language, 2, "M")]).values())
        == 12
    )
    assert (
        sum(row[2] for row in _print_record_rows(book[_week_sheet_name(language, 2, "F")]).values())
        == 18
    )
    assert "Machine Learning" in _text(book[_full_sheet_name(language, "M")])
    assert "Machine Learning" not in _text(book[_week_sheet_name(language, 2, "M")])


@pytest.mark.parametrize("language", ["en", "ar"])
def test_weekly_sheets_obey_department_and_gender_filters(language):
    data = _weekly_fixture()
    female = _book(data, language=language, genders=["F"])
    details_name = "التفاصيل" if language == "ar" else "Details"
    assert female.sheetnames == [
        details_name,
        _full_sheet_name(language, "F"),
        _week_sheet_name(language, 1, "F"),
        _week_sheet_name(language, 2, "F"),
    ]
    assert "Programming I" not in "\n".join(_text(sheet) for sheet in female)
    cs = _book(data, language=language, departments=["cs"])
    assert cs.sheetnames == [
        details_name,
        _full_sheet_name(language, "M"),
        _full_sheet_name(language, "F"),
        _week_sheet_name(language, 1, "F"),
    ]
    weekly_rows = _print_record_rows(cs[_week_sheet_name(language, 1, "F")])
    assert sum(row[2] for row in weekly_rows.values()) == 17
    assert "Mathematics" not in "\n".join(_text(sheet) for sheet in cs)
    assert "CS Only" in _text(cs[_full_sheet_name(language, "M")])
    unknown = _book(data, language=language, genders=["U"])
    assert unknown.sheetnames == [details_name, _full_sheet_name(language, "U")]


@pytest.mark.parametrize("language", ["en", "ar"])
def test_unknown_day_labels_and_overflow_are_not_guessed_into_weekly_views(language):
    data = make_department_data()
    for slot in data["slots"]:
        slot["day"] = ["Orientation", "W7 custom session", "Special session"][slot["index"]]
    for entry in data["schedule"]:
        if entry["day"] != "OVERFLOW":
            entry["day"] = data["slots"][entry["slot_index"]]["day"]
    book = _book(data, language=language)
    details_name = "التفاصيل" if language == "ar" else "Details"
    assert book.sheetnames == [details_name] + [
        _full_sheet_name(language, gender) for gender in ("M", "F", "U")
    ]
    assert (
        sum(
            sum(row[2] for row in _print_record_rows(sheet).values())
            for sheet in _full_print_sheets(book)
        )
        == 41
    )


@pytest.mark.parametrize("language", ["en", "ar"])
def test_bare_sunday_rollover_continues_from_an_explicit_week_anchor(language):
    data = make_department_data()
    labels = ["w3 Thursday", "Sunday", "Monday"]
    for slot in data["slots"]:
        slot["day"] = labels[slot["index"]]
    for entry in data["schedule"]:
        if entry["day"] != "OVERFLOW":
            entry["day"] = labels[entry["slot_index"]]
    book = _book(data, language=language)
    expected = [
        _week_sheet_name(language, week, gender) for week, gender in [(3, "F"), (4, "M"), (4, "F")]
    ]
    assert book.sheetnames[4:] == expected
    assert sum(row[2] for row in _print_record_rows(book[expected[0]]).values()) == 12
    assert sum(row[2] for row in _print_record_rows(book[expected[1]]).values()) == 16
    assert sum(row[2] for row in _print_record_rows(book[expected[2]]).values()) == 6


@pytest.mark.parametrize("language", ["en", "ar"])
def test_oversized_explicit_week_prefix_stays_full_only_until_a_valid_anchor(language):
    data = make_department_data()
    labels = ["W" + "9" * 5000 + " Thu", "Sun", "W2 Mon"]
    for slot in data["slots"]:
        slot["day"] = labels[slot["index"]]
    for entry in data["schedule"]:
        if entry["day"] != "OVERFLOW":
            entry["day"] = labels[entry["slot_index"]]
    # Move mathematics to the next valid anchor, keeping the preceding Sunday
    # unclassified after the unusable explicit week number.
    data["schedule"][2].update(day=labels[2], slot_index=2)
    book = _book(data, language=language)
    assert book.sheetnames[4:] == [
        _week_sheet_name(language, 2, "M"),
        _week_sheet_name(language, 2, "F"),
    ]
    assert all(len(name) <= 31 for name in book.sheetnames)
    assert (
        sum(row[2] for row in _print_record_rows(book[_week_sheet_name(language, 2, "M")]).values())
        == 4
    )
    assert (
        sum(row[2] for row in _print_record_rows(book[_week_sheet_name(language, 2, "F")]).values())
        == 6
    )
    assert "Foundations" in _text(book[_full_sheet_name(language, "F")])
    assert "Programming I" in _text(book[_full_sheet_name(language, "M")])


# ── Direction marks, never isolates ────────────────────────────

_ISOLATES = tuple(chr(code) for code in range(0x2066, 0x206A))
_MARKS = LRM + RLM
_DATED = {"Sun": "2026-09-20", "Mon": "2026-09-21", "Tue": "2026-09-22"}
_ARABIC_DAYS = {"Sun": "W1-الأحد", "Mon": "W1-الاثنين", "Tue": "W1-الثلاثاء"}


def _arabic_day_fixture():
    data = make_department_data()
    for slot in data["slots"]:
        slot["day"] = _ARABIC_DAYS[slot["day"]]
    for entry in data["schedule"]:
        entry["day"] = _ARABIC_DAYS.get(entry["day"], entry["day"])
    return data


def _long_session_fixture():
    # One Monday session long enough to spill over pages, so its band is
    # repeated at the top of each continuation page.
    data = make_department_data()
    sections = [_section(100 + index, f"M{index:03}", "M", {"AI2": 1}) for index in range(40)]
    entry = data["schedule"][1]
    entry["enrolled_count"] = 40
    entry["programs"] = ["AI2"]
    entry["rooms"] = [_room(f"M-{200 + index}", section) for index, section in enumerate(sections)]
    course = data["operations_snapshot"]["courses"]["CS111 (2)"]
    course["sections"] = deepcopy(sections)
    course["program_counts"] = [{"program": "AI2", "gender": "M", "student_count": 40}]
    data["section_enrollment"]["CS111 (2)"] = deepcopy(sections)
    return data


# Every shape a print band takes: dated and undated, weekly sheets, Arabic day
# labels (a date after an Arabic day printed reversed before the marks), and a
# session repeated on continuation pages.
_BAND_CASES = {
    "dated": (make_department_data, _DATED),
    "undated": (make_department_data, None),
    "weekly": (_weekly_fixture, None),
    "arabic-days-dated": (
        _arabic_day_fixture,
        {_ARABIC_DAYS[day]: value for day, value in _DATED.items()},
    ),
    "long-session": (_long_session_fixture, None),
}


def _all_profile_workbooks(data, language, dates):
    profiles = [profile["id"] for profile in department_export_options(data)["departments"]]
    payload = {"departments": profiles, "genders": ["M", "F", "U"], "language": language}
    if dates:
        payload["dates"] = dates
    content, filename, _ = _export(data, **payload)
    if filename.endswith(".zip"):
        with ZipFile(BytesIO(content)) as outer:
            workbooks = [outer.read(name) for name in outer.namelist()]
    else:
        workbooks = [content]
    assert len(workbooks) == len(profiles) >= 2
    return workbooks


@pytest.mark.parametrize("case", sorted(_BAND_CASES))
def test_arabic_department_files_carry_no_unicode_isolates(case):
    # Excel prints U+2066..U+2069 as visible LRI/PDI boxes and still reverses
    # the run. No part and no cell of any profile, cohort or sheet holds one.
    make, dates = _BAND_CASES[case]
    for raw in _all_profile_workbooks(make(), "ar", dates):
        with ZipFile(BytesIO(raw)) as archive:
            for part in archive.namelist():
                text = archive.read(part).decode("utf-8", errors="replace")
                assert not any(mark in text for mark in _ISOLATES), part
        book = load_workbook(BytesIO(raw), rich_text=True)
        for sheet in book:
            for row in sheet.iter_rows():
                for cell in row:
                    value = "" if cell.value is None else str(cell.value)
                    assert not any(mark in value for mark in _ISOLATES), (
                        sheet.title,
                        cell.coordinate,
                    )


def _bands(sheet):
    return [
        cell.value
        for row in sheet
        for cell in row
        if isinstance(cell.value, str) and "  |  " in cell.value
    ]


@pytest.mark.parametrize("language", ["en", "ar"])
@pytest.mark.parametrize("case", sorted(_BAND_CASES))
def test_every_print_band_marks_each_left_to_right_field(case, language):
    # Checked band by band on every print sheet after Details (full and
    # weekly, including bands repeated on continuation pages) of every profile
    # and cohort. Arabic: the day, an entered date and the period each read
    # LRM + run + RLM; empty runs and "date not entered" stay bare. English:
    # no marks, except one LRM closing an Arabic day label so the date and
    # time after it keep European digits.
    make, dates = _BAND_CASES[case]
    data = make()
    days = {slot["day"] for slot in data["slots"]} | {"OVERFLOW"}
    entered = set((dates or {}).values())
    arabic = language == "ar"
    undated = "التاريخ غير محدد" if arabic else "Date not entered"
    # Spelled out, not built with the helpers under test.
    wrap = (lambda run: LRM + run + RLM) if arabic else str
    arabic_days = set(_ARABIC_DAYS.values())
    weekly_sheets = dated_bands = repeated_bands = closed_arabic_days = 0
    for raw in _all_profile_workbooks(data, language, dates):
        book = load_workbook(BytesIO(raw))
        for sheet in book.worksheets[1:]:
            bands = _bands(sheet)
            assert bands, sheet.title
            weekly_sheets += sheet.title.startswith(("الأسبوع", "Week"))
            repeated_bands += len(bands) - len(set(bands))
            for band in bands:
                day, date_label, period = band.split("  |  ")
                core_day = day.strip(_MARKS)
                assert core_day in days, band
                closed = core_day + LRM if core_day in arabic_days else core_day
                assert day == (wrap(core_day) if arabic else closed), band
                closed_arabic_days += (not arabic) and day != core_day
                if date_label != undated:
                    dated_bands += 1
                    assert date_label.strip(_MARKS) in entered, band
                    assert date_label == wrap(date_label.strip(_MARKS)), band
                core = period.strip(_MARKS)
                assert core in {"08:00-10:00", ""}, band
                assert period == (wrap(core) if core else ""), band
                if not arabic:
                    assert RLM not in band and LRM not in date_label + period, band
    assert weekly_sheets > 0
    assert dated_bands > 0 if entered else dated_bands == 0
    assert repeated_bands > 0 if case == "long-session" else True
    assert closed_arabic_days > 0 if (case == "arabic-days-dated" and not arabic) else True
