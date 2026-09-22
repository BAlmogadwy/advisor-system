"""A checked manual placement is the same placement that Save and Excel publish."""

from collections import Counter
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from core import exam_views
from core.models import Course, ExamTimetableRun, ProgrammeRequirement, Room, Student, StudentCourse
from core.services import exam_timetable
from tests.exam_source_factory import scraped_exam_registration
from tests.test_exam_export_source import FUNDAMENTALS, PROGRAMMING
from tests.test_exam_export_source import export_client as export_client
from tests.test_exam_export_source import export_courses as export_courses

pytestmark = pytest.mark.django_db
DAYS = ["Sun", "Mon", "Tue"]
PERIODS = ["08:00-10:00", "13:00-15:00"]


def _post(client, route, payload):
    response = client.post(reverse(route), payload, content_type="application/json")
    assert response.status_code == 200, response.content
    return response.json()


def _placements(data):
    return {
        entry["course_code"]: (entry["day"], entry["period"], entry["slot_index"])
        for entry in data["schedule"]
    }


def _stable_qa(data):
    qa = deepcopy(data["qa"])
    # Check and Save capture at different instants; all substantive QA agrees.
    qa["enrolment_snapshot"].pop("snapshot_timestamp", None)
    return qa


def _export_and_reconcile(client, data, tmp_path):
    response = client.get(reverse("exam_timetable_export", args=[data["run_id"]]))
    assert response.status_code == 200, response.content
    content = b"".join(response.streaming_content)
    response.close()
    assert not list(tmp_path.glob("*.xlsx"))
    with ZipFile(BytesIO(content)) as archive:
        assert archive.testzip() is None
    book = load_workbook(BytesIO(content))
    expected_names = {
        "Schedule",
        "Schedule (M)",
        "Schedule (F)",
        "Courses",
        "Students (M)",
        "Students (F)",
        "QA Summary",
    }
    if data["assign_rooms"]:
        expected_names |= {"Room Assignments", "Invigilators"}
    assert set(book.sheetnames) == expected_names
    rows = list(book["Courses"].values)
    courses = {row[0]: dict(zip(rows[0], row, strict=True)) for row in rows[1:]}
    assert set(courses) == set(data["courses"])
    pinned = {pin["course_code"] for pin in data["pinned"]}
    for entry in data["schedule"]:
        row = courses[entry["course_code"]]
        assert row["Course Name"] == entry["course_name"]
        assert row["Source Course Code"] == entry["source_course_code"]
        assert row["Enrolled Students"] == entry["enrolled_count"]
        assert (row["Day"], row["Period"], row["Slot Index"]) == (
            entry["day"],
            entry["period"],
            entry["slot_index"],
        )
        assert row["Fixed Time"] == ("Yes" if entry["course_code"] in pinned else "No")
        for sheet_name in ("Schedule", "Schedule (F)"):
            sheet = book[sheet_name]
            period_column = next(cell.column for cell in sheet[1] if cell.value == entry["period"])
            cells = [
                str(row[period_column - 1].value or "")
                for row in sheet.iter_rows(min_row=2)
                if row[0].value == entry["day"]
            ]
            assert any(
                entry["course_name"] in text and entry["course_code"] in text for text in cells
            )

    values = list(book["Students (F)"].values)
    female_rows = {row[0]: dict(zip(values[0], row, strict=True)) for row in values[1:]}
    assert female_rows["TOTAL EXAM ENROLMENTS"]["Exam Enrolments"] == 5
    for entry in data["schedule"]:
        assert female_rows[entry["course_code"]]["Exam Enrolments"] == entry["enrolled_count"]
        assert female_rows[entry["course_code"]]["Enrolled Sections"] == 1
        assert female_rows[entry["course_code"]]["Missing Section Enrolments"] == 0
        assert female_rows[entry["course_code"]]["Ambiguous Section Enrolments"] == 0
    assert book["Students (M)"]["A2"].value == "No M enrolments in this timetable"

    metrics = {
        row[0]: row[1]
        for row in book["QA Summary"].values
        if row[0]
        in {
            "Unique Students",
            "Total Courses",
            "Fixed Exams",
            "Same-Slot Conflicts",
            "Schedule Violations",
            "Approved-Only Student Clashes",
            "Hard Student Clashes",
            "Unassigned Room Groups",
        }
    }
    assert metrics["Unique Students"] == 3
    assert metrics["Total Courses"] == 3
    assert metrics["Fixed Exams"] == len(pinned)
    assert metrics["Same-Slot Conflicts"] == data["qa"]["conflict_count"]
    assert metrics["Schedule Violations"] == data["qa"]["schedule_violation_count"]
    assert metrics["Approved-Only Student Clashes"] == 0
    assert metrics["Hard Student Clashes"] == data["qa"]["conflict_count"]

    expected_rooms = [
        (
            entry["course_code"],
            entry["day"],
            entry["period"],
            room["section"],
            room["room_code"],
            room["student_count"],
            room["gender"],
        )
        for entry in data["schedule"]
        for room in entry.get("rooms", [])
    ]
    if data["assign_rooms"]:
        values = list(book["Room Assignments"].values)
        room_rows = [
            dict(zip(values[0], row, strict=True)) for row in values[1:] if row[2] in courses
        ]
        assert Counter(
            (
                row["Course"],
                row["Day"],
                row["Period"],
                row["Section"],
                row["Room"],
                row["Students"],
                row["Gender"],
            )
            for row in room_rows
        ) == Counter(expected_rooms)
        assert sum(row["Students"] for row in room_rows) == 5
        assert all(row["Students"] <= row["Capacity"] for row in room_rows)
        assert all(row["Section Mapping"] == "Recorded" for row in room_rows)
        values = list(book["Invigilators"].values)
        detail_header = next(
            i for i, row in enumerate(values) if row[:3] == ("Day", "Period", "Course")
        )
        invigilators = [
            dict(zip(values[detail_header], row, strict=True))
            for row in values[detail_header + 1 :]
            if row[2] in courses
        ]
        assert Counter(
            (
                row["Course"],
                row["Day"],
                row["Period"],
                row["Section"],
                row["Room"],
                row["Students"],
                row["Gender"],
            )
            for row in invigilators
        ) == Counter(expected_rooms)
        assert all(row["Invigilators"] == 1 for row in invigilators)
        assert sum(row["Invigilators"] for row in invigilators) == 5
    else:
        assert expected_rooms == []
        assert metrics["Unassigned Room Groups"] == 0
        assert any(
            "Room assignment not requested" in str(cell.value)
            for row in book["Schedule"]
            for cell in row
        )
    book.close()


@pytest.mark.parametrize("assign_rooms", [False, True])
def test_check_save_export_agree_without_pinning_manual_moves(
    export_courses,  # noqa: F811 — use the imported shared pytest fixture.
    export_client,  # noqa: F811 — use the imported shared pytest fixture.
    monkeypatch,
    tmp_path,
    assign_rooms,
):
    companion = Course.objects.create(course_code="EX222", credit_hours=3)
    ProgrammeRequirement.objects.create(
        program="AI",
        course_code="EX222",
        course_name="Second exam",
        programme_term=1,
    )
    for student in Student.objects.filter(program="AI"):
        StudentCourse.objects.create(student=student, course=companion, status="studying")
        scraped_exam_registration(student, companion)
    for number in range(4):
        Room.objects.create(
            room_code=f"F-{number}", capacity=1, section="F", building="Exam College"
        )
    monkeypatch.setattr(exam_timetable, "RUNTIME_DIR", tmp_path)
    scope = {"programs": ["AI", "AI2"], "sections": ["F"]}
    _, metadata = exam_timetable.build_enrolled_sets_with_meta(**scope)
    code_by_name = {entry["course_name"]: code for code, entry in metadata.items()}
    programming = code_by_name[PROGRAMMING]
    fundamentals = code_by_name[FUNDAMENTALS]
    initial = {
        "label": "Manual changes export",
        "days": DAYS,
        "periods": PERIODS,
        "max_per_day": 2,
        **scope,
        "assign_rooms": assign_rooms,
        "selected_courses": sorted(metadata),
        "selected_course_entries": [
            {"course_code": code, **entry} for code, entry in metadata.items()
        ],
        "pinned": [],
        "thin_conflict_threshold": 0,
    }
    built = _post(export_client, "exam_timetable_build", initial)
    original_run = ExamTimetableRun.objects.get(pk=built["run_id"])
    original_json = original_run.result_json
    changed = deepcopy(built["schedule"])
    for entry in changed:
        entry["day"], entry["period"] = (
            ("Tue", PERIODS[1]) if entry["course_code"] == fundamentals else ("Sun", PERIODS[0])
        )
        # Deliberately stale indexes cannot override an explicit manual placement.
        entry["slot_index"] = 999
    expected_placements = {
        programming: ("Sun", PERIODS[0], 0),
        "EX222": ("Sun", PERIODS[0], 0),
        fundamentals: ("Tue", PERIODS[1], 5),
    }
    assert _placements(built) != expected_placements
    payload = {
        **initial,
        "previous_run_id": built["run_id"],
        "base_schedule": changed,
        "programs": [],
        "sections": [],
        "editor_revision": 7,
    }

    def scheduling_is_forbidden(*args, **kwargs):
        pytest.fail("Check and Save must evaluate the submitted placements without scheduling")

    with monkeypatch.context() as patch:
        patch.setattr(exam_views, "schedule", scheduling_is_forbidden)
        patch.setattr(exam_timetable, "schedule", scheduling_is_forbidden)
        patch.setattr(exam_timetable, "_rebalance_invigilators_pass", scheduling_is_forbidden)
        # The evaluator holds its own reference to the pass; patch that too.
        patch.setattr(
            "core.services.exam_evaluation._rebalance_invigilators_pass",
            scheduling_is_forbidden,
        )
        count_before = ExamTimetableRun.objects.count()
        checked = _post(export_client, "exam_timetable_draft_impact", payload)
        assert ExamTimetableRun.objects.count() == count_before
        assert "run_id" not in checked
        assert checked["editor_revision"] == 7
        assert checked["input_fingerprint"]
        assert _placements(checked) == expected_placements
        assert checked["pinned"] == []
        assert checked["enrollment_scope"] == scope
        assert checked["students_count"] == 3
        assert checked["courses_count"] == 3
        assert checked["assign_rooms"] is assign_rooms
        assert checked["qa"]["conflict_count"] == 2
        assert checked["qa"]["bucket_day_violations_count"] == 1
        assert checked["qa"]["manual_override_count"] == 3
        assert checked["qa"]["schedule_violation_count"] == 3
        assert checked["qa"]["approved_thin_conflict_count"] == 0
        assert checked["qa"]["hard_conflict_count"] == 2
        assert checked["qa"]["rebalance_moves"] == 0
        assert checked["qa"]["multi_sitting_sections"] == (2 if assign_rooms else 0)
        assert {
            entry["course_name"]
            for entry in checked["schedule"]
            if entry["source_course_code"] == "CS111"
        } == {PROGRAMMING, FUNDAMENTALS}
        if assign_rooms:
            assert checked["qa"]["rooms"]["total_demand"] == 5
            assert not checked["qa"]["rooms"]["unassigned_room_sections"]
            assert checked["qa"]["building_footprint"]
        else:
            assert checked["rooms_count"] == 0
            assert all(not entry.get("rooms") for entry in checked["schedule"])
        original_run.refresh_from_db()
        assert original_run.result_json == original_json
        saved = _post(
            export_client,
            "exam_timetable_build",
            {
                **payload,
                "mode": "save_loaded_changes",
                "expected_input_fingerprint": checked["input_fingerprint"],
            },
        )
        assert ExamTimetableRun.objects.count() == count_before + 1
    assert saved["input_fingerprint"] == checked["input_fingerprint"]
    assert saved["schedule"] == checked["schedule"]
    assert _stable_qa(saved) == _stable_qa(checked)
    assert saved["primary_status"] == checked["primary_status"]
    assert saved["status_flags"] == checked["status_flags"]
    assert saved["pinned"] == []

    detail_response = export_client.get(reverse("exam_timetable_detail", args=[saved["run_id"]]))
    assert detail_response.status_code == 200
    detail = detail_response.json()
    assert detail["schedule"] == saved["schedule"]
    assert detail["qa"] == saved["qa"]
    _export_and_reconcile(export_client, detail, tmp_path)

    # Optimization is an explicit next operation. Only the pin the registrar
    # now selects is fixed; the earlier manual moves were never implicit pins.
    fixed = {"course_code": programming, "day": "Sun", "period": PERIODS[0]}
    optimized = _post(
        export_client,
        "exam_timetable_build",
        {
            **payload,
            "mode": "optimize_loaded",
            "previous_run_id": saved["run_id"],
            "base_schedule": saved["schedule"],
            "pinned": [fixed],
        },
    )
    assert optimized["pinned"] == [fixed]
    assert _placements(optimized)[programming] == expected_placements[programming]
    assert _placements(optimized)["EX222"][0] != "Sun"
    assert optimized["qa"]["conflict_count"] == 0
    assert optimized["qa"]["bucket_day_violations_count"] == 0
    assert optimized["enrollment_scope"] == scope
    assert optimized["assign_rooms"] is assign_rooms
    assert optimized["students_count"] == 3
    _export_and_reconcile(export_client, optimized, tmp_path)
