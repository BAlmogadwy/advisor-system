"""H15: a cross-course overlap is critical only when it blocks students.

"Shared" counts demand -- students who NEED both courses -- not students seated
in both sections. With sibling sections those students can usually take a
clash-free pair, so the overlap costs choice, not a timetable. Only when no
placed section of one course fits round a placed section of the other are they
blocked, and only then does the overlap block publish. Every reader (badges,
next actions, the publish gate, the student-impact panel) takes that one answer
from ``detect_board_conflicts``.
"""

from __future__ import annotations

import pytest

from core.models import (
    DeliveryBoard,
    ScenarioStudentCourseRequest,
    SectionPlacement,
    TermSection,
    TermSectionMeeting,
    TimetableScenario,
)
from core.services.timetable_workspace import (
    build_scenario_builder_actions,
    check_publish_readiness,
    compute_affected_students,
    detect_board_conflicts,
    get_scenario_boards_summary,
)

pytestmark = pytest.mark.django_db

NETS = "AI201::NEURAL_NETWORKS"
ALGEBRA = "MATH201::LINEAR_ALGEBRA"


def _board() -> DeliveryBoard:
    scenario = TimetableScenario.objects.create(
        academic_year="1448", term="2", name="H15 blocked overlaps"
    )
    return DeliveryBoard.objects.create(
        scenario=scenario, label="Term 4", nominal_term=4, program="AI", display_order=1
    )


def _place(board, key, section, day, start="08:00", end="09:15"):
    code = key.split("::", 1)[0]
    ts, _ = TermSection.objects.get_or_create(
        scenario=board.scenario,
        course_key=key,
        section=section,
        defaults={"course_code": code, "course_number": code},
    )
    TermSectionMeeting.objects.create(term_section=ts, day=day, start_time=start, end_time=end)
    return SectionPlacement.objects.create(
        board=board, term_section=ts, day=day, start_time=start, end_time=end
    )


def _need(board, keys: tuple[str, ...], students: int) -> None:
    for key in keys:
        for sid in range(1, students + 1):
            ScenarioStudentCourseRequest.objects.create(
                scenario=board.scenario,
                student_id=sid,
                course_key=key,
                course_code=key.split("::", 1)[0],
                primary_term=4,
            )


def _need_both(board, students: int) -> None:
    _need(board, (NETS, ALGEBRA), students)


def _clash_actions(board) -> list[str]:
    actions = build_scenario_builder_actions(board.scenario_id)["actions"]
    return [a["severity"] for a in actions if a["kind"] == "student_time_clash"]


def test_overlap_students_route_round_through_a_sibling_is_a_warning() -> None:
    board = _board()
    _need_both(board, 25)
    _place(board, NETS, "S1", "SUN")
    _place(board, ALGEBRA, "S1", "SUN")
    # NEURAL NETWORKS S2 + LINEAR ALGEBRA S1 is a clash-free pair for all 25.
    _place(board, NETS, "S2", "MON")

    conflicts = detect_board_conflicts(board.id)

    assert [
        (o["shared_students"], o["blocked_students"], o["severity"]) for o in conflicts["overlaps"]
    ] == [(25, 0, "warning")]
    assert conflicts["summary"] == {"critical": 0, "warning": 1, "info": 0}
    readiness = check_publish_readiness(board.scenario_id)
    assert readiness["ready"] is True, readiness["blockers"]
    assert readiness["warnings"] == ["Board 'Term 4': 1 warnings"]
    assert _clash_actions(board) == ["warn"]
    [badge] = get_scenario_boards_summary(board.scenario_id)
    assert (badge["critical"], badge["warning"]) == (0, 1)
    impact = compute_affected_students(board.id)
    assert (impact["affected_count"], impact["blocked_count"]) == (25, 0)


@pytest.mark.parametrize("sibling", ["none", "clashes_with_another_meeting"])
def test_overlap_leaving_students_no_clash_free_pair_is_critical(sibling) -> None:
    board = _board()
    _need_both(board, 25)
    _place(board, NETS, "S1", "SUN")
    _place(board, ALGEBRA, "S1", "SUN")
    if sibling == "clashes_with_another_meeting":
        # LINEAR ALGEBRA S1 also meets on Monday, where the only sibling sits:
        # a section's whole week counts, not just the placement that clashed.
        _place(board, ALGEBRA, "S1", "MON")
        _place(board, NETS, "S2", "MON")

    conflicts = detect_board_conflicts(board.id)

    overlaps = conflicts["overlaps"]
    assert overlaps and all(
        (o["blocked_students"], o["severity"]) == (25, "critical") for o in overlaps
    )
    assert conflicts["summary"]["critical"] == len(overlaps)
    readiness = check_publish_readiness(board.scenario_id)
    assert readiness["ready"] is False
    assert f"Board 'Term 4': {len(overlaps)} critical conflicts" in readiness["blockers"]
    assert set(_clash_actions(board)) == {"block"}
    impact = compute_affected_students(board.id)
    assert impact["blocked_count"] == 25 * len(overlaps)
    assert {d["resolvable"] for d in impact["overlap_details"]} == {False}


def test_blocked_overlap_below_the_threshold_stays_a_warning() -> None:
    board = _board()
    _need_both(board, 5)
    _place(board, NETS, "S1", "SUN")
    _place(board, ALGEBRA, "S1", "SUN")

    [overlap] = detect_board_conflicts(board.id)["overlaps"]

    assert (overlap["blocked_students"], overlap["severity"]) == (5, "warning")


def _exported_overlap_lines(board) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    from core.services.timetable_export import export_scenario_xlsx

    exported = export_scenario_xlsx(board.scenario_id)
    try:
        sheet = load_workbook(exported)[f"Term {board.nominal_term}"]
        lines = [
            (str(cell.value), str(cell.font.color.rgb)[-6:])
            for row in sheet.iter_rows()
            for cell in row
            if str(cell.value or "").startswith(("OVERLAP", "Conflicts:"))
        ]
    finally:
        exported.unlink(missing_ok=True)
    # The heading first; the overlaps under it in whatever order they were found.
    return lines[:1] + sorted(lines[1:])


CRITICAL_LINE = ("OVERLAP: AI201-S1 vs MATH201-S1", "C03030")
WARNING_LINE = ("OVERLAP (non-blocking): AI201-S1 vs MATH201-S1", "D97706")
MIXED_WARNING_LINE = ("OVERLAP (non-blocking): CS301-S1 vs CS302-S1", "D97706")


@pytest.mark.parametrize(
    "case, expected",
    [
        ("warning_only", [("Conflicts:", "D97706"), WARNING_LINE]),
        ("critical_only", [("Conflicts:", "C03030"), CRITICAL_LINE]),
        # One critical conflict makes the heading red even beside a warning.
        ("critical_and_warning", [("Conflicts:", "C03030"), CRITICAL_LINE, MIXED_WARNING_LINE]),
    ],
)
def test_export_board_sheet_writes_each_overlap_at_its_own_severity(case, expected) -> None:
    board = _board()
    _need_both(board, 25)
    _place(board, NETS, "S1", "SUN")
    _place(board, ALGEBRA, "S1", "SUN")
    if case == "warning_only":
        _place(board, NETS, "S2", "MON")
    if case == "critical_and_warning":
        _need(board, ("CS301::ALPHA", "CS302::BETA"), 25)
        _place(board, "CS301::ALPHA", "S1", "TUE")
        _place(board, "CS302::BETA", "S1", "TUE")
        _place(board, "CS301::ALPHA", "S2", "WED")

    assert _exported_overlap_lines(board) == expected[:1] + sorted(expected[1:])


def test_same_course_overlap_stays_critical_and_blocks_publish() -> None:
    # Two sections of one course at one time double-book the course whatever
    # the demand. Students need only this course here: no pair needs both, so
    # nothing but the same-course rule can make the overlap critical.
    board = _board()
    _need(board, (NETS,), 25)
    _place(board, NETS, "S1", "SUN")
    _place(board, NETS, "S2", "SUN")

    conflicts = detect_board_conflicts(board.id)

    [overlap] = conflicts["overlaps"]
    assert (overlap["severity"], overlap["same_course"], overlap["blocked_students"]) == (
        "critical",
        True,
        None,
    )
    assert conflicts["summary"]["critical"] == 1
    readiness = check_publish_readiness(board.scenario_id)
    assert readiness["ready"] is False
    assert "Board 'Term 4': 1 critical conflicts" in readiness["blockers"]
    assert _clash_actions(board) == ["block"]
    # A student takes one section of a course, so the pair affects no one.
    impact = compute_affected_students(board.id)
    assert (impact["affected_count"], impact["blocked_count"]) == (0, 0)
    assert impact["overlap_details"] == []
