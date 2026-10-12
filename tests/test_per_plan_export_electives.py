"""The per-plan export shows this term's electives on the plan term of their slot.

A plan lists elective slots (AI1, AI2, AI3), never the courses. For 1448 T2 the
department mapped AI2 -> AI441 and AI3 -> AI462, Build gave them sections on the
Term 10 board, and the export still filed them under "not in this plan": no
term sheet showed them and the slots read "missing" (scenario 8 online,
2026-10-12). These tests pin where a mapped elective lands, what the slot's
coverage row says, and what stays untouched.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.models import (
    DeliveryBoard,
    ElectiveCourse,
    ElectiveTermMapping,
    ProgrammeRequirement,
    ScenarioSectionBudget,
    SectionPlacement,
    TermSection,
    TimetableScenario,
)
from core.services.timetable_per_plan_export import (
    export_scenario_per_plan,
    get_plan_elective_terms,
    get_plan_term_map,
)

pytestmark = pytest.mark.django_db


def _plan(program: str, rows: list[tuple[str, int]]) -> None:
    for code, term in rows:
        ProgrammeRequirement.objects.create(
            program=program, course_code=code, course_name=code, programme_term=term
        )


def _map(slot: str, code: str, *, programme: str = "AI", year: str = "1448", term: int = 2):
    elective, _ = ElectiveCourse.objects.get_or_create(
        programme=programme, course_code=code, defaults={"course_name": f"{code} name"}
    )
    ElectiveTermMapping.objects.create(
        academic_year=year, term=term, programme=programme, placeholder_code=slot, elective=elective
    )


def _scenario(program: str = "AI", board_term: int = 10) -> tuple[TimetableScenario, DeliveryBoard]:
    scenario = TimetableScenario.objects.create(
        academic_year="1448", term="2", name=f"{program} electives"
    )
    board = DeliveryBoard.objects.create(
        scenario=scenario, label=f"Term {board_term}", nominal_term=board_term, program=program
    )
    return scenario, board


def _place(board: DeliveryBoard, code: str, day: str = "SUN", start: str = "09:00") -> None:
    ScenarioSectionBudget.objects.create(
        scenario=board.scenario,
        course_code=code,
        course_key=code,
        department=code[:2],
        programme_term=board.nominal_term,
        credit_hours=3,
        planned_sections=1,
        max_per_section=40,
        total_demand=30,
    )
    section = TermSection.objects.create(
        scenario=board.scenario,
        course_code=code,
        course_number=code,
        course_key=code,
        course_name=f"{code} name",
        section="S1",
        source_tag="tw_auto",
    )
    end = {"09:00": "10:15", "10:30": "11:45", "13:00": "14:15"}[start]
    SectionPlacement.objects.create(
        board=board, term_section=section, day=day, start_time=start, end_time=end
    )


def _workbook(scenario: TimetableScenario):
    path, _name, is_zip = export_scenario_per_plan(scenario.id)
    assert not is_zip
    try:
        return load_workbook(path)
    finally:
        Path(path).unlink(missing_ok=True)


def _sheet_text(ws) -> str:
    return "\n".join(
        str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None
    )


def _coverage_rows(wb) -> dict[str, tuple]:
    """Plan table rows of the coverage sheet, by course code."""
    ws = wb["Plan Coverage"]
    rows = {}
    for values in ws.iter_rows(min_row=5, values_only=True):
        if not values[0] or not str(values[0]).startswith("T"):
            break
        rows[values[1]] = values
    return rows


def _other_courses(wb) -> list[str]:
    ws = wb["Plan Coverage"]
    seen_header, codes = False, []
    for values in ws.iter_rows(values_only=True):
        if values[0] and str(values[0]).startswith("Other courses placed"):
            seen_header = True
            continue
        if seen_header and values[1]:
            codes.append(values[1])
    return codes


# ── where a mapped elective lands ──────────────────────────────


def test_a_mapped_elective_is_on_its_slots_term_sheet() -> None:
    _plan("AI", [("AI401", 10), ("AI2", 10), ("AI3", 10)])
    _map("AI2", "AI441")
    _map("AI3", "AI462")
    scenario, board = _scenario()
    _place(board, "AI401", "SUN", "09:00")
    _place(board, "AI441", "MON", "10:30")
    _place(board, "AI462", "TUE", "13:00")

    wb = _workbook(scenario)

    term_10 = _sheet_text(wb["Term 10"])
    assert "AI441" in term_10
    assert "AI462" in term_10
    assert _other_courses(wb) == []


def test_the_slot_row_is_covered_by_the_course_that_fills_it() -> None:
    _plan("AI", [("AI2", 10), ("AI3", 10)])
    _map("AI2", "AI441")
    _map("AI3", "AI462")
    scenario, board = _scenario()
    _place(board, "AI441")

    rows = _coverage_rows(_workbook(scenario))

    assert rows["AI2"][4] == "matched"
    assert rows["AI2"][5].startswith("via AI441")
    assert rows["AI3"][4] == "missing"  # AI462 has no section in this scenario
    assert "AI441" not in rows  # one row per plan requirement, not per course


def test_without_a_mapping_the_course_stays_outside_the_plan() -> None:
    _plan("AI", [("AI401", 10), ("AI2", 10)])
    scenario, board = _scenario()
    _place(board, "AI401", "SUN", "09:00")
    _place(board, "AI441", "MON", "10:30")

    wb = _workbook(scenario)

    assert "AI441" not in _sheet_text(wb["Term 10"])
    assert _other_courses(wb) == ["AI441"]
    assert _coverage_rows(wb)["AI2"][4] == "missing"


def test_a_mapping_for_another_term_is_ignored() -> None:
    _plan("AI", [("AI401", 10), ("AI2", 10)])
    _map("AI2", "AI441", term=1)
    scenario, board = _scenario()
    _place(board, "AI401", "SUN", "09:00")
    _place(board, "AI441", "MON", "10:30")

    assert _other_courses(_workbook(scenario)) == ["AI441"]


def test_a_second_cohort_plan_uses_its_base_programmes_mapping_and_its_own_slot_term() -> None:
    _plan("AI2", [("AI201", 9), ("AI2", 9)])
    _map("AI2", "AI441", programme="AI")
    scenario, board = _scenario(program="AI2", board_term=9)
    _place(board, "AI201", "SUN", "09:00")
    _place(board, "AI441", "MON", "10:30")

    wb = _workbook(scenario)

    assert "AI441" in _sheet_text(wb["Term 9"])
    assert _coverage_rows(wb)["AI2"][4] == "matched"


# ── the plan-term map itself ───────────────────────────────────


def test_an_elective_filling_two_slots_takes_the_earlier_term() -> None:
    _plan("DS", [("DS1", 8), ("DS3", 10)])
    _map("DS1", "DS482", programme="DS")
    _map("DS3", "DS482", programme="DS")
    scenario, _board = _scenario(program="DS", board_term=8)
    plan_map = get_plan_term_map("DS")

    terms, fills = get_plan_elective_terms("DS", scenario, {"DS482"}, plan_map)

    assert terms == {"DS482": 8}
    assert fills == {"DS1": ["DS482"], "DS3": ["DS482"]}


def test_a_course_already_in_the_plan_keeps_its_own_term() -> None:
    _plan("AI", [("AI2", 10), ("AI441", 7)])
    _map("AI2", "AI441")
    scenario, _board = _scenario()
    plan_map = get_plan_term_map("AI")

    terms, fills = get_plan_elective_terms("AI", scenario, {"AI441"}, plan_map)

    assert terms == {}
    assert fills == {}


def test_only_courses_the_scenario_has_are_added() -> None:
    _plan("AI", [("AI2", 10)])
    _map("AI2", "AI441")
    _map("AI2", "AI499")
    scenario, _board = _scenario()

    terms, fills = get_plan_elective_terms("AI", scenario, {"AI441"}, get_plan_term_map("AI"))

    assert terms == {"AI441": 10}
    assert fills == {"AI2": ["AI441"]}
