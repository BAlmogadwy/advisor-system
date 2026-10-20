"""One teacher per section on screen and on paper (D1, D2, D3, I21).

The names written into the meetings, the drawer's ``teacher``, the export
grid and its Instructors sheet all come from the one map the engine charges,
so they agree with each other and with the engine before any reconcile. Each
request reads the stored rows once, whatever reads the map inside it.
"""

from __future__ import annotations

import json

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext

from core.models import SectionPlacement, TermSection, TermSectionMeeting
from core.services import section_teacher_assignment as sta
from tests.one_teacher_support import (
    FLAG_ON,
    budget,
    instructor,
    link,
    place,
    requirement,
    scenario,
    section,
)


@pytest.fixture
def flags(settings):
    for name, value in FLAG_ON.items():
        setattr(settings, name, value)
    return settings


def _admin_client():
    from tests.test_one_teacher_pipeline import _admin_client as client

    return client()


def _fixture():
    """AI201 is taught in AI by Dr A and in DS by Dr B, three sections at one
    time: two get a teacher each and one is to be decided (``clash``). AI299
    has no links and keeps the registrar's free text. The meeting names are
    stale, as before a reconcile."""
    shared = requirement("AI", "AI201", "SHARED COURSE")
    requirement("DS", "AI201", "SHARED COURSE")
    free = requirement("AI", "AI299", "NO LINKS COURSE")
    sc, board = scenario(programs=("AI", "DS"))
    link("AI", "AI201", instructor("Dr A"))
    link("DS", "AI201", instructor("Dr B"))
    budget(sc, shared, 3)
    budget(sc, free, 1)
    for label, room in (("S1", "R1"), ("S2", "R2"), ("S3", "R3")):
        place(board, section(sc, shared, label), [("SUN", "09:00", "10:15")], room=room)
    place(board, section(sc, free, "S1"), [("MON", "09:00", "10:15")], instructor_name="Free Text")
    sta.prepare_section_teachers(sc.id, phase="layout")
    TermSectionMeeting.objects.filter(term_section__course_key=shared).update(instructor="stale")
    return sc, board, shared, free


def _expected(sc) -> dict[int, str]:
    """What every surface must show per term section: the teacher's name, ""
    for a section to be decided; a no-links section is absent."""
    out = {}
    for ts_id, row in sta.section_teacher_display(sc.id).items():
        if row["status"] in ("engine", "pinned", "frozen"):
            out[ts_id] = row["name"]
        elif row["status"] == "tbd":
            out[ts_id] = ""
    return out


def _meeting_names(sc) -> dict[int, set[str]]:
    out: dict[int, set[str]] = {}
    for ts_id, name in TermSectionMeeting.objects.filter(term_section__scenario=sc).values_list(
        "term_section_id", "instructor"
    ):
        out.setdefault(ts_id, set()).add(name)
    return out


def _section_instructor_reads(ctx) -> int:
    return sum(
        1
        for q in ctx.captured_queries
        if "section_instructors" in q["sql"] and q["sql"].lstrip().upper().startswith("SELECT")
    )


# ── D1: names ─────────────────────────────────────────────────────────────


@pytest.mark.django_db
def test_d1_each_section_shows_its_own_teacher_and_a_tbd_section_none(flags) -> None:
    from core.services.course_instructor_assignment import (
        apply_primary_instructor,
        reconcile_scenario_instructors,
    )

    sc, board, shared, free = _fixture()
    expected = _expected(sc)
    shared_ids = list(
        TermSection.objects.filter(scenario=sc, course_key=shared).values_list("id", flat=True)
    )
    # Dr B, whom the first-programme rule never named, teaches one section.
    assert sorted(expected[i] for i in shared_ids) == ["", "Dr A", "Dr B"]

    reconcile_scenario_instructors(sc)

    names = _meeting_names(sc)
    assert {i: names[i] for i in shared_ids} == {i: {expected[i]} for i in shared_ids}
    [free_id] = TermSection.objects.filter(course_key=free).values_list("id", flat=True)
    assert names[free_id] == {"Free Text"}  # no links: the registrar's text stays

    TermSectionMeeting.objects.filter(term_section__course_key=shared).update(instructor="stale")
    for ts in TermSection.objects.filter(scenario=sc):
        apply_primary_instructor(ts, sc, board, ts.course_code)

    assert _meeting_names(sc) == names


# ── D2 / I21: the drawer's teacher, one read per request ─────────────────


@pytest.mark.django_db
def test_d2_board_detail_gives_each_placement_its_teacher_from_one_read(flags, monkeypatch) -> None:
    sc, board, shared, free = _fixture()
    expected = sta.section_teacher_display(sc.id)
    loads: list[int] = []
    real = sta.load_assignment

    def counted(scenario_id, **kwargs):
        loads.append(scenario_id)
        return real(scenario_id, **kwargs)

    monkeypatch.setattr(sta, "load_assignment", counted)
    client = _admin_client()

    with CaptureQueriesContext(connection) as ctx:
        response = client.get(f"/ops/tw/boards/{board.id}/")

    assert response.status_code == 200
    placements = response.json()["placements"]
    assert len(placements) == 4
    for p in placements:
        assert p["teacher"] == expected[p["term_section_id"]]
    statuses = sorted(p["teacher"]["status"] for p in placements)
    assert statuses == ["engine", "engine", "no_links", "tbd"]
    assert sorted(
        (p["teacher"]["name"], p["teacher"]["source"])
        for p in placements
        if p["teacher"]["status"] == "engine"
    ) == [("Dr A", "engine"), ("Dr B", "engine")]  # stored rows, not a read-time pick
    [tbd] = [p["teacher"] for p in placements if p["teacher"]["status"] == "tbd"]
    assert tbd == {"id": None, "name": None, "status": "tbd", "source": None, "reason": "clash"}
    assert loads == [sc.id]  # one map for the whole board, not one per placement
    assert _section_instructor_reads(ctx) == 1


@pytest.mark.django_db
def test_d2_a_teacher_no_row_stores_yet_is_shown_as_derived(flags) -> None:
    from core.models import SectionInstructor

    sc, _board, _shared, _free = _fixture()
    row = SectionInstructor.objects.filter(scenario=sc, source="engine").order_by("pk").first()
    ts = TermSection.objects.get(scenario=sc, course_key=row.course_key, section=row.section)
    row.delete()  # the read path's closure still gives the section a teacher

    shown = sta.section_teacher_display(sc.id)[ts.id]

    assert shown["status"] == "engine" and shown["name"] and shown["source"] == "derived"


@pytest.mark.django_db
def test_d2_without_the_rule_placements_have_no_teacher_key(flags) -> None:
    sc, board, _shared, _free = _fixture()
    flags.TIMETABLE_ONE_TEACHER_PER_SECTION = False

    with CaptureQueriesContext(connection) as ctx:
        response = _admin_client().get(f"/ops/tw/boards/{board.id}/")

    assert response.status_code == 200
    assert all("teacher" not in p for p in response.json()["placements"])
    assert _section_instructor_reads(ctx) == 0


@pytest.mark.django_db
def test_i21_a_drag_reads_the_stored_rows_once(flags) -> None:
    sc, board, shared, free = _fixture()
    moved = SectionPlacement.objects.get(term_section__course_key=free)
    client = _admin_client()

    with CaptureQueriesContext(connection) as ctx:
        response = client.post(
            f"/ops/tw/placements/{moved.id}/move/",
            data=json.dumps({"day": "TUE", "start_time": "09:00", "end_time": "10:15"}),
            content_type="application/json",
        )

    assert response.status_code == 200, response.content
    # Validation, the backstop and the display each read the map; the rows once.
    assert _section_instructor_reads(ctx) == 1
    assert response.json()["placement"]["teacher"]["status"] == "no_links"


@pytest.mark.django_db
def test_i21_a_write_inside_a_nested_scope_reaches_the_rest_of_the_request(flags) -> None:
    from dataclasses import replace

    sc, _board, _shared, _free = _fixture()
    with sta.request_scope():
        before = sta.load_assignment(sc.id)
        engine = sorted(k for k, s in before.status.items() if s == "engine")
        assert len(engine) == 2  # Dr A and Dr B on the two clashing sections
        swapped = {engine[0]: before.teacher[engine[1]], engine[1]: before.teacher[engine[0]]}
        with sta.assignment_scope(sc.id):  # e.g. a service that opens its own scope
            sta.persist_assignment(sc.id, replace(before, teacher={**before.teacher, **swapped}))

        after = sta.load_assignment(sc.id)

    assert {key: after.teacher[key] for key in engine} == swapped


@pytest.mark.django_db
def test_i21_a_nested_scope_shares_the_requests_read(flags) -> None:
    sc, _board, _shared, _free = _fixture()

    with CaptureQueriesContext(connection) as ctx, sta.request_scope():
        with sta.assignment_scope(sc.id):  # a service that opens its own scope first
            sta.load_assignment(sc.id)
        sta.load_assignment(sc.id)

    assert _section_instructor_reads(ctx) == 1


@pytest.mark.django_db
def test_i21_an_english_pattern_switch_reads_the_rows_once(flags) -> None:
    from core.services.timetable_section_policy import create_policy_section_bundle

    sc, board = scenario()
    link("AI", "ENG101", instructor("Dr E"))
    english = TermSection.objects.create(
        scenario=sc, course_code="ENG101", course_key="ENG101", section="M1"
    )
    rows = create_policy_section_bundle(
        board, english, day="SUN", start_time="09:00", end_time="12:00"
    )
    sta.prepare_section_teachers(sc.id, phase="layout")
    client = _admin_client()

    with CaptureQueriesContext(connection) as ctx:
        response = client.post(
            f"/ops/tw/placements/{rows[0].id}/switch-pattern/",
            data=json.dumps({"pattern_id": "ENG_INTENSIVE_V1_PM"}),
            content_type="application/json",
        )

    assert response.status_code == 200, response.content
    placements = response.json()["placements"]
    assert len(placements) == 4
    assert {p["teacher"]["name"] for p in placements} == {"Dr E"}
    # Four validations, the backstop and the display: one read of the rows.
    assert _section_instructor_reads(ctx) == 1


@pytest.mark.django_db
@pytest.mark.parametrize("planned", [False, True])
def test_i21_a_create_reads_the_rows_once_and_again_after_naming_the_section(
    flags, planned
) -> None:
    sc, board, _shared, free = _fixture()
    other = requirement("AI", "AI202", "OTHER COURSE")
    link("AI", "AI202", instructor("Dr C"))
    budget(sc, other, 1)
    client = _admin_client()
    meeting = {"day": "TUE", "start_time": "09:00", "end_time": "10:15", "room": "R9"}
    if planned:
        url = "/ops/tw/placements/create-planned/"
        payload = {
            "board_id": board.id,
            "course_code": "AI202",
            "course_key": other,
            "course_name": "OTHER COURSE",
            "section_label": "S1",
            "capacity": 30,
            "meetings": [meeting],
        }
    else:
        ts = section(sc, other, "S1")
        TermSectionMeeting.objects.create(term_section=ts, **meeting)
        url = "/ops/tw/placements/create/"
        payload = {"board_id": board.id, "term_section_id": ts.id, **meeting}

    with CaptureQueriesContext(connection) as ctx:
        response = client.post(url, data=json.dumps(payload), content_type="application/json")

    assert response.status_code == 201, response.content
    assert response.json()["placement"]["teacher"]["name"] == "Dr C"
    # Validation, naming, the backstop and the display share the rows: one read
    # before the section's row is written, one after.
    assert _section_instructor_reads(ctx) == 2


@pytest.mark.django_db
@pytest.mark.parametrize(
    "url", ["/ops/tw/scenarios/{id}/export.xlsx", "/ops/tw/scenarios/{id}/export-per-plan"]
)
def test_i21_an_export_reads_the_stored_rows_once(flags, url) -> None:
    sc, _board, _shared, _free = _fixture()
    client = _admin_client()

    with CaptureQueriesContext(connection) as ctx:
        response = client.get(url.format(id=sc.id))
        b"".join(response.streaming_content) if response.streaming else response.content

    assert response.status_code == 200
    assert _section_instructor_reads(ctx) == 1


# ── D3: the export agrees with the drawer before any reconcile ───────────


def _cells(path) -> dict[str, list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    return {
        ws.title: [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
        for ws in workbook.worksheets
    }


def _check_export(sheets: dict[str, list[str]], sc, shared: str) -> None:
    expected = _expected(sc)
    grid = [text for title, texts in sheets.items() if title != "Instructors" for text in texts]
    for ts in TermSection.objects.filter(scenario=sc, course_key=shared):
        lines = [  # grid cells: code and section, then the teacher, then the room
            text.split("\n")
            for text in grid
            if "\n" in text and text.split("\n")[0].startswith(f"{ts.course_code} {ts.section}")
        ]
        assert lines, f"{ts.section} is not on the grid"
        for cell in lines:
            assert cell[1] == (expected[ts.id] or "TBD"), cell
    assert not any("stale" in text for text in grid)
    teachers = sheets["Instructors"]
    assert not any("stale" in text for text in teachers)
    for ts in TermSection.objects.filter(scenario=sc, course_key=shared):
        if expected[ts.id]:
            assert any(expected[ts.id] in text for text in teachers)
    # The section to be decided is under nobody; the no-links one keeps its name.
    [tbd] = [ts for ts in TermSection.objects.filter(course_key=shared) if not expected[ts.id]]
    assert not any(f"{tbd.course_code} {tbd.section}" in text for text in teachers)
    assert any("Free Text" in text for text in teachers)


@pytest.mark.django_db
def test_d3_the_export_names_the_drawers_teachers_before_any_reconcile(flags) -> None:
    from core.services.timetable_export import export_scenario_xlsx

    sc, _board, shared, _free = _fixture()

    _check_export(_cells(export_scenario_xlsx(sc.id)), sc, shared)


@pytest.mark.django_db
def test_d3_the_per_plan_export_names_the_drawers_teachers_before_any_reconcile(flags) -> None:
    import zipfile

    from core.services.timetable_per_plan_export import export_scenario_per_plan

    sc, _board, shared, _free = _fixture()

    path, _name, is_zip = export_scenario_per_plan(sc.id)

    if not is_zip:
        _check_export(_cells(path), sc, shared)
        return
    with zipfile.ZipFile(path) as bundle:
        names = bundle.namelist()
        assert names
        for name in names:
            target = path.with_name(name)
            target.write_bytes(bundle.read(name))
            _check_export(_cells(target), sc, shared)
