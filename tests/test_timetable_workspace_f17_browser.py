"""F17 in Chromium (stage 3, B2): under the production flags, a scenario with
English sections that have no one complete weekly pattern opens; such a
section's teacher reads "—" with its reason (en and ar); both exports download
on the page with the notice first; a failed export is a toast in words and no
tab ever opens. The jsdom suites pin each branch; this runs the real server and
the real browser download once.
"""

from __future__ import annotations

import io
import os
from pathlib import Path

os.environ.setdefault("DJANGO_ALLOW_ASYNC_UNSAFE", "1")

from unittest import mock

import pytest
from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.test import Client, override_settings
from openpyxl import load_workbook

from core.models import SectionPlacement
from tests.tw_incomplete_english_support import IncompleteEnglishScenario
from tests.tw_scenario_seating_support import PRODUCTION

pytest.importorskip("playwright.sync_api")

from tests.timetable_workspace_browser_support import WorkspaceBrowserMixin  # noqa: E402

REASON = {
    "en": (
        "Teacher not shown: this English section does not have one complete weekly "
        "pattern on every board. Remove it and place it again."
    ),
    "ar": (
        "لم يُعرض المدرّس: ليس لهذه الشعبة الإنجليزية نمط أسبوعي كامل واحد على كل "
        "اللوحات. احذفها وضعها من جديد."
    ),
}
#: Term 2's cards: ENG102 S1 and S2 (four meetings each), ENG101 S1 (three),
#: AI201 S1 and S2.
TERM2_CARDS = 13


def _holder(held: list):
    """A route handler that keeps each request unanswered in ``held``."""

    def hold(route) -> None:
        held.append(route)

    return hold


@override_settings(**PRODUCTION)
class IncompleteEnglishBrowserTests(WorkspaceBrowserMixin, StaticLiveServerTestCase):
    def setUp(self) -> None:  # the mixin's scenario is not this one
        user = get_user_model().objects.create_superuser(
            username="workspace-f17", email="workspace-f17@example.test", password="unused"
        )
        client = Client()
        client.force_login(user)
        self.session = client.cookies[settings.SESSION_COOKIE_NAME].value
        self.fx = IncompleteEnglishScenario(name="F17 degrade")
        self.path = f"/timetable-workspace/?scenario={self.fx.sc.id}&board={self.fx.b1.id}"

    def _selection_teacher(self, page, placement_id: int) -> tuple[str, str]:
        page.click(f'.tw-card[data-placement-id="{placement_id}"]')
        page.wait_for_selector("#twRightBody .tw-sel-field")
        return page.evaluate(
            """() => {
              const field = [...document.querySelectorAll('#twRightBody .tw-sel-field')]
                .find(f => ['Teacher', 'المدرس'].includes(f.querySelector('dt').textContent));
              const dd = field.querySelector('dd');
              return [dd.firstChild.textContent, dd.querySelector('.tw-sel-hint')?.textContent || ''];
            }"""
        )

    def test_the_board_opens_and_an_unknown_teacher_reads_as_a_dash_and_why(self) -> None:
        broken = SectionPlacement.objects.filter(term_section=self.fx.s1, board=self.fx.b1)
        legacy = SectionPlacement.objects.filter(term_section=self.fx.legacy).first()
        known = SectionPlacement.objects.get(term_section=self.fx.a1)
        for language in ("en", "ar"):
            with self.subTest(language=language):
                page, sent, errors = self._load(language, "light", 1440, 790, cards=TERM2_CARDS)
                for placement in (broken.first(), legacy):
                    self.assertEqual(
                        self._selection_teacher(page, placement.id), ["—", REASON[language]]
                    )
                self.assertEqual(self._selection_teacher(page, known.id), ["Dr A", ""])
                self.assertNotIn(
                    "Registrar Name", page.inner_text("#twRightBody"), "never the meeting's name"
                )
                self.assertEqual(errors, [])
                self.assertEqual(sent, [], "reading writes nothing")

    def test_both_exports_download_on_the_page_and_a_failure_is_said_in_words(self) -> None:
        page, _sent, errors = self._load("en", "light", 1440, 790, cards=TERM2_CARDS)
        context = page.context
        notices = []

        def sheets_of(download) -> list[str]:
            # Playwright saves without an extension; openpyxl reads by name.
            return load_workbook(io.BytesIO(Path(download.path()).read_bytes())).sheetnames

        with page.expect_download() as per_plan:
            page.click("#twExport")
        self.assertEqual(per_plan.value.suggested_filename, "AI.xlsx")
        self.assertEqual(sheets_of(per_plan.value)[0], "Notice")
        page.wait_for_function("!document.getElementById('twExport').hasAttribute('aria-busy')")

        page.click("#twExportMenu")
        with page.expect_download() as combined:
            page.click('.tw-export-item[data-mode="combined"]')
        self.assertTrue(combined.value.suggested_filename.startswith("timetable_F17_degrade_"))
        self.assertEqual(sheets_of(combined.value)[0], "Notice")

        page.on("download", lambda d: notices.append(f"download {d.suggested_filename}"))
        with mock.patch(
            "core.services.timetable_export.export_scenario_xlsx", side_effect=KeyError("x")
        ):
            page.wait_for_function("!document.getElementById('twExport').hasAttribute('aria-busy')")
            page.click("#twExportMenu")
            page.click('.tw-export-item[data-mode="combined"]')
            toast = page.wait_for_selector(".ux-toast.error")
            text = toast.inner_text()
        self.assertIn("Export failed", text)
        self.assertIn("The export could not be built.", text)
        self.assertEqual(notices, [], "an error saves nothing")
        self.assertEqual(len(context.pages), 1, "no tab ever opens")
        self.assertEqual(errors, [])

    def test_the_busy_export_button_moves_nothing_in_any_window(self) -> None:
        """While an export is built the button reads "Exporting…" (longer than
        "Export", in both languages); the toolbar wraps, so a longer label could
        push the chrome down. At every window of the record's matrix, in the
        wide font, the busy label changes no height and scrolls nothing."""
        measure = """() => {
          const main = document.getElementById('main-content');
          const button = document.getElementById('twExport');
          return {
            panes: Math.round(document.querySelector('.tw-workspace').getBoundingClientRect().top),
            button: Math.round(button.getBoundingClientRect().height),
            sideways: document.documentElement.scrollWidth > innerWidth + 0.5
              || main.scrollWidth > main.clientWidth + 0.5,
            busy: button.getAttribute('aria-busy'),
            label: button.textContent.trim(),
          };
        }"""
        problems = {}
        for language in ("en", "ar"):
            for width, height in ((1366, 650), (1440, 790), (1536, 730), (1920, 1080)):
                page, _sent, errors = self._load(language, "light", width, height, TERM2_CARDS)
                held: list = []  # the export request, held until measured
                page.route("**/export-per-plan*", _holder(held))
                before = page.evaluate(measure)
                page.click("#twExport")
                page.wait_for_function(
                    "document.getElementById('twExport').getAttribute('aria-busy') === 'true'"
                )
                during = page.evaluate(measure)
                self.assertEqual(
                    during["label"], "جارٍ التصدير…" if language == "ar" else "Exporting…"
                )
                for route in held:
                    route.abort()
                page.wait_for_function(
                    "!document.getElementById('twExport').hasAttribute('aria-busy')"
                )
                key = f"{language} {width}x{height}"
                changed = {
                    k: (before[k], during[k])
                    for k in ("panes", "button", "sideways")
                    if before[k] != during[k]
                }
                if changed or during["sideways"] or errors:
                    problems[key] = {**changed, "errors": errors}
                page.context.close()
        self.assertEqual(problems, {})
