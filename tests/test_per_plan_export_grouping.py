"""Per-plan export groups follow who sits in each section (owner 2026-10-21).

A programme's term sheet packed sections into groups by time fit alone, so the
AI2 file of scenario 38 paired AI236 S1 (eleven DS2 students and one AI2) with
IS124 S2 (all AI2): a week no AI2 student has. The groups now read the plan
lens's seating, the evaluator with the "mainly for programme" marks:

* R1 a section is main for P when it is marked for P or at least 70% P;
* R2 no course disappears: failing R1 the section seating most P is main,
  labelled mixed, and a course no P student sits in keeps today's sections;
* R3 every main group holds one main section of every course;
* R4 other sections seating some P follow in their own block with the count,
  and a section seating no P student is not in P's file;
* R7 the lens, and the seater inside it, run once per export.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pytest
from django.test import override_settings
from openpyxl import load_workbook

from core.models import (
    DeliveryBoard,
    ProgrammeRequirement,
    ScenarioSectionBudget,
    SectionPlacement,
    TermSection,
    TermSectionMeeting,
    TimetableScenario,
)
from core.services import timetable_per_plan_export as per_plan
from core.services import timetable_plan_lens
from core.services.timetable_autoplace import DEFAULT_SLOTS
from core.services.timetable_per_plan_export import (
    ROLE_MAIN,
    ROLE_MIXED,
    ROLE_OTHER,
    SectionSeat,
    _time_fit_groups,
    classify_course_sections,
    export_scenario_per_plan,
)
from tests._section_programme_fixture import build_marked_scenario

END = {"09:00": "10:15", "10:30": "11:45", "13:00": "14:15"}
LABEL = re.compile(r"^[A-Z]+\d+ S\d+")
BLOCK_HEADS = ("Group ", "Other group ", "Other sections with")


# ── pure lens rows ─────────────────────────────────────────────────────────


def _row(*, mark: str = "", plans: tuple[str, ...] = (), **seated: int) -> dict:
    """One lens section row: who the evaluator seats there, its mark and the
    demand allocation (``filter_plans``)."""
    return {
        "actual_plans": dict(seated),
        "actual_total": sum(seated.values()),
        "preferred_programme": mark,
        "filter_plans": list(plans),
        "shared": False,
    }


def _lens(rows: dict[int, dict]) -> dict:
    return {"sections": {str(ts_id): row for ts_id, row in rows.items()}}


def _roles(seats: dict[int, SectionSeat]) -> dict[int, str]:
    return {ts_id: seat.role for ts_id, seat in seats.items()}


# ── R1 / R2: which sections are main ───────────────────────────────────────


def test_a_section_marked_for_the_programme_is_main_whatever_its_share() -> None:
    lens = _lens({1: _row(AI2=5, DS2=5), 2: _row(mark="AI2", AI2=4, DS2=8)})

    assert classify_course_sections("AI2", [1, 2], lens) == {
        1: SectionSeat(ROLE_OTHER, 5, 10),
        2: SectionSeat(ROLE_MAIN, 4, 12),  # main, not "mixed": the mark decides
    }


def test_seventy_percent_is_main_and_sixty_nine_is_not() -> None:
    lens = _lens({1: _row(AI2=7, DS2=3), 2: _row(AI2=69, DS2=31), 3: _row(AI2=14, DS2=6)})

    assert _roles(classify_course_sections("AI2", [1, 2, 3], lens)) == {
        1: ROLE_MAIN,  # 7 of 10
        2: ROLE_OTHER,  # 69 of 100, beside a main section
        3: ROLE_MAIN,  # 14 of 20
    }


def test_a_section_marked_for_another_programme_is_still_main_at_seventy_percent() -> None:
    lens = _lens({1: _row(mark="DS2", AI2=8, DS2=2), 2: _row(AI2=2, DS2=8)})

    assert classify_course_sections("AI2", [1, 2], lens) == {
        1: SectionSeat(ROLE_MAIN, 8, 10),
        2: SectionSeat(ROLE_OTHER, 2, 10),
    }
    # ... and main in DS2's file too, by its mark: one section, two files.
    assert _roles(classify_course_sections("DS2", [1, 2], lens)) == {1: ROLE_MAIN, 2: ROLE_MAIN}


def test_with_no_main_section_the_one_seating_most_students_is_mixed() -> None:
    lens = _lens({1: _row(AI2=18, DS2=11), 2: _row(AI2=10, DS2=5)})

    assert classify_course_sections("AI2", [1, 2], lens) == {
        1: SectionSeat(ROLE_MIXED, 18, 29),  # 62%, but the most AI2 students
        2: SectionSeat(ROLE_OTHER, 10, 15),  # 67%, a higher share of fewer
    }
    # Equal counts: the higher share, then grid order.
    tie = _lens({1: _row(AI2=6, DS2=6), 2: _row(AI2=6, DS2=3), 3: _row(AI2=6, DS2=3)})
    assert _roles(classify_course_sections("AI2", [1, 2, 3], tie)) == {
        1: ROLE_OTHER,
        2: ROLE_MIXED,
        3: ROLE_OTHER,
    }


def test_a_course_no_programme_student_sits_in_keeps_todays_sections() -> None:
    lens = _lens(
        {1: _row(plans=("AI2",)), 2: _row(plans=("DS2",), DS2=25), 3: _row(), 4: _row(CS2=3)}
    )

    # The lens allocation decides, as before seating counted: S2 is DS2's.
    # Id 5 has no lens row (another scenario's section) and passes, as before.
    assert classify_course_sections("AI2", [1, 2, 3, 4, 5], lens) == {
        1: SectionSeat(ROLE_MAIN, 0, 0),
        3: SectionSeat(ROLE_MAIN, 0, 0),
        4: SectionSeat(ROLE_MAIN, 0, 3),
        5: SectionSeat(ROLE_MAIN, 0, 0),
    }


def test_a_marked_section_counts_even_in_a_course_nobody_sits_in() -> None:
    lens = _lens({1: _row(plans=("AI2",)), 2: _row(mark="AI2", plans=("DS2",), DS2=25)})

    assert _roles(classify_course_sections("AI2", [1, 2], lens)) == {1: ROLE_MAIN, 2: ROLE_MAIN}


# ── R3: the time fit, pure ─────────────────────────────────────────────────


def test_least_overlap_still_decides_placement_and_repeats() -> None:
    sun, mon, wed = 0b0001, 0b0010, 0b1000
    options = {
        "A": [("A1", sun), ("A2", mon), ("A3", sun)],
        "B": [("B1", sun), ("B2", wed)],
    }

    assert _time_fit_groups(options, repeat=True) == [
        ["A1", "B2"],  # B1 clashes with A1 on Sunday
        ["A2", "B1"],  # B's only unused section
        ["A3", "B2"],  # a repeat still avoids the clash
    ]
    # Without repeats a course whose sections are used sits the group out.
    assert _time_fit_groups(options, repeat=False) == [["A1", "B2"], ["A2", "B1"], ["A3"]]


# ── workbook helpers ───────────────────────────────────────────────────────


def _scenario(*terms: int, programs: str = "AI2,DS2") -> tuple[TimetableScenario, list]:
    scenario = TimetableScenario.objects.create(
        academic_year="1448", term="2", name="Seated groups"
    )
    boards = [
        DeliveryBoard.objects.create(
            scenario=scenario,
            label=f"Term {term}",
            nominal_term=term,
            program=programs,
            display_order=term,
        )
        for term in terms or (2,)
    ]
    return scenario, boards


def _course(
    board: DeliveryBoard, code: str, *sections: tuple[str, str, str], programs=("AI2", "DS2")
) -> dict[str, TermSection]:
    """One plan course on ``board``'s term: ``sections`` as (label, day, start)."""
    for program in programs:
        ProgrammeRequirement.objects.create(
            program=program, course_code=code, course_name=code, programme_term=board.nominal_term
        )
    ScenarioSectionBudget.objects.create(
        scenario=board.scenario,
        course_code=code,
        course_key=code,
        course_name=code,
        department=code[:2],
        programme_term=board.nominal_term,
        credit_hours=3,
        planned_sections=len(sections),
        max_per_section=40,
    )
    out = {}
    for label, day, start in sections:
        section = TermSection.objects.create(
            scenario=board.scenario,
            course_code=code,
            course_number=code,
            course_key=code,
            course_name=code,
            section=label,
            source_tag="tw_auto",
        )
        SectionPlacement.objects.create(
            board=board, term_section=section, day=day, start_time=start, end_time=END[start]
        )
        out[label] = section
    return out


def _seat(monkeypatch, rows: dict[TermSection, dict]) -> None:
    """Export with the real lens, its seating replaced by ``rows``."""
    real = per_plan.build_scenario_plan_lens

    def lens(scenario_id: int) -> dict:
        built = real(scenario_id)
        for section, row in rows.items():
            built["sections"][str(section.id)].update(row)
        return built

    monkeypatch.setattr(per_plan, "build_scenario_plan_lens", lens)


def _export(scenario: TimetableScenario) -> dict:
    """``{programme: workbook}``."""
    path, name, is_zip = export_scenario_per_plan(scenario.id)
    try:
        if not is_zip:
            return {name.removesuffix(".xlsx"): load_workbook(path)}
        with zipfile.ZipFile(path) as bundle:
            return {
                member.removesuffix(".xlsx"): load_workbook(io.BytesIO(bundle.read(member)))
                for member in bundle.namelist()
            }
    finally:
        Path(path).unlink(missing_ok=True)


def _grid_cells(ws) -> list[tuple[int, list[str]]]:
    """``(row, lines)`` of every filled day×slot cell. ``splitlines`` keeps it
    CRLF-safe: without lxml a newline may come back as CRLF."""
    out = []
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=1 + len(DEFAULT_SLOTS)):
        for cell in row:
            if isinstance(cell.value, str) and cell.value:
                out.append((cell.row, cell.value.splitlines()))
    return out


def _blocks(ws) -> list[tuple[str, set[str]]]:
    """The sheet's grids in order: ``(heading, sections drawn under it)``; the
    heading is "" for grids before any heading."""
    heads = {
        row[0].row: row[0].value
        for row in ws.iter_rows(min_row=2, max_col=1)
        if isinstance(row[0].value, str) and row[0].value.startswith(BLOCK_HEADS)
    }
    blocks: list[tuple[str, set[str]]] = [("", set())]
    cells = _grid_cells(ws)
    for row_num in range(2, ws.max_row + 1):
        if row_num in heads:
            blocks.append((heads[row_num], set()))
        for cell_row, lines in cells:
            if cell_row == row_num:
                blocks[-1][1].update(m.group(0) for line in lines if (m := LABEL.match(line)))
    return [block for block in blocks if block[0] or block[1]]


def _cell_lines(ws, label: str) -> list[list[str]]:
    return [lines for _row, lines in _grid_cells(ws) if lines[0].startswith(label)]


def _coverage_status(wb) -> dict[str, str]:
    ws = wb["Plan Coverage"]
    rows = {}
    for values in ws.iter_rows(min_row=5, values_only=True):
        if not values[0] or not str(values[0]).startswith("T"):
            break
        rows[values[1]] = values[4]
    return rows


def _scenario_38_like(monkeypatch, courses=("CS112", "AI236", "IS124", "ENGL214")):
    """The design's worked example: AI236 S1 is DS2's (one AI2 in it), S2 is
    AI2's; IS124 splits cleanly; CS112 has two AI2 sections and a DS2 one;
    ENGL214 has no section 70% AI2. No two sections share a time."""
    scenario, [board] = _scenario()
    table = {
        "CS112": (
            (("S1", "SUN", "09:00"), _row(AI2=28, DS2=2)),
            (("S2", "SUN", "10:30"), _row(AI2=2, DS2=28)),
            (("S3", "MON", "09:00"), _row(AI2=25, CS2=3)),
        ),
        "AI236": (
            (("S1", "TUE", "09:00"), _row(mark="DS2", AI2=1, DS2=11)),
            (("S2", "TUE", "10:30"), _row(mark="AI2", AI2=12)),
        ),
        "IS124": (
            (("S1", "WED", "09:00"), _row(DS2=25)),
            (("S2", "WED", "10:30"), _row(AI2=24)),
        ),
        "ENGL214": (
            (("S1", "THU", "09:00"), _row(AI2=18, DS2=11)),
            (("S2", "THU", "10:30"), _row(AI2=4, DS2=26)),
        ),
    }
    sections, rows = {}, {}
    for code in courses:
        made = _course(board, code, *(where for where, _seating in table[code]))
        for (label, _day, _start), row in table[code]:
            sections[f"{code} {label}"] = made[label]
            rows[made[label]] = row
    _seat(monkeypatch, rows)
    return scenario, sections


# ── R3: cohort groups ──────────────────────────────────────────────────────


@pytest.mark.django_db
def test_every_main_group_holds_one_section_of_every_course(monkeypatch) -> None:
    scenario, [board] = _scenario()
    cs = _course(board, "CS112", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"))
    ai = _course(board, "AI236", ("S1", "TUE", "09:00"))
    is_ = _course(board, "IS124", ("S1", "WED", "09:00"))
    _seat(
        monkeypatch,
        {
            cs["S1"]: _row(AI2=20),
            cs["S2"]: _row(AI2=20),
            ai["S1"]: _row(AI2=30),
            is_["S1"]: _row(AI2=30),
        },
    )

    assert _blocks(_export(scenario)["AI2"]["Term 2"]) == [
        ("Group 1 — 3 courses", {"CS112 S1", "AI236 S1", "IS124 S1"}),
        ("Group 2 — 3 courses", {"CS112 S2", "AI236 S1", "IS124 S1"}),
    ]


@pytest.mark.django_db
def test_group_count_is_the_largest_main_section_count(monkeypatch) -> None:
    scenario, [board] = _scenario()
    cs = _course(
        board, "CS112", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"), ("S3", "TUE", "09:00")
    )
    ai = _course(board, "AI236", ("S1", "WED", "09:00"), ("S2", "THU", "09:00"))
    _seat(
        monkeypatch,
        {
            cs["S1"]: _row(AI2=28, DS2=2),
            cs["S2"]: _row(AI2=2, DS2=28),
            cs["S3"]: _row(AI2=1, DS2=29),
            ai["S1"]: _row(AI2=25),
            ai["S2"]: _row(AI2=24),
        },
    )

    # Two main groups (AI236's two), not three (CS112's sections in all).
    assert _blocks(_export(scenario)["AI2"]["Term 2"]) == [
        ("Group 1 — 2 courses", {"CS112 S1", "AI236 S1"}),
        ("Group 2 — 2 courses", {"CS112 S1", "AI236 S2"}),
        ("Other sections with AI2 students", set()),
        ("Other group 1 — 1 courses", {"CS112 S2"}),
        ("Other group 2 — 1 courses", {"CS112 S3"}),
    ]


@pytest.mark.django_db
def test_the_ai2_file_never_pairs_a_ds2_section_with_ai2_sections(monkeypatch) -> None:
    scenario, _sections = _scenario_38_like(monkeypatch)

    blocks = _blocks(_export(scenario)["AI2"]["Term 2"])

    assert blocks == [
        ("Group 1 — 4 courses", {"CS112 S1", "AI236 S2", "IS124 S2", "ENGL214 S1"}),
        ("Group 2 — 4 courses", {"CS112 S3", "AI236 S2", "IS124 S2", "ENGL214 S1"}),
        ("Other sections with AI2 students", {"AI236 S1", "CS112 S2", "ENGL214 S2"}),
    ]


# ── R4: other sections ─────────────────────────────────────────────────────


@pytest.mark.django_db
def test_other_sections_follow_the_main_groups_with_their_count(monkeypatch) -> None:
    scenario, _sections = _scenario_38_like(monkeypatch, courses=("AI236", "IS124"))

    ws = _export(scenario)["AI2"]["Term 2"]

    # One main group, still headed: the other block follows it.
    assert _blocks(ws) == [
        ("Group 1 — 2 courses", {"AI236 S2", "IS124 S2"}),
        ("Other sections with AI2 students", {"AI236 S1"}),
    ]
    [cell] = _cell_lines(ws, "AI236 S1")
    assert cell[-1] == "AI2: 1 of 12"
    [main_cell] = _cell_lines(ws, "AI236 S2")
    assert not any("AI2:" in line for line in main_cell)


@pytest.mark.django_db
def test_a_section_with_no_programme_students_is_not_in_the_file(monkeypatch) -> None:
    scenario, _sections = _scenario_38_like(monkeypatch)

    workbooks = _export(scenario)

    assert not _cell_lines(workbooks["AI2"]["Term 2"], "IS124 S1")
    assert _cell_lines(workbooks["DS2"]["Term 2"], "IS124 S1")


@pytest.mark.django_db
def test_a_section_the_lens_gave_another_programme_still_shows_where_we_sit(
    monkeypatch,
) -> None:
    scenario, [board] = _scenario()
    cs = _course(board, "CS112", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"))
    engl = _course(board, "ENGL214", ("S1", "TUE", "09:00"))
    _seat(
        monkeypatch,
        {
            cs["S1"]: _row(plans=("AI2",), AI2=28, DS2=2),
            cs["S2"]: _row(plans=("DS2",), AI2=2, DS2=28),
            engl["S1"]: _row(plans=("DS2",), AI2=3, DS2=27),
        },
    )

    wb = _export(scenario)["AI2"]

    assert _blocks(wb["Term 2"]) == [
        ("Group 1 — 2 courses", {"CS112 S1", "ENGL214 S1"}),
        ("Other sections with AI2 students", {"CS112 S2"}),
    ]
    # Three AI2 students sit in ENGL214 through DS2's section: not "missing".
    assert _coverage_status(wb)["ENGL214"] == "matched"


@pytest.mark.django_db
def test_a_mixed_cell_keeps_course_teacher_room_then_the_label(monkeypatch) -> None:
    scenario, sections = _scenario_38_like(monkeypatch)
    engl = sections["ENGL214 S1"]
    SectionPlacement.objects.filter(term_section=engl).update(room="B12")
    TermSectionMeeting.objects.create(
        term_section=engl, day="THU", start_time="09:00", end_time="10:15", instructor="Dr Huda"
    )

    cells = _cell_lines(_export(scenario)["AI2"]["Term 2"], "ENGL214 S1")

    assert cells == [["ENGL214 S1", "Dr Huda", "B12", "mixed · AI2: 18 of 29"]] * 2  # 2 groups


@pytest.mark.django_db
def test_each_programme_file_groups_by_its_own_seating(monkeypatch) -> None:
    scenario, [board] = _scenario()
    ai = _course(board, "AI236", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"))
    _seat(
        monkeypatch,
        {ai["S1"]: _row(AI2=1, DS2=11), ai["S2"]: _row(AI2=11, DS2=1)},
    )

    workbooks = _export(scenario)

    assert _blocks(workbooks["AI2"]["Term 2"]) == [
        ("Group 1 — 1 courses", {"AI236 S2"}),
        ("Other sections with AI2 students", {"AI236 S1"}),
    ]
    assert _blocks(workbooks["DS2"]["Term 2"]) == [
        ("Group 1 — 1 courses", {"AI236 S1"}),
        ("Other sections with DS2 students", {"AI236 S2"}),
    ]
    [ds2_cell] = _cell_lines(workbooks["DS2"]["Term 2"], "AI236 S2")
    assert ds2_cell[-1] == "DS2: 1 of 12"


# ── a course's place is its main sections' ─────────────────────────────────


def _sidebar_status(ws) -> dict[str, str]:
    """The term sheet's course list: ``{course: Status}``."""
    [head] = [cell for cell in ws[1] if cell.value == "Course"]
    out = {}
    for values in ws.iter_rows(min_row=2, min_col=head.column, values_only=True):
        if not values[0]:
            break
        out[values[0]] = values[6]
    return out


def _cs112_on_two_boards(monkeypatch, *, s2_room: str = ""):
    """CS112 is AI2's at T2 and DS2's at T4: S1 on the T2 board (Monday 09:00)
    seats 28 AI2, S2 on the T4 board (Monday 10:30) 28 DS2 and 2 AI2."""
    scenario, boards = _scenario(2, 4)
    for program, term in (("AI2", 2), ("DS2", 4)):
        ProgrammeRequirement.objects.create(
            program=program, course_code="CS112", course_name="CS112", programme_term=term
        )
    ScenarioSectionBudget.objects.create(
        scenario=scenario,
        course_code="CS112",
        course_key="CS112",
        course_name="CS112",
        department="CS",
        programme_term=2,
        credit_hours=3,
        planned_sections=2,
        max_per_section=40,
    )
    sections = {}
    for label, board, start, room in (
        ("S1", boards[0], "09:00", ""),
        ("S2", boards[1], "10:30", s2_room),
    ):
        sections[label] = TermSection.objects.create(
            scenario=scenario,
            course_code="CS112",
            course_number="CS112",
            course_key="CS112",
            course_name="CS112",
            section=label,
            source_tag="tw_auto",
        )
        SectionPlacement.objects.create(
            board=board,
            term_section=sections[label],
            day="MON",
            start_time=start,
            end_time=END[start],
            room=room,
        )
    _seat(
        monkeypatch,
        {
            sections["S1"]: _row(plans=("AI2",), AI2=28),
            sections["S2"]: _row(plans=("DS2",), AI2=2, DS2=28),
        },
    )
    return scenario


def _misplaced_fill(cell) -> bool:
    return str(cell.fill.start_color.rgb).endswith("FFF3CD")


@pytest.mark.django_db
def test_an_other_section_on_another_board_leaves_the_course_in_place(monkeypatch) -> None:
    scenario = _cs112_on_two_boards(monkeypatch)

    workbooks = _export(scenario)
    wb = workbooks["AI2"]
    ws = wb["Term 2"]

    # AI2's own section is on its plan's board: two AI2 students in DS2's
    # section on the T4 board do not make the course misplaced.
    assert _coverage_status(wb)["CS112"] == "matched"
    assert _sidebar_status(ws)["CS112"] == "OK"
    [(main_row, main_lines)] = [c for c in _grid_cells(ws) if c[1][0].startswith("CS112 S1")]
    assert main_lines == ["CS112 S1"]
    main_cell = next(c for c in ws[main_row] if c.value == "CS112 S1")
    assert not _misplaced_fill(main_cell)
    # The other section still says where it is: on its own board, not the course's.
    [(other_row, other_lines)] = [c for c in _grid_cells(ws) if c[1][0].startswith("CS112 S2")]
    assert other_lines == ["CS112 S2 (on T4)", "AI2: 2 of 30"]
    assert _misplaced_fill(next(c for c in ws[other_row] if str(c.value).startswith("CS112 S2")))
    # DS2's file, the mirror: its section matches, AI2's (no DS2 in it) is not there.
    assert _coverage_status(workbooks["DS2"])["CS112"] == "matched"
    assert not _cell_lines(workbooks["DS2"]["Term 4"], "CS112 S1")


@pytest.mark.django_db
def test_a_course_searched_down_to_an_other_section_is_still_placed(monkeypatch) -> None:
    # Only S2's room matches. S1 is read first, before the course is known to
    # match, so the search drops it and keeps S2 alone.
    scenario = _cs112_on_two_boards(monkeypatch, s2_room="B12")

    path, name, is_zip = export_scenario_per_plan(scenario.id, search="AI2 B12")
    try:
        assert not is_zip and name.endswith(".xlsx")
        wb = load_workbook(path)
    finally:
        Path(path).unlink(missing_ok=True)

    # The search keeps only DS2's section: the course is placed, on T4 by it.
    assert _blocks(wb["Term 2"]) == [("Other sections with AI2 students", {"CS112 S2"})]
    assert _coverage_status(wb)["CS112"] == "misplaced"
    assert _sidebar_status(wb["Term 2"])["CS112"] == "On T4"


# ── R7: one seating per export ─────────────────────────────────────────────


@pytest.mark.django_db
def test_the_plan_lens_and_the_seater_run_once_per_export(monkeypatch) -> None:
    scenario, (term_2, term_3) = _scenario(2, 3)
    _course(term_2, "AI236", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"))
    _course(term_3, "CS112", ("S1", "TUE", "09:00"), ("S2", "WED", "09:00"))
    calls = {"lens": 0, "seater": 0}
    real_lens = per_plan.build_scenario_plan_lens
    real_seater = timetable_plan_lens._section_assignment_evidence

    def lens(*args, **kwargs):
        calls["lens"] += 1
        return real_lens(*args, **kwargs)

    def seater(*args, **kwargs):
        calls["seater"] += 1
        return real_seater(*args, **kwargs)

    monkeypatch.setattr(per_plan, "build_scenario_plan_lens", lens)
    monkeypatch.setattr(timetable_plan_lens, "_section_assignment_evidence", seater)

    workbooks = _export(scenario)

    assert sorted(workbooks) == ["AI2", "DS2"]
    assert all(wb.sheetnames[1:3] == ["Term 2", "Term 3"] for wb in workbooks.values())
    assert calls == {"lens": 1, "seater": 1}


# ── the real seating, with the flag on and off ─────────────────────────────


def _two_programme_marked_scenario(mark: str):
    """The PR #158 board (six AI2 and six DS students, AI236 S2 marked) with
    both programmes' plans, so the export writes an AI2 and a DS file."""
    scenario, board, placements = build_marked_scenario(mark=mark)
    board.program = "AI2,DS"
    board.save(update_fields=["program"])
    for program in ("AI2", "DS"):
        for code in ("AI236", "IS124"):
            ProgrammeRequirement.objects.create(
                program=program, course_code=code, course_name=code, programme_term=2
            )
    return scenario, placements


@pytest.mark.django_db
@override_settings(TIMETABLE_SECTION_PROGRAMME_PREFERENCE=True)
def test_with_the_flag_on_marks_and_seating_come_from_the_evaluator() -> None:
    scenario, _placements = _two_programme_marked_scenario(mark="AI2")

    workbooks = _export(scenario)

    # The grouped seating: AI236 S2 is AI2's six, S1 is DS's six; IS124 seats
    # all twelve, half each, so it is mixed in both files.
    ai2 = workbooks["AI2"]["Term 2"]
    assert _blocks(ai2) == [("", {"AI236 S2", "IS124 S1"})]
    assert _cell_lines(ai2, "IS124 S1") == [["IS124 S1", "R1", "mixed · AI2: 6 of 12"]]
    ds = workbooks["DS"]["Term 2"]
    assert _blocks(ds) == [("", {"AI236 S1", "IS124 S1"})]
    assert _cell_lines(ds, "IS124 S1") == [["IS124 S1", "R1", "mixed · DS: 6 of 12"]]

    # The lens carries the marks too: marked for AI2, IS124's one section is
    # main in AI2's file at half its seats, and still mixed in DS's.
    TermSection.objects.filter(scenario=scenario, course_key="IS124").update(
        preferred_programme="AI2"
    )
    workbooks = _export(scenario)
    assert _cell_lines(workbooks["AI2"]["Term 2"], "IS124 S1") == [["IS124 S1", "R1"]]
    assert _cell_lines(workbooks["DS"]["Term 2"], "IS124 S1") == [
        ["IS124 S1", "R1", "mixed · DS: 6 of 12"]
    ]


@pytest.mark.django_db
@override_settings(TIMETABLE_SECTION_PROGRAMME_PREFERENCE=False)
def test_with_the_flag_off_a_stored_mark_changes_nothing_in_the_export() -> None:
    scenario, _placements = _two_programme_marked_scenario(mark="AI2")

    def term_2_values() -> list:
        ws = _export(scenario)["AI2"]["Term 2"]
        return [cell.value for row in ws.iter_rows() for cell in row]

    marked = term_2_values()
    TermSection.objects.filter(scenario=scenario).update(preferred_programme="")

    assert term_2_values() == marked
