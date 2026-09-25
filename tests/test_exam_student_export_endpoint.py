"""Student export endpoints: access, gates, the fail-closed audit and the download."""

from __future__ import annotations

import json
import re
from io import BytesIO

import pytest
from django.contrib.auth.models import Group
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from core import exam_student_export_views as views
from core.middleware import ExamCommitteeAccessMiddleware
from core.models import AuditLog, ExamTimetableRun, StudentTermSection
from core.services import exam_student_export as export
from core.services.audit import log_audit_event, validate_hash_chain
from core.services.rbac import ROLE_ADVISOR, ROLE_EXAM_COMMITTEE, ROLE_STUDENT, ensure_role_groups
from tests.exam_student_export_fixture import (
    ALL_IDS,
    build_population,
    build_saved_run,
    save_payload,
    saved_payload,
    student_name,
)

pytestmark = pytest.mark.django_db

VIEWS = ("exam_student_export_preflight", "exam_student_export")


@pytest.fixture(autouse=True)
def _fast_passwords(settings):
    settings.PASSWORD_HASHERS = ["django.contrib.auth.hashers.MD5PasswordHasher"]


@pytest.fixture
def run():
    build_population()
    return build_saved_run()


def _user(model, username, role=None, **kwargs):
    ensure_role_groups()
    user = model.objects.create_user(username=username, password="x-Pass-1234", **kwargs)
    if role:
        user.groups.add(Group.objects.get(name=role))
    return user


@pytest.fixture
def committee(client, django_user_model):
    client.force_login(_user(django_user_model, "exam.committee1", ROLE_EXAM_COMMITTEE))
    return client


def _url(run_id, view="exam_student_export"):
    return reverse(view, args=[run_id])


def _post(client, run_id, payload, view="exam_student_export"):
    return client.post(_url(run_id, view), json.dumps(payload), content_type="application/json")


def _body(response):
    return b"".join(response.streaming_content) if response.streaming else response.content


PAYLOAD = {"scope": {"kind": "all"}, "language": "en", "one_file_per_group": False}


def _tokens(text: str) -> set[str]:
    return set(re.findall(r"[A-Za-z0-9.\-]+", text))


# ── Access ─────────────────────────────────────────────────────


def test_both_views_are_on_the_committee_allow_list():
    assert set(VIEWS) <= ExamCommitteeAccessMiddleware.allowed_views


def test_committee_and_super_admin_can_preflight_and_download(run, committee, django_user_model):
    assert _post(committee, run.pk, {}, VIEWS[0]).status_code == 200
    response = _post(committee, run.pk, PAYLOAD)
    assert response.status_code == 200, response.content
    admin = Client()
    admin.force_login(django_user_model.objects.create_superuser(username="root-admin"))
    assert _post(admin, run.pk, PAYLOAD).status_code == 200


@pytest.mark.parametrize("role", [ROLE_ADVISOR, ROLE_STUDENT])
def test_other_roles_are_refused_before_any_roster_is_built(
    run, client, django_user_model, role, monkeypatch
):
    monkeypatch.setattr(views, "build_roster_model", lambda *a, **k: pytest.fail("roster built"))
    client.force_login(_user(django_user_model, f"user-{role}", role))
    for view in VIEWS:
        assert _post(client, run.pk, PAYLOAD, view).status_code == 403


def test_anonymous_requests_are_sent_to_login(run, client):
    for view in VIEWS:
        response = _post(client, run.pk, PAYLOAD, view)
        assert response.status_code == 302 and "login" in response.url


def test_posts_need_csrf_and_nothing_answers_get(run, django_user_model):
    strict = Client(enforce_csrf_checks=True)
    strict.force_login(_user(django_user_model, "strict-committee", ROLE_EXAM_COMMITTEE))
    for view in VIEWS:
        assert _post(strict, run.pk, PAYLOAD, view).status_code == 403
        assert strict.get(_url(run.pk, view)).status_code == 405


# ── The page that offers it ────────────────────────────────────


def test_the_dialog_is_arabic_only_when_arabic_is_asked_for(committee):
    """The page renders English unless the request asks for Arabic.

    An Arabic assertion against an unforced page would assert nothing, so the
    same page is asked for both ways and each must hold its own words only.
    """
    english = committee.get(reverse("exam_timetable_page")).content.decode()
    arabic = committee.get(
        reverse("exam_timetable_page"), HTTP_ACCEPT_LANGUAGE="ar"
    ).content.decode()
    words = {
        "title": ("Export student data to Excel", "تصدير بيانات الطلاب إلى Excel"),
        "button": (">Student data</button>", ">بيانات الطلاب</button>"),
        "link": ("Export student data…", "تصدير بيانات الطلاب…"),
        "privacy": (
            "Contains student names and IDs. Share only with exam staff.",
            "يحتوي على أسماء الطلاب وأرقامهم الجامعية. شاركه مع منسوبي الاختبارات فقط.",
        ),
        "matches": ("Student lists match this timetable", "قوائم الطلاب مطابقة لهذا الجدول"),
    }
    for name, (en, ar) in words.items():
        assert en in english and ar not in english, name
        assert ar in arabic and en not in arabic, name
    assert '<html lang="ar" dir="rtl"' in arabic
    assert '<html lang="en" dir="ltr"' in english
    for page in (english, arabic):
        assert "js/exam-student-export.js?v=" in page
        assert 'dir="auto"' not in page


# ── Gates and validation ───────────────────────────────────────


def test_missing_run_is_404(committee):
    for view in VIEWS:
        response = _post(committee, 987654, PAYLOAD, view)
        assert response.status_code == 404


def test_rebuild_required_and_term_mismatch_are_409_with_nothing_audited(run, committee):
    data = saved_payload(run)
    for rows in data["section_enrollment"].values():
        for row in rows:
            row["term"] = "2"
    save_payload(run, data)
    for view in VIEWS:
        response = _post(committee, run.pk, PAYLOAD, view)
        assert response.status_code == 409 and response.json()["code"] == "lists_term_mismatch"
        # The terms themselves, so the dialog can say them in Arabic too.
        assert response.json()["live_term"] == ["1448", "1"]
        assert response.json()["saved_term"] == ["1448", "2"]
    data["section_enrollment"] = {}
    save_payload(run, data)
    response = _post(committee, run.pk, PAYLOAD)
    assert response.status_code == 409
    assert response.json() == {
        "ok": False,
        "code": "rebuild_required",
        "error": "Rebuild and save this timetable to export student data.",
    }
    assert not AuditLog.objects.exists()


def test_no_imported_lists_is_409_lists_unavailable(run, committee):
    StudentTermSection.objects.all().delete()
    response = _post(committee, run.pk, PAYLOAD)
    assert response.status_code == 409 and response.json()["code"] == "lists_unavailable"


@pytest.mark.parametrize(
    "body, code, field",
    [
        ({"scope": {"kind": "moon"}}, "invalid_options", "scope.kind"),
        ({**PAYLOAD, "student_ids": [4401001]}, "invalid_options", "student_ids"),
        (
            {**PAYLOAD, "scope": {"kind": "course", "exam": "CS101"}, "programs": ["AI"]},
            "empty_scope",
            "scope",
        ),
    ],
)
def test_bad_options_are_400_named_against_their_field(run, committee, body, code, field):
    response = _post(committee, run.pk, body)
    assert response.status_code == 400
    assert (response.json()["code"], response.json()["field"]) == (code, field)
    assert not AuditLog.objects.exists()


def test_oversized_or_malformed_bodies_are_400(run, committee):
    # Otherwise valid: the download ignores pickers, so only the size cap can refuse it.
    huge = {**PAYLOAD, "pickers": {"padding": "x" * 40000}}
    response = _post(committee, run.pk, huge)
    assert response.status_code == 400 and response.json()["field"] == "body"
    small = {**PAYLOAD, "pickers": {"padding": "x"}}
    assert _post(committee, run.pk, small).status_code == 200
    response = committee.post(_url(run.pk), "{not json", content_type="application/json")
    assert response.status_code == 400 and response.json()["field"] == "body"


def test_a_busy_export_slot_answers_503(run, committee, monkeypatch):
    monkeypatch.setattr(views, "EXPORT_SLOT_WAIT_SECONDS", 0)
    assert views._EXPORT_SLOT.acquire()
    try:
        response = _post(committee, run.pk, PAYLOAD)
    finally:
        views._EXPORT_SLOT.release()
    assert response.status_code == 503 and response.json()["code"] == "export_slot_busy"


# ── The audited download ───────────────────────────────────────


def test_download_is_audited_first_and_carries_its_reference_everywhere(run, committee):
    response = _post(committee, run.pk, PAYLOAD)
    assert response.status_code == 200
    content = _body(response)
    entry = AuditLog.objects.get(action="exam_timetable.export_students")
    reference = "EXR-" + entry.entry_hash[:8].upper()
    assert response["X-Export-Reference"] == reference
    assert response["Cache-Control"] == "private, no-store"
    assert response["X-Content-Type-Options"] == "nosniff"
    disposition = response["Content-Disposition"]
    assert disposition.isascii() and f"_r{run.pk}_en_{reference[4:]}.xlsx" in disposition
    book = load_workbook(BytesIO(content))
    about = {row[0]: row[1] for row in book["About"].iter_rows(values_only=True) if row[0]}
    assert about["Reference"].startswith(reference)
    assert book.properties.subject == reference
    info = {row[0]: row[1] for row in book["File info"].iter_rows(values_only=True) if row[0]}
    assert info["audit_entry_hash16"] == entry.entry_hash[:16]
    details = json.loads(entry.details_json)
    assert details["run_id"] == run.pk and details["files"][0]["rows"] > 0
    assert details["files"][0]["data_sha256"] == info["data_sha256"]
    assert entry.actor_username == "exam.committee1" and entry.actor_role == ROLE_EXAM_COMMITTEE
    assert not {str(sid) for sid in ALL_IDS} & _tokens(entry.details_json)
    assert not any(student_name(sid) in entry.details_json for sid in ALL_IDS)
    assert validate_hash_chain()["ok"]


def test_audit_failure_means_503_and_no_file(run, committee, monkeypatch):
    def refuse(*args, **kwargs):
        raise RuntimeError("audit table unavailable")

    monkeypatch.setattr(AuditLog.objects, "create", refuse)
    rendered = []
    monkeypatch.setattr(views, "render_export", lambda *a, **k: rendered.append(1))
    response = _post(committee, run.pk, PAYLOAD)
    assert response.status_code == 503
    assert response["Content-Type"] == "application/json"
    assert response.json()["code"] == "audit_unavailable"
    assert rendered == [] and "Content-Disposition" not in response


def test_existing_audit_callers_stay_fail_open(rf, monkeypatch, django_user_model):
    monkeypatch.setattr(
        AuditLog.objects, "create", lambda **kw: (_ for _ in ()).throw(RuntimeError())
    )
    request = rf.post("/anything/")
    request.user = django_user_model.objects.create_user(username="plain")
    assert log_audit_event(request, action="db.something", status="success") is None


def test_a_render_failure_after_the_audit_is_recorded_and_sends_nothing(
    run, committee, monkeypatch
):
    def explode(*args, **kwargs):
        raise RuntimeError("disk full")

    monkeypatch.setattr(views, "render_export", explode)
    response = _post(committee, run.pk, PAYLOAD)
    assert response.status_code == 500 and response.json()["code"] == "export_failed"
    audited = AuditLog.objects.get(action="exam_timetable.export_students")
    failed = AuditLog.objects.get(action="exam_timetable.export_students_failed")
    reference = "EXR-" + audited.entry_hash[:8].upper()
    assert json.loads(failed.details_json) == {"run_id": run.pk, "reference": reference}
    assert response.json()["reference"] == reference


def test_one_file_per_group_downloads_a_zip(run, committee):
    response = _post(
        committee, run.pk, {**PAYLOAD, "groups": ["M", "F"], "one_file_per_group": True}
    )
    assert response.status_code == 200
    assert response["Content-Type"] == export.ZIP_TYPE
    assert ".zip" in response["Content-Disposition"]


# ── Preflight: counts only, not audited ────────────────────────


def test_preflight_counts_without_ids_or_names_and_is_not_audited(run, committee):
    response = _post(
        committee,
        run.pk,
        {
            "scope": {"kind": "day", "day": "Sun"},
            "groups": ["M", "F"],
            "language": "en",
            "pickers": {"exam": "CS101", "slot_index": 0, "room_code": "M-B", "day": "Sun"},
        },
        VIEWS[0],
    )
    assert response.status_code == 200
    body = response.json()
    assert body["mode"] == "sync" and body["check"]["status"] == "matches"
    assert body["check"]["sections_matching"] == body["check"]["sections_total"]
    counts = body["counts"]
    assert counts["scope_rows"]["course"] == 16 and counts["scope_rows"]["room"] == 20
    assert (
        counts["groups"]["U"] == 1
        and counts["rows"] == counts["groups"]["M"] + counts["groups"]["F"]
    )
    assert [f["gender"] for f in body["files"]] == ["M", "F"]
    assert all(f["name"].endswith("_<REF>.xlsx") for f in body["files"])
    assert body["download_name"].endswith("_<REF>.zip")
    text = json.dumps(body)
    assert not {str(sid) for sid in ALL_IDS} & _tokens(text)
    assert not any(student_name(sid) in text for sid in ALL_IDS)
    assert response["Cache-Control"] == "private, no-store"
    assert not AuditLog.objects.exists()


def test_preflight_without_a_scope_returns_the_check_and_pickers(run, committee):
    StudentTermSection.objects.filter(term_section__section="F2").delete()
    body = _post(committee, run.pk, {}, VIEWS[0]).json()
    assert body["check"]["status"] == "changed"
    assert [(c["exam"], c["section"], c["membership"]) for c in body["check"]["changed"]] == [
        ("CS101", "F2", "gone")
    ]
    assert "counts" not in body
    exams = {exam["code"]: exam for exam in body["choices"]["exams"]}
    assert set(exams) == {"MATH101", "IS201", "CS101", "PHYS103 (1)", "PHYS103 (2)"}
    assert {d["id"] for d in body["choices"]["departments"]} == {"ai-ds", "cs", "is"}
    assert ExamTimetableRun.objects.filter(pk=run.pk).exists()
