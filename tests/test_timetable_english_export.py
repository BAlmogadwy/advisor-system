"""Continuous English lectures remain visible in every exported timetable."""

from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from core.models import (
    DeliveryBoard,
    Room,
    SectionPlacement,
    TermSection,
    TermSectionMeeting,
    TimetableScenario,
)
from core.services.timetable_autoplace import DEFAULT_SLOTS
from core.services.timetable_export import (
    _english_lecture,
    _merge_english_spans,
    export_scenario_xlsx,
)


def _placement(code="ENG101", start="09:00", end="12:00"):
    return SimpleNamespace(
        term_section=SimpleNamespace(course_code=code, course_key=code, course_number=code),
        start_time=start,
        end_time=end,
    )


def test_english_span_keeps_conflicting_cell_text_and_exact_times():
    ws = Workbook().active
    ws.cell(2, 2, "ENG101 S1")
    ws.cell(2, 3, "ENG101 S1\nAI301 S1")
    ws.cell(2, 4, "ENG101 S1")
    _merge_english_spans(ws, 2, DEFAULT_SLOTS, [_placement()])
    assert "B2:D2" in {str(span) for span in ws.merged_cells.ranges}
    assert "AI301 S1" in ws["B2"].value
    assert "09:00–12:00 · continuous lecture" in ws["B2"].value
    assert ws["B2"].fill.fgColor.rgb.endswith("F4CCCC")
    assert ws.row_dimensions[2].height >= 58


def test_ordinary_long_meetings_do_not_become_english_lectures():
    ws = Workbook().active
    ws.cell(2, 2, "AI301 Lab")
    _merge_english_spans(ws, 2, DEFAULT_SLOTS, [_placement("AI301")])
    assert not ws.merged_cells.ranges
    assert not _english_lecture(_placement("ENG103"))
    assert _english_lecture(_placement("ENG102"))


@pytest.mark.django_db
@pytest.mark.parametrize("start,end", [("09:00", "12:00"), ("13:00", "16:00")])
def test_complete_english_pattern_exports_without_room_occupancy(start, end):
    scenario = TimetableScenario.objects.create(academic_year="1448", term="1", name="English")
    board = DeliveryBoard.objects.create(
        scenario=scenario, label="T1", nominal_term=1, program="AI"
    )
    Room.objects.create(room_code="English space", room_type="lecture", capacity=40)
    Room.objects.create(room_code="Ordinary room", room_type="lecture", capacity=40)
    section = TermSection.objects.create(
        scenario=scenario,
        course_code="ENG101",
        course_number="101",
        course_key="ENG101",
        course_name="English",
        section="S1",
        source_tag="test",
    )
    for day in ["SUN", "MON", "WED", "THU"]:
        TermSectionMeeting.objects.create(
            term_section=section,
            day=day,
            start_time=start,
            end_time=end,
            instructor="Dr English",
        )
        SectionPlacement.objects.create(
            board=board,
            term_section=section,
            day=day,
            start_time=start,
            end_time=end,
            room="English space",
        )
    ordinary = TermSection.objects.create(
        scenario=scenario,
        course_code="AI301",
        course_key="AI301",
        course_name="AI",
        section="S1",
    )
    SectionPlacement.objects.create(
        board=board,
        term_section=ordinary,
        day="TUE",
        start_time="16:00",
        end_time="17:15",
        room="Ordinary room",
    )
    exported = export_scenario_xlsx(scenario.id)
    try:
        workbook = load_workbook(exported)
        for sheet_name in ["Term 1", "Instructors"]:
            sheet = workbook[sheet_name]
            lecture_cells = [
                cell
                for row in sheet.iter_rows()
                for cell in row
                if cell.value and "continuous lecture" in str(cell.value)
            ]
            assert len(lecture_cells) == 4, sheet_name
            assert all(f"{start}–{end}" in cell.value for cell in lecture_cells)
            assert all("No room required" in cell.value for cell in lecture_cells)
            assert all(
                any(cell.coordinate in span for span in sheet.merged_cells.ranges)
                for cell in lecture_cells
            )
            assert all("Lab" not in cell.value for cell in lecture_cells)
        room_values = [
            str(cell.value) for row in workbook["Rooms"].iter_rows() for cell in row if cell.value
        ]
        assert any("AI301" in value for value in room_values)
        assert not any("ENG101" in value or "English space" in value for value in room_values)
    finally:
        exported.unlink(missing_ok=True)
