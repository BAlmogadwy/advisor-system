"""Saved-run exports must preserve identities, constraints and unresolved work."""

import json
import zipfile
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.models import ExamTimetableRun
from core.services import exam_timetable
from core.services.exam_run_schema import STATUS_DERIVATION_VERSION, stamp_schema_version

pytestmark = pytest.mark.django_db


def _room(section, room_code, count, *, merged_from=None):
    return {
        "section": section,
        "gender": "F",
        "room_code": room_code,
        "student_count": count,
        "room_capacity": 30 if room_code != "UNASSIGNED" else 0,
        "merged_from": merged_from or [section],
        "building": "College A",
        "floor": "1",
    }


@pytest.fixture
def saved_data():
    schedule = [
        {
            "course_code": "CS111 (1)",
            "source_course_code": "CS111",
            "course_identity": "CS111|fundamentals of programming",
            "course_name": "Fundamentals of Programming",
            "programs": ["DS"],
            "enrolled_count": 20,
            "is_online": False,
            "day": "Sun",
            "period": "08:00-10:00",
            "slot_index": 0,
            # Two enrolment sections are merged into one room allocation.
            "rooms": [_room("F1+F2", "F-101", 20, merged_from=["F1", "F2"])],
        },
        {
            "course_code": "CS111 (2)",
            "source_course_code": "CS111",
            "course_identity": "CS111|programming i",
            "course_name": "Programming I",
            "programs": ["AI"],
            "enrolled_count": 40,
            "is_online": True,
            "day": "Mon",
            "period": "08:00-10:00",
            "slot_index": 1,
            # One enrolment section is split over two room allocations.
            "rooms": [_room("F3/1", "F-102", 20), _room("F3/2", "UNASSIGNED", 20)],
        },
        {
            "course_code": "GS101",
            "source_course_code": "GS101",
            "course_identity": "GS101|general studies",
            "course_name": "General Studies",
            "programs": ["AI"],
            "enrolled_count": 3,
            "is_online": False,
            "day": "OVERFLOW",
            "period": "",
            "slot_index": 2,
            "rooms": [],
        },
    ]
    return {
        "status": "ok",
        "enrollment_source": "scraper_timetable",
        "schedule": schedule,
        "slots": [
            {"index": 0, "day": "Sun", "period": "08:00-10:00"},
            {"index": 1, "day": "Mon", "period": "08:00-10:00"},
        ],
        "courses": [e["course_code"] for e in schedule],
        "courses_count": 3,
        "students_count": 43,
        "enrollment_scope": {"programs": ["AI", "DS"], "sections": ["F"]},
        "pinned": [{"course_code": "CS111 (2)", "day": "Mon", "period": "08:00-10:00"}],
        "assign_rooms": True,
        "credit_map": {"CS111 (1)": 4, "CS111 (2)": 3, "GS101": 2},
        "section_enrollment": {
            "CS111 (1)": [
                {"section": "F1", "gender": "F", "student_count": 12},
                {"section": "F2", "gender": "F", "student_count": 8},
            ],
            "CS111 (2)": [{"section": "F3", "gender": "F", "student_count": 40}],
            "GS101": [{"section": "F4", "gender": "F", "student_count": 3}],
        },
        "qa": {
            "total_students": 43,
            "total_courses": 3,
            "conflict_count": 1,
            "same_slot_conflicts": [
                {"student_id": 90, "slot_index": 0, "courses": ["CS111 (1)", "CS111 (2)"]}
            ],
            "overload_details": [
                {"student_id": 90, "day": "Sun", "courses": [{"code": "CS111 (1)"}]}
            ],
            "heavy_day_details": [
                {"student_id": 90, "day": "Sun", "courses": [{"code": "CS111 (2)"}]}
            ],
        },
    }


def _export(data, tmp_path, monkeypatch, label="Export integrity"):
    monkeypatch.setattr(exam_timetable, "RUNTIME_DIR", tmp_path)
    run = ExamTimetableRun.objects.create(
        label=label,
        result_json=json.dumps(stamp_schema_version(data)),
    )
    path = exam_timetable.export_exam_timetable_xlsx(run.pk)
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
    return load_workbook(path), path


def _records(sheet):
    rows = list(sheet.values)
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def _text(sheet):
    return "\n".join(str(cell.value) for row in sheet for cell in row if cell.value is not None)


def test_identity_and_fixed_online_markers_survive_every_export_view(
    saved_data, tmp_path, monkeypatch
):
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    courses = {r["Course Code"]: r for r in _records(book["Courses"])}
    assert courses["CS111 (1)"]["Course Name"] == "Fundamentals of Programming"
    assert courses["CS111 (2)"]["Course Name"] == "Programming I"
    assert courses["CS111 (2)"]["Fixed Time"] == "Yes"
    assert courses["CS111 (2)"]["Online"] == "Yes"
    assert courses["CS111 (1)"]["Programmes"] == "DS"
    assert courses["CS111 (2)"]["Programmes"] == "AI"
    assert (
        courses["CS111 (1)"]["Source Course Code"]
        == courses["CS111 (2)"]["Source Course Code"]
        == "CS111"
    )
    for name in ("Schedule", "Schedule (F)"):
        text = _text(book[name])
        assert "Fundamentals of Programming" in text
        assert "Programming I" in text
        assert "[FIXED]" in text and "[ONLINE]" in text
    for name in ("Students (F)", "Room Assignments", "Invigilators", "QA Summary"):
        assert "Programming I" in _text(book[name])
    assert "CS111" not in _text(book["Schedule (M)"])


def test_gender_counts_use_enrolment_sections_and_include_overflow(
    saved_data, tmp_path, monkeypatch
):
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    rows = {r["Course Code"]: r for r in _records(book["Students (F)"])}
    assert rows["CS111 (1)"]["Enrolled Sections"] == 2
    assert rows["CS111 (2)"]["Enrolled Sections"] == 1
    assert rows["CS111 (1)"]["Exam Enrolments"] == 20
    assert rows["CS111 (2)"]["Exam Enrolments"] == 40
    assert rows["GS101"]["Day"] == "OVERFLOW"
    assert rows["GS101"]["Exam Enrolments"] == 3
    assert rows["TOTAL EXAM ENROLMENTS"]["Exam Enrolments"] == 63
    qa = _text(book["QA Summary"])
    assert "Unique Students\n43" in qa
    assert "Sections in Scope\nF" in qa
    assert "Programmes in Scope\nAI, DS" in qa
    for name in ("Schedule", "Schedule (F)"):
        text = _text(book[name])
        assert "UNSCHEDULED EXAMS" in text
        assert "GS101 — General Studies" in text
        assert "UNASSIGNED" in text


@pytest.mark.parametrize("assign_rooms", [False, True])
def test_gender_views_keep_current_enrolments_when_no_rooms_exist(
    saved_data, tmp_path, monkeypatch, assign_rooms
):
    saved_data["assign_rooms"] = assign_rooms
    for entry in saved_data["schedule"]:
        entry["rooms"] = []
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    text = _text(book["Schedule (F)"])
    assert "Programming I" in text
    assert "pre-dates" not in text
    assert ("UNASSIGNED" if assign_rooms else "Room assignment not requested") in text
    assert "TOTAL EXAM ENROLMENTS\n" in _text(book["Students (F)"])


def test_export_preserves_unassigned_room_rows_and_correct_invigilator_demand(
    saved_data, tmp_path, monkeypatch
):
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    rooms = _records(book["Room Assignments"])
    assert len(rooms) == 3
    assert sum(r["Students"] for r in rooms) == 60
    missing = next(r for r in rooms if r["Room"] == "UNASSIGNED")
    assert missing["Students"] == 20
    assert missing["Course Name"] == "Programming I"
    assert missing["Fixed Time"] == "Yes"
    # Only the two assigned physical rooms require invigilators.
    invigilators = list(book["Invigilators"].values)
    assert any(row[:4] == ("TOTAL", 0, 2, 2) for row in invigilators)
    assert "UNASSIGNED" not in _text(book["Invigilators"])


def _mapped_section(key, label, count, *, status="mapped"):
    return {
        "section_key": key,
        "section": label,
        "term_section_id": int(key.split(":")[1]) if key.startswith("term-section:") else None,
        "mapping_status": status,
        "student_count": count,
        "gender": "F",
    }


def test_official_sections_and_room_groups_remain_separate_in_saved_export(
    saved_data, tmp_path, monkeypatch
):
    first, second, overflow = saved_data["schedule"]
    first_sections = [
        _mapped_section("term-section:1", "F01", 12),
        _mapped_section("term-section:2", "F1", 8),
        _mapped_section("unmapped:F:missing", "", 3, status="missing"),
        _mapped_section("unmapped:F:ambiguous", "", 2, status="ambiguous"),
    ]
    second_section = _mapped_section("term-section:3", "F3/1", 40)
    saved_data["section_enrollment"] = {
        first["course_code"]: first_sections,
        second["course_code"]: [second_section],
        overflow["course_code"]: [_mapped_section("term-section:4", "F4", 3)],
    }
    first["enrolled_count"] = 25
    first["rooms"] = [
        {**_room("F01+F1", "F-101", 20), "section_parts": first_sections[:2], "room_group": ""},
        {**_room("", "F-103", 3), "section_parts": [first_sections[2]], "room_group": ""},
        {**_room("", "UNASSIGNED", 2), "section_parts": [first_sections[3]], "room_group": ""},
    ]
    second["rooms"] = [
        {
            **_room("F3/1", room_code, 20),
            "section_parts": [
                {
                    **second_section,
                    "student_count": 20,
                    "room_group_index": index,
                    "room_group_count": 2,
                }
            ],
            "room_group": f"{index}/2",
        }
        for index, room_code in enumerate(["F-102", "F-104"], 1)
    ]
    saved_data["qa"]["section_mapping"] = {
        "mapped_sections": 4,
        "mapped_enrollments": 63,
        "missing_enrollments": 3,
        "ambiguous_enrollments": 2,
        "details": [
            {
                **first_sections[2],
                "course_code": first["course_code"],
                "reason": "No registered section",
            },
            {
                **first_sections[3],
                "course_code": first["course_code"],
                "reason": "Multiple registered sections",
            },
        ],
    }
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    rooms = _records(book["Room Assignments"])
    assert len(rooms) == 5
    assert sum(row["Students"] for row in rooms) == 65
    assert rooms[0]["Section"] == "F01 (12); F1 (8)"
    assert rooms[0]["Room Group"] is None
    assert rooms[1]["Section"] == "Section not recorded"
    assert rooms[1]["Section Mapping"] == "Section not recorded"
    assert rooms[2]["Section"] == "Ambiguous section"
    assert rooms[2]["Section Mapping"] == "Ambiguous section"
    split_rows = [row for row in rooms if row["Course"] == second["course_code"]]
    assert [row["Section"] for row in split_rows] == ["F3/1", "F3/1"]
    assert [row["Room Group"] for row in split_rows] == ["1/2", "2/2"]
    assert all(row["Section Mapping"] == "Recorded" for row in split_rows)
    assert all(row["Course Name"] == second["course_name"] for row in split_rows)
    for index, row in enumerate(rooms, 2):
        cell = book["Room Assignments"].cell(index, 9)
        assert cell.number_format == "0%"
        if row["Room"] == "UNASSIGNED":
            assert cell.value is None
        else:
            assert cell.data_type == "n"
            assert cell.value == pytest.approx(row["Students"] / row["Capacity"])
    student_rows = {row["Course Code"]: row for row in _records(book["Students (F)"])}
    student = student_rows[first["course_code"]]
    assert student["Enrolled Sections"] == 2
    assert student["Exam Enrolments"] == 25
    assert student["Missing Section Enrolments"] == 3
    assert student["Ambiguous Section Enrolments"] == 2
    assert student_rows["TOTAL EXAM ENROLMENTS"]["Exam Enrolments"] == 68
    assert student_rows["TOTAL EXAM ENROLMENTS"]["Missing Section Enrolments"] == 3
    assert student_rows["TOTAL EXAM ENROLMENTS"]["Ambiguous Section Enrolments"] == 2
    invig_values = list(book["Invigilators"].values)
    detail_index = next(
        i for i, row in enumerate(invig_values) if row[:4] == ("Day", "Period", "Course", "Type")
    )
    invig_rows = [
        dict(zip(invig_values[detail_index], row, strict=True))
        for row in invig_values[detail_index + 1 :]
    ]
    assert len(invig_rows) == 4
    assert sum(row["Invigilators"] for row in invig_rows) == 4
    assert invig_rows[0]["Section"] == rooms[0]["Section"]
    assert [row["Room Group"] for row in invig_rows[-2:]] == ["1/2", "2/2"]
    qa_text = _text(book["QA Summary"])
    assert "Missing Section Enrolments\n3" in qa_text
    assert "Ambiguous Section Enrolments\n2" in qa_text
    assert "No registered section" in qa_text
    assert "Multiple registered sections" in qa_text


def test_distinct_section_keys_count_even_when_official_labels_match(
    saved_data, tmp_path, monkeypatch
):
    saved_data["section_enrollment"]["CS111 (1)"] = [
        _mapped_section("term-section:1", "F01", 12),
        _mapped_section("term-section:2", "F01", 8),
    ]
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    rows = {row["Course Code"]: row for row in _records(book["Students (F)"])}
    assert rows["CS111 (1)"]["Enrolled Sections"] == 2


def test_room_qa_average_utilization_is_an_excel_number(saved_data, tmp_path, monkeypatch):
    saved_data["qa"]["rooms"] = {"avg_utilization": 0.625}
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    sheet = book["Room Assignments"]
    row = next(row for row in sheet if row[0].value == "Avg Utilization")
    assert row[1].value == 0.625
    assert row[1].data_type == "n"
    assert row[1].number_format == "0.0%"


def test_saved_literals_are_never_interpreted_as_excel_formulas(saved_data, tmp_path, monkeypatch):
    saved_data["schedule"][0]["course_name"] = '=HYPERLINK("https://invalid.example")'
    book, path = _export(saved_data, tmp_path, monkeypatch, label="=1+1")
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    assert "=1+1" in _text(book["QA Summary"])
    assert saved_data["schedule"][0]["course_name"] in _text(book["Courses"])
    assert list(Path(tmp_path).glob("*.xlsx")) == [path]


def test_export_does_not_modify_persisted_run(saved_data, tmp_path, monkeypatch):
    _, path = _export(saved_data, tmp_path, monkeypatch)
    run = ExamTimetableRun.objects.get()
    original = run.result_json
    second_path = exam_timetable.export_exam_timetable_xlsx(run.pk)
    assert second_path != path
    with zipfile.ZipFile(second_path) as archive:
        assert archive.testzip() is None
    run.refresh_from_db()
    assert run.result_json == original


def test_export_works_while_previous_download_is_open(saved_data, tmp_path, monkeypatch):
    _, first_path = _export(saved_data, tmp_path, monkeypatch)
    run = ExamTimetableRun.objects.get()
    with first_path.open("rb") as active_download:
        next_path = exam_timetable.export_exam_timetable_xlsx(run.pk)
        assert next_path != first_path
        with zipfile.ZipFile(active_download) as archive:
            assert archive.testzip() is None
        with zipfile.ZipFile(next_path) as archive:
            assert archive.testzip() is None


def test_export_failure_cleans_partial_artifact(saved_data, tmp_path, monkeypatch):
    from openpyxl import Workbook

    monkeypatch.setattr(exam_timetable, "RUNTIME_DIR", tmp_path)
    run = ExamTimetableRun.objects.create(
        label="Failed export",
        result_json=json.dumps(stamp_schema_version(saved_data)),
    )

    def fail_save(self, filename):
        Path(filename).write_bytes(b"partial zip")
        raise OSError("Storage unavailable")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(OSError, match="Storage unavailable"):
        exam_timetable.export_exam_timetable_xlsx(run.pk)
    assert list(tmp_path.iterdir()) == []


def test_dense_schedule_uses_separate_course_rows_with_repeated_day_labels(
    saved_data, tmp_path, monkeypatch
):
    entries = []
    enrolment = {}
    for index in range(24):
        entry = deepcopy(saved_data["schedule"][0])
        entry.update(
            course_code=f"DENSE{index:02}",
            course_name=f"Distinct Course Name {index:02}",
            source_course_code=f"DENSE{index:02}",
            course_identity=f"dense{index:02}",
            rooms=[_room("F1", f"F-{index:03}", 20)],
        )
        entries.append(entry)
        enrolment[entry["course_code"]] = [{"section": "F1", "gender": "F", "student_count": 20}]
    saved_data.update(schedule=entries, section_enrollment=enrolment, pinned=[])
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    for name in ("Schedule", "Schedule (F)"):
        sheet = book[name]
        assert not sheet.merged_cells.ranges
        assert [sheet.cell(row, 1).value for row in range(2, 26)] == ["Sun"] * 24
        assert sheet["A26"].value == "Mon"
        for index in range(24):
            value = sheet.cell(index + 2, 2).value
            assert f"DENSE{index:02}" in value
            assert f"Distinct Course Name {index:02}" in value
            assert f"F-{index:03}" in value
            assert value.count("DENSE") == 1
        assert all((row.height or 0) <= 409.5 for row in sheet.row_dimensions.values())


def test_large_single_course_continues_without_losing_name_or_room_details(
    saved_data, tmp_path, monkeypatch
):
    entry = saved_data["schedule"][0]
    long_name = "Complete course name retained across continuation cells " * 40
    entry["course_name"] = long_name
    entry["rooms"] = [_room(f"F{index}", f"F-ROOM-{index:04}", 20) for index in range(150)]
    saved_data.update(schedule=[entry], pinned=[])
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    for name in ("Schedule", "Schedule (F)"):
        sheet = book[name]
        cells = [str(row[1].value) for row in sheet.iter_rows(min_row=2) if row[1].value]
        assert len(cells) > 1
        assert "[CONTINUED]" in cells[1]
        details = "".join(value.split("\n", 1)[1] for value in cells)
        assert long_name in details
        for index in range(150):
            assert f"F-ROOM-{index:04}" in details
        assert all((row.height or 0) <= 409.5 for row in sheet.row_dimensions.values())


def test_invigilator_headers_and_rule_notes_are_visible(saved_data, tmp_path, monkeypatch):
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    sheet = book["Invigilators"]
    for row in sheet:
        if row[0].value == "Day":
            assert row[0].fill.fgColor.rgb == "000A8E6E"
            assert row[0].font.bold
    assert "A2:M2" in {str(merged) for merged in sheet.merged_cells.ranges}
    assert "A3:M3" in {str(merged) for merged in sheet.merged_cells.ranges}
    assert sheet.row_dimensions[2].height >= 30
    for sheet in book:
        assert all((row.height or 0) <= 409.5 for row in sheet.row_dimensions.values())


@pytest.mark.parametrize("empty_gender", ["M", "F"])
def test_empty_gender_students_sheet_explains_scope_exclusion(
    saved_data, tmp_path, monkeypatch, empty_gender
):
    populated_gender = "F" if empty_gender == "M" else "M"
    for sections in saved_data["section_enrollment"].values():
        for section in sections:
            section["gender"] = populated_gender
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    sheet = book[f"Students ({empty_gender})"]
    assert sheet["A2"].value == f"No {empty_gender} enrolments in this timetable"
    assert sheet["E2"].value == sheet["F2"].value == 0
    assert f"No {populated_gender} enrolments" not in _text(book[f"Students ({populated_gender})"])


@pytest.mark.parametrize("student_count, expected_invigilators", [(2, 1), (30, 2)])
def test_unknown_course_prefix_is_labeled_with_actual_staffing_rule(
    saved_data,
    tmp_path,
    monkeypatch,
    student_count,
    expected_invigilators,
):
    entry = saved_data["schedule"][0]
    entry.update(course_code="QREX336", rooms=[_room("F1", "F-101", student_count)])
    saved_data.update(schedule=[entry], pinned=[])
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    sheet = book["Invigilators"]
    assert "Other unlisted course prefixes use this department rule" in sheet["A2"].value
    detail = next(row for row in sheet.values if len(row) > 2 and row[2] == "QREX336")
    assert detail[3] == "Other (department rule)"
    assert detail[8] == expected_invigilators


@pytest.mark.parametrize(
    "primary_status, expected_label",
    [
        ("clean", "Clean"),
        ("clean_with_approved_thin_conflicts", "Clean (with approved tiny-course clashes)"),
        ("requires_room_action", "Requires room action"),
        ("requires_section_review", "Teaching sections need review"),
        ("contains_overflow", "Contains overflow exams"),
        ("contains_manual_override", "Contains schedule violations"),
        ("contains_workload_warnings", "Contains workload warnings"),
        ("infeasible", "Infeasible"),
        ("new_review_needed", "new review needed"),
    ],
)
def test_qa_status_uses_readable_screen_labels(
    saved_data,
    tmp_path,
    monkeypatch,
    primary_status,
    expected_label,
):
    saved_data.update(
        primary_status=primary_status,
        status_derivation_version=STATUS_DERIVATION_VERSION,
        status_flags=[
            "approved_thin_conflicts",
            "room_action_required",
            "section_mapping_incomplete",
            "overflow",
            "manual_override",
            "multi_sitting_required",
            "legacy_incomplete_qa",
            "daily_limit_exceeded",
            "heavy_credit_day",
            "new_review_needed",
        ],
    )
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    sheet = book["QA Summary"]
    # Keep the same metric positions: only presentation text changes.
    assert sheet["A4"].value == "Status"
    assert sheet["B4"].value == expected_label
    assert sheet["A5"].value == "Status Flags"
    assert sheet["B5"].value == (
        "approved thin conflicts, room action required, teaching-section mapping incomplete, overflow, schedule violations, "
        "multi-sitting required, legacy / incomplete QA data, daily exam limit exceeded, "
        "heavy credit days, new review needed"
    )
    assert "_" not in sheet["B4"].value
    assert "_" not in sheet["B5"].value


@pytest.mark.parametrize("source", [None, "studying", "fallback_studying", "manual"])
def test_export_requires_exact_scraped_population_provenance(
    saved_data, tmp_path, monkeypatch, source
):
    if source is None:
        saved_data.pop("enrollment_source")
    else:
        saved_data["enrollment_source"] = source
    with pytest.raises(ValueError, match="Load Courses from actual scraped timetables"):
        _export(saved_data, tmp_path, monkeypatch)
    assert not list(tmp_path.glob("*.xlsx"))


def test_export_print_settings_keep_wide_tables_together(saved_data, tmp_path, monkeypatch):
    book, _ = _export(saved_data, tmp_path, monkeypatch)
    for sheet in book:
        assert sheet.page_setup.orientation == "landscape"
        assert str(sheet.page_setup.paperSize) == sheet.PAPERSIZE_A3
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.page_setup.fitToHeight == 0
        assert sheet.sheet_properties.pageSetUpPr.fitToPage
        assert str(sheet.print_area)
    for name in ("Courses", "Students (M)", "Students (F)", "Room Assignments"):
        assert book[name].print_title_rows == "$1:$1"
    assert book["QA Summary"].freeze_panes == "B2"
    invigilators = book["Invigilators"]
    assert invigilators.freeze_panes is None
    assert invigilators.sheet_view.pane is None
    assert all(selection.pane is None for selection in invigilators.sheet_view.selection)
