"""Department export endpoints keep authorization and saved evidence boundaries."""

import json
from io import BytesIO
from zipfile import ZipFile

import pytest
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from core.models import ExamTimetableRun
from core.services.exam_run_schema import EXAM_RUN_SCHEMA_VERSION
from tests.test_exam_department_export import XLSX_TYPE, make_department_data

pytestmark = pytest.mark.django_db


@pytest.fixture
def saved_department_run():
    return ExamTimetableRun.objects.create(
        label="Department export endpoint audit",
        result_json=json.dumps(make_department_data()),
    )


@pytest.fixture
def department_admin(client, django_user_model):
    client.force_login(django_user_model.objects.create_superuser(username="department-admin"))
    return client


def _url(run, export=False):
    return reverse("exam_department_export" if export else "exam_department_options", args=[run.pk])


def _payload(**changes):
    return {"departments": ["ai-ds"], "genders": ["M", "F"], "language": "en", **changes}


def _download(client, run, payload=None):
    return client.post(
        _url(run, True), _payload() if payload is None else payload, content_type="application/json"
    )


def _body(response):
    return b"".join(response.streaming_content) if response.streaming else response.content


def test_options_and_workbook_download_agree_on_saved_departments(
    department_admin, saved_department_run
):
    options = department_admin.get(_url(saved_department_run))
    assert options.status_code == 200
    departments = {item["id"]: item for item in options.json()["departments"]}
    assert departments["ai-ds"]["course_count"] == 6
    assert departments["ai-ds"]["student_sittings"] == 41
    response = _download(department_admin, saved_department_run)
    assert response.status_code == 200, response.content
    assert response["Content-Type"] == XLSX_TYPE
    assert ".xlsx" in response["Content-Disposition"]
    content = _body(response)
    with ZipFile(BytesIO(content)) as archive:
        assert archive.testzip() is None
        assert "xl/workbook.xml" in archive.namelist()
    book = load_workbook(BytesIO(content))
    assert "Details" in book.sheetnames
    response.close()


def test_anonymous_and_non_superadmin_cannot_read_or_export_departments(
    client, django_user_model, saved_department_run
):
    for export in (False, True):
        response = (
            _download(client, saved_department_run)
            if export
            else client.get(_url(saved_department_run))
        )
        assert response.status_code in (302, 403)
    client.force_login(
        django_user_model.objects.create_user(username="department-reader", is_staff=True)
    )
    assert client.get(_url(saved_department_run)).status_code == 403
    assert _download(client, saved_department_run).status_code == 403


def test_export_requires_csrf_and_correct_methods(
    department_admin, saved_department_run, django_user_model
):
    strict_client = Client(enforce_csrf_checks=True)
    strict_client.force_login(django_user_model.objects.get(username="department-admin"))
    assert _download(strict_client, saved_department_run).status_code == 403
    assert department_admin.get(_url(saved_department_run, True)).status_code == 405
    assert department_admin.post(_url(saved_department_run)).status_code == 405


def test_missing_run_is_404_for_options_and_download(department_admin):
    missing = ExamTimetableRun(pk=987654321)
    assert department_admin.get(_url(missing)).status_code == 404
    assert _download(department_admin, missing).status_code == 404


@pytest.mark.parametrize(
    "sentinel", ["future_schema", "future_version_unrenderable", "unrenderable"]
)
def test_unrenderable_saved_run_cannot_emit_department_options_or_workbook(
    department_admin, saved_department_run, sentinel
):
    data = json.loads(saved_department_run.result_json)
    if sentinel == "future_schema":
        data["schema_version"] = EXAM_RUN_SCHEMA_VERSION + 1
    else:
        data["status"] = sentinel
    saved_department_run.result_json = json.dumps(data)
    saved_department_run.save(update_fields=["result_json"])
    for response in (
        department_admin.get(_url(saved_department_run)),
        _download(department_admin, saved_department_run),
    ):
        assert response.status_code == 409
        assert response["Content-Type"].startswith("application/json")
        assert response.json()["code"] == "operations_snapshot_required"
        assert response.json()["error"]
        assert "Content-Disposition" not in response


@pytest.mark.parametrize(
    "changes",
    [
        {"departments": []},
        {"departments": ["unknown-department"]},
        {"departments": "ai-ds"},
        {"departments": ["ai-ds", "ai-ds"]},
        {"genders": []},
        {"genders": ["X"]},
        {"genders": "F"},
        {"genders": ["F", "F"]},
        {"language": "fr"},
        {"language": []},
        {"language": {}},
        {"unknown_option": True},
        {"dates": "2026-09-20"},
        {"dates": {"Sun": "2026-02-30"}},
        {"dates": {"Sun": "20/09/2026"}},
        {"dates": {"Sun": "not-a-date"}},
        {"dates": {"Wed": "2026-09-23"}},
        {"dates": {"Sun": "2026-09-20", "Mon": "2026-09-20"}},
        {"dates": {"Sun": "2026-09-21", "Mon": "2026-09-20"}},
    ],
)
def test_invalid_selection_or_date_is_rejected_before_export(
    department_admin, saved_department_run, changes
):
    response = _download(department_admin, saved_department_run, _payload(**changes))
    assert response.status_code == 400, response.content
    assert response["Content-Type"].startswith("application/json")
    assert response.json()["error"]


@pytest.mark.parametrize("body", ["{", "[]", "null", '"a string"'])
def test_malformed_json_body_returns_400(department_admin, saved_department_run, body):
    response = department_admin.post(
        _url(saved_department_run, True), data=body, content_type="application/json"
    )
    assert response.status_code == 400, response.content


@pytest.mark.parametrize(
    "problem",
    [
        "no_snapshot",
        "old_source",
        "unknown_snapshot_source",
        "future_snapshot",
        "boolean_version",
        "missing_course",
        "wrong_identity",
        "invalid_identity_type",
        "negative_section",
        "boolean_section_count",
        "wrong_program_total",
        "wrong_gender_total",
        "invalid_gender_type",
        "section_program_mismatch",
        "duplicate_section",
        "invalid_original_sections_type",
        "invalid_original_section_row",
        "changed_section_label",
        "changed_section_id",
        "changed_mapping_status",
        "changed_mapping_source",
    ],
)
def test_missing_or_tampered_saved_evidence_returns_rebuild_conflict(
    department_admin, saved_department_run, problem
):
    data = json.loads(saved_department_run.result_json)
    snapshot = data["operations_snapshot"]
    course = snapshot["courses"]["CS111 (1)"]
    section = course["sections"][0]
    if problem == "no_snapshot":
        data.pop("operations_snapshot")
    elif problem == "old_source":
        data["enrollment_source"] = "plan"
    elif problem == "unknown_snapshot_source":
        snapshot["enrollment_source"] = "unverified"
    elif problem == "future_snapshot":
        snapshot["version"] = 999
    elif problem == "boolean_version":
        snapshot["version"] = True
    elif problem == "missing_course":
        del snapshot["courses"]["CS111 (1)"]
    elif problem == "wrong_identity":
        course["course_identity"] = "CS111|different canonical course"
    elif problem == "invalid_identity_type":
        course["course_identity"] = data["schedule"][0]["course_identity"] = ["invalid identity"]
    elif problem == "negative_section":
        section["student_count"] = -1
    elif problem == "boolean_section_count":
        section["student_count"] = True
    elif problem == "wrong_program_total":
        course["program_counts"][0]["student_count"] += 1
    elif problem == "wrong_gender_total":
        course["program_counts"][0]["gender"] = "M"
    elif problem == "invalid_gender_type":
        section["gender"] = []
    elif problem == "section_program_mismatch":
        section["program_counts"]["AI"] += 1
    elif problem == "duplicate_section":
        course["sections"].append(dict(section))
    elif problem == "invalid_original_sections_type":
        data["section_enrollment"] = []
    elif problem == "invalid_original_section_row":
        data["section_enrollment"]["CS111 (1)"] = [42]
    elif problem == "changed_section_label":
        section["section"] = "F88"
    elif problem == "changed_section_id":
        section["term_section_id"] = 99999
    elif problem == "changed_mapping_status":
        section["mapping_status"] = "missing"
    elif problem == "changed_mapping_source":
        section["mapping_source"] = "plan"
    saved_department_run.result_json = json.dumps(data)
    saved_department_run.save(update_fields=["result_json"])
    for response in (
        department_admin.get(_url(saved_department_run)),
        _download(department_admin, saved_department_run),
    ):
        assert response.status_code == 409, response.content
        assert response["Content-Type"].startswith("application/json")
        assert response.json()["error"]
        assert "Content-Disposition" not in response


def test_departments_export_reads_no_live_enrolment_or_instructor_tables(
    department_admin, saved_department_run
):
    from django.db import connection
    from django.test.utils import CaptureQueriesContext

    with CaptureQueriesContext(connection) as queries:
        assert department_admin.get(_url(saved_department_run)).status_code == 200
        response = _download(department_admin, saved_department_run)
        assert response.status_code == 200
        assert load_workbook(BytesIO(_body(response)))["Details"].max_row > 1
    sql = "\n".join(query["sql"].lower() for query in queries)
    # Authentication and the saved run may be read, but the original source
    # can have changed since the registrar saved this schedule.
    for table in (
        "students",
        "student_term_sections",
        "term_sections",
        "term_section_meetings",
        "programme_requirements",
        "student_courses",
    ):
        assert f'"{table}"' not in sql
