"""Copying history duplicates saved evidence without rebuilding an exam run."""

import json
from datetime import timedelta
from io import BytesIO

import pytest
from django.db import connection
from django.test import Client
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core import exam_views
from core.models import AuditLog, ExamTimetableRun
from core.services.exam_department_export import export_department_workbooks
from core.services.exam_run_schema import EXAM_RUN_SCHEMA_VERSION, load_normalised_run
from tests.test_exam_department_export import _records, make_department_data

pytestmark = pytest.mark.django_db


@pytest.fixture
def copy_source():
    data = make_department_data()
    data.update(
        {
            "run_id": 987654,
            "label": "Stored internal label",
            "created_at": "2020-01-01T00:00:00+00:00",
            "future_metadata": {"note": "دليل محفوظ", "flags": [True, None, {"untouched": 7}]},
        }
    )
    run = ExamTimetableRun.objects.create(
        label="Original saved timetable",
        result_json=json.dumps(data, ensure_ascii=False, indent=3) + "\n\n",
    )
    ExamTimetableRun.objects.filter(pk=run.pk).update(created_at=timezone.now() - timedelta(days=1))
    run.refresh_from_db()
    return run


@pytest.fixture
def copy_admin(client, django_user_model):
    client.force_login(django_user_model.objects.create_superuser(username="copy-admin"))
    return client


def _copy(client, source, payload=None, **request_options):
    return client.post(
        reverse("exam_timetable_copy", args=[source.pk]),
        {} if payload is None else payload,
        content_type="application/json",
        **request_options,
    )


def test_copy_creates_an_independent_history_row_preserving_raw_snapshot(copy_admin, copy_source):
    source_text = copy_source.result_json
    source_created = copy_source.created_at
    source_label = copy_source.label
    source_snapshot = load_normalised_run(copy_source)

    response = _copy(copy_admin, copy_source, {"label": "  Working copy  "})

    assert response.status_code == 201, response.content
    body = response.json()
    copied = ExamTimetableRun.objects.get(pk=body["run_id"])
    assert body["ok"] is True
    assert body["source_run_id"] == copy_source.pk
    assert copied.pk != copy_source.pk
    assert copied.label == body["label"] == "Working copy"
    assert copied.created_at > source_created
    assert body["created_at"] == copied.created_at.isoformat()
    assert copied.result_json == source_text
    for key, value in source_snapshot.items():
        if key not in {"run_id", "label", "created_at", "ok", "source_run_id"}:
            assert body[key] == value
    copy_source.refresh_from_db()
    assert (copy_source.result_json, copy_source.created_at, copy_source.label) == (
        source_text,
        source_created,
        source_label,
    )

    audit = AuditLog.objects.get(action="exam_timetable.copy_run")
    assert audit.actor_username == "copy-admin"
    assert audit.status == "success"
    assert audit.method == "POST"
    assert json.loads(audit.details_json) == {
        "source_run_id": copy_source.pk,
        "run_id": copied.pk,
        "label": "Working copy",
    }


def test_editing_and_deleting_a_copy_leaves_the_original_unchanged(copy_admin, copy_source):
    original = (copy_source.label, copy_source.result_json, copy_source.created_at)
    response = _copy(copy_admin, copy_source, {"label": "Editable copy"})
    assert response.status_code == 201
    copied = ExamTimetableRun.objects.get(pk=response.json()["run_id"])
    edited = json.loads(copied.result_json)
    edited["schedule"][0]["day"] = "Thu"
    edited["pinned"] = []
    copied.result_json = json.dumps(edited)
    copied.label = "Renamed copy"
    copied.save(update_fields=["result_json", "label"])
    copy_source.refresh_from_db()
    assert (copy_source.label, copy_source.result_json, copy_source.created_at) == original

    deleted = copy_admin.post(
        reverse("exam_timetable_delete", args=[copied.pk]),
        {"confirm": "DELETE"},
        content_type="application/json",
    )
    assert deleted.status_code == 200
    assert not ExamTimetableRun.objects.filter(pk=copied.pk).exists()
    copy_source.refresh_from_db()
    assert (copy_source.label, copy_source.result_json, copy_source.created_at) == original


def test_duplicate_names_are_allowed_and_copies_appear_newest_first(copy_admin, copy_source):
    first = _copy(copy_admin, copy_source, {"label": copy_source.label})
    second = _copy(copy_admin, copy_source, {"label": copy_source.label})
    assert first.status_code == second.status_code == 201
    first_id, second_id = first.json()["run_id"], second.json()["run_id"]
    assert len({first_id, second_id, copy_source.pk}) == 3
    assert ExamTimetableRun.objects.filter(label=copy_source.label).count() == 3
    history = copy_admin.get(reverse("exam_timetable_list")).json()
    assert [run["id"] for run in history["runs"][:3]] == [second_id, first_id, copy_source.pk]


def test_history_uses_new_id_to_order_copies_with_equal_timestamps(copy_admin, copy_source):
    response = _copy(copy_admin, copy_source)
    assert response.status_code == 201
    copied_id = response.json()["run_id"]
    ExamTimetableRun.objects.filter(pk=copied_id).update(created_at=copy_source.created_at)
    history = copy_admin.get(reverse("exam_timetable_list")).json()
    assert [run["id"] for run in history["runs"][:2]] == [copied_id, copy_source.pk]


@pytest.mark.parametrize("language,suffix", [("en", " — Copy"), ("ar", " — نسخة")])
def test_omitted_label_uses_localized_suffix_and_truncates_only_the_base(
    copy_admin, copy_source, language, suffix
):
    copy_source.label = "م" * 120
    copy_source.save(update_fields=["label"])
    response = _copy(copy_admin, copy_source, HTTP_ACCEPT_LANGUAGE=language)
    assert response.status_code == 201
    assert response.json()["label"] == copy_source.label[: 120 - len(suffix)] + suffix
    assert len(response.json()["label"]) == 120


@pytest.mark.parametrize("label", ["x", "س" * 120])
def test_boundary_label_lengths_are_accepted(copy_admin, copy_source, label):
    response = _copy(copy_admin, copy_source, {"label": label})
    assert response.status_code == 201
    assert response.json()["label"] == label


@pytest.mark.parametrize(
    "payload",
    [
        {"label": ""},
        {"label": " \t\n "},
        {"label": "x" * 121},
        {"label": "س" * 121},
        {"label": None},
        {"label": True},
        {"label": 42},
        {"label": []},
        {"label": {}},
        {"label": "Before\x00after"},
        {"label": "Before\nafter"},
        {"label": "Before\x7fafter"},
        {"label": "\ud800"},
        {"label": "Copy", "unknown": True},
    ],
)
def test_invalid_labels_or_unknown_fields_do_not_create_history_rows(
    copy_admin, copy_source, payload
):
    count = ExamTimetableRun.objects.count()
    response = _copy(copy_admin, copy_source, payload)
    assert response.status_code == 400
    assert response.json()["error"]
    assert ExamTimetableRun.objects.count() == count
    assert not AuditLog.objects.filter(action="exam_timetable.copy_run", status="success").exists()


@pytest.mark.parametrize("body", [b"{", b"[]", b"null", b'"text"', b"42", b'{"label":"\xff"}'])
def test_malformed_or_non_object_requests_return_400_without_copy(copy_admin, copy_source, body):
    response = copy_admin.post(
        reverse("exam_timetable_copy", args=[copy_source.pk]),
        data=body,
        content_type="application/json",
    )
    assert response.status_code == 400
    assert ExamTimetableRun.objects.count() == 1


def test_deep_json_within_body_limit_returns_400_without_copy_or_audit(copy_admin, copy_source):
    body = b'{"label":' + b"[" * 1600 + b"0" + b"]" * 1600 + b"}"
    assert len(body) < 4096
    audit_count = AuditLog.objects.count()
    response = copy_admin.post(
        reverse("exam_timetable_copy", args=[copy_source.pk]),
        data=body,
        content_type="application/json",
    )
    assert response.status_code == 400
    assert response.json()["error"]
    assert ExamTimetableRun.objects.count() == 1
    assert AuditLog.objects.count() == audit_count


def test_copy_request_body_limit_is_enforced_at_4096_bytes(copy_admin, copy_source):
    raw = b'{"label":"Bounded copy"}'
    at_limit = raw + b" " * (4096 - len(raw))
    url = reverse("exam_timetable_copy", args=[copy_source.pk])
    accepted = copy_admin.post(url, at_limit, content_type="application/json")
    assert accepted.status_code == 201
    rejected = copy_admin.post(url, at_limit + b" ", content_type="application/json")
    assert rejected.status_code == 400
    assert ExamTimetableRun.objects.count() == 2


@pytest.mark.parametrize(
    "problem",
    [
        "future_schema",
        "future_version_unrenderable",
        "unrenderable",
        "feasibility_error",
        "corrupt_json",
        "null_json",
    ],
)
def test_unrenderable_or_unsuccessful_source_cannot_be_copied(copy_admin, copy_source, problem):
    data = json.loads(copy_source.result_json)
    if problem == "future_schema":
        data["schema_version"] = EXAM_RUN_SCHEMA_VERSION + 1
    elif problem in {"future_version_unrenderable", "unrenderable", "feasibility_error"}:
        data["status"] = problem
    copy_source.result_json = (
        "{" if problem == "corrupt_json" else "null" if problem == "null_json" else json.dumps(data)
    )
    copy_source.save(update_fields=["result_json"])
    response = _copy(copy_admin, copy_source)
    assert response.status_code == 409
    assert response.json()["error"]
    assert response.json()["code"] == "run_not_copyable"
    assert ExamTimetableRun.objects.count() == 1
    assert not AuditLog.objects.filter(action="exam_timetable.copy_run", status="success").exists()


def test_copy_requires_superadmin_login_and_post(client, django_user_model, copy_source):
    assert _copy(client, copy_source).status_code in (302, 403)
    client.force_login(django_user_model.objects.create_user(username="copy-staff", is_staff=True))
    assert _copy(client, copy_source).status_code == 403
    client.force_login(django_user_model.objects.create_superuser(username="copy-method-admin"))
    assert client.get(reverse("exam_timetable_copy", args=[copy_source.pk])).status_code == 405
    assert ExamTimetableRun.objects.count() == 1


def test_copy_keeps_normal_csrf_protection(django_user_model, copy_source, settings):
    client = Client(enforce_csrf_checks=True)
    client.force_login(django_user_model.objects.create_superuser(username="copy-csrf-admin"))
    assert _copy(client, copy_source).status_code == 403
    assert ExamTimetableRun.objects.count() == 1
    token = "a" * 32
    client.cookies[settings.CSRF_COOKIE_NAME] = token
    assert _copy(client, copy_source, HTTP_X_CSRFTOKEN=token).status_code == 201


def test_missing_source_is_404(copy_admin):
    response = _copy(copy_admin, ExamTimetableRun(pk=987654321))
    assert response.status_code == 404
    assert ExamTimetableRun.objects.count() == 0


def test_copy_does_not_schedule_check_or_query_live_membership(
    copy_admin, copy_source, monkeypatch
):
    def forbidden(*args, **kwargs):
        pytest.fail("History copy attempted to rebuild or evaluate saved evidence")

    for name in (
        "build_exam_timetable",
        "evaluate_exam_schedule",
        "_loaded_request_context",
        "build_enrolled_sets_with_meta",
        "schedule",
        "run_multistart",
        "build_conflict_graph",
    ):
        monkeypatch.setattr(exam_views, name, forbidden)
    with CaptureQueriesContext(connection) as queries:
        response = _copy(copy_admin, copy_source)
        assert response.status_code == 201
    sql = "\n".join(query["sql"].lower() for query in queries)
    for table in (
        "students",
        "student_term_sections",
        "term_sections",
        "term_section_meetings",
        "programme_requirements",
        "student_courses",
    ):
        assert f'"{table}"' not in sql


def test_copy_preserves_department_export_rows_and_identity_keys(copy_admin, copy_source):
    response = _copy(copy_admin, copy_source, {"label": "Department working copy"})
    assert response.status_code == 201
    copied = ExamTimetableRun.objects.get(pk=response.json()["run_id"])
    payload = {"departments": ["ai-ds"], "genders": ["M", "F", "U"], "language": "en"}
    original_content, _, _ = export_department_workbooks(copy_source, payload)
    copied_content, _, _ = export_department_workbooks(copied, payload)
    original_book = load_workbook(BytesIO(original_content))
    copied_book = load_workbook(BytesIO(copied_content))
    assert _records(copied_book["Details"]) == _records(original_book["Details"])
    copied_detail = copy_admin.get(reverse("exam_timetable_detail", args=[copied.pk])).json()
    original_detail = copy_admin.get(reverse("exam_timetable_detail", args=[copy_source.pk])).json()
    for field in (
        "schedule",
        "operations_snapshot",
        "section_enrollment",
        "enrollment_scope",
        "schema_version",
        "pinned",
        "future_metadata",
    ):
        assert copied_detail[field] == original_detail[field]


def test_copy_preserves_legacy_compatible_payload_without_inventing_new_snapshot(
    copy_admin, copy_source
):
    legacy = json.loads(copy_source.result_json)
    legacy.pop("operations_snapshot")
    legacy.pop("enrollment_source")
    legacy["schema_version"] = 1
    copy_source.result_json = json.dumps(legacy, indent=2)
    copy_source.save(update_fields=["result_json"])
    response = _copy(copy_admin, copy_source)
    assert response.status_code == 201
    copied = ExamTimetableRun.objects.get(pk=response.json()["run_id"])
    assert copied.result_json == copy_source.result_json
    assert response.json()["operations_snapshot"] is None
