"""The per-plan export's workbooks as digests, for the 6b move (§7.4, PV1).

Not collected (no ``test_`` prefix). ``SCENARIOS`` rebuilds every scenario the
export's own tests build (``test_per_plan_export_grouping.py``,
``test_per_plan_export_electives.py``, the off-grid case of
``test_timetable_slot_config.py``) plus a few that make the SHARED search and
out-of-plan sections count. ``scenario_digests`` exports each with every
search in ``SEARCHES`` and reduces each sheet of each file to a digest of its
canonical lines (``sheet_lines``: names, merges, column widths, and each
cell's value, fill and font), so an oracle recorded before the move
(``tests/fixtures/per_plan_export_digests.json``, written by
``tests/record_per_plan_export_digests.py`` with lxml blocked) holds the
workbooks without holding a dump of them.

The lines are read back through openpyxl, so the writer it used (lxml or its
own) does not count, and a newline written as CRLF (Windows, no lxml) reads as
LF. To see a difference the digests only name, dump both sides in full
(``TW_PER_PLAN_DUMP=<file>`` with the recorder, on the base and on the branch)
and run ``python tests/per_plan_export_digests_support.py <base> <branch>``:
it prints the first line that differs.
"""

from __future__ import annotations

import hashlib
import io
import json
import sys
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

#: Every search the oracle holds: none, the SHARED token alone, a programme
#: token, a programme with a text token, SHARED with a token that is a
#: programme on some scenarios and text on the others.
SEARCHES = ("", "SHARED", "AI2", "AI2 CS1", "SHARED DS")

#: The flags the oracle is recorded and replayed with (the export reads the
#: teacher and mark flags through the lens and the teacher texts).
FLAGS = {
    "TIMETABLE_INSTRUCTOR_LINKS_ENABLED": True,
    "TIMETABLE_TIERED_OBJECTIVE_ENABLED": True,
    "TIMETABLE_ONE_TEACHER_PER_SECTION": False,
    "TIMETABLE_INSTRUCTOR_DAY_OFF_POLICY": False,
    "TIMETABLE_SECTION_PROGRAMME_PREFERENCE": False,
}


# ── canonical lines ─────────────────────────────────────────────────────────


def _color(color: Any) -> str:
    if color is None:
        return "-"
    return f"{color.type}:{color.value}:{color.tint}"


def _value(value: Any) -> Any:
    if isinstance(value, str):
        return value.replace("\r\n", "\n")
    if value is None or isinstance(value, bool | int | float):
        return value
    return str(value)


def sheet_lines(ws) -> list[str]:
    """One sheet as lines: its name, merges, widths, then every cell that has
    a value or a style, with its value, fill and font."""
    lines = [f"sheet {ws.title}"]
    lines += [f"merge {rng}" for rng in sorted(str(r) for r in ws.merged_cells.ranges)]
    lines += [
        f"width {key} {dim.width}"
        for key, dim in sorted(ws.column_dimensions.items())
        if dim.width is not None
    ]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            fill, font = cell.fill, cell.font
            lines.append(
                json.dumps(
                    [
                        cell.coordinate,
                        _value(cell.value),
                        [
                            getattr(fill, "fill_type", None),
                            _color(getattr(fill, "start_color", None)),
                            _color(getattr(fill, "end_color", None)),
                        ],
                        [font.name, font.sz, font.b, font.i, font.u, _color(font.color)],
                    ],
                    ensure_ascii=False,
                )
            )
    return lines


def _digest(lines: list[str]) -> str:
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()[:16]


def export_lines(scenario_id: int, search: str) -> dict[str, Any]:
    """One export: ``{"name", "files": {file: {sheet: lines}}}`` (files and
    sheets in order), or ``{"error": message}`` when the export refuses."""
    from openpyxl import load_workbook

    from core.services.timetable_per_plan_export import export_scenario_per_plan

    try:
        path, name, is_zip = export_scenario_per_plan(scenario_id, search=search)
    except RuntimeError as exc:
        return {"error": str(exc)}
    try:
        data = Path(path).read_bytes()
    finally:
        Path(path).unlink(missing_ok=True)
    if is_zip:
        with zipfile.ZipFile(io.BytesIO(data)) as bundle:
            members = [(member, bundle.read(member)) for member in bundle.namelist()]
    else:
        members = [(name, data)]
    files: dict[str, dict[str, list[str]]] = {}
    for member, content in members:
        workbook = load_workbook(io.BytesIO(content))
        files[member] = {ws.title: sheet_lines(ws) for ws in workbook.worksheets}
    return {"name": name, "files": files}


def digests_of(lines: dict[str, Any]) -> dict[str, Any]:
    """``export_lines`` with each sheet's lines reduced to its digest."""
    if "error" in lines:
        return lines
    return {
        "name": lines["name"],
        "files": {
            member: {title: _digest(sheet) for title, sheet in sheets.items()}
            for member, sheets in lines["files"].items()
        },
    }


def scenario_exports(scenario_id: int) -> dict[str, dict[str, Any]]:
    """Every search's ``export_lines`` for one scenario."""
    return {search: export_lines(scenario_id, search) for search in SEARCHES}


def first_difference(recorded: Any, now: Any, path: tuple[str, ...] = ()) -> str | None:
    """The path of the first place two digest trees differ, or ``None``."""
    if isinstance(recorded, dict) and isinstance(now, dict):
        if list(recorded) != list(now):
            return f"{' / '.join(path) or '(top)'}: keys {list(recorded)} != {list(now)}"
        for key in recorded:
            found = first_difference(recorded[key], now[key], (*path, repr(key)))
            if found:
                return found
        return None
    if recorded != now:
        return f"{' / '.join(path)}: recorded {recorded!r}, now {now!r}"
    return None


# ── the scenarios ───────────────────────────────────────────────────────────


def _grouping():
    from tests import test_per_plan_export_grouping as grouping

    return grouping


def _electives():
    from tests import test_per_plan_export_electives as electives

    return electives


def _every_main_group(monkeypatch, settings):
    g = _grouping()
    scenario, [board] = g._scenario()
    cs = g._course(board, "CS112", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"))
    ai = g._course(board, "AI236", ("S1", "TUE", "09:00"))
    is_ = g._course(board, "IS124", ("S1", "WED", "09:00"))
    g._seat(
        monkeypatch,
        {
            cs["S1"]: g._row(AI2=20),
            cs["S2"]: g._row(AI2=20),
            ai["S1"]: g._row(AI2=30),
            is_["S1"]: g._row(AI2=30),
        },
    )
    return scenario


def _group_count(monkeypatch, settings):
    g = _grouping()
    scenario, [board] = g._scenario()
    cs = g._course(
        board, "CS112", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"), ("S3", "TUE", "09:00")
    )
    ai = g._course(board, "AI236", ("S1", "WED", "09:00"), ("S2", "THU", "09:00"))
    g._seat(
        monkeypatch,
        {
            cs["S1"]: g._row(AI2=28, DS2=2),
            cs["S2"]: g._row(AI2=2, DS2=28),
            cs["S3"]: g._row(AI2=1, DS2=29),
            ai["S1"]: g._row(AI2=25),
            ai["S2"]: g._row(AI2=24),
        },
    )
    return scenario


def _scenario_38(monkeypatch, settings):
    return _grouping()._scenario_38_like(monkeypatch)[0]


def _scenario_38_two_courses(monkeypatch, settings):
    return _grouping()._scenario_38_like(monkeypatch, courses=("AI236", "IS124"))[0]


def _allocated_elsewhere(monkeypatch, settings):
    g = _grouping()
    scenario, [board] = g._scenario()
    cs = g._course(board, "CS112", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"))
    engl = g._course(board, "ENGL214", ("S1", "TUE", "09:00"))
    g._seat(
        monkeypatch,
        {
            cs["S1"]: g._row(plans=("AI2",), AI2=28, DS2=2),
            cs["S2"]: g._row(plans=("DS2",), AI2=2, DS2=28),
            engl["S1"]: g._row(plans=("DS2",), AI2=3, DS2=27),
        },
    )
    return scenario


def _mixed_cell(monkeypatch, settings):
    from core.models import SectionPlacement, TermSectionMeeting

    scenario, sections = _grouping()._scenario_38_like(monkeypatch)
    engl = sections["ENGL214 S1"]
    SectionPlacement.objects.filter(term_section=engl).update(room="B12")
    TermSectionMeeting.objects.create(
        term_section=engl, day="THU", start_time="09:00", end_time="10:15", instructor="Dr Huda"
    )
    return scenario


def _each_programme_file(monkeypatch, settings):
    g = _grouping()
    scenario, [board] = g._scenario()
    ai = g._course(board, "AI236", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"))
    g._seat(monkeypatch, {ai["S1"]: g._row(AI2=1, DS2=11), ai["S2"]: g._row(AI2=11, DS2=1)})
    return scenario


def _cs112_two_boards(monkeypatch, settings):
    return _grouping()._cs112_on_two_boards(monkeypatch)


def _cs112_two_boards_room(monkeypatch, settings):
    return _grouping()._cs112_on_two_boards(monkeypatch, s2_room="B12")


def _real_lens_two_boards(monkeypatch, settings):
    g = _grouping()
    scenario, (term_2, term_3) = g._scenario(2, 3)
    g._course(term_2, "AI236", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"))
    g._course(term_3, "CS112", ("S1", "TUE", "09:00"), ("S2", "WED", "09:00"))
    return scenario


def _marked(mark_on: bool, *, is124_marked: bool = False, unmarked: bool = False):
    def build(monkeypatch, settings):
        from core.models import TermSection

        settings.TIMETABLE_SECTION_PROGRAMME_PREFERENCE = mark_on
        scenario, _placements = _grouping()._two_programme_marked_scenario(mark="AI2")
        if is124_marked:
            TermSection.objects.filter(scenario=scenario, course_key="IS124").update(
                preferred_programme="AI2"
            )
        if unmarked:
            TermSection.objects.filter(scenario=scenario).update(preferred_programme="")
        return scenario

    return build


def _off_grid(monkeypatch, settings):
    from core.models import SectionPlacement, TermSection

    g = _grouping()
    scenario, [board] = g._scenario()
    sections = g._course(board, "CS112", ("S1", "SUN", "09:00"))
    off = TermSection.objects.create(
        scenario=scenario,
        course_code="CS112",
        course_number="CS112",
        course_key="CS112",
        course_name="CS112",
        section="S2",
        source_tag="tw_auto",
    )
    SectionPlacement.objects.create(
        board=board, term_section=off, day="MON", start_time="18:30", end_time="19:45"
    )
    g._seat(monkeypatch, {sections["S1"]: g._row(AI2=20), off: g._row(AI2=20)})
    return scenario


def _electives_scenario(build: Callable[[Any], Any]):
    def scenario(monkeypatch, settings):
        return build(_electives())

    return scenario


def _e_mapped_on_slot_term(e):
    e._plan("AI", [("AI401", 10), ("AI2", 10), ("AI3", 10)])
    e._map("AI2", "AI441")
    e._map("AI3", "AI462")
    scenario, board = e._scenario()
    e._place(board, "AI401", "SUN", "09:00")
    e._place(board, "AI441", "MON", "10:30")
    e._place(board, "AI462", "TUE", "13:00")
    return scenario


def _e_slot_row_covered(e):
    e._plan("AI", [("AI2", 10), ("AI3", 10)])
    e._map("AI2", "AI441")
    e._map("AI3", "AI462")
    scenario, board = e._scenario()
    e._place(board, "AI441")
    return scenario


def _e_without_mapping(e, *, other_term: bool = False):
    e._plan("AI", [("AI401", 10), ("AI2", 10)])
    if other_term:
        e._map("AI2", "AI441", term=1)
    scenario, board = e._scenario()
    e._place(board, "AI401", "SUN", "09:00")
    e._place(board, "AI441", "MON", "10:30")
    return scenario


def _e_second_cohort(e):
    e._plan("AI2", [("AI201", 9), ("AI2", 9)])
    e._map("AI2", "AI441", programme="AI")
    scenario, board = e._scenario(program="AI2", board_term=9)
    e._place(board, "AI201", "SUN", "09:00")
    e._place(board, "AI441", "MON", "10:30")
    return scenario


def _e_two_slots(e):
    e._plan("DS", [("DS1", 8), ("DS3", 10)])
    e._map("DS1", "DS482", programme="DS")
    e._map("DS3", "DS482", programme="DS")
    return e._scenario(program="DS", board_term=8)[0]


def _e_later_slot_board(e):
    from core.models import DeliveryBoard

    e._plan("DS", [("DS101", 8), ("DS1", 8), ("DS3", 10)])
    e._map("DS1", "DS482", programme="DS")
    e._map("DS3", "DS482", programme="DS")
    scenario, board_8 = e._scenario(program="DS", board_term=8)
    board_10 = DeliveryBoard.objects.create(
        scenario=scenario, label="Term 10", nominal_term=10, program="DS"
    )
    e._place(board_8, "DS101", "SUN", "09:00")
    e._place(board_10, "DS482", "MON", "10:30")
    return scenario


def _e_search_slot(e):
    e._plan("AI", [("AI401", 10), ("AI2", 10)], names={"AI2": "PROGRAM ELECTIVE COURSE II"})
    e._map("AI2", "AI441")
    scenario, board = e._scenario()
    e._place(board, "AI401", "SUN", "09:00")
    e._place(board, "AI441", "MON", "10:30")
    return scenario


def _e_course_in_plan(e, names: dict[str, str] | None = None):
    e._plan("AI", [("AI2", 10), ("AI441", 7)], names=names)
    e._map("AI2", "AI441")
    return e._scenario()[0]


def _e_only_scenario_courses(e):
    e._plan("AI", [("AI2", 10)])
    e._map("AI2", "AI441")
    e._map("AI2", "AI499")
    return e._scenario()[0]


def _e_board_none_of_slots(e):
    e._plan("AI", [("AI2", 10)])
    e._map("AI2", "AI441")
    scenario, board_8 = e._scenario(board_term=8)
    e._place(board_8, "AI441", "MON", "10:30")
    return scenario


def _section(scenario, code: str, label: str):
    from core.models import TermSection

    return TermSection.objects.create(
        scenario=scenario,
        course_code=code,
        course_number=code,
        course_key=code,
        course_name=code,
        section=label,
        source_tag="tw_auto",
    )


def _shared_roles(monkeypatch, settings):
    """SHARED changes roles, not only membership: CS112 S2 (shared) is AI2's
    other section, and its mixed one once SHARED leaves it alone; ENGL214 S1
    (shared) is mixed for both; IS124 S1 (shared, nobody seated, allocated
    to AI2) is AI2's main and in no DS2 sheet under SHARED. MATH101 is in no
    plan: M1 allocated to AI2 (not shared), M2 unallocated (shared); MATH103,
    in no plan either, has one section allocated to AI2, not shared, so
    SHARED drops it from the coverage sheet. MATH102 is another scenario's
    section placed here: no lens row."""
    from core.models import SectionPlacement, TimetableScenario

    g = _grouping()
    scenario, [board] = g._scenario()
    cs = g._course(board, "CS112", ("S1", "SUN", "09:00"), ("S2", "MON", "09:00"))
    engl = g._course(board, "ENGL214", ("S1", "TUE", "09:00"), ("S2", "TUE", "10:30"))
    is_ = g._course(board, "IS124", ("S1", "WED", "09:00"), ("S2", "WED", "10:30"))
    outside = {}
    for code, label, day, start in (
        ("MATH101", "M1", "THU", "09:00"),
        ("MATH101", "M2", "THU", "10:30"),
        ("MATH103", "M3", "SUN", "10:30"),
    ):
        outside[label] = _section(scenario, code, label)
        SectionPlacement.objects.create(
            board=board,
            term_section=outside[label],
            day=day,
            start_time=start,
            end_time=g.END[start],
        )
    elsewhere = TimetableScenario.objects.create(academic_year="1448", term="2", name="Elsewhere")
    SectionPlacement.objects.create(
        board=board,
        term_section=_section(elsewhere, "MATH102", "S1"),
        day="SUN",
        start_time="13:00",
        end_time="14:15",
    )

    def shared(row: dict) -> dict:
        return {**row, "shared": True}

    g._seat(
        monkeypatch,
        {
            cs["S1"]: g._row(AI2=28),
            cs["S2"]: shared(g._row(AI2=2, DS2=28)),
            engl["S1"]: shared(g._row(AI2=18, DS2=11)),
            engl["S2"]: g._row(AI2=4, DS2=26),
            is_["S1"]: shared(g._row(plans=("AI2",))),
            is_["S2"]: g._row(plans=("DS2",), DS2=25),
            outside["M1"]: g._row(plans=("AI2",), AI2=3),
            outside["M2"]: shared(g._row(AI2=1, DS2=1)),
            outside["M3"]: g._row(plans=("AI2",), AI2=5),
        },
    )
    return scenario


def _elective_two_sheets(monkeypatch, settings):
    """DS482 fills DS1 (T8) and DS3 (T10): S1 on the T8 board seats 10 DS of
    30, alone on its sheet (mixed there); S2 on the T10 board seats 25 DS (main
    there). Judged as one course, S1 would be S2's other section."""
    from core.models import DeliveryBoard, ScenarioSectionBudget, SectionPlacement

    e, g = _electives(), _grouping()
    e._plan("DS", [("DS101", 8), ("DS1", 8), ("DS3", 10)])
    e._map("DS1", "DS482", programme="DS")
    e._map("DS3", "DS482", programme="DS")
    scenario, board_8 = e._scenario(program="DS", board_term=8)
    board_10 = DeliveryBoard.objects.create(
        scenario=scenario, label="Term 10", nominal_term=10, program="DS", display_order=1
    )
    e._place(board_8, "DS101", "SUN", "09:00")
    ScenarioSectionBudget.objects.create(
        scenario=scenario,
        course_code="DS482",
        course_key="DS482",
        department="DS",
        programme_term=10,
        credit_hours=3,
        planned_sections=2,
        max_per_section=40,
        total_demand=55,
    )
    s1, s2 = _section(scenario, "DS482", "S1"), _section(scenario, "DS482", "S2")
    SectionPlacement.objects.create(
        board=board_8, term_section=s1, day="MON", start_time="10:30", end_time="11:45"
    )
    SectionPlacement.objects.create(
        board=board_10, term_section=s2, day="TUE", start_time="10:30", end_time="11:45"
    )
    g._seat(monkeypatch, {s1: g._row(DS=10, AI=20), s2: g._row(DS=25)})
    return scenario


SCENARIOS: dict[str, Callable[[Any, Any], Any]] = {
    "grouping/every_main_group": _every_main_group,
    "grouping/group_count": _group_count,
    "grouping/scenario_38": _scenario_38,
    "grouping/scenario_38_two_courses": _scenario_38_two_courses,
    "grouping/allocated_elsewhere": _allocated_elsewhere,
    "grouping/mixed_cell": _mixed_cell,
    "grouping/each_programme_file": _each_programme_file,
    "grouping/cs112_two_boards": _cs112_two_boards,
    "grouping/cs112_two_boards_room": _cs112_two_boards_room,
    "grouping/real_lens_two_boards": _real_lens_two_boards,
    "grouping/marked_flag_on": _marked(True),
    "grouping/marked_flag_on_is124_marked": _marked(True, is124_marked=True),
    "grouping/marked_flag_off": _marked(False),
    "grouping/marked_flag_off_unmarked": _marked(False, unmarked=True),
    "slot_config/off_grid": _off_grid,
    "electives/mapped_on_slot_term": _electives_scenario(_e_mapped_on_slot_term),
    "electives/slot_row_covered": _electives_scenario(_e_slot_row_covered),
    "electives/without_mapping": _electives_scenario(_e_without_mapping),
    "electives/mapping_other_term": _electives_scenario(
        lambda e: _e_without_mapping(e, other_term=True)
    ),
    "electives/second_cohort": _electives_scenario(_e_second_cohort),
    "electives/two_slots": _electives_scenario(_e_two_slots),
    "electives/later_slot_board": _electives_scenario(_e_later_slot_board),
    "electives/search_slot": _electives_scenario(_e_search_slot),
    "electives/course_in_plan": _electives_scenario(_e_course_in_plan),
    "electives/plan_course_by_code": _electives_scenario(
        lambda e: _e_course_in_plan(e, names={"AI441": "Computer Vision"})
    ),
    "electives/only_scenario_courses": _electives_scenario(_e_only_scenario_courses),
    "electives/board_none_of_slots": _electives_scenario(_e_board_none_of_slots),
    "extra/shared_roles": _shared_roles,
    "extra/elective_two_sheets": _elective_two_sheets,
}


def apply_flags(settings) -> None:
    for name, value in FLAGS.items():
        setattr(settings, name, value)


def build(name: str, monkeypatch, settings) -> int:
    """Build one scenario with the oracle's flags; its id."""
    apply_flags(settings)
    return SCENARIOS[name](monkeypatch, settings).id


# ── the dump diff ───────────────────────────────────────────────────────────


def _first_line_difference(base: dict, branch: dict) -> str:
    for name in base:
        for search in base[name]:
            left, right = base[name][search], branch.get(name, {}).get(search)
            if left == right:
                continue
            if right is None or "error" in left or "error" in right:
                return f"{name} / {search!r}: {left.get('error')!r} vs {right!r}"
            for member, sheets in left["files"].items():
                for title, lines in sheets.items():
                    other = right["files"].get(member, {}).get(title)
                    if other == lines:
                        continue
                    if other is None:
                        return f"{name} / {search!r} / {member} / {title}: missing on the branch"
                    for idx, (a, b) in enumerate(zip(lines, other, strict=False)):
                        if a != b:
                            return f"{name} / {search!r} / {member} / {title} line {idx}:\n  base   {a}\n  branch {b}"
                    return f"{name} / {search!r} / {member} / {title}: {len(lines)} vs {len(other)} lines"
            return (
                f"{name} / {search!r}: file names {list(left['files'])} vs {list(right['files'])}"
            )
    return "no difference"


if __name__ == "__main__":
    base_dump, branch_dump = (
        json.loads(Path(arg).read_text(encoding="utf-8")) for arg in sys.argv[1:3]
    )
    print(_first_line_difference(base_dump, branch_dump))
