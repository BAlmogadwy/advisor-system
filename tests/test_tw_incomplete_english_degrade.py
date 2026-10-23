"""F17 degrades at the source (stage 3, B2; plan D-4, evidence N-3, N-9).

With one teacher per section on (production), an English section whose board
copies do not share one complete weekly pattern used to raise out of the
one-teacher map: the board read, every check a manual edit runs, and both
exports answered a raw 500 (QA 662 and 664). Now a display read collects that
section as ``unknown`` (a dash and a reason) and every other section keeps its
teacher; the exports print a dash and say why on a first sheet in English and
Arabic; an export that still fails answers JSON in words. Every engine still
raises.

All under the three production flags (render.yaml), as the QA probes ran.
"""

from __future__ import annotations

import io
import json
import logging
import zipfile

import pytest

from core.models import (
    SectionPlacement,
    TermSection,
    TermSectionMeeting,
)
from core.services import section_teacher_assignment as sta
from core.services.timetable_optimizer_v2 import (
    IncompletePolicyPattern,
    build_section_states_for_scenario,
)
from tests.one_teacher_support import (
    budget,
    engine_rows,
    instructor,
    link,
    requirement,
    scenario,
    section,
)
from tests.tw_incomplete_english_support import PM
from tests.tw_incomplete_english_support import IncompleteEnglishScenario as Fx
from tests.tw_scenario_seating_support import PRODUCTION

#: What an export prints for a teacher it does not show (an em dash).
DASH = chr(0x2014)
UNKNOWN = {
    "id": None,
    "name": None,
    "status": "unknown",
    "source": None,
    "reason": "incomplete_english_pattern",
}


@pytest.fixture
def production(settings):
    for name, value in PRODUCTION.items():
        setattr(settings, name, value)
    return settings


def _admin_client():
    from tests.test_one_teacher_pipeline import _admin_client as client

    return client()


def _teachers_on(client, board) -> dict[int, list]:
    response = client.get(f"/ops/tw/boards/{board.id}/")
    assert response.status_code == 200, response.content
    out: dict[int, list] = {}
    for p in response.json()["placements"]:
        out.setdefault(p["term_section_id"], []).append(p["teacher"])
    return out


def _one(rows: list) -> dict:
    assert rows and all(row == rows[0] for row in rows), rows
    return rows[0]


# ── The board read: the broken sections are unknown, the rest keep theirs ──


@pytest.mark.django_db
def test_the_board_read_answers_200_and_degrades_only_the_broken_sections(production) -> None:
    fx = Fx()
    client = _admin_client()

    term2, term4 = _teachers_on(client, fx.b1), _teachers_on(client, fx.b2)

    assert _one(term2[fx.s1.id]) == UNKNOWN
    assert _one(term4[fx.s1.id]) == UNKNOWN
    assert _one(term2[fx.legacy.id]) == UNKNOWN, "never the registrar's meeting name"
    assert _one(term2[fx.s2.id])["name"] == "Dr E"
    assert _one(term2[fx.a1.id])["name"] == "Dr A"
    assert _one(term2[fx.a2.id])["name"] == "Dr A"
    # Dr E holds S1 (unknown) and S2: the English cap leaves S3 to be decided.
    assert _one(term4[fx.s3.id]) == {
        "id": None,
        "name": None,
        "status": "tbd",
        "source": None,
        "reason": "d10",
    }


@pytest.mark.django_db
def test_a_consistent_scenario_reads_as_before(production) -> None:
    """No unknown section, no notice: the degrade changes nothing else."""
    fx = Fx(broken=False)

    rows = sta.section_teacher_display(fx.sc.id)

    assert {row["status"] for row in rows.values()} <= {"engine", "tbd"}
    assert rows[fx.s1.id]["name"] == "Dr E"
    assert sta.load_assignment(fx.sc.id).incomplete == ()


@pytest.mark.django_db
def test_reads_write_nothing(production) -> None:
    fx = Fx()
    client = _admin_client()
    rows, names = engine_rows(fx.sc), _meeting_names(fx.sc)

    _teachers_on(client, fx.b1)
    client.get(f"/ops/tw/boards/{fx.b1.id}/conflicts/")
    client.get(f"/ops/tw/scenarios/{fx.sc.id}/readiness/")

    assert engine_rows(fx.sc) == rows
    assert _meeting_names(fx.sc) == names


def _meeting_names(sc) -> dict[int, set[str]]:
    out: dict[int, set[str]] = {}
    for ts_id, name in TermSectionMeeting.objects.filter(term_section__scenario=sc).values_list(
        "term_section_id", "instructor"
    ):
        out.setdefault(ts_id, set()).add(name)
    return out


# ── Every check a manual edit runs, and the edits themselves ─────────────


@pytest.mark.django_db
@pytest.mark.parametrize(
    "tail", ["conflicts/", "summary/", "capacity/", "unplaced/"], ids=lambda t: t.strip("/")
)
def test_the_boards_other_reads_answer_200(production, tail) -> None:
    fx = Fx()
    client = _admin_client()

    for board in (fx.b1, fx.b2):
        assert client.get(f"/ops/tw/boards/{board.id}/{tail}").status_code == 200


@pytest.mark.django_db
def test_readiness_answers_200_and_publishing_stays_blocked(production) -> None:
    fx = Fx()

    response = _admin_client().get(f"/ops/tw/scenarios/{fx.sc.id}/readiness/")

    assert response.status_code == 200
    readiness = response.json()["readiness"]
    assert readiness["ready"] is False
    invalid = {
        (item["params"]["course_code"], item["params"]["section"])
        for item in readiness["blocker_items"]
        if item["code"] == "invalid_section_pattern"
    }
    assert invalid == {("ENG102", "S1"), ("ENG101", "S1")}


@pytest.mark.django_db
def test_the_move_and_room_previews_answer_200(production) -> None:
    fx = Fx()
    client = _admin_client()
    english = SectionPlacement.objects.filter(term_section=fx.s1).first()
    plain = SectionPlacement.objects.get(term_section=fx.a1)

    for placement in (english, plain):
        for tail in ("slot-candidates/?verdict=1&offer=manual", "room-candidates/?audit=1"):
            response = client.get(f"/ops/tw/placements/{placement.id}/{tail}")
            assert response.status_code == 200, (tail, response.content[:200])


@pytest.mark.django_db
def test_create_planned_answers_201_and_leaves_the_broken_sections_as_they_are(
    production, caplog
) -> None:
    caplog.set_level(logging.WARNING)
    fx = Fx()
    other = requirement("AI", "AI202", "OTHER COURSE")
    link("AI", "AI202", instructor("Dr C"))
    budget(fx.sc, other, 1)
    names = _meeting_names(fx.sc)
    before = SectionPlacement.objects.filter(board__scenario=fx.sc).count()

    response = _admin_client().post(
        "/ops/tw/placements/create-planned/",
        data=json.dumps(
            {
                "board_id": fx.b1.id,
                "course_code": "AI202",
                "course_key": other,
                "course_name": "OTHER COURSE",
                "section_label": "S1",
                "capacity": 30,
                "meetings": [{"day": "TUE", "start_time": "13:00", "end_time": "14:15"}],
            }
        ),
        content_type="application/json",
    )

    assert response.status_code == 201, response.content
    assert response.json()["placement"]["teacher"]["name"] == "Dr C"
    # The clash check after the write could not read the broken section's
    # meetings: it says so (never a clean board), without a traceback.
    backstop = response.json()["instructor_clash_backstop"]
    assert (backstop["checked"], backstop["reason"]) == (False, "incomplete_english_pattern")
    assert not [r for r in caplog.records if r.name.endswith("instructor_backstop") and r.exc_info]
    assert SectionPlacement.objects.filter(board__scenario=fx.sc).count() == before + 1
    rows = engine_rows(fx.sc)
    # The broken section's stored teacher is kept (not re-picked, not dropped).
    assert rows[(fx.eng.upper(), "S1")] == "Dr E"
    assert rows[(fx.eng.upper(), "S2")] == "Dr E"
    assert rows[(other.upper(), "S1")] == "Dr C"
    after = _meeting_names(fx.sc)
    for ts in (fx.s1, fx.legacy, fx.s2):
        assert after[ts.id] == names[ts.id], ts


@pytest.mark.django_db
def test_create_move_and_switch_answer_in_words(production) -> None:
    fx = Fx()
    client = _admin_client()
    spare = section(fx.sc, fx.ai, "S3")
    TermSectionMeeting.objects.create(
        term_section=spare, day="TUE", start_time="13:00", end_time="14:15", room="R1"
    )

    created = client.post(
        "/ops/tw/placements/create/",
        data=json.dumps(
            {
                "board_id": fx.b1.id,
                "term_section_id": spare.id,
                "day": "TUE",
                "start_time": "13:00",
                "end_time": "14:15",
                "room": "R1",
            }
        ),
        content_type="application/json",
    )
    assert created.status_code == 201, created.content
    assert created.json()["placement"]["teacher"]["name"] == "Dr A"

    english = TermSection.objects.create(
        scenario=fx.sc, course_code="ENG102", course_key="ENG102", section="M9"
    )
    bundle = client.post(
        "/ops/tw/placements/create/",
        data=json.dumps(
            {
                "board_id": fx.b2.id,
                "term_section_id": english.id,
                "meetings": [{"day": d, "start_time": s, "end_time": e} for d, s, e in PM],
            }
        ),
        content_type="application/json",
    )
    assert bundle.status_code == 201, bundle.content

    plain = SectionPlacement.objects.get(term_section=fx.a1)
    moved = client.post(
        f"/ops/tw/placements/{plain.id}/move/",
        data=json.dumps({"day": "TUE", "start_time": "08:00", "end_time": "09:15"}),
        content_type="application/json",
    )
    assert moved.status_code == 200, moved.content
    assert moved.json()["placement"]["teacher"]["name"] == "Dr A"

    s2 = SectionPlacement.objects.filter(term_section=fx.s2).first()
    switched = client.post(
        f"/ops/tw/placements/{s2.id}/switch-pattern/",
        data=json.dumps({"pattern_id": "ENG_INTENSIVE_V1_AM"}),
        content_type="application/json",
    )
    assert switched.status_code == 200, switched.content
    assert {p["teacher"]["status"] for p in switched.json()["placements"]} == {"engine"}

    # Switching the broken section itself is a refusal in words, as before.
    s1 = SectionPlacement.objects.filter(term_section=fx.s1).first()
    refused = client.post(
        f"/ops/tw/placements/{s1.id}/switch-pattern/",
        data=json.dumps({"pattern_id": "ENG_INTENSIVE_V1_AM"}),
        content_type="application/json",
    )
    assert refused.status_code == 400
    assert refused.json()["error"]["code"] == "COURSE_SECTION_PATTERN_REQUIRED"


@pytest.mark.django_db
def test_student_blockers_refuse_in_words(production) -> None:
    fx = Fx()

    response = _admin_client().get(f"/ops/tw/scenarios/{fx.sc.id}/student-blockers/")

    assert response.status_code == 409
    error = response.json()["error"]
    assert error["code"] == "INCOMPLETE_ENGLISH_PATTERN"
    assert "ENG102 S1" in error["message"]
    assert error["details"]["board_ids"] == [fx.b1.id, fx.b2.id]


# ── The exports: a dash, the notice, and words for a failure ───────────────


def _workbooks(response) -> list:
    from openpyxl import load_workbook

    data = b"".join(response.streaming_content)
    if response["Content-Type"] == "application/zip":
        with zipfile.ZipFile(io.BytesIO(data)) as bundle:
            return [load_workbook(io.BytesIO(bundle.read(name))) for name in bundle.namelist()]
    return [load_workbook(io.BytesIO(data))]


def _grid_blocks(wb) -> list[list[str]]:
    """Each section's lines in each grid cell of the term sheets (a cell holds
    several sections apart by ``---``); the code and section first."""
    return [
        block.split("\n")
        for ws in wb.worksheets
        if ws.title.startswith("Term")
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and "\n" in cell.value
        for block in cell.value.split("\n---\n")
    ]


def _teacher_lines(books, code_section: str) -> set[str]:
    """The teacher line of every grid block of ``code_section`` in ``books``:
    the first line after the code that is not an English time line."""
    out = set()
    for wb in books:
        for lines in _grid_blocks(wb):
            if lines[0].split(" (on T")[0] == code_section:
                rest = [line for line in lines[1:] if "lecture" not in line]
                out.add(rest[0] if rest else "")
    return out


@pytest.mark.django_db
@pytest.mark.parametrize("kind", ["export.xlsx", "export-per-plan"])
def test_both_exports_answer_200_with_dashes_and_the_notice(production, kind) -> None:
    fx = Fx()

    response = _admin_client().get(f"/ops/tw/scenarios/{fx.sc.id}/{kind}")

    assert response.status_code == 200
    books = _workbooks(response)
    assert books
    for wb in books:
        assert wb.sheetnames[0] == "Notice" and wb.active.title == "Notice"
        en, ar = wb["Notice"]["A1"].value, wb["Notice"]["A3"].value
        assert en.startswith(
            "Teachers not shown for: ENG101 S1 (Term 2), ENG102 S1 (Term 2, Term 4). "
            "These English sections do not have one complete weekly pattern on every board."
        ), en
        assert "\u200eENG102 S1\u200f (\u200eTerm 2\u200f\u060c \u200eTerm 4\u200f)" in ar
        assert ar.startswith("لم يُعرض المدرّسان لـ:"), ascii(ar)
        instructors = [
            str(cell.value)
            for row in wb["Instructors"].iter_rows()
            for cell in row
            if cell.value is not None
        ]
        assert not any("\u2014" == text for text in instructors), "a dash is nobody"
        assert not any("ENG102 S1" in text or "ENG101 S1" in text for text in instructors)
        assert any("ENG102 S2" in text for text in instructors)
    # Each section's teacher line on the grid (ENG101 is in no plan of the
    # programme, so only the combined workbook draws it).
    assert _teacher_lines(books, "ENG102 S1") == {DASH}
    assert _teacher_lines(books, "ENG101 S1") == ({DASH} if kind == "export.xlsx" else set())
    assert _teacher_lines(books, "AI201 S1") == {"Dr A"}
    assert _teacher_lines(books, "ENG102 S2") == {"Dr E"}
    assert _teacher_lines(books, "ENG102 S3") == {"TBD"}


@pytest.mark.django_db
@pytest.mark.parametrize("kind", ["export.xlsx", "export-per-plan"])
def test_a_consistent_scenarios_exports_have_no_notice(production, kind) -> None:
    fx = Fx(broken=False)

    response = _admin_client().get(f"/ops/tw/scenarios/{fx.sc.id}/{kind}")

    assert response.status_code == 200
    for wb in _workbooks(response):
        assert "Notice" not in wb.sheetnames


def test_the_notice_agrees_in_number_in_both_languages() -> None:
    from core.services.timetable_export import teacher_notice_lines

    row = {"course_code": "ENG102", "section": "S1", "boards": ["Term 2"]}
    one = teacher_notice_lines([row])
    two = teacher_notice_lines([row, {**row, "section": "S2"}])
    many = teacher_notice_lines([row, {**row, "section": "S2"}, {**row, "section": "S3"}])
    assert one[0].startswith("Teacher not shown for: ENG102 S1 (Term 2). This English section")
    assert one[1].startswith("لم يُعرض المدرّس لـ:") and "احذفها" in one[1]
    assert two[0].startswith("Teachers not shown for:") and "These English sections" in two[0]
    assert two[1].startswith("لم يُعرض المدرّسان لـ:") and "هاتين الشعبتين" in two[1]
    assert many[1].startswith("لم يُعرض المدرّسون لـ:") and "هذه الشعب" in many[1]
    # Codes and labels stay in reading order inside the Arabic sentence.
    assert "\u200eENG102 S3\u200f (\u200eTerm 2\u200f)" in many[1]
    assert not any(ch in many[1] for ch in "\u2066\u2067\u2068\u2069"), "Excel paints isolates"


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("url", "target", "error"),
    [
        ("export.xlsx", "core.services.timetable_export.export_scenario_xlsx", KeyError("x")),
        (
            "export-per-plan",
            "core.services.timetable_per_plan_export.export_scenario_per_plan",
            ValueError("x"),
        ),
    ],
)
def test_an_export_that_fails_answers_json_in_words_and_is_logged(
    production, monkeypatch, caplog, url, target, error
) -> None:
    fx = Fx(broken=False)

    def boom(*_args, **_kwargs):
        raise error

    module, name = target.rsplit(".", 1)
    monkeypatch.setattr(f"{module}.{name}", boom)
    with caplog.at_level(logging.ERROR, logger="core.timetable_workspace_views"):
        response = _admin_client().get(f"/ops/tw/scenarios/{fx.sc.id}/{url}")

    assert response.status_code == 500
    assert response["Content-Type"] == "application/json"
    body = response.json()
    assert body["error"]["code"] == "EXPORT_FAILED"
    assert body["error"]["message"].startswith("The export could not be built.")
    assert any(r.exc_info for r in caplog.records), "logged with its traceback"


@pytest.mark.django_db
def test_the_per_plan_exports_own_refusals_keep_their_words(production) -> None:
    sc, _board = scenario(board_program="")

    response = _admin_client().get(f"/ops/tw/scenarios/{sc.id}/export-per-plan")

    assert response.status_code == 500
    assert response.json()["error"] == {
        "code": "EXPORT_FAILED",
        "message": "Scenario has no boards with a configured program",
    }


# ── The engine path still raises ─────────────────────────────────────────


@pytest.mark.django_db
def test_every_engine_read_still_raises(production) -> None:
    from core.services.timetable_pr4_instructor import build_section_instructor_ids

    fx = Fx()

    with pytest.raises(IncompletePolicyPattern) as raised:
        build_section_states_for_scenario(fx.sc.id)
    assert raised.value.term_section_id == fx.s1.id
    assert raised.value.board_ids == (fx.b1.id, fx.b2.id)
    assert (raised.value.course_code, raised.value.section) == ("ENG102", "S1")
    assert (
        str(raised.value) == "English board copies must each use the same complete weekly pattern"
    )
    for engine in (
        lambda: sta.load_assignment(fx.sc.id),
        lambda: build_section_instructor_ids(fx.sc),
        lambda: sta.prepare_section_teachers(fx.sc.id, phase="layout"),
        lambda: sta.prepare_section_teachers(fx.sc.id, phase="ensure"),
    ):
        with pytest.raises(IncompletePolicyPattern):
            engine()


@pytest.mark.django_db
def test_collect_leaves_out_only_the_broken_sections(production) -> None:
    fx = Fx()
    found: list = []

    states = build_section_states_for_scenario(fx.sc.id, incomplete="collect", collected=found)

    assert {e.term_section_id for e in found} == {fx.s1.id, fx.legacy.id}
    assert {s.section_id for s in states} == {
        f"{fx.eng}_S2",
        f"{fx.eng}_S3",
        f"{fx.ai}_S1",
        f"{fx.ai}_S2",
    }
    with pytest.raises(TypeError):
        build_section_states_for_scenario(fx.sc.id, incomplete="collect")
    with pytest.raises(ValueError, match="incomplete must be one of"):
        build_section_states_for_scenario(fx.sc.id, incomplete="skip")


@pytest.mark.django_db
def test_a_different_value_error_still_surfaces(production, monkeypatch) -> None:
    """Only the incomplete-pattern error is collected; any other is a bug."""
    from core.services import timetable_optimizer_v2 as v2

    fx = Fx(broken=False)

    def boom(*_args, **_kwargs):
        raise ValueError("a different failure")

    monkeypatch.setattr(v2, "_validated_policy_placement_pattern", boom)
    with pytest.raises(ValueError, match="a different failure"):
        sta.section_teacher_display(fx.sc.id)


@pytest.mark.django_db
def test_one_requests_memo_serves_both_modes(production) -> None:
    fx = Fx()

    with sta.request_scope():
        collected = sta.load_assignment(fx.sc.id, incomplete="collect")
        assert collected.status[f"{fx.eng}_S1"] == "unknown"
        with pytest.raises(IncompletePolicyPattern):
            sta.load_assignment(fx.sc.id)


@pytest.mark.django_db
def test_no_persist_drops_an_unknown_sections_row(production) -> None:
    """Even a full persist (a run's) keeps the row a display read could not
    check: the section's meetings are not known."""
    fx = Fx()
    assignment = sta.load_assignment(fx.sc.id, incomplete="collect")

    sta.persist_assignment(fx.sc.id, assignment, only_missing=False, incomplete="collect")

    assert engine_rows(fx.sc)[(fx.eng.upper(), "S1")] == "Dr E"


@pytest.mark.django_db
def test_a_held_row_is_neither_a_seed_nor_reported_stale(production) -> None:
    """The map's report says which stored rows it ignored (and the next
    persist drops them): an unknown section's row is held, not stale."""
    fx = Fx()

    report = sta.load_assignment(fx.sc.id, incomplete="collect").report

    assert report["stale_rows_ignored"] == 0
    assert report["ignored_row_keys"] == []
    assert f"{fx.eng}_S1" not in report["closure_keys"]
