"""Section Planning sizes courses the way the timetable builder will.

- Declared limits hold in every view: the all-programmes view (the default)
  sizes a course by the lowest limit any programme in scope declares, through
  the same function the builder uses for a pooled scenario.
- COE is one of our departments: its courses take the local 25/40 rules, and
  the page reads that list from the server.
- Male and female students are never planned together: every course is sized
  for each cohort, Total = M + F in rows, KPIs, summary and export, and students
  with no recorded gender are reported apart, never pooled.
- A resolved elective is planned for its slot, with the slot's declared limit,
  and slot demand that becomes no course is reported.
"""

from __future__ import annotations

import io
import json
from collections import Counter
from pathlib import Path

import pytest
from django.contrib.auth.models import Group, User
from django.test import Client

from core.authz import _rate_buckets
from core.models import (
    Course,
    ElectiveCourse,
    ElectiveTermMapping,
    ProgrammeRequirement,
    Student,
)
from core.services.rbac import (
    ROLE_GENERAL_ADVISOR,
    ensure_role_groups,
    ensure_scope_schema,
    set_user_scope,
)
from core.services.reporting import clear_aggregate_cache
from core.services.section_planning import (
    LOCAL_DEPARTMENTS,
    compute_section_plan,
    lowest_declared_capacities,
)

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def _fresh_state():
    _rate_buckets.clear()
    clear_aggregate_cache()
    yield
    _rate_buckets.clear()
    clear_aggregate_cache()


@pytest.fixture
def planner(client: Client) -> Client:
    ensure_role_groups()
    ensure_scope_schema()
    user, _ = User.objects.get_or_create(username="sp-sizing")
    user.groups.clear()
    user.groups.add(Group.objects.get(name=ROLE_GENERAL_ADVISOR))
    set_user_scope(user.id, advisor_id="", departments="")
    client.force_login(user)
    return client


def _requirement(
    program: str, code: str, cap: int | None, *, credits: int = 3, name: str = ""
) -> None:
    ProgrammeRequirement.objects.create(
        program=program,
        course_code=code,
        course_name=name or f"{code} NAME",
        credit_hours=credits,
        type="Mandatory",
        programme_term=5,
        max_capacity=cap,
    )


def _students(program: str, first_id: int, count: int, *, section: str = "M") -> list[int]:
    ids = list(range(first_id, first_id + count))
    Student.objects.bulk_create(
        [
            Student(student_id=sid, registration_no=str(sid), program=program, section=section)
            for sid in ids
        ]
    )
    return ids


def _recommend(monkeypatch: pytest.MonkeyPatch, by_program: dict[str, list[str]]) -> None:
    """Every student of a programme is recommended that programme's courses."""

    def single(student_ids, program, _year, _term, **_kw):
        return {sid: list(by_program.get(program, [])) for sid in student_ids}

    def multi(student_ids, _year, _term, **_kw):
        programs = dict(
            Student.objects.filter(student_id__in=student_ids).values_list("student_id", "program")
        )
        return {sid: list(by_program.get(programs[sid], [])) for sid in student_ids}

    monkeypatch.setattr("core.services.reporting.batch_recommend", single)
    monkeypatch.setattr("core.services.reporting.batch_recommend_multi_program", multi)


def _generate(client: Client, **body) -> dict:
    response = client.post(
        "/ops/section-planning/generate/",
        json.dumps({"year": 1448, "semester": 1, **body}),
        content_type="application/json",
    )
    assert response.status_code == 200, response.content
    return response.json()


def _row(plan: list[dict], code: str) -> dict:
    rows = [row for row in plan if row["course_code"] == code]
    assert len(rows) == 1, rows
    return rows[0]


def _export_workbook(client: Client, **body):
    from openpyxl import load_workbook

    response = client.post(
        "/ops/section-planning/export/",
        json.dumps({"year": 1448, "semester": 1, **body}),
        content_type="application/json",
    )
    assert response.status_code == 200
    return load_workbook(io.BytesIO(b"".join(response.streaming_content)), data_only=False)


def _sheet_rows(sheet) -> list[dict]:
    headers = [cell.value for cell in sheet[2]]
    return [
        dict(zip(headers, [cell.value for cell in row], strict=False))
        for row in sheet.iter_rows(min_row=3)
        if row[0].value is not None
    ]


def _export_rows(client: Client, **body) -> list[dict]:
    return _sheet_rows(_export_workbook(client, **body).worksheets[0])


# ── declared limits in every view ─────────────────────────────────


def test_lowest_declared_capacities_takes_the_most_restrictive_programme() -> None:
    _requirement("AI", "CS211", 30)
    _requirement("DS", "CS211", 25)
    _requirement("CS", "CS211", 20)  # not in scope: must not count
    _requirement("AI", "AI491", 5)
    _requirement("DS", "AI491", None)
    _requirement("AI", "MATH203", 0)  # a zero is not a declaration
    metadata = {
        "CS211::DATA_STRUCTURES": {"course_code": "CS211"},
        "AI491": {"course_code": "AI491"},
        "MATH203": {"course_code": "MATH203"},
        "PHYS103": {},
    }

    caps = lowest_declared_capacities(["AI", "DS"], metadata)

    assert caps == {"CS211::DATA_STRUCTURES": 25, "AI491": 5}
    assert lowest_declared_capacities([], metadata) == {}


@pytest.fixture
def graduation_project(monkeypatch: pytest.MonkeyPatch) -> None:
    """31 AI students need AI491, which AI limits to 5 per section; DS limits nothing."""
    Course.objects.create(
        course_code="AI491", department="AI", credit_hours=2, description="GRADUATION PROJECT I"
    )
    Course.objects.create(
        course_code="DS321", department="DS", credit_hours=3, description="DS321 NAME"
    )
    _requirement("AI", "AI491", 5, credits=2, name="GRADUATION PROJECT I")
    _requirement("AI2", "AI491", None, credits=2, name="GRADUATION PROJECT I")
    _requirement("DS", "DS321", None)
    _students("AI", 441000001, 16, section="M")
    _students("AI", 442000001, 15, section="F")
    _students("DS", 443000001, 4, section="M")
    _recommend(monkeypatch, {"AI": ["AI491"], "DS": ["DS321"]})


def test_ai491_has_the_same_max_in_its_programme_and_in_all_programmes(
    planner: Client, graduation_project
) -> None:
    programme_view = _row(_generate(planner, program="AI")["plan"], "AI491")
    default_view = _row(_generate(planner)["plan"], "AI491")

    assert programme_view["max_per_section"] == 5
    assert default_view["max_per_section"] == programme_view["max_per_section"]
    assert default_view["num_sections"] == programme_view["num_sections"]


def test_the_all_programmes_export_carries_the_declared_limit(
    planner: Client, graduation_project
) -> None:
    row = next(r for r in _export_rows(planner) if r["Course"] == "AI491")
    assert row["Max/Section"] == 5


def test_the_builder_and_section_planning_size_a_shared_course_alike(
    planner: Client, monkeypatch: pytest.MonkeyPatch
) -> None:
    from core.services.timetable_generate import generate_workspace_scenario

    Course.objects.create(
        course_code="CS211", department="CS", credit_hours=3, description="DATA STRUCTURES"
    )
    _requirement("AI", "CS211", 30, name="DATA STRUCTURES")
    _requirement("DS", "CS211", 35, name="DATA STRUCTURES")
    ids = _students("AI", 451000001, 40) + _students("DS", 452000001, 40)
    _recommend(monkeypatch, {"AI": ["CS211"], "DS": ["CS211"]})
    monkeypatch.setattr("core.services.timetable_generate.get_student_ids", lambda **_kw: ids)
    monkeypatch.setattr(
        "core.services.recommender_batch.batch_recommend_multi_program",
        lambda student_ids, _year, _term: {sid: ["CS211"] for sid in student_ids},
    )

    built = generate_workspace_scenario(1448, 1, ["AI", "DS"], section="M", run_autoplace=False)
    planned = _generate(planner, program="AI,DS")

    builder_row = _row(built["section_plan"], "CS211")
    planning_row = _row(planned["combined_plan"], "CS211")
    assert builder_row["max_per_section"] == 30
    assert planning_row["max_per_section"] == builder_row["max_per_section"]
    assert planning_row["num_sections"] == builder_row["num_sections"]


# ── COE is one of our departments ────────────────────────────────


def test_coe_courses_take_the_local_rules() -> None:
    Course.objects.create(course_code="COE211", department="COE", credit_hours=4)
    Course.objects.create(course_code="COE321", department="", credit_hours=3)
    Course.objects.create(course_code="MATH203", department="MATH", credit_hours=4)

    plan = {
        row["course_code"]: row
        for row in compute_section_plan(Counter({"COE211": 60, "COE321": 60, "MATH203": 60}))
    }

    assert "COE" in LOCAL_DEPARTMENTS
    assert plan["COE211"]["max_per_section"] == 25  # local, 4 credits
    assert plan["COE211"]["num_sections"] == 3
    assert (
        plan["COE321"]["max_per_section"] == 40
    )  # local, other credits (department from the code)
    assert plan["COE321"]["is_external"] is False
    assert plan["MATH203"]["max_per_section"] == 50  # a service department stays external


def test_a_coe_course_flagged_external_keeps_the_external_rule() -> None:
    """COE joining our departments does not reinterpret the per-course flag."""
    Course.objects.create(course_code="COE466", department="COE", credit_hours=3, is_external=True)

    (row,) = compute_section_plan(Counter({"COE466": 60}))

    assert row["max_per_section"] == 50


def test_the_page_is_given_the_servers_department_list(planner: Client) -> None:
    import re

    html = planner.get("/section-planning/").content.decode("utf-8")
    embedded = re.search(
        r'<script id="spLocalDepartments" type="application/json">(.*?)</script>', html
    )

    assert embedded, "the page must be handed the server's list"
    assert json.loads(embedded.group(1)) == sorted(LOCAL_DEPARTMENTS)
    js = (Path(__file__).resolve().parents[1] / "static/js/page-section-planning.js").read_text(
        encoding="utf-8"
    )
    assert "'CYB'" not in js and '"CYB"' not in js, "the page must not keep a list of its own"


# ── male and female are never planned together ────────────────────


@pytest.fixture
def mixed_cohorts(monkeypatch: pytest.MonkeyPatch) -> None:
    """AI: 12 men and 12 women need AI331 (4 cr, local: 25 a section); 3 students
    have no recorded gender, and one of them alone needs AI352."""
    Course.objects.create(
        course_code="AI331", department="AI", credit_hours=4, description="MACHINE LEARNING"
    )
    Course.objects.create(
        course_code="AI352", department="AI", credit_hours=3, description="AI352 NAME"
    )
    _requirement("AI", "AI331", None, credits=4, name="MACHINE LEARNING")
    _requirement("AI", "AI352", None)
    _students("AI", 461000001, 12, section="M")
    _students("AI", 462000001, 12, section=" f ")  # recorded, only badly spelt
    unknown = _students("AI", 463000001, 3, section="")
    by_student = {sid: ["AI331"] for sid in Student.objects.values_list("student_id", flat=True)}
    by_student[unknown[0]] = ["AI331", "AI352"]
    by_student[unknown[1]] = []
    by_student[unknown[2]] = ["AI352"]

    def single(student_ids, _program, _year, _term, **_kw):
        return {sid: list(by_student[sid]) for sid in student_ids}

    def multi(student_ids, _year, _term, **_kw):
        return {sid: list(by_student[sid]) for sid in student_ids}

    monkeypatch.setattr("core.services.reporting.batch_recommend", single)
    monkeypatch.setattr("core.services.reporting.batch_recommend_multi_program", multi)


@pytest.mark.parametrize("program", ["AI", None])
def test_male_and_female_are_sized_apart_and_total_is_m_plus_f(
    planner: Client, mixed_cohorts, program
) -> None:
    data = _generate(planner, **({"program": program} if program else {}))
    row = _row(data["plan"], "AI331")

    # Pooled, 24 students fit one 25-seat section; apart, each cohort needs one.
    assert (row["male_students"], row["female_students"]) == (12, 12)
    assert (row["male_sections"], row["female_sections"]) == (1, 1)
    assert row["num_sections"] == 2
    assert row["total_students"] == 24
    assert row["max_per_section"] == 25
    assert data["cohorts"] == {"M": 12, "F": 12, "no_gender": 3}
    assert data["student_count"] == 27
    summary = data["summary"]
    assert (summary["male_sections"], summary["female_sections"], summary["total_sections"]) == (
        1,
        1,
        2,
    )
    (ai,) = summary["departments"]
    assert (ai["male_sections"], ai["female_sections"], ai["sections"]) == (1, 1, 2)


def test_students_with_no_recorded_gender_are_reported_never_pooled(
    planner: Client, mixed_cohorts
) -> None:
    data = _generate(planner, program="AI")

    ai331 = _row(data["plan"], "AI331")
    assert ai331["unknown_students"] == 1
    assert ai331["total_students"] == 24, "a no-gender student is in neither cohort"
    # A course only no-gender students need is listed, but plans no section.
    ai352 = _row(data["plan"], "AI352")
    assert (ai352["unknown_students"], ai352["num_sections"], ai352["status"]) == (
        2,
        0,
        "no_gender",
    )
    assert data["summary"]["total_courses"] == 1
    assert data["summary"]["total_sections"] == 2
    assert data["summary"]["no_gender"] == {"students": 3, "seat_demand": 3, "courses": 2}


def test_a_section_field_in_the_request_changes_nothing(planner: Client, mixed_cohorts) -> None:
    """The free-text Section filter is gone: both cohorts are always planned."""
    data = _generate(planner, program="AI", section="M")

    assert data["cohorts"] == {"M": 12, "F": 12, "no_gender": 3}
    assert _row(data["plan"], "AI331")["female_sections"] == 1
    html = planner.get("/section-planning/").content.decode("utf-8")
    assert 'id="spSection"' not in html


def test_several_programmes_are_pooled_by_cohort(
    planner: Client, monkeypatch: pytest.MonkeyPatch
) -> None:
    """AI and DS share CS211 (limits 30 and 35): men and women each pool across programmes."""
    Course.objects.create(
        course_code="CS211", department="CS", credit_hours=3, description="DATA STRUCTURES"
    )
    _requirement("AI", "CS211", 30, name="DATA STRUCTURES")
    _requirement("DS", "CS211", 35, name="DATA STRUCTURES")
    _students("AI", 471000001, 20, section="M")
    _students("DS", 472000001, 20, section="M")
    _students("AI", 473000001, 5, section="F")
    _recommend(monkeypatch, {"AI": ["CS211"], "DS": ["CS211"]})

    data = _generate(planner, program="AI,DS")

    row = _row(data["combined_plan"], "CS211")
    assert row["programs"] == ["AI", "DS"]
    assert row["max_per_section"] == 30
    assert (row["male_students"], row["female_students"]) == (40, 5)
    assert (row["male_sections"], row["female_sections"], row["num_sections"]) == (2, 1, 3)
    assert data["combined_summary"]["total_sections"] == 3
    per_programme = {p["program"]: _row(p["plan"], "CS211") for p in data["programs"]}
    assert (per_programme["AI"]["male_sections"], per_programme["AI"]["female_sections"]) == (1, 1)
    assert (per_programme["DS"]["max_per_section"], per_programme["DS"]["num_sections"]) == (35, 1)
    # The Department Summary adds up to the KPI (it is the pooled plan's).
    assert sum(d["sections"] for d in data["combined_summary"]["departments"]) == 3


def test_same_code_different_plan_names_stay_two_courses(
    planner: Client, monkeypatch: pytest.MonkeyPatch
) -> None:
    Course.objects.create(
        course_code="CS111", department="CS", credit_hours=4, description="GLOBAL NAME"
    )
    _requirement("AI", "CS111", 20, credits=4, name="PROGRAMMING I")
    _requirement("AI2", "CS111", None, credits=4, name="FUNDAMENTALS OF PROGRAMMING")
    _requirement("DS", "CS111", 30, credits=4, name="PROGRAMMING I")
    _students("AI", 481000001, 10)
    _students("AI2", 482000001, 4)
    _students("DS", 483000001, 15)
    _recommend(monkeypatch, {"AI": ["CS111"], "AI2": ["CS111"], "DS": ["CS111"]})

    data = _generate(planner, program="AI,AI2,DS")

    rows = {
        row["course_name"]: row for row in data["combined_plan"] if row["course_code"] == "CS111"
    }
    assert set(rows) == {"PROGRAMMING I", "FUNDAMENTALS OF PROGRAMMING"}, "the plan's own names win"
    assert rows["PROGRAMMING I"]["programs"] == ["AI", "DS"]
    assert rows["PROGRAMMING I"]["total_students"] == 25
    assert rows["PROGRAMMING I"]["max_per_section"] == 20
    assert rows["PROGRAMMING I"]["num_sections"] == 2
    assert rows["FUNDAMENTALS OF PROGRAMMING"]["programs"] == ["AI2"]
    assert rows["FUNDAMENTALS OF PROGRAMMING"]["total_students"] == 4


def test_the_export_carries_male_and_female(planner: Client, mixed_cohorts) -> None:
    workbook = _export_workbook(planner, program="AI")
    sections, summary = workbook.worksheets[0], workbook.worksheets[1]
    rows = {row["Course"]: row for row in _sheet_rows(sections)}

    ai331 = rows["AI331"]
    assert (ai331["Male students"], ai331["Female students"]) == (12, 12)
    assert ai331["No gender recorded"] == 1
    assert ai331["Max/Section"] == 25
    r = next(i for i, row in enumerate(sections.iter_rows(min_row=3), 3) if row[2].value == "AI331")
    assert ai331["Students"] == f"=G{r}+H{r}"
    assert ai331["Male sections"] == f"=IF(G{r}>0,CEILING(G{r}/J{r},1),0)"
    assert ai331["Female sections"] == f"=IF(H{r}>0,CEILING(H{r}/J{r},1),0)"
    assert ai331["Sections"] == f"=K{r}+L{r}"
    assert rows["AI352"]["Status"] == "No gender recorded"

    labels = {summary.cell(row=r, column=c).value: (r, c) for r in (3, 4) for c in (1, 3, 5)}
    r, c = labels["Male Sections"]
    assert summary.cell(row=r, column=c + 1).value == "=SUM(Sections!K3:K4)"
    r, c = labels["Female Sections"]
    assert summary.cell(row=r, column=c + 1).value == "=SUM(Sections!L3:L4)"
    r, c = labels["Total Sections"]
    assert summary.cell(row=r, column=c + 1).value == "=SUM(Sections!M3:M4)"
    r, c = labels["No gender recorded"]
    assert summary.cell(row=r, column=c + 1).value == 3


def test_the_multi_programme_export_writes_the_pooled_plan_first(
    planner: Client, monkeypatch: pytest.MonkeyPatch
) -> None:
    Course.objects.create(
        course_code="CS211", department="CS", credit_hours=3, description="DATA STRUCTURES"
    )
    _requirement("AI", "CS211", 30, name="DATA STRUCTURES")
    _requirement("DS", "CS211", 35, name="DATA STRUCTURES")
    _students("AI", 491000001, 20, section="M")
    _students("DS", 492000001, 20, section="F")
    _recommend(monkeypatch, {"AI": ["CS211"], "DS": ["CS211"]})

    workbook = _export_workbook(planner, program="AI,DS")

    assert workbook.sheetnames[:2] == ["Sections-All", "Summary-All"]
    assert "Sections-AI" in workbook.sheetnames and "Sections-DS" in workbook.sheetnames
    (pooled,) = _sheet_rows(workbook["Sections-All"])
    assert (pooled["Male students"], pooled["Female students"], pooled["Max/Section"]) == (
        20,
        20,
        30,
    )
    assert pooled["Name"] == "AI, DS - DATA STRUCTURES"


# ── a resolved elective takes its slot's declared limit ──────────


def _slot(program: str, code: str, cap: int | None) -> None:
    ProgrammeRequirement.objects.create(
        program=program,
        course_code=code,
        course_name=f"DEPARTMENT ELECTIVE {code[-1]}",
        credit_hours=3,
        type="Program Elective",
        programme_term=7,
        max_capacity=cap,
    )


def _publish(program: str, slot: str, code: str, *, prerequisites: str = "") -> None:
    elective = ElectiveCourse.objects.create(
        programme=program,
        course_code=code,
        course_name=f"{code} ELECTIVE",
        credit_hours=3,
        prerequisites_csv=prerequisites,
    )
    ElectiveTermMapping.objects.create(
        academic_year="1448", term=1, programme=program, placeholder_code=slot, elective=elective
    )


@pytest.fixture
def electives(monkeypatch: pytest.MonkeyPatch):
    """AI1 (30 seats) resolves to AI463, a course flagged external; AI2 is not
    published this term; AI3 resolves only to a course nobody is eligible for."""
    Course.objects.create(course_code="AI463", department="AI", credit_hours=3, is_external=True)
    _slot("AI", "AI1", 30)
    _slot("AI", "AI2", 30)
    _slot("AI", "AI3", None)
    _publish("AI", "AI1", "AI463")
    _publish("AI", "AI3", "AI464", prerequisites="AI999")
    men = _students("AI", 501000001, 70, section="M")
    women = _students("AI", 502000001, 52, section="F")
    unpublished = _students("AI", 503000001, 5, section="M")
    blocked = _students("AI", 504000001, 3, section="F")
    wants = {
        **{sid: ["AI1"] for sid in men + women},
        **{sid: ["AI2"] for sid in unpublished},
        **{sid: ["AI3"] for sid in blocked},
    }

    def single(student_ids, _program, _year, _term, **_kw):
        return {sid: list(wants[sid]) for sid in student_ids}

    def multi(student_ids, _year, _term, **_kw):
        return {sid: list(wants[sid]) for sid in student_ids}

    monkeypatch.setattr("core.services.reporting.batch_recommend", single)
    monkeypatch.setattr("core.services.reporting.batch_recommend_multi_program", multi)


@pytest.mark.parametrize("program", ["AI", None])
def test_a_resolved_elective_takes_its_slots_declared_limit(
    planner: Client, electives, program
) -> None:
    data = _generate(planner, **({"program": program} if program else {}))
    row = _row(data["plan"], "AI463")

    assert row["is_external"] is True
    assert row["max_per_section"] == 30, "the AI1 slot's limit, not the external 50"
    assert (row["male_sections"], row["female_sections"]) == (3, 2)  # 70/30, 52/30
    assert row["limit_source"] == "slot"
    assert row["slots"] == ["AI1"]


def test_a_slot_that_declares_nothing_leaves_the_course_to_its_own_rule(
    planner: Client, electives
) -> None:
    ProgrammeRequirement.objects.filter(program="AI", course_code="AI1").update(max_capacity=None)

    row = _row(_generate(planner, program="AI")["plan"], "AI463")

    assert row["max_per_section"] == 50
    assert row["limit_source"] == "rule"
    assert row["slots"] == ["AI1"]


def test_a_draft_on_the_slot_is_a_what_if_for_its_elective(planner: Client, electives) -> None:
    row = _row(_generate(planner, program="AI", course_overrides={"AI1": 20})["plan"], "AI463")

    assert row["max_per_section"] == 20
    assert row["limit_source"] == "draft"
    assert ProgrammeRequirement.objects.get(program="AI", course_code="AI1").max_capacity == 30


def test_slot_demand_that_becomes_no_course_is_reported(planner: Client, electives) -> None:
    data = _generate(planner, program="AI")

    assert data["electives"] == {
        "dropped": [
            {"program": "AI", "slot": "AI2", "reason": "not_published", "students": 5},
            {"program": "AI", "slot": "AI3", "reason": "no_eligible_course", "students": 3},
        ],
        "dropped_total": 8,
    }
    assert {row["course_code"] for row in data["plan"]} == {"AI463"}


def test_the_panel_lists_the_elective_under_its_slot(planner: Client, electives) -> None:
    response = planner.get(
        "/ops/section-planning/courses/", {"program": "AI", "year": 1448, "semester": 1}
    )
    courses = {c["course_code"]: c for c in response.json()["courses"]}

    assert courses["AI1"]["slot_electives"] == [
        {"program": "AI", "status": "ready", "courses": ["AI463"]}
    ]
    assert courses["AI2"]["slot_electives"] == [
        {"program": "AI", "status": "not_published", "courses": []}
    ]
    without_term = planner.get("/ops/section-planning/courses/", {"program": "AI"}).json()
    assert {c["course_code"]: c["slot_electives"] for c in without_term["courses"]}["AI1"] == []


def test_the_export_names_the_slot_and_the_limits_source(planner: Client, electives) -> None:
    workbook = _export_workbook(planner, program="AI")
    (row,) = [r for r in _sheet_rows(workbook.worksheets[0]) if r["Course"] == "AI463"]

    assert row["Max/Section"] == 30
    assert row["Limit from"] == "Slot AI1 limit"
    assert row["Fills slot"] == "AI1"
    note = workbook.worksheets[1].cell(row=6, column=1).value
    assert note.startswith("Elective-slot demand that became no course: 8")
    assert "AI AI2: not published for this term (5)" in note
