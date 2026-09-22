"""Room allocation stays reproducible across build, fixed checks and saves.

All persisted rows in this module belong to pytest's isolated database.
"""

from collections import Counter
from copy import deepcopy

import pytest
from django.core.cache import cache
from django.urls import reverse
from openpyxl import load_workbook

from core.models import (
    Course,
    ExamTimetableRun,
    ProgrammeRequirement,
    Room,
    Student,
    StudentCourse,
    StudentTermSection,
    TermSection,
)
from core.services import exam_input_fingerprint
from core.services.exam_evaluation import evaluate_exam_schedule
from core.services.exam_timetable import (
    _build_qa,
    _build_room_qa,
    apply_thin_conflict_policy,
    attach_exam_relaxation_qa,
    build_conflict_graph,
    build_exam_timetable,
    check_room_feasibility,
    export_exam_timetable_xlsx,
)
from tests.exam_source_factory import scraped_exam_registration

pytestmark = pytest.mark.django_db
DAYS = ["Sun", "Mon", "Tue"]
PERIODS = ["08:00-10:00", "13:00-15:00"]


def _evaluate(source, schedule=None, pinned=None):
    return evaluate_exam_schedule(
        days=DAYS,
        periods=PERIODS,
        max_per_day=source["qa"]["max_per_day"],
        schedule_raw=deepcopy(schedule if schedule is not None else source["schedule"]),
        selected_courses=source["courses"],
        assign_rooms=source["assign_rooms"],
        seed=source["seed"],
        thin_conflict_threshold=0,
        programs=source["enrollment_scope"]["programs"],
        sections=source["enrollment_scope"]["sections"],
        pinned=deepcopy(source["pinned"] if pinned is None else pinned),
    )


def _placements(result):
    return {
        row["course_identity"]: (row["day"], row["period"], row["slot_index"])
        for row in result["schedule"]
    }


def _room_rows(result):
    # Room-list order is deliberate output, so retain it in these assertions.
    return {row["course_identity"]: row["rooms"] for row in result["schedule"]}


def _report(result):
    qa = deepcopy(result["qa"])
    qa.pop("rebalance_moves", None)  # Historical optimizer activity, not current QA.
    qa["enrolment_snapshot"].pop("snapshot_timestamp", None)
    return qa


def _visible_report(result):
    """Match every substantive field used by the frontend report signature."""
    return {
        "qa": _report(result),
        **{
            key: result.get(key)
            for key in [
                "courses_count",
                "students_count",
                "conflicts_count",
                "conflicts",
                "credit_map",
                "section_enrollment",
                "primary_status",
                "status_flags",
            ]
        },
        "schedule": sorted(
            [
                {
                    "identity": row["course_identity"],
                    "day": row["day"],
                    "period": row["period"],
                    "rooms": row["rooms"],
                }
                for row in result["schedule"]
            ],
            key=lambda row: row["identity"],
        ),
    }


def _assert_room_constraints(result):
    occupied = {}
    room_totals = Counter()
    for course in result["schedule"]:
        if course["day"] == "OVERFLOW":
            continue
        demand = Counter(
            {
                (row["section_key"], row["gender"]): row["student_count"]
                for row in result["section_enrollment"][course["course_code"]]
            }
        )
        seated = Counter()
        official = {
            (row["section_key"], row["gender"]): row["section"]
            for row in result["section_enrollment"][course["course_code"]]
        }
        for room in course["rooms"]:
            assert (
                sum(part["student_count"] for part in room["section_parts"])
                == room["student_count"]
            )
            for part in room["section_parts"]:
                key = (part["section_key"], part["gender"])
                assert part["section"] == official[key]
                assert part["gender"] == room["gender"]
                seated[key] += part["student_count"]
            if room["room_code"] == "UNASSIGNED":
                continue
            assert room["student_count"] <= room["room_capacity"]
            assert Room.objects.get(room_code=room["room_code"]).section == room["gender"]
            key = (course["slot_index"], room["room_code"])
            assert key not in occupied or occupied[key] == course["course_identity"]
            occupied[key] = course["course_identity"]
            room_totals[key] += room["student_count"]
            assert room_totals[key] <= room["room_capacity"]
        assert seated == demand


@pytest.fixture
def tied_inventory():
    students = []
    for start, count, program, gender in [
        (100, 12, "AI", "F"),
        (200, 12, "DS", "F"),
        (300, 8, "AI", "M"),
    ]:
        students.extend(
            Student(student_id=sid, program=program, section=gender)
            for sid in range(start, start + count)
        )
    Student.objects.bulk_create(students)
    roster = {student.student_id: student for student in students}
    for code, credits in [("CS301", 4), ("CS111", 4), ("DS113", 3), ("CS201", 3)]:
        course = Course.objects.create(
            course_code=code, description=f"Course {code}", credit_hours=credits
        )
        for program in ["AI", "DS"]:
            name = f"{program} Programming" if code == "CS111" else course.description
            ProgrammeRequirement.objects.create(
                program=program,
                course_code=code,
                course_name=name,
                programme_term={"CS111": 1, "DS113": 2, "CS201": 3, "CS301": 4}[code],
            )
        for sid, student in roster.items():
            if code == "CS111" and sid % 3 == 0:
                continue
            label = (
                f"{student.section}{1 + sid % 2}/أ+{code}"
                if code == "DS113"
                else f"{student.section}11"
            )
            scraped_exam_registration(student, course, section_label=label)
    for gender, capacities in [("F", [12, 12, 8, 5]), ("M", [10, 10, 5])]:
        for index, capacity in enumerate(capacities):
            Room.objects.create(
                room_code=f"{gender}{index}",
                capacity=capacity,
                section=gender,
                building=f"B{index % 2}",
                floor="1",
            )
    # Neither studying nor future/manual section evidence can join this population.
    extra = Course.objects.create(course_code="FAKE999", description="Not an actual exam")
    StudentCourse.objects.create(student=roster[100], course=extra, status="studying")
    future = TermSection.objects.create(
        course_key=extra.course_code, course_code=extra.course_code, section="F99"
    )
    StudentTermSection.objects.create(
        student_id=100,
        term_section=future,
        academic_year="1449",
        term="1",
        source="registration_plan_1449_t1",
    )


@pytest.mark.parametrize("seed", [None, 17, 1392698829])
@pytest.mark.parametrize("rebalance", [False, True])
def test_build_check_repeated_check_keep_identical_rooms_and_report(
    tied_inventory, seed, rebalance
):
    built = build_exam_timetable(
        label="Stable rooms",
        days=DAYS,
        periods=PERIODS,
        programs=["DS", "AI"],
        sections=["M", "F"],
        seed=seed,
        rebalance_invigilators=rebalance,
        persist=False,
    )
    assert ExamTimetableRun.objects.count() == 0
    assert built["enrollment_source"] == "scraper_timetable"
    assert "FAKE999" not in built["courses"]
    variants = [row for row in built["schedule"] if row["source_course_code"] == "CS111"]
    assert len({row["course_identity"] for row in variants}) == 2
    checked = _evaluate(built, list(reversed(built["schedule"])))
    repeated = _evaluate(checked)
    for current in [checked, repeated]:
        assert _placements(current) == _placements(built)
        assert current["input_fingerprint"] == built["input_fingerprint"]
        assert current["section_enrollment"] == built["section_enrollment"]
        assert _room_rows(current) == _room_rows(built)
        assert _report(current) == _report(built)
        assert _visible_report(current) == _visible_report(built)
        assert (current["primary_status"], current["status_flags"]) == (
            built["primary_status"],
            built["status_flags"],
        )
        _assert_room_constraints(current)


@pytest.fixture
def moving_inventory():
    next_sid = 1000
    for code, groups in [
        ("DS113", [("F1/أ", 26), ("F2+ب", 23)]),
        ("CS248", [("F1", 48)]),
        ("CS118", [("F01", 47)]),
        ("CS999", [("F1", 12)]),
    ]:
        course = Course.objects.create(course_code=code, description=code, credit_hours=3)
        for label, size in groups:
            students = [
                Student(student_id=sid, program="AI", section="F")
                for sid in range(next_sid, next_sid + size)
            ]
            Student.objects.bulk_create(students)
            for student in students:
                scraped_exam_registration(student, course, section_label=label)
            next_sid += size
    for capacity in [54, 26, 23]:
        Room.objects.create(room_code=f"F{capacity}", capacity=capacity, section="F", building="B1")


def test_manual_move_reallocates_both_periods_without_moving_other_exams_and_save_is_identical(
    moving_inventory, client, django_user_model, monkeypatch, tmp_path
):
    cache.clear()
    monkeypatch.setattr("core.authz._rate_buckets", {})
    client.force_login(django_user_model.objects.create_superuser(username="room-stability"))
    pins = [
        {"course_code": "DS113", "day": "Sun", "period": PERIODS[0]},
        {"course_code": "CS248", "day": "Sun", "period": PERIODS[1]},
        {"course_code": "CS118", "day": "Sun", "period": PERIODS[1]},
        {"course_code": "CS999", "day": "Tue", "period": PERIODS[0]},
    ]
    built = build_exam_timetable(
        label="Source and destination", days=DAYS, periods=PERIODS, pinned=pins, seed=17
    )
    before = {row["course_code"]: row for row in built["schedule"]}
    assert [room["room_code"] for room in before["DS113"]["rooms"]] == ["F54"]
    assert len(before["CS118"]["rooms"]) == 2
    moved = deepcopy(built["schedule"])
    next(row for row in moved if row["course_code"] == "CS248").update(
        day="Sun", period=PERIODS[0], slot_index=0
    )
    kept_pins = [pin for pin in built["pinned"] if pin["course_code"] != "CS248"]
    with pytest.raises(ValueError, match="pin|fixed|Pinned"):
        _evaluate(built, moved)

    def forbidden(*args, **kwargs):
        raise AssertionError("A fixed placement workflow called a scheduling optimizer")

    # Resolve the view module before replacing the service function. Its
    # imported scheduler alias must be patched and restored by monkeypatch,
    # not permanently captured from this test during the first URL import.
    monkeypatch.setattr("core.exam_views.schedule", forbidden)
    monkeypatch.setattr("core.services.exam_timetable.schedule", forbidden)
    monkeypatch.setattr("core.services.exam_timetable._rebalance_invigilators_pass", forbidden)
    # The evaluator holds its own reference to the pass; patch that binding too.
    monkeypatch.setattr("core.services.exam_evaluation._rebalance_invigilators_pass", forbidden)
    checked = _evaluate(built, moved, kept_pins)
    after = {row["course_code"]: row for row in checked["schedule"]}
    assert {room["room_code"] for room in after["DS113"]["rooms"]} == {"F26", "F23"}
    assert [room["room_code"] for room in after["CS248"]["rooms"]] == ["F54"]
    assert [room["room_code"] for room in after["CS118"]["rooms"]] == ["F54"]
    assert after["CS999"] == before["CS999"], "An unaffected period keeps its exact allocation"
    assert _placements(checked) == _placements({"schedule": moved})
    assert checked["pinned"] == kept_pins
    assert checked["input_fingerprint"] == built["input_fingerprint"]
    assert ExamTimetableRun.objects.count() == 1
    _assert_room_constraints(checked)
    payload = {
        "label": "Saved source and destination",
        "mode": "save_loaded_changes",
        "previous_run_id": built["run_id"],
        "base_schedule": moved,
        "selected_courses": built["courses"],
        "days": DAYS,
        "periods": PERIODS,
        "max_per_day": 2,
        "assign_rooms": True,
        "pinned": kept_pins,
        "expected_input_fingerprint": checked["input_fingerprint"],
    }
    response = client.post(
        reverse("exam_timetable_build"), payload, content_type="application/json"
    )
    assert response.status_code == 200, response.content
    saved = response.json()
    assert ExamTimetableRun.objects.count() == 2
    assert _placements(saved) == _placements(checked)
    assert _room_rows(saved) == _room_rows(checked)
    assert _report(saved) == _report(checked)
    assert _visible_report(saved) == _visible_report(checked)
    assert saved["section_enrollment"] == checked["section_enrollment"]
    assert saved["input_fingerprint"] == checked["input_fingerprint"]
    monkeypatch.setattr("core.services.exam_timetable.RUNTIME_DIR", tmp_path)
    workbook = load_workbook(export_exam_timetable_xlsx(saved["run_id"]))
    sheet = workbook["Room Assignments"]
    exported = {
        (row[2], row[6]): row
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[2] in saved["courses"]
    }
    for entry in saved["schedule"]:
        for room in entry["rooms"]:
            row = exported[(entry["course_code"], room["room_code"])]
            assert row[:2] == (entry["day"], entry["period"])
            assert row[3] == room["section"]
            assert row[4:8] == (
                room["gender"],
                room["student_count"],
                room["room_code"],
                room["room_capacity"],
            )
            assert row[10] == ("No" if entry["course_code"] == "CS248" else "Yes")
            assert (row[13] or "") == room.get("room_group", "")
    assert len(exported) == sum(len(entry["rooms"]) for entry in saved["schedule"])
    workbook.close()


def test_qa_order_is_independent_of_course_student_section_and_room_traversal():
    enrolled = {"ZZ101": {8, 1}, "AA101": {2, 8, 1}, "BB101": {2, 1}}
    schedule = [
        {
            "course_code": code,
            "course_identity": code,
            "day": "Sun",
            "period": PERIODS[0],
            "slot_index": 0,
            "rooms": [
                {"room_code": "UNASSIGNED", "section": label, "student_count": 1, "gender": "F"}
                for label in ["F2", "F1"]
            ],
        }
        for code in enrolled
    ]
    buckets = {("AI", 1): set(enrolled)}
    first = _build_qa(
        enrolled,
        schedule,
        max_per_day=1,
        plan_term_buckets=buckets,
        credit_map={code: 4 for code in enrolled},
    )
    shuffled = deepcopy(list(reversed(schedule)))
    for row in shuffled:
        row["rooms"].reverse()
    second = _build_qa(
        dict(reversed(list(enrolled.items()))),
        shuffled,
        max_per_day=1,
        plan_term_buckets=buckets,
        credit_map={code: 4 for code in enrolled},
    )
    assert first == second
    assert build_conflict_graph(enrolled) == build_conflict_graph(
        dict(reversed(list(enrolled.items())))
    )
    _, adjacency = build_conflict_graph(enrolled)
    _, thin_courses = apply_thin_conflict_policy(enrolled, adjacency, 10)
    attach_exam_relaxation_qa(first, enrolled, schedule, 10, thin_courses)
    attach_exam_relaxation_qa(second, enrolled, shuffled, 10, thin_courses)
    assert first == second, "Approved tiny-course clash reports also retain their exact row order"
    assert _build_room_qa(schedule, []) == _build_room_qa(shuffled, [])
    sections = {
        code: [
            {"section": label, "section_key": label, "student_count": 30, "gender": "F"}
            for label in ["F2", "F1"]
        ]
        for code in enrolled
    }
    rooms = [{"room_code": "F10", "section": "F", "capacity": 10}]
    assert check_room_feasibility(sections, rooms) == check_room_feasibility(
        {code: list(reversed(rows)) for code, rows in reversed(list(sections.items()))}, rooms
    )


def _fingerprint_inputs():
    return dict(
        result={
            "enrollment_source": "scraper_timetable",
            "credit_map": {"CS111": 3},
            "section_enrollment": {
                "CS111": [
                    {"section_key": "a", "membership_fingerprint": "aaa", "student_count": 2},
                    {"section_key": "b", "membership_fingerprint": "bbb", "student_count": 1},
                ]
            },
            "buckets_summary": [
                {"program": "DS", "courses": ["DS113", "CS111"]},
                {"program": "AI", "courses": ["CS111"]},
            ],
            "enrollment_scope": {"programs": ["DS", "AI"], "sections": ["M", "F"]},
            "assign_rooms": True,
            "seed": 17,
        },
        enrolled_sets={"CS111": {1, 2, 3}},
        course_meta={"CS111": {"course_identity": "cs111-programming", "programs": ["DS", "AI"]}},
        student_attribution=[
            {"student_id": 2, "program": "AI", "section": "F"},
            {"student_id": 1, "program": "DS", "section": "F"},
        ],
        rooms=[
            {"room_code": "F2", "capacity": 12, "section": "F", "building": "B2"},
            {"room_code": "F1", "capacity": 12, "section": "F", "building": "B1"},
        ],
        days=DAYS,
        periods=PERIODS,
        max_per_day=2,
        thin_conflict_threshold=0,
    )


def test_fingerprint_ignores_unordered_input_traversal_without_mutating_source():
    inputs = _fingerprint_inputs()
    original = deepcopy(inputs)
    expected = exam_input_fingerprint.fingerprint_exam_inputs(**inputs)
    assert inputs == original
    for value in [
        inputs["rooms"],
        inputs["student_attribution"],
        inputs["course_meta"]["CS111"]["programs"],
        inputs["result"]["section_enrollment"]["CS111"],
        inputs["result"]["buckets_summary"],
        *inputs["result"]["enrollment_scope"].values(),
    ]:
        value.reverse()
    for bucket in inputs["result"]["buckets_summary"]:
        bucket["courses"].reverse()
    assert exam_input_fingerprint.fingerprint_exam_inputs(**inputs) == expected


@pytest.mark.parametrize(
    "change",
    [
        "policy",
        "source",
        "room_capacity",
        "room_gender",
        "room_building",
        "membership",
        "identity",
        "seed",
        "days",
        "periods",
    ],
)
def test_fingerprint_invalidates_real_input_or_policy_changes(change, monkeypatch):
    inputs = _fingerprint_inputs()
    before = exam_input_fingerprint.fingerprint_exam_inputs(**inputs)
    if change == "policy":
        monkeypatch.setattr(
            exam_input_fingerprint,
            "ROOM_ALLOCATION_POLICY_VERSION",
            exam_input_fingerprint.ROOM_ALLOCATION_POLICY_VERSION + 1,
        )
    elif change == "source":
        inputs["result"]["enrollment_source"] = "different_source"
    elif change.startswith("room_"):
        key = change.removeprefix("room_")
        inputs["rooms"][0]["section" if key == "gender" else key] = {
            "capacity": 13,
            "gender": "M",
            "building": "B3",
        }[key]
    elif change == "membership":
        inputs["result"]["section_enrollment"]["CS111"][0]["membership_fingerprint"] = (
            "changed-membership"
        )
    elif change == "identity":
        inputs["course_meta"]["CS111"]["course_identity"] = "cs111-fundamentals"
    elif change == "seed":
        inputs["result"]["seed"] = 18
    else:
        inputs[change] = list(reversed(inputs[change]))
    assert exam_input_fingerprint.fingerprint_exam_inputs(**inputs) != before


def _evaluate_board(schedule_raw, source, *, rebalance):
    return evaluate_exam_schedule(
        days=DAYS,
        periods=PERIODS,
        max_per_day=source["qa"]["max_per_day"],
        schedule_raw=deepcopy(schedule_raw),
        selected_courses=source["courses"],
        assign_rooms=True,
        seed=source["seed"],
        thin_conflict_threshold=0,
        programs=source["enrollment_scope"]["programs"],
        sections=source["enrollment_scope"]["sections"],
        pinned=[],
        rebalance_invigilators=rebalance,
    )


def test_a_rebalanced_board_survives_a_fixed_time_check(tied_inventory):
    """The Optimise path moves exams, then a Check must reproduce that board.

    Optimise re-solves every placement and so runs the invigilator post-pass -
    without it, day totals on a real board went from a spread of 2 to 45. That
    pass is the only way a non-build path can move an exam, so the Build/Check
    agreement this module defends has to survive it. The other tests here use a
    board that is already flat, where the pass exits early without moving
    anything and proves nothing about this.
    """
    built = build_exam_timetable(
        label="Rebalanced",
        days=DAYS,
        periods=PERIODS,
        programs=["DS", "AI"],
        sections=["M", "F"],
        seed=None,
        rebalance_invigilators=False,
        persist=False,
    )
    # Crowd the timetable onto Sunday so the post-pass has real work to do.
    skewed = deepcopy(built["schedule"])
    for index, row in enumerate(sorted(skewed, key=lambda r: r["course_identity"])):
        slot = 0 if index < len(skewed) - 1 else 2
        row["slot_index"] = slot
        row["day"] = DAYS[slot // len(PERIODS)]
        row["period"] = PERIODS[slot % len(PERIODS)]

    optimised = _evaluate_board(skewed, built, rebalance=True)
    assert optimised["qa"]["rebalance_moves"] > 0, (
        "The fixture no longer exercises the post-pass; this test would prove nothing."
    )

    checked = _evaluate_board(optimised["schedule"], built, rebalance=False)
    assert _placements(checked) == _placements(optimised)
    assert _room_rows(checked) == _room_rows(optimised)
    assert _report(checked) == _report(optimised)
    assert _visible_report(checked) == _visible_report(optimised)
    assert checked["input_fingerprint"] == optimised["input_fingerprint"]
    assert ExamTimetableRun.objects.count() == 0
