"""Exam student data as Excel: one schema, rows chosen by scope, every row verified.

The workbook is the committee's student-by-exam list, rebuilt from the live
registrar lists and checked group by group against the saved timetable (see
``exam_rosters``). It is written with openpyxl's write-only mode: the student
fact table can hold 20,000+ rows and normal mode costs ~100 MB beside a
CP-SAT solve on a 512 MB instance.

Rules this module enforces
--------------------------
* **One schema.** Sheets, column order and Table names never depend on the
  scope; only the row counts do. Two layouts: Full and Summaries only (no
  names or IDs).
* **Allowed student fields only** (``STUDENT_EXPORT_FIELDS``): Student ID,
  Name, Program, Section, Exam room and the Clash / Same day flags. Department
  derives from the programme and Group is the sitting's cohort - timetable
  facts, not new personal data. No level column exists anywhere.
* **The scope picks rows; it never changes what a row says.** Flags and every
  "(timetable)" count come from the whole run; ID ranges and "in this file"
  figures come from the file's own rows, so a CS file never shows another
  programme's IDs, not even as the end of a range.
* **Every difference is listed where it applies.** Sections and the
  ChangeLog list each section of the file's exams, groups and programmes that
  differs from the save - its students or its programme counts - whether or
  not it still has a row in the file, so check 11 can always be reconciled.
  "Every section matches" is said only when the whole timetable does.
* **Excel safety.** IDs are integers formatted ``0`` (never scientific or
  ``#,##0``); a string starting with ``=`` is forced to text; every string
  loses what XML 1.0 cannot hold (U+FFFE, U+FFFF too); no merged cells,
  no hidden anything, sheet names are fixed strings; ``ws.max_row`` is never
  used (it is quadratic in write-only mode); the write-only spool files, which
  hold student rows on local disk, are removed even when writing fails.
* **Deterministic data sheets.** Nothing volatile is written outside About,
  File info, the page header and the document properties, so the same run,
  lists and options give identical data sheets and the same ``data_sha256``.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import warnings
from collections import Counter, defaultdict
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time
from io import BytesIO
from typing import Any
from zipfile import ZIP_STORED, ZipFile

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.packaging.custom import IntProperty, StringProperty
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet._writer import ALL_TEMP_FILES
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.services.exam_department_export import (
    _DEPARTMENTS,
    ExamDateError,
    day_weekday,
    parse_exam_dates,
)
from core.services.exam_rosters import (
    BASIS_NO_SEAT,
    BASIS_NOT_SCHEDULED,
    BASIS_SPLIT,
    BASIS_UNASSIGNED,
    BASIS_WHOLE,
    CHANGED,
    GENDERS,
    GONE,
    MATCHES,
    NEW,
    UNSEATED_BASES,
    ExamFacts,
    RoomPart,
    RosterModel,
    SectionGroup,
    SittingFlags,
    natural_key,
)
from core.services.student_sections import arabic_term_section_course_names

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ZIP_TYPE = "application/zip"
FORMAT = "exam_student_data_xlsx"
FORMAT_VERSION = 1

#: Every student field a file may carry (owner decision 3: no level). Anything
#: else written is a timetable fact or an integrity marker.
STUDENT_EXPORT_FIELDS = (
    "student_id",
    "name",
    "program",
    "section",
    "exam_room",
    "clash",
    "same_day",
)

SCOPE_KINDS = ("section", "course", "room", "period", "day", "all")
ROWS_OPTIONS = ("all", "flagged")
CONTENTS_OPTIONS = ("full", "summary")
LANGUAGES = ("ar", "en")
PREPARED_FOR_MAX = 80

SEAT_RULE = "student_id_ascending_fill_saved_part_sizes"
FLAG_RULE = "same_slot|same_day_other_slot;whole_run;overflow_excluded"

NAVY = "213A59"
SLATE = "53657A"
TEAL = "227F86"
MIST = "8796A8"
DANGER = "B42332"
DANGER_SOFT = "FFF3F3"
WARN = "85530A"
WARN_SOFT = "FFF8E6"
OK_TEXT = "087F72"
OK_SOFT = "EDF7F5"
BODY = "20334D"

# Excel does not implement the Unicode 6.3 isolates (U+2066..U+2069): it paints
# them as visible LRI/PDI boxes and still reverses the digits (measured through
# Excel's own PDF export). A left-to-right mark before a code, time or ID run
# and a right-to-left mark after it keep 2026-09-24 00:10 and 08:00-10:00 in
# reading order inside Arabic text, with nothing visible.
_LRM, _RLM = "\u200e", "\u200f"
_JOIN = " · "


class ExportOptionsError(ValueError):
    """A request option is invalid; answered 400 against the named field."""

    code = "invalid_options"

    def __init__(self, message: str, *, field: str = "", day: str = "") -> None:
        self.field = field
        # For ``field == "dates"``: the day whose date is refused, when one is.
        self.day = day
        super().__init__(message)


class EmptyScope(ExportOptionsError):
    code = "empty_scope"


# ── Vocabulary (EN, AR) ────────────────────────────────────────

_WORDS: dict[str, tuple[str, str]] = {
    "yes": ("Yes", "نعم"),
    "no": ("No", "لا"),
    "M": ("Male", "طلاب"),
    "F": ("Female", "طالبات"),
    "U": ("Not recorded", "غير مسجل"),
    BASIS_WHOLE: ("Whole section", "الشعبة كاملة"),
    BASIS_SPLIT: ("Split by student ID", "مقسّمة حسب الرقم الجامعي"),
    BASIS_NO_SEAT: ("No seat", "بلا مقعد"),
    BASIS_UNASSIGNED: ("Room not assigned", "لم تُحدَّد قاعة"),
    BASIS_NOT_SCHEDULED: ("Not scheduled", "غير مجدول"),
    MATCHES: ("Matches", "مطابقة"),
    CHANGED: ("Changed", "تغيّرت"),
    NEW: ("New since save", "جديدة بعد الحفظ"),
    GONE: ("Gone since save", "غير موجودة الآن"),
    "section_changed": ("Section changed", "تغيّرت الشعبة"),
    "mix_changed": ("Program mix changed", "تغيّر توزيع البرامج"),
    "exam_missing": ("Exam no longer in the lists", "الاختبار لم يعد في القوائم"),
    "mapped": ("Mapped", "مسجلة"),
    "missing": ("Not recorded", "غير مسجلة"),
    "ambiguous": ("Unclear", "غير محددة"),
    "clash": ("Clash", "تعارض"),
    "same_day": ("Same day", "اليوم نفسه"),
    "ok": ("OK", "مطابق"),
    "differs": ("Differs", "مختلف"),
    "info": ("Info", "معلومة"),
    "other_department": ("Other", "أخرى"),
    "room_not_assigned": ("Not assigned", "لم تُحدَّد"),
    "total": ("Total", "المجموع"),
}

_WEEKDAY_NAMES = (
    ("Monday", "الاثنين"),
    ("Tuesday", "الثلاثاء"),
    ("Wednesday", "الأربعاء"),
    ("Thursday", "الخميس"),
    ("Friday", "الجمعة"),
    ("Saturday", "السبت"),
    ("Sunday", "الأحد"),
)

_ROLE_NAMES = {
    "EXAM_COMMITTEE": ("Exam Committee", "لجنة الاختبارات"),
    "SUPER_ADMIN": ("Super admin", "مدير النظام"),
}


def _w(key: str, lang: str) -> str:
    en, ar = _WORDS[key]
    return ar if lang == "ar" else en


def _pick(lang: str, en: str, ar: str) -> str:
    return ar if lang == "ar" else en


def _iso(text: object, lang: str) -> str:
    """Keep a code, time or ID run in reading order inside an Arabic sentence."""
    return f"{_LRM}{text}{_RLM}" if lang == "ar" else str(text)


def _department(program: str) -> tuple[str, str, str]:
    for key, english, arabic, members in _DEPARTMENTS:
        if program in members:
            return key, english, arabic
    return "", _WORDS["other_department"][0], _WORDS["other_department"][1]


def _department_label(program: str, lang: str) -> str:
    _key, english, arabic = _department(program)
    return arabic if lang == "ar" else english


def _department_rank(program: str) -> int:
    for index, (_key, _en, _ar, members) in enumerate(_DEPARTMENTS):
        if program in members:
            return index
    return len(_DEPARTMENTS)


# ── Sheet and column specifications ────────────────────────────


@dataclass(frozen=True)
class Col:
    key: str
    en: str
    ar: str
    kind: str  # id | int | count | date | time | pct | delta | text | enum
    width: float
    meaning_en: str
    meaning_ar: str
    example: str = ""
    width_ar: float | None = None

    def header(self, lang: str) -> str:
        return self.ar if lang == "ar" else self.en


@dataclass(frozen=True)
class TableSpec:
    name: str
    columns: tuple[Col, ...]
    one_row_en: str
    one_row_ar: str

    def headers(self, lang: str) -> list[str]:
        return [column.header(lang) for column in self.columns]

    def index(self, key: str) -> int:
        return next(i for i, column in enumerate(self.columns) if column.key == key)

    def letter(self, key: str) -> str:
        return get_column_letter(self.index(key) + 1)


@dataclass(frozen=True)
class SheetSpec:
    key: str
    en: str
    ar: str
    tab: str
    tables: tuple[TableSpec, ...]
    freeze: str | None = None
    paper: int = 9  # A4; A3 is 8
    landscape: bool = True
    in_summary: bool = True

    def title(self, lang: str) -> str:
        return self.ar if lang == "ar" else self.en


# fmt: off
_ID = ("Student ID", "الرقم الجامعي", "id", 11, "The student's university ID.", "الرقم الجامعي للطالب.", "7 digits")
_NAME = ("Name", "الاسم", "text", 36, "The name as recorded by the registrar.", "الاسم كما سجلته عمادة القبول.", "AS RECORDED")
_PROGRAM = ("Program", "البرنامج", "text", 9, "The student's programme.", "برنامج الطالب.", "CS")
_DEPT = ("Department", "القسم", "enum", 22, "The department the programme belongs to.", "القسم الذي يتبعه البرنامج.", "Computer Science")
_GROUP = ("Group", "الفئة", "enum", 10, "Male or Female cohort of the exam sitting.", "فئة جلسة الاختبار: طلاب أو طالبات.", "Female")
_EXAM = ("Exam", "الاختبار", "text", 13, "The exam as shown on the timetable.", "الاختبار كما يظهر في الجدول.", "PHYS103 (2)")
_CODE = ("Course code", "رمز المقرر", "text", 10, "The registrar course code.", "رمز المقرر لدى عمادة القبول.", "PHYS103")
_CNAME = ("Course name", "اسم المقرر", "text", 32, "The course name.", "اسم المقرر.", "GENERAL PHYSICS")
_SECTION = ("Section", "الشعبة", "text", 9, "The teaching section; Not recorded or Unclear when the lists name none or several.", "الشعبة الدراسية؛ غير مسجلة أو غير محددة إذا لم تذكر القوائم شعبة أو ذكرت أكثر من شعبة.", "M7")
_DAYNO = ("Day no.", "ترتيب اليوم", "int", 8, "Position of the exam day in the timetable; empty when not scheduled.", "ترتيب يوم الاختبار في الجدول؛ فارغ إذا لم يُجدول.", "1")
_DAY = ("Day", "اليوم", "text", 10, "The timetable day label.", "تسمية اليوم في الجدول.", "W1-Sun")
_WEEKDAY = ("Weekday", "يوم الأسبوع", "enum", 11, "The weekday of the exam day.", "يوم الأسبوع.", "Sunday")
_DATE = ("Date", "التاريخ", "date", 11, "The exam date entered for this export; empty when not set.", "تاريخ الاختبار المُدخل لهذا التصدير؛ فارغ إذا لم يُحدد.", "2026-12-13")
_PERIOD = ("Period", "الفترة", "text", 12, "The exam period.", "فترة الاختبار.", "08:00-10:00")
_START = ("Start", "وقت البدء", "time", 7, "The start time of the period.", "وقت بدء الفترة.", "08:00")
_ONLINE = ("Online", "عن بُعد", "enum", 7, "Yes when the exam is online (online exams still hold rooms).", "نعم إذا كان الاختبار عن بُعد (وتبقى له قاعات).", "No")
_BUILDING = ("Building", "المبنى", "text", 9, "The building as saved with the room.", "المبنى كما حُفظ مع القاعة.", "172")
_FLOOR = ("Floor", "الدور", "int", 6, "The floor when recorded as a number.", "الدور إذا سُجل رقماً.", "1")


def _c(key: str, spec: tuple, **overrides: Any) -> Col:
    en, ar, kind, width, meaning_en, meaning_ar, example = spec
    values = {
        "key": key,
        "en": en,
        "ar": ar,
        "kind": kind,
        "width": width,
        "meaning_en": meaning_en,
        "meaning_ar": meaning_ar,
        "example": example,
    }
    values.update(overrides)
    return Col(**values)


def _col(
    key: str,
    en: str,
    ar: str,
    kind: str,
    width: float,
    meaning_en: str,
    meaning_ar: str,
    example: str = "",
) -> Col:
    return Col(key, en, ar, kind, width, meaning_en, meaning_ar, example)


STUDENT_EXAMS = TableSpec(
    "StudentExams",
    (
        _c("student_id", _ID),
        _c("name", _NAME),
        _c("program", _PROGRAM),
        _c("department", _DEPT, width_ar=30),
        _c("group", _GROUP),
        _c("exam", _EXAM),
        _c("course_code", _CODE),
        _c("course_name", _CNAME),
        _c("section", _SECTION),
        _c("day_no", _DAYNO),
        _col("day", "Day", "اليوم", "text", 10, "The timetable day label; Not scheduled when the timetable could not place the exam.", "تسمية اليوم في الجدول؛ غير مجدول إذا تعذر وضع الاختبار.", "W1-Sun"),
        _c("weekday", _WEEKDAY),
        _c("date", _DATE),
        _c("period", _PERIOD),
        _c("start", _START),
        _col("exam_room", "Exam room", "قاعة الاختبار", "text", 11, "The room the student sits in; empty unless seated.", "القاعة التي يجلس فيها الطالب؛ فارغة إن لم يكن له مقعد.", "172FA001"),
        _c("building", _BUILDING),
        _c("floor", _FLOOR),
        _col("room_basis", "Room basis", "أساس القاعة", "enum", 20, "How the room was decided: Whole section, Split by student ID, No seat, Room not assigned or Not scheduled.", "كيف حُددت القاعة: الشعبة كاملة، مقسّمة حسب الرقم الجامعي، بلا مقعد، لم تُحدَّد قاعة، غير مجدول.", "Whole section"),
        _c("online", _ONLINE),
        _col("clash", "Clash", "تعارض", "enum", 7, "Yes when the student has another exam in the same period.", "نعم إذا كان للطالب اختبار آخر في الفترة نفسها.", "No"),
        _col("clash_with", "Clash with", "يتعارض مع", "text", 16, "The other exams in the same period (codes only).", "الاختبارات الأخرى في الفترة نفسها (الرموز فقط).", "CS211"),
        _col("same_day", "Same day", "اختبار آخر في اليوم نفسه", "enum", 9, "Yes when the student has another exam that day in a different period.", "نعم إذا كان للطالب اختبار آخر في اليوم نفسه في فترة أخرى.", "Yes"),
        _col("same_day_with", "Same day with", "في اليوم نفسه مع", "text", 22, "The other exams that day, with their start times.", "الاختبارات الأخرى في ذلك اليوم مع أوقات بدئها.", "13:00 GS111"),
        _col("exams_that_day", "Exams that day", "اختبارات ذلك اليوم", "int", 8, "The student's exams that day in the whole timetable.", "عدد اختبارات الطالب في ذلك اليوم في الجدول كاملاً.", "2"),
        _col("change", "Change since save", "تغيّر بعد الحفظ", "enum", 16, "Section changed or New since save when the lists differ from the saved timetable.", "تغيّرت الشعبة أو جديدة بعد الحفظ إذا اختلفت القوائم عن الجدول المحفوظ.", ""),
        _col("first_that_day", "First exam that day", "أول اختبار في اليوم", "int", 8, "1 on the student's earliest row in this file that day, else 0; its sum counts students per day.", "1 في أول سطر للطالب في هذا الملف في ذلك اليوم، وإلا 0؛ مجموعه يعد الطلاب في اليوم.", "1"),
    ),
    "one student sitting one exam",
    "طالب واحد في اختبار واحد",
)

EXAM_STUDENTS = TableSpec(
    "ExamStudents",
    (
        _c("student_id", _ID),
        _c("name", _NAME),
        _c("program", _PROGRAM),
        _c("department", _DEPT, width_ar=30),
        _c("group", _GROUP),
        _col("exams_timetable", "Exams (timetable)", "الاختبارات (الجدول)", "int", 10, "The student's exams in the whole timetable.", "اختبارات الطالب في الجدول كاملاً.", "6"),
        _col("exams_file", "Exams (file)", "الاختبارات (الملف)", "int", 10, "The student's rows in this file.", "أسطر الطالب في هذا الملف.", "2"),
        _col("days_file", "Exam days (file)", "أيام الاختبار (الملف)", "int", 10, "Distinct exam days among this file's rows.", "أيام الاختبار المختلفة في أسطر هذا الملف.", "1"),
        _col("clashes_timetable", "Clashes (timetable)", "التعارضات (الجدول)", "int", 10, "Periods with two or more of the student's exams, whole timetable.", "الفترات التي فيها اختباران أو أكثر للطالب في الجدول كاملاً.", "0"),
        _col("days_2plus_timetable", "Days with 2+ exams (timetable)", "أيام بها اختباران أو أكثر (الجدول)", "int", 12, "Days with two or more of the student's exams, whole timetable.", "الأيام التي فيها اختباران أو أكثر للطالب في الجدول كاملاً.", "1"),
        _col("most_in_day_timetable", "Most exams in a day (timetable)", "أكثر اختبارات في يوم (الجدول)", "int", 11, "The most exams the student has on one day, whole timetable.", "أكبر عدد اختبارات للطالب في يوم واحد في الجدول كاملاً.", "2"),
        _col("first_exam", "First exam (file)", "أول اختبار (الملف)", "text", 10, "Day of the earliest exam in this file.", "يوم أول اختبار في هذا الملف.", "W1-Sun"),
        _col("first_date", "First exam date (file)", "تاريخ أول اختبار (الملف)", "date", 11, "Date of the earliest exam in this file; empty when not set.", "تاريخ أول اختبار في هذا الملف؛ فارغ إذا لم يُحدد.", "2026-12-13"),
        _col("last_exam", "Last exam (file)", "آخر اختبار (الملف)", "text", 10, "Day of the latest exam in this file.", "يوم آخر اختبار في هذا الملف.", "W1-Sun"),
        _col("last_date", "Last exam date (file)", "تاريخ آخر اختبار (الملف)", "date", 11, "Date of the latest exam in this file; empty when not set.", "تاريخ آخر اختبار في هذا الملف؛ فارغ إذا لم يُحدد.", "2026-12-13"),
        _col("unseated_file", "Without a seat (file)", "دون مقعد (الملف)", "int", 10, "Rows in this file with No seat, Room not assigned or Not scheduled.", "أسطر هذا الملف بلا مقعد أو دون قاعة أو غير مجدولة.", "0"),
        _col("schedule_file", "Exam schedule (file)", "جدول الاختبارات (الملف)", "text", 90, "The student's exams in this file: day, start, exam, section, room.", "اختبارات الطالب في هذا الملف: اليوم، وقت البدء، الاختبار، الشعبة، القاعة.", "W1-Sun 08:00 COE211 F4 205FCL005"),
    ),
    "one student",
    "طالب واحد",
)

EXAM_FLAGS = TableSpec(
    "ExamFlags",
    (
        _c("student_id", _ID),
        _c("name", _NAME),
        _c("program", _PROGRAM),
        _c("group", _GROUP),
        _c("day_no", _DAYNO),
        _c("day", _DAY),
        _c("date", _DATE),
        _col("flag", "Flag", "التنبيه", "enum", 10, "Clash when two exams share a period that day, else Same day.", "تعارض إذا اشترك اختباران في فترة ذلك اليوم، وإلا اليوم نفسه.", "Same day"),
        _col("exams_that_day", "Exams that day", "اختبارات ذلك اليوم", "int", 8, "The student's exams that day in the whole timetable.", "عدد اختبارات الطالب في ذلك اليوم في الجدول كاملاً.", "2"),
        _col("exams", "Exams", "الاختبارات", "text", 60, "That day's exams with start times (codes only).", "اختبارات ذلك اليوم مع أوقات البدء (الرموز فقط).", "08:00 STAT301 (1) · 13:00 MATH204"),
    ),
    "one student on one day with two or more exams",
    "طالب واحد في يوم له فيه اختباران أو أكثر",
)

_RANGES_MEANING = ("Rooms in part order with the ID range of this file's students in each; summaries-only files show counts only.", "القاعات بترتيب الأجزاء مع نطاق أرقام طلاب هذا الملف في كل قاعة؛ ملفات الملخصات تعرض الأعداد فقط.")

EXAM_SECTIONS = TableSpec(
    "ExamSections",
    (
        _c("exam", _EXAM),
        _c("course_code", _CODE),
        _c("course_name", _CNAME, width=30),
        _c("section", _SECTION),
        _c("group", _GROUP),
        _c("day_no", _DAYNO),
        _c("day", _DAY),
        _c("date", _DATE),
        _c("period", _PERIOD),
        _c("online", _ONLINE),
        _col("record", "Section record", "سجل الشعبة", "enum", 12, "Mapped, Not recorded or Unclear.", "مسجلة، غير مسجلة، غير محددة.", "Mapped"),
        _col("saved", "Students at save", "الطلاب عند الحفظ", "int", 10, "Students in this section when the timetable was saved (what the exam card shows).", "طلاب الشعبة عند حفظ الجدول (كما في بطاقة الاختبار).", "38"),
        _col("now", "Students now", "الطلاب الآن", "int", 10, "Students in this section in the current lists.", "طلاب الشعبة في القوائم الحالية.", "40"),
        _col("difference", "Difference", "الفرق", "delta", 8, "Students now minus students at save.", "الطلاب الآن ناقص الطلاب عند الحفظ.", "+2"),
        _col("membership", "Membership", "العضوية", "enum", 14, "Matches, Changed, New since save or Gone since save, checked against the saved checksum.", "مطابقة، تغيّرت، جديدة بعد الحفظ، غير موجودة الآن، بالتحقق من البصمة المحفوظة.", "Matches"),
        _col("program_mix", "Program mix", "توزيع البرامج", "enum", 11, "Whether the students per programme match the saved timetable.", "هل يطابق عدد الطلاب لكل برنامج الجدول المحفوظ.", "Matches"),
        _col("selected_saved", "Selected programs at save", "البرامج المختارة عند الحفظ", "int", 12, "Saved students of this file's programmes in this section.", "طلاب برامج هذا الملف في الشعبة عند الحفظ.", "38"),
        _col("rows_file", "Rows in this file", "الأسطر في هذا الملف", "int", 10, "This section's rows in this file.", "أسطر هذه الشعبة في هذا الملف.", "40"),
        _col("no_seat", "No seat", "بلا مقعد", "int", 8, "Rows in this file without a seat.", "أسطر هذا الملف دون مقعد.", "2"),
        _col("clash_students", "Clash students", "طلاب لديهم تعارض", "int", 9, "Rows in this file with a clash.", "أسطر هذا الملف التي فيها تعارض.", "0"),
        _col("same_day_students", "Same-day students", "طلاب لديهم اختبار آخر في اليوم", "int", 10, "Rows in this file with another exam that day.", "أسطر هذا الملف التي لها اختبار آخر في اليوم نفسه.", "3"),
        _col("room_count", "Room count", "عدد القاعات", "int", 8, "Rooms the saved timetable gave this section.", "القاعات التي خصصها الجدول المحفوظ للشعبة.", "2"),
        _col("programs_file", "Programs in this file", "البرامج في هذا الملف", "text", 24, "This file's students of the section by programme.", "طلاب الشعبة في هذا الملف حسب البرنامج.", "CS2 22 · IS2 18"),
        _col("ranges", "Rooms and ID ranges", "القاعات ونطاقات الأرقام", "text", 70, *_RANGES_MEANING, "172FB001 (1 of 2): 20 · No seat (2)"),
    ),
    "one exam section and group",
    "شعبة واحدة وفئة واحدة في اختبار",
)

EXAM_ROOMS = TableSpec(
    "ExamRooms",
    (
        _c("day_no", _DAYNO),
        _c("day", _DAY),
        _c("date", _DATE),
        _c("period", _PERIOD),
        _c("start", _START),
        _col("room", "Room", "القاعة", "text", 11, "The room; Not assigned for students the timetable could not room.", "القاعة؛ لم تُحدَّد للطلاب الذين لم يجد لهم الجدول قاعة.", "205FCL005"),
        _c("building", _BUILDING),
        _c("floor", _FLOOR),
        _c("group", _GROUP),
        _col("capacity", "Capacity", "السعة", "int", 8, "The room's seats.", "مقاعد القاعة.", "37"),
        _col("seated_saved", "Seated at save", "المقاعد عند الحفظ", "int", 10, "Students the saved timetable put in the room.", "الطلاب الذين وضعهم الجدول المحفوظ في القاعة.", "37"),
        _col("seated_now", "Seated now", "المقاعد الآن", "int", 10, "Students seated in the whole room now by the seat rule.", "الطلاب في القاعة كلها الآن حسب قاعدة المقاعد.", "37"),
        _col("free", "Free seats", "المقاعد الشاغرة", "int", 9, "Capacity minus seated now; negative means over capacity.", "السعة ناقص المقاعد الآن؛ القيمة السالبة تعني تجاوز السعة.", "0"),
        _col("use", "Use", "الإشغال", "pct", 7, "Seated now divided by capacity.", "المقاعد الآن مقسومة على السعة.", "100%"),
        _col("exams", "Exams", "الاختبارات", "int", 7, "Exams held in the room in this period.", "الاختبارات في القاعة في هذه الفترة.", "1"),
        _c("online", _ONLINE),
        _col("rows_file", "Rows in this file", "الأسطر في هذا الملف", "int", 10, "This file's students seated in the room.", "طلاب هذا الملف في القاعة.", "37"),
        _col("sections", "Sections in room", "الشعب في القاعة", "text", 44, "Sections seated in the room now, with counts.", "الشعب في القاعة الآن مع أعدادها.", "COE211 F4 (20) + F5 (17)"),
    ),
    "one room in one period",
    "قاعة واحدة في فترة واحدة",
)

DAY_SUMMARY = TableSpec(
    "DaySummary",
    (
        _c("day_no", _DAYNO),
        _c("day", _DAY),
        _c("weekday", _WEEKDAY),
        _c("date", _DATE),
        _col("exams", "Exams", "الاختبارات", "int", 8, "Exams with rows in this file that day.", "الاختبارات التي لها أسطر في هذا الملف في ذلك اليوم.", "16"),
        _col("sections", "Sections", "الشعب", "int", 9, "Section groups with rows in this file that day.", "مجموعات الشعب التي لها أسطر في هذا الملف في ذلك اليوم.", "80"),
        _col("rooms", "Rooms used", "القاعات المستخدمة", "int", 9, "Rooms seating this file's students that day (room-periods).", "القاعات التي يجلس فيها طلاب هذا الملف في ذلك اليوم (قاعة لكل فترة).", "52"),
        _col("seats", "Seats", "المقاعد", "count", 9, "Capacity of those rooms.", "سعة تلك القاعات.", "1,690"),
        _col("sittings", "Sittings", "الجلسات", "count", 9, "Rows in this file that day.", "أسطر هذا الملف في ذلك اليوم.", "2,210"),
        _col("students", "Students", "الطلاب", "int", 9, "Distinct students in this file that day.", "الطلاب المختلفون في هذا الملف في ذلك اليوم.", "2,050"),
        _col("male", "Male", "طلاب", "int", 8, "Distinct male students.", "الطلاب (الذكور) المختلفون.", "890"),
        _col("female", "Female", "طالبات", "int", 8, "Distinct female students.", "الطالبات المختلفات.", "1,160"),
        _col("two_plus", "Students with 2+ exams", "طلاب لديهم اختباران أو أكثر", "int", 12, "Students in this file with two or more exams that day (whole timetable).", "طلاب هذا الملف الذين لهم اختباران أو أكثر في ذلك اليوم (الجدول كاملاً).", "160"),
        _col("clash_students", "Clash students", "طلاب لديهم تعارض", "int", 9, "Students in this file with a clash that day.", "طلاب هذا الملف الذين لديهم تعارض في ذلك اليوم.", "0"),
    ),
    "one exam day",
    "يوم اختبار واحد",
)

PERIOD_SUMMARY = TableSpec(
    "PeriodSummary",
    (
        _c("day_no", _DAYNO),
        _c("day", _DAY),
        _c("date", _DATE),
        _c("period", _PERIOD),
        _c("start", _START),
        _col("exams", "Exams", "الاختبارات", "int", 8, "Exams with rows in this file in the period.", "الاختبارات التي لها أسطر في هذا الملف في الفترة.", "5"),
        _col("rooms", "Rooms used", "القاعات المستخدمة", "int", 9, "Rooms seating this file's students in the period.", "القاعات التي يجلس فيها طلاب هذا الملف في الفترة.", "18"),
        _col("seats", "Seats", "المقاعد", "count", 9, "Capacity of those rooms.", "سعة تلك القاعات.", "610"),
        _col("sittings", "Sittings", "الجلسات", "count", 9, "Rows in this file in the period.", "أسطر هذا الملف في الفترة.", "590"),
        _col("students", "Students", "الطلاب", "int", 9, "Distinct students in this file in the period.", "الطلاب المختلفون في هذا الملف في الفترة.", "590"),
        _col("clash_students", "Clash students", "طلاب لديهم تعارض", "int", 9, "Students in this file with a clash in the period.", "طلاب هذا الملف الذين لديهم تعارض في الفترة.", "0"),
        _col("use", "Use", "الإشغال", "pct", 7, "Students seated now in those whole rooms divided by their capacity.", "الطلاب في تلك القاعات كلها الآن مقسوماً على سعتها.", "97%"),
    ),
    "one exam period",
    "فترة اختبار واحدة",
)

_PROGRAM_DAYS_FIXED = (
    _c("program", _PROGRAM),
    _c("department", _DEPT, width_ar=30),
    _col("students", "Students", "الطلاب", "int", 9, "Distinct students of the programme in this file.", "طلاب البرنامج المختلفون في هذا الملف.", "324"),
)
_PROGRAM_DAYS_DAY = _col("day_students", "", "", "int", 9, "Distinct students of the programme sitting that day in this file.", "طلاب البرنامج الذين لهم اختبار في ذلك اليوم في هذا الملف.", "120")
_PROGRAM_DAYS_NS = _col("not_scheduled", "Not scheduled", "غير مجدول", "int", 10, "Distinct students of the programme with an exam the timetable could not place.", "طلاب البرنامج الذين لهم اختبار لم يُجدول.", "0")

EXPORT_CHECKS = TableSpec(
    "ExportChecks",
    (
        _col("check", "Check", "التحقق", "text", 44, "What is compared.", "ما تجري مقارنته.", "Courses"),
        _col("saved", "Saved (timetable screen)", "المحفوظ (شاشة الجدول)", "int", 14, "The value the saved timetable shows.", "القيمة التي يعرضها الجدول المحفوظ.", "172"),
        _col("now", "Now / this file", "الآن / هذا الملف", "int", 14, "The value from the current lists or this file.", "القيمة من القوائم الحالية أو من هذا الملف.", "172"),
        _col("result", "Result", "النتيجة", "enum", 9, "OK, Differs or Info.", "مطابق، مختلف، معلومة.", "OK"),
        _col("where", "Where on the screen", "مكانها في الشاشة", "text", 30, "Where to find the saved value on the timetable page.", "مكان القيمة المحفوظة في صفحة الجدول.", "Full summary › Courses"),
        _col("note", "Note", "ملاحظة", "text", 50, "What the check covers.", "ما يغطيه التحقق.", "" ),
    ),
    "one check",
    "تحقق واحد",
)

CHANGE_LOG = TableSpec(
    "ChangeLog",
    (
        _c("exam", _EXAM),
        _c("section", _SECTION),
        _c("group", _GROUP),
        _col("saved", "Students at save", "الطلاب عند الحفظ", "int", 10, "Students when the timetable was saved.", "الطلاب عند حفظ الجدول.", "38"),
        _col("now", "Students now", "الطلاب الآن", "int", 10, "Students in the current lists.", "الطلاب في القوائم الحالية.", "40"),
        _col("difference", "Difference", "الفرق", "delta", 8, "Students now minus students at save.", "الطلاب الآن ناقص الطلاب عند الحفظ.", "+2"),
        _col("what", "What changed", "ما الذي تغيّر", "enum", 22, "Changed, New since save, Gone since save, Program mix changed or Exam no longer in the lists.", "تغيّرت، جديدة بعد الحفظ، غير موجودة الآن، تغيّر توزيع البرامج، الاختبار لم يعد في القوائم.", "Changed"),
        _col("effect", "Effect in this file", "الأثر في هذا الملف", "text", 50, "What the change does to this file's rows.", "أثر التغيير على أسطر هذا الملف.", ""),
        _col("todo", "What to do", "الإجراء", "text", 44, "How to bring the timetable up to date.", "كيف يُحدَّث الجدول.", ""),
    ),
    "one changed section",
    "شعبة واحدة تغيّرت",
)

FILE_INFO = TableSpec(
    "FileInfo",
    (
        _col("key", "Key", "المفتاح", "text", 30, "A fixed English key.", "مفتاح ثابت بالإنجليزية.", "run_id"),
        _col("value", "Value", "القيمة", "text", 60, "The value.", "القيمة.", "484"),
        _col("meaning", "Meaning", "المعنى", "text", 60, "What the key means.", "معنى المفتاح.", "" ),
    ),
    "one fact about this file",
    "معلومة واحدة عن هذا الملف",
)

COLUMN_GUIDE = TableSpec(
    "ColumnGuide",
    (
        _col("sheet", "Sheet", "الورقة", "text", 22, "The sheet.", "الورقة.", "Students"),
        _col("column", "Column", "العمود", "text", 30, "The column header.", "عنوان العمود.", "Student ID"),
        _col("meaning", "Meaning", "المعنى", "text", 70, "What the column holds.", "ما يحتويه العمود.", ""),
        _col("type", "Type", "النوع", "text", 16, "The kind of value.", "نوع القيمة.", "Number"),
        _col("example", "Example", "مثال", "text", 26, "An illustrative value.", "قيمة توضيحية.", "" ),
    ),
    "one column of one sheet",
    "عمود واحد في ورقة",
)

SHEET_GUIDE = TableSpec(
    "SheetGuide",
    (
        _col("sheet", "Sheet", "الورقة", "text", 34, "", "", ""),
        _col("rows", "Rows", "الأسطر", "int", 14, "", "", ""),
        _col("one_row", "One row is…", "السطر الواحد…", "text", 90, "", "", ""),
    ),
    "",
    "",
)

# fmt: on
SHEETS: tuple[SheetSpec, ...] = (
    SheetSpec("about", "About", "حول الملف", NAVY, (SHEET_GUIDE,), landscape=False),
    SheetSpec(
        "student_exams",
        "Student exams",
        "اختبارات الطلاب",
        NAVY,
        (STUDENT_EXAMS,),
        "C2",
        paper=8,
        in_summary=False,
    ),
    SheetSpec("students", "Students", "الطلاب", NAVY, (EXAM_STUDENTS,), "C2", in_summary=False),
    SheetSpec("flags", "Flags", "التنبيهات", NAVY, (EXAM_FLAGS,), "C2", in_summary=False),
    SheetSpec("sections", "Sections", "الشعب", SLATE, (EXAM_SECTIONS,), "E2"),
    SheetSpec("rooms", "Rooms", "القاعات", SLATE, (EXAM_ROOMS,), "G2"),
    SheetSpec("by_day", "By day", "حسب اليوم", TEAL, (DAY_SUMMARY, PERIOD_SUMMARY)),
    SheetSpec("program_days", "Programs by day", "البرامج حسب الأيام", TEAL, ()),
    SheetSpec(
        "checks", "Checks and changes", "المطابقة والتغييرات", SLATE, (EXPORT_CHECKS, CHANGE_LOG)
    ),
    SheetSpec("file_info", "File info", "بيانات الملف", MIST, (FILE_INFO, COLUMN_GUIDE)),
)
SHEET_BY_KEY = {sheet.key: sheet for sheet in SHEETS}

FILE_INFO_KEYS = (
    "format",
    "format_version",
    "generator_commit",
    "run_id",
    "run_label",
    "run_saved_at_utc",
    "run_status",
    "run_input_fingerprint16",
    "enrollment_source",
    "academic_year",
    "term",
    "lists_code_saved",
    "lists_code_now",
    "lists_match",
    "sections_total",
    "sections_matching",
    "sections_changed",
    "sections_new",
    "sections_gone",
    "lists_checked_at_utc",
    "scope_kind",
    "scope_value",
    "programs",
    "groups",
    "file_group",
    "file_part",
    "rows_option",
    "contents",
    "language",
    "exam_dates_source",
    "prepared_for",
    "rows",
    "students",
    "exams",
    "sections_in_file",
    "data_sha256",
    "seat_rule",
    "flag_rule",
    "generated_at_utc",
    "generated_by",
    "reference",
    "audit_entry_hash16",
)

_FILE_INFO_MEANING: dict[str, tuple[str, str]] = {
    "format": ("The file format.", "صيغة الملف."),
    "format_version": ("The format version.", "إصدار الصيغة."),
    "generator_commit": (
        "The application version that wrote the file.",
        "إصدار التطبيق الذي كتب الملف.",
    ),
    "run_id": ("The saved exam timetable.", "الجدول المحفوظ."),
    "run_label": ("The saved timetable's name.", "اسم الجدول المحفوظ."),
    "run_saved_at_utc": (
        "When the timetable was saved (UTC).",
        "وقت حفظ الجدول (بالتوقيت العالمي).",
    ),
    "run_status": ("The saved timetable's status.", "حالة الجدول المحفوظ."),
    "run_input_fingerprint16": (
        "The start of the timetable's input checksum.",
        "بداية بصمة مدخلات الجدول.",
    ),
    "enrollment_source": ("Where the student lists come from.", "مصدر قوائم الطلاب."),
    "academic_year": ("The academic year of the lists.", "العام الجامعي للقوائم."),
    "term": ("The term of the lists.", "الفصل الدراسي للقوائم."),
    "lists_code_saved": (
        "Checksum of every section list when the timetable was saved.",
        "بصمة قوائم الشعب عند حفظ الجدول.",
    ),
    "lists_code_now": (
        "Checksum of every section list now; equal to the saved code exactly when nothing changed.",
        "بصمة قوائم الشعب الآن؛ تساوي البصمة المحفوظة فقط إذا لم يتغيّر شيء.",
    ),
    "lists_match": (
        "Whether every section's students match the saved timetable (the two lists codes are equal); programme counts are checked on Checks and changes.",
        "هل يطابق طلاب كل شعبة الجدول المحفوظ (تتساوى بصمتا القوائم)؛ أعداد البرامج تُطابَق في صفحة المطابقة والتغييرات.",
    ),
    "sections_total": (
        "Section groups in the saved timetable and the lists.",
        "مجموعات الشعب في الجدول المحفوظ والقوائم.",
    ),
    "sections_matching": (
        "Section groups identical to the saved timetable.",
        "مجموعات الشعب المطابقة للجدول المحفوظ.",
    ),
    "sections_changed": (
        "Section groups whose students changed.",
        "مجموعات الشعب التي تغيّر طلابها.",
    ),
    "sections_new": ("Section groups new since the save.", "مجموعات الشعب الجديدة بعد الحفظ."),
    "sections_gone": ("Section groups gone since the save.", "مجموعات الشعب غير الموجودة الآن."),
    "lists_checked_at_utc": (
        "When the lists were read (UTC).",
        "وقت قراءة القوائم (بالتوقيت العالمي).",
    ),
    "scope_kind": ("Which exams the file covers.", "الاختبارات التي يغطيها الملف."),
    "scope_value": (
        "The exam, section, room, period or day chosen.",
        "الاختبار أو الشعبة أو القاعة أو الفترة أو اليوم المختار.",
    ),
    "programs": ("The programmes chosen; empty means all.", "البرامج المختارة؛ الفراغ يعني الكل."),
    "groups": ("The groups chosen.", "الفئات المختارة."),
    "file_group": (
        "The group of this file when there is one file per group.",
        "فئة هذا الملف عند وجود ملف لكل فئة.",
    ),
    "file_part": (
        "This file's place among the files of one export.",
        "ترتيب هذا الملف بين ملفات التصدير.",
    ),
    "rows_option": (
        "Every student, or only students with a flag.",
        "كل الطلاب أو الطلاب ذوو التنبيهات فقط.",
    ),
    "contents": (
        "Full, or summaries only without names or IDs.",
        "كامل، أو ملخصات فقط دون أسماء أو أرقام.",
    ),
    "language": ("The file's language.", "لغة الملف."),
    "exam_dates_source": ("Where the exam dates come from.", "مصدر تواريخ الاختبارات."),
    "prepared_for": ("Who the file was prepared for.", "الجهة التي أُعد لها الملف."),
    "rows": ("Student-exam rows in this file.", "أسطر الطلاب والاختبارات في هذا الملف."),
    "students": ("Distinct students in this file.", "الطلاب المختلفون في هذا الملف."),
    "exams": ("Exams in this file.", "الاختبارات في هذا الملف."),
    "sections_in_file": ("Section groups in this file.", "مجموعات الشعب في هذا الملف."),
    "data_sha256": (
        "sha256 of the Student exams rows (Sections rows in a summaries-only file); recompute it to detect edits.",
        "بصمة sha256 لأسطر اختبارات الطلاب (أسطر الشعب في ملف الملخصات)؛ أعد حسابها لكشف أي تعديل.",
    ),
    "seat_rule": (
        "How students of a split section are seated.",
        "كيفية توزيع طلاب الشعبة المقسّمة على القاعات.",
    ),
    "flag_rule": ("How Clash and Same day are decided.", "كيفية تحديد التعارض واليوم نفسه."),
    "generated_at_utc": ("When the file was made (UTC).", "وقت إنشاء الملف (بالتوقيت العالمي)."),
    "generated_by": ("Who made the file.", "منشئ الملف."),
    "reference": (
        "Quote it with any question about this file.",
        "اذكره مع أي استفسار عن هذا الملف.",
    ),
    "audit_entry_hash16": ("The start of the audit record's hash.", "بداية بصمة سجل التدقيق."),
}

_TYPE_NAMES = {
    "id": ("Student ID (number)", "رقم جامعي (عدد)"),
    "int": ("Whole number", "عدد صحيح"),
    "count": ("Whole number", "عدد صحيح"),
    "date": ("Date", "تاريخ"),
    "time": ("Time", "وقت"),
    "pct": ("Percent", "نسبة مئوية"),
    "delta": ("Change (+/−)", "فرق (+/−)"),
    "text": ("Text", "نص"),
    "enum": ("One of a fixed list", "قيمة من قائمة ثابتة"),
}

_NUMBER_FORMATS = {
    "id": "0",
    "date": "yyyy-mm-dd",
    "time": "hh:mm",
    "pct": "0%",
    "delta": "+0;-0;0",
    "count": "#,##0",
}


# ── Options ────────────────────────────────────────────────────


@dataclass(frozen=True)
class ExportOptions:
    scope_kind: str
    exam: str = ""
    section_key: str = ""
    gender: str = ""
    slot_index: int | None = None
    room_code: str = ""
    day: str = ""
    programs: tuple[str, ...] = ()
    groups: tuple[str, ...] = GENDERS
    one_file_per_group: bool = False
    rows: str = "all"
    contents: str = "full"
    language: str = "ar"
    dates: dict[str, date] = field(default_factory=dict)
    prepared_for: str = ""


_OPTION_KEYS = {
    "scope",
    "programs",
    "groups",
    "one_file_per_group",
    "rows",
    "contents",
    "language",
    "dates",
    "prepared_for",
    "pickers",
    "known_choices",
}
_SCOPE_FIELDS = {
    "section": ("exam", "section_key", "gender"),
    "course": ("exam",),
    "room": ("slot_index", "room_code"),
    "period": ("slot_index",),
    "day": ("day",),
    "all": (),
}


def available_programs(model: RosterModel) -> list[str]:
    """Programmes a filter may name: in the saved run or the live lists."""
    programs = {facts.program for facts in model.students.values()}
    for group in model.groups:
        programs.update(group.saved_program_counts)
    programs.discard("")
    return sorted(programs, key=lambda program: (_department_rank(program), program))


def scheduled_slots(model: RosterModel) -> dict[int, ExamFacts]:
    """One representative exam per scheduled slot, for its day and period."""
    slots: dict[int, ExamFacts] = {}
    for exam in model.exams.values():
        if exam.scheduled:
            slots.setdefault(exam.slot_index, exam)
    return dict(sorted(slots.items()))


def _scope_values(raw: object, kind: str, model: RosterModel) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise ExportOptionsError("Choose which exams to export.", field="scope")
    values: dict[str, Any] = {}
    for name in _SCOPE_FIELDS[kind]:
        value = raw.get(name)
        if name == "slot_index":
            if isinstance(value, bool) or not isinstance(value, int):
                raise ExportOptionsError("Choose a period.", field=f"scope.{name}")
        elif not isinstance(value, str) or not value.strip():
            raise ExportOptionsError("Complete the exam choice.", field=f"scope.{name}")
        values[name] = value
    if "exam" in values and values["exam"] not in model.exams:
        raise ExportOptionsError("Choose an exam from this timetable.", field="scope.exam")
    if kind == "section":
        keys = {(g.section_key, g.gender) for g in model.groups if g.exam == values["exam"]}
        if (values["section_key"], values["gender"]) not in keys:
            raise ExportOptionsError("Choose a section of this exam.", field="scope.section_key")
    if kind == "room" and (values["slot_index"], values["room_code"]) not in model.saved.rooms:
        raise ExportOptionsError("Choose a room used in that period.", field="scope.room_code")
    if kind == "period" and values["slot_index"] not in scheduled_slots(model):
        raise ExportOptionsError("Choose a period of this timetable.", field="scope.slot_index")
    if kind == "day" and values["day"] not in model.saved.days:
        raise ExportOptionsError("Choose a day of this timetable.", field="scope.day")
    return values


def parse_export_options(
    payload: object, model: RosterModel, *, require_scope: bool = True
) -> ExportOptions | None:
    """Validate a request body. Only IDs of timetable facts travel; never students."""
    if not isinstance(payload, dict):
        raise ExportOptionsError("Export options must be an object.", field="body")
    unknown = set(payload) - _OPTION_KEYS
    if unknown:
        raise ExportOptionsError("Unknown export option.", field=sorted(unknown)[0])
    scope = payload.get("scope")
    if scope is None and not require_scope:
        return None
    if not isinstance(scope, dict) or scope.get("kind") not in SCOPE_KINDS:
        raise ExportOptionsError("Choose which exams to export.", field="scope.kind")
    kind = scope["kind"]
    if set(scope) - {"kind", *_SCOPE_FIELDS[kind]}:
        raise ExportOptionsError("Unknown exam choice.", field="scope")
    values = _scope_values(scope, kind, model)

    programs = payload.get("programs", [])
    available = set(available_programs(model))
    if (
        not isinstance(programs, list)
        or any(not isinstance(program, str) or program not in available for program in programs)
        or len(set(programs)) != len(programs)
    ):
        raise ExportOptionsError(
            "Choose programmes of this timetable without duplicates.", field="programs"
        )
    if set(programs) == available:
        programs = []
    groups = payload.get("groups", list(GENDERS))
    if (
        not isinstance(groups, list)
        or not groups
        or any(not isinstance(group, str) or group not in GENDERS for group in groups)
        or len(set(groups)) != len(groups)
    ):
        raise ExportOptionsError("Tick at least one group.", field="groups")
    one_file = payload.get("one_file_per_group", len(groups) > 1)
    if not isinstance(one_file, bool):
        raise ExportOptionsError(
            "One file per group must be on or off.", field="one_file_per_group"
        )
    rows = payload.get("rows", "all")
    if rows not in ROWS_OPTIONS:
        raise ExportOptionsError("Choose every student or flagged students only.", field="rows")
    contents = payload.get("contents", "full")
    if contents not in CONTENTS_OPTIONS:
        raise ExportOptionsError("Choose Full or Summaries only.", field="contents")
    language = payload.get("language", "ar")
    if language not in LANGUAGES:
        raise ExportOptionsError("Choose Arabic or English for the file.", field="language")
    try:
        dates = parse_exam_dates(payload.get("dates", {}), model.saved.days)
    except ExamDateError as exc:
        raise ExportOptionsError(str(exc), field="dates", day=exc.day) from exc
    prepared_for = payload.get("prepared_for", "")
    if (
        not isinstance(prepared_for, str)
        or len(prepared_for.strip()) > PREPARED_FOR_MAX
        or re.search(r"[\x00-\x1f\x7f]", prepared_for)
        or XML_ILLEGAL_RE.search(prepared_for)
    ):
        raise ExportOptionsError(
            f"Prepared for must be plain text of {PREPARED_FOR_MAX} characters or fewer.",
            field="prepared_for",
        )
    return ExportOptions(
        scope_kind=kind,
        exam=values.get("exam", ""),
        section_key=values.get("section_key", ""),
        gender=values.get("gender", ""),
        slot_index=values.get("slot_index"),
        room_code=values.get("room_code", ""),
        day=values.get("day", ""),
        programs=tuple(sorted(programs, key=lambda p: (_department_rank(p), p))),
        groups=tuple(g for g in GENDERS if g in groups),
        one_file_per_group=one_file,
        rows=rows,
        contents=contents,
        language=language,
        dates=dates,
        prepared_for=" ".join(prepared_for.split()),
    )


# ── Sittings: the rows a file is made of ───────────────────────


@dataclass(frozen=True)
class Sitting:
    student_id: int
    group: SectionGroup
    exam: ExamFacts
    part: RoomPart | None
    basis: str
    seat_rank: int
    flags: SittingFlags | None


def all_sittings(model: RosterModel) -> list[Sitting]:
    """Every live student sitting of the run, in total file order (= seat order)."""
    sittings: list[Sitting] = []
    for group in model.groups:
        exam = model.exams[group.exam]
        rank = {part: index for index, part in enumerate(group.parts)}
        ordered = sorted(
            group.members,
            key=lambda sid: (
                rank.get(group.seats.get(sid), len(group.parts)),  # type: ignore[arg-type]
                sid,
            ),
        )
        for sid in ordered:
            part = group.seats.get(sid)
            sittings.append(
                Sitting(
                    student_id=sid,
                    group=group,
                    exam=exam,
                    part=part,
                    basis=model.basis(group, sid),
                    seat_rank=rank.get(part, len(group.parts)),  # type: ignore[arg-type]
                    flags=model.flags.sittings.get((sid, group.exam)),
                )
            )
    return sittings


def _in_scope(sitting: Sitting, options: ExportOptions) -> bool:
    exam = sitting.exam
    kind = options.scope_kind
    if kind == "all":
        return True
    if kind == "course":
        return exam.code == options.exam
    if kind == "section":
        return (
            exam.code == options.exam
            and sitting.group.section_key == options.section_key
            and sitting.group.gender == options.gender
        )
    if kind == "period":
        return exam.scheduled and exam.slot_index == options.slot_index
    if kind == "day":
        return exam.scheduled and exam.day == options.day
    if kind == "room":
        return (
            exam.scheduled
            and exam.slot_index == options.slot_index
            and sitting.basis in {BASIS_WHOLE, BASIS_SPLIT}
            and sitting.part is not None
            and sitting.part.room_code == options.room_code
        )
    return False


def _selected(
    sitting: Sitting, options: ExportOptions, model: RosterModel, *, check_groups: bool = True
) -> bool:
    if not _in_scope(sitting, options):
        return False
    if check_groups and sitting.group.gender not in options.groups:
        return False
    if options.programs and model.students[sitting.student_id].program not in options.programs:
        return False
    return not (options.rows == "flagged" and not model.flags.flagged(sitting.student_id))


def _group_in_scope(group: SectionGroup, options: ExportOptions, model: RosterModel) -> bool:
    """Whether a section group belongs to the file's exam scope (for groups with no rows)."""
    exam = model.exams[group.exam]
    kind = options.scope_kind
    if kind == "all":
        return True
    if kind == "course":
        return exam.code == options.exam
    if kind == "section":
        return (exam.code, group.section_key, group.gender) == (
            options.exam,
            options.section_key,
            options.gender,
        )
    if kind == "period":
        return exam.scheduled and exam.slot_index == options.slot_index
    if kind == "day":
        return exam.scheduled and exam.day == options.day
    return any(part.room_code == options.room_code for part in group.parts) and (
        exam.slot_index == options.slot_index
    )


# ── Tokens and names ───────────────────────────────────────────


def _token(text: str) -> str:
    """An ASCII file-name token: "PHYS103 (2)" -> "PHYS103-2", ":" dropped.

    Every other run of characters outside ``A-Za-z0-9-`` becomes one hyphen,
    so an Arabic-only label yields "" and the caller supplies an ASCII fallback.
    """
    text = str(text).replace(":", "")
    text = re.sub(r"[^A-Za-z0-9-]+", "-", text)
    return re.sub(r"-{2,}", "-", text).strip("-")


def _slot_token(exam: ExamFacts, options: ExportOptions) -> str:
    when = options.dates.get(exam.day)
    day = when.isoformat() if when else (_token(exam.day) or f"day{exam.day_no}")
    start = exam.start.strftime("%H%M") if exam.start else f"slot{exam.slot_index}"
    return f"{day}-{start}"


def scope_token(options: ExportOptions, model: RosterModel) -> str:
    kind = options.scope_kind
    if kind == "all":
        return "all"
    if kind in {"course", "section"}:
        exam = _token(options.exam) or "exam"
        if kind == "course":
            return exam
        group = next(
            g
            for g in model.groups
            if (g.exam, g.section_key, g.gender)
            == (options.exam, options.section_key, options.gender)
        )
        label = _token(group.section) if group.mapping_status == "mapped" else ""
        if not label:
            label = ("notrecorded" if group.mapping_status == "missing" else "unclear") + (
                f"-{group.gender}"
            )
        return f"{exam}-{label}"
    if kind == "day":
        when = options.dates.get(options.day)
        return when.isoformat() if when else (_token(options.day) or "day")
    exam = scheduled_slots(model)[options.slot_index]  # type: ignore[index]
    slot = _slot_token(exam, options)
    if kind == "period":
        return slot
    return f"{_token(options.room_code) or 'room'}_{slot}"


def programs_token(options: ExportOptions, model: RosterModel) -> str:
    if not options.programs:
        return ""
    available = set(available_programs(model))
    for key, _en, _ar, members in _DEPARTMENTS:
        present = available.intersection(members)
        if present and set(options.programs) == present:
            return f"dept-{key}"
    return "-".join(_token(program) or "p" for program in options.programs)


def file_stem(options: ExportOptions, model: RosterModel, gender: str | None = None) -> str:
    """The ASCII name without its reference: ``exam_students_<scope>..._<lang>[_G]``."""
    parts = ["exam_students", scope_token(options, model)]
    programs = programs_token(options, model)
    if programs:
        parts.append(programs)
    if options.rows == "flagged":
        parts.append("flagged")
    if options.contents == "summary":
        parts.append("summary")
    parts.append(f"r{model.saved.run_id}")
    parts.append(options.language)
    if gender:
        parts.append(gender)
    return "_".join(parts)


def reference_token(reference: str) -> str:
    return reference.removeprefix("EXR-")


# ── Preparing a file: every table's rows, before anything is written ──


@dataclass
class PreparedFile:
    gender: str | None
    part_no: int
    part_total: int
    stem: str
    sittings: list[Sitting]
    groups: list[SectionGroup]
    tables: dict[str, list[tuple]]
    day_headers: tuple[str, ...]
    rows: int
    students: int
    exams: int
    sections: int
    data_sha256: str
    no_changes: bool

    def name(self, reference: str) -> str:
        return f"{self.stem}_{reference_token(reference)}.xlsx"


@dataclass
class PreparedExport:
    model: RosterModel
    options: ExportOptions
    files: list[PreparedFile]
    generated_by: str
    generated_role: str
    generated_at: datetime
    zip_stem: str

    @property
    def rows(self) -> int:
        return sum(item.rows for item in self.files)

    def filename(self, reference: str) -> str:
        if len(self.files) == 1:
            return self.files[0].name(reference)
        return f"{self.zip_stem}_{reference_token(reference)}.zip"

    def audit_details(self) -> dict:
        """What the audit row records: options and counts, never a student."""
        options = self.options
        model = self.model
        return {
            "run_id": model.saved.run_id,
            "scope": {"kind": options.scope_kind, "value": scope_value(options, model)},
            "programs": list(options.programs),
            "groups": list(options.groups),
            "one_file_per_group": options.one_file_per_group,
            "rows_option": options.rows,
            "contents": options.contents,
            "language": options.language,
            "prepared_for": options.prepared_for,
            "exam_dates_source": "entered" if options.dates else "not_set",
            "check": {
                "status": "matches" if model.unchanged else "changed",
                "sections_changed": len(model.changed_groups),
                "program_mix_changed": _mix_only_changes(model),
                "exams_missing": len(model.missing_exams),
                "no_seat": whole_run_no_seat(model),
            },
            "lists_code_saved": model.lists_code_saved,
            "lists_code_now": model.lists_code_now,
            "files": [
                {
                    "name": item.stem,
                    "rows": item.rows,
                    "students": item.students,
                    "data_sha256": item.data_sha256,
                }
                for item in self.files
            ],
        }


def scope_value(options: ExportOptions, model: RosterModel) -> str:
    kind = options.scope_kind
    if kind == "all":
        return ""
    if kind == "course":
        return options.exam
    if kind == "section":
        return f"{options.exam}|{options.section_key}|{options.gender}"
    if kind == "day":
        return options.day
    exam = scheduled_slots(model)[options.slot_index]  # type: ignore[index]
    slot = f"{exam.day} {exam.period}"
    return slot if kind == "period" else f"{options.room_code} {slot}"


def whole_run_no_seat(model: RosterModel) -> int:
    return sum(
        1
        for group in model.groups
        for sid in group.members
        if model.basis(group, sid) == BASIS_NO_SEAT
    )


def _mix_only_changes(model: RosterModel) -> int:
    """Groups with the saved students whose per-programme counts moved."""
    return sum(1 for g in model.groups if g.membership == MATCHES and g.program_mix != MATCHES)


def _data_sha256(rows: Iterable[tuple]) -> str:
    digest = hashlib.sha256()
    for row in rows:
        digest.update(("\x1f".join(_hash_text(value) for value in row) + "\n").encode("utf-8"))
    return digest.hexdigest()


def _hash_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return repr(value)
    return str(value)


def _floor(value: str) -> int | None:
    text = str(value or "").strip()
    return int(text) if re.fullmatch(r"-?\d{1,4}", text) else None


def _section_label(group: SectionGroup, lang: str) -> str:
    if group.mapping_status == "mapped" and group.section:
        return group.section
    return _w("ambiguous" if group.mapping_status == "ambiguous" else "missing", lang)


def _weekday(exam_day: str, when: date | None, lang: str) -> str | None:
    index = when.weekday() if when else day_weekday(exam_day)
    if index is None:
        return None
    return _WEEKDAY_NAMES[index][1 if lang == "ar" else 0]


def _yes(value: bool, lang: str) -> str:
    return _w("yes" if value else "no", lang)


def _id_range(ids: list[int]) -> str:
    low, high = min(ids), max(ids)
    return f"{low} ({len(ids)})" if low == high else f"{low}–{high} ({len(ids)})"


@dataclass
class _Context:
    model: RosterModel
    options: ExportOptions
    lang: str
    arabic_names: dict[str, str]
    room_now: Counter
    room_sections: dict[tuple[int, str], Counter]
    unassigned_now: Counter
    exams_per_student: Counter

    def course_name(self, exam: ExamFacts) -> str:
        if self.lang == "ar":
            return self.arabic_names.get(exam.source_code) or exam.name
        return exam.name

    def date(self, exam: ExamFacts) -> date | None:
        return self.options.dates.get(exam.day) if exam.scheduled else None


def _context(model: RosterModel, options: ExportOptions) -> _Context:
    room_now: Counter = Counter()
    room_sections: dict[tuple[int, str], Counter] = defaultdict(Counter)
    unassigned_now: Counter = Counter()
    exams_per_student: Counter = Counter()
    for group in model.groups:
        exam = model.exams[group.exam]
        for sid in group.members:
            exams_per_student[sid] += 1
            basis = model.basis(group, sid)
            part = group.seats.get(sid)
            if basis in {BASIS_WHOLE, BASIS_SPLIT} and part is not None:
                room_now[(exam.slot_index, part.room_code)] += 1
                room_sections[(exam.slot_index, part.room_code)][
                    (group.exam, group.section_key, group.gender)
                ] += 1
            elif basis == BASIS_UNASSIGNED and exam.scheduled and part is not None:
                unassigned_now[(exam.slot_index, group.gender)] += 1
    arabic_names = (
        arabic_term_section_course_names(sorted({e.source_code for e in model.exams.values()}))
        if options.language == "ar"
        else {}
    )
    return _Context(
        model=model,
        options=options,
        lang=options.language,
        arabic_names=arabic_names,
        room_now=room_now,
        room_sections=room_sections,
        unassigned_now=unassigned_now,
        exams_per_student=exams_per_student,
    )


def prepare_export(
    model: RosterModel,
    options: ExportOptions,
    *,
    generated_by: str,
    generated_role: str,
    now: datetime | None = None,
) -> PreparedExport:
    """Compute every file's rows and ``data_sha256`` - all before the audit row."""
    context = _context(model, options)
    everything = all_sittings(model)
    chosen = [s for s in everything if _selected(s, options, model)]
    if not chosen:
        raise EmptyScope(
            "No students sit exams in this choice. Change the exams, programmes or groups.",
            field="scope",
        )
    if options.one_file_per_group:
        genders = [g for g in GENDERS if any(s.group.gender == g for s in chosen)]
        splits: list[tuple[str | None, list[Sitting]]] = [
            (gender, [s for s in chosen if s.group.gender == gender]) for gender in genders
        ]
    else:
        splits = [(None, chosen)]
    files = []
    for part_no, (gender, sittings) in enumerate(splits, 1):
        file_groups = (gender,) if gender else options.groups
        files.append(
            _prepare_file(
                context,
                sittings,
                file_groups,
                gender=gender,
                part_no=part_no,
                part_total=len(splits),
                stem=file_stem(options, model, gender if options.one_file_per_group else None),
            )
        )
    return PreparedExport(
        model=model,
        options=options,
        files=files,
        generated_by=generated_by,
        generated_role=generated_role,
        generated_at=now or timezone.now(),
        zip_stem=file_stem(options, model),
    )


def _prepare_file(
    context: _Context,
    sittings: list[Sitting],
    file_groups: tuple[str, ...],
    *,
    gender: str | None,
    part_no: int,
    part_total: int,
    stem: str,
) -> PreparedFile:
    model, options = context.model, context.options
    groups_with_rows: dict[tuple[str, str, str], list[Sitting]] = defaultdict(list)
    for sitting in sittings:
        group = sitting.group
        groups_with_rows[(group.exam, group.section_key, group.gender)].append(sitting)
    file_groups_list = [
        group
        for group in model.groups
        if (group.exam, group.section_key, group.gender) in groups_with_rows
        or _changed_in_file(group, options, model, file_groups)
    ]
    tables: dict[str, list[tuple]] = {}
    exam_rows = _student_exam_rows(context, sittings)
    tables["StudentExams"] = exam_rows
    tables["ExamStudents"] = _student_rows(context, sittings)
    tables["ExamFlags"] = _flag_rows(context, sittings)
    tables["ExamSections"] = _section_rows(context, file_groups_list, groups_with_rows)
    tables["ExamRooms"] = _room_rows(context, sittings)
    tables["DaySummary"] = _day_rows(context, sittings)
    tables["PeriodSummary"] = _period_rows(context, sittings)
    day_headers, program_rows = _program_day_rows(context, sittings)
    tables["ProgramDays"] = program_rows
    tables["ExportChecks"] = _check_rows(context, sittings, file_groups)
    change_rows = _change_rows(context, file_groups_list)
    tables["ChangeLog"] = change_rows
    digest_table = "StudentExams" if options.contents == "full" else "ExamSections"
    return PreparedFile(
        gender=gender,
        part_no=part_no,
        part_total=part_total,
        stem=stem,
        sittings=sittings,
        groups=file_groups_list,
        tables=tables,
        day_headers=day_headers,
        rows=len(sittings),
        students=len({s.student_id for s in sittings}),
        exams=len({s.exam.code for s in sittings}),
        sections=len(file_groups_list),
        data_sha256=_data_sha256(tables[digest_table]),
        no_changes=not change_rows,
    )


def _changed_in_file(
    group: SectionGroup, options: ExportOptions, model: RosterModel, file_groups: tuple[str, ...]
) -> bool:
    """A section that differs from the save and belongs to this file, rows or not.

    Its students may all have left the file's programmes (or it is gone), so
    it has no row here; it is still listed on Sections and in the ChangeLog,
    because check 11 counts its saved students and the reader must see why
    the rows differ. The same scope, group and programme tests as check 11.
    """
    if group.membership == MATCHES and group.program_mix == MATCHES:
        return False
    if group.gender not in file_groups or not _group_in_scope(group, options, model):
        return False
    return not options.programs or any(
        group.saved_program_counts.get(program) or group.live_program_counts.get(program)
        for program in options.programs
    )


def _student_exam_rows(context: _Context, sittings: list[Sitting]) -> list[tuple]:
    model, lang = context.model, context.lang
    seen_days: set[tuple[int, str]] = set()
    rows = []
    for sitting in sittings:
        exam, group, sid = sitting.exam, sitting.group, sitting.student_id
        student = model.students[sid]
        seated = sitting.basis in {BASIS_WHOLE, BASIS_SPLIT} and sitting.part is not None
        flags = sitting.flags
        first = None
        if exam.scheduled:
            first = 0 if (sid, exam.day) in seen_days else 1
            seen_days.add((sid, exam.day))
        change = ""
        if group.membership == CHANGED:
            change = _w("section_changed", lang)
        elif group.membership == NEW:
            change = _w(NEW, lang)
        rows.append(
            (
                sid,
                student.name,
                student.program,
                _department_label(student.program, lang),
                _w(group.gender, lang),
                exam.code,
                exam.source_code,
                context.course_name(exam),
                _section_label(group, lang),
                exam.day_no,
                exam.day if exam.scheduled else _w(BASIS_NOT_SCHEDULED, lang),
                _weekday(exam.day, context.date(exam), lang) if exam.scheduled else None,
                context.date(exam),
                exam.period or None,
                exam.start,
                sitting.part.room_code if seated else None,  # type: ignore[union-attr]
                (sitting.part.building or None) if seated else None,  # type: ignore[union-attr]
                _floor(sitting.part.floor) if seated else None,  # type: ignore[union-attr]
                _w(sitting.basis, lang),
                _yes(exam.is_online, lang),
                _yes(bool(flags and flags.clash), lang),
                _JOIN.join(flags.clash_with) if flags and flags.clash else None,
                _yes(bool(flags and flags.same_day), lang),
                _JOIN.join(f"{start} {code}" for start, code in flags.same_day_with)
                if flags and flags.same_day
                else None,
                flags.exams_that_day if flags else None,
                change or None,
                first,
            )
        )
    return rows


def _schedule_text(sitting: Sitting, lang: str) -> str:
    exam = sitting.exam
    when = exam.day if exam.scheduled else _w(BASIS_NOT_SCHEDULED, lang)
    start = exam.start.strftime("%H:%M") if exam.start else ""
    seated = sitting.basis in {BASIS_WHOLE, BASIS_SPLIT} and sitting.part is not None
    room = sitting.part.room_code if seated else _w(sitting.basis, lang)  # type: ignore[union-attr]
    return " ".join(
        part for part in (when, start, exam.code, _section_label(sitting.group, lang), room) if part
    )


def _student_rows(context: _Context, sittings: list[Sitting]) -> list[tuple]:
    model, lang = context.model, context.lang
    by_student: dict[int, list[Sitting]] = defaultdict(list)
    for sitting in sittings:
        by_student[sitting.student_id].append(sitting)
    rows = []
    for sid in sorted(by_student):
        own = by_student[sid]
        student = model.students[sid]
        scheduled = [s for s in own if s.exam.scheduled]
        days = model.flags.days.get(sid, {})
        first = scheduled[0].exam if scheduled else None
        last = scheduled[-1].exam if scheduled else None
        rows.append(
            (
                sid,
                student.name,
                student.program,
                _department_label(student.program, lang),
                _w(own[0].group.gender, lang),
                context.exams_per_student[sid],
                len(own),
                len({s.exam.day for s in scheduled}),
                model.flags.clash_periods(sid),
                sum(1 for exams in days.values() if len(exams) >= 2),
                max((len(exams) for exams in days.values()), default=0),
                first.day if first else None,
                context.date(first) if first else None,
                last.day if last else None,
                context.date(last) if last else None,
                sum(1 for s in own if s.basis in UNSEATED_BASES),
                _JOIN.join(_schedule_text(s, lang) for s in own),
            )
        )
    return rows


def _flag_rows(context: _Context, sittings: list[Sitting]) -> list[tuple]:
    model, lang = context.model, context.lang
    in_file: dict[int, set[str]] = defaultdict(set)
    group_of: dict[int, str] = {}
    for sitting in sittings:
        group_of.setdefault(sitting.student_id, sitting.group.gender)
        if sitting.exam.scheduled:
            in_file[sitting.student_id].add(sitting.exam.day)
    day_no = {day: index for index, day in enumerate(model.saved.days, 1)}
    records = []
    for sid, days in in_file.items():
        for day in days:
            exams = model.flags.days[sid][day]
            if len(exams) < 2:
                continue
            clash = model.flags.day_has_clash(sid, day)
            records.append((not clash, day_no[day], sid, day, clash, exams))
    records.sort(key=lambda item: item[:3])
    rows = []
    for _not_clash, number, sid, day, clash, exams in records:
        student = model.students[sid]
        rows.append(
            (
                sid,
                student.name,
                student.program,
                _w(group_of[sid], lang),
                number,
                day,
                context.options.dates.get(day),
                _w("clash" if clash else "same_day", lang),
                len(exams),
                _JOIN.join(f"{start} {code}" for _slot, code, start in exams),
            )
        )
    return rows


def _part_text(part: RoomPart, position: int, total: int, lang: str) -> str:
    """A room part as "M-B (2 of 2)", numbered within its own section group.

    Never the saved ``room_group_index``/``room_group_count``: the build
    numbers those per section_key across BOTH cohorts, so a label with no
    gender shared by male and female students would read "(2 of 2)" for a
    group that sits whole in one room. The position is the seat order.
    """
    label = part.room_code if part.assigned else _w(BASIS_UNASSIGNED, lang)
    of = _pick(lang, "of", "من")
    return f"{label} ({position} {of} {total})"


def _ranges_text(group: SectionGroup, rows: list[Sitting], lang: str, full: bool) -> str:
    pieces = []
    for position, part in enumerate(group.parts, 1):
        ids = sorted(
            s.student_id
            for s in rows
            if s.part == part and s.basis in {BASIS_WHOLE, BASIS_SPLIT, BASIS_UNASSIGNED}
        )
        if ids:
            pieces.append(
                f"{_part_text(part, position, len(group.parts), lang)}: "
                + (_id_range(ids) if full else str(len(ids)))
            )
    for basis in (BASIS_NO_SEAT, BASIS_UNASSIGNED, BASIS_NOT_SCHEDULED):
        count = sum(
            1
            for s in rows
            if s.basis == basis and not (basis == BASIS_UNASSIGNED and s.part is not None)
        )
        if count:
            pieces.append(f"{_w(basis, lang)} ({count})")
    return _JOIN.join(pieces)


def _section_rows(
    context: _Context,
    groups: list[SectionGroup],
    rows_by_group: dict[tuple[str, str, str], list[Sitting]],
) -> list[tuple]:
    model, options, lang = context.model, context.options, context.lang
    full = options.contents == "full"
    rows = []
    for group in groups:
        exam = model.exams[group.exam]
        rows_in = rows_by_group.get((group.exam, group.section_key, group.gender), [])
        by_program = Counter(model.students[s.student_id].program for s in rows_in)
        selected_saved = sum(
            count
            for program, count in group.saved_program_counts.items()
            if not options.programs or program in options.programs
        )
        rows.append(
            (
                exam.code,
                exam.source_code,
                context.course_name(exam),
                _section_label(group, lang),
                _w(group.gender, lang),
                exam.day_no,
                exam.day if exam.scheduled else _w(BASIS_NOT_SCHEDULED, lang),
                context.date(exam),
                exam.period or None,
                _yes(exam.is_online, lang),
                _w(group.mapping_status, lang),
                group.saved_count,
                len(group.members),
                len(group.members) - group.saved_count,
                _w(group.membership, lang),
                _w(group.program_mix, lang),
                selected_saved,
                len(rows_in),
                sum(1 for s in rows_in if s.basis == BASIS_NO_SEAT),
                sum(1 for s in rows_in if s.flags and s.flags.clash),
                sum(1 for s in rows_in if s.flags and s.flags.same_day),
                len(group.assigned_parts),
                _JOIN.join(
                    f"{program or '—'} {count}"
                    for program, count in sorted(by_program.items(), key=lambda i: (-i[1], i[0]))
                )
                or None,
                _ranges_text(group, rows_in, lang, full) or None,
            )
        )
    return rows


def _room_rows(context: _Context, sittings: list[Sitting]) -> list[tuple]:
    model, lang = context.model, context.lang
    in_file: Counter = Counter()
    unassigned_in_file: Counter = Counter()
    exams_by_slot: dict[int, ExamFacts] = {}
    for sitting in sittings:
        exam = sitting.exam
        if not exam.scheduled:
            continue
        exams_by_slot.setdefault(exam.slot_index, exam)
        if sitting.basis in {BASIS_WHOLE, BASIS_SPLIT} and sitting.part is not None:
            in_file[(exam.slot_index, sitting.part.room_code)] += 1
        elif sitting.basis == BASIS_UNASSIGNED and sitting.part is not None:
            unassigned_in_file[(exam.slot_index, sitting.group.gender)] += 1
    records = []
    for (slot, code), count in in_file.items():
        room = model.saved.rooms[(slot, code)]
        exam = exams_by_slot[slot]
        seated_now = context.room_now[(slot, code)]
        capacity = room["capacity"]
        sections = context.room_sections[(slot, code)]
        by_exam: dict[str, list[str]] = defaultdict(list)
        for (exam_code, key, gender), number in sorted(
            sections.items(),
            key=lambda item: (
                natural_key(item[0][0]),
                natural_key(_group_label(model, item[0], lang)),
            ),
        ):
            by_exam[exam_code].append(
                f"{_group_label(model, (exam_code, key, gender), lang)} ({number})"
            )
        records.append(
            (
                (
                    exam.day_no or 0,
                    exam.start or time.max,
                    slot,
                    0,
                    natural_key(room["building"]),
                    _floor(room["floor"]) or 0,
                    natural_key(code),
                ),
                (
                    exam.day_no,
                    exam.day,
                    context.date(exam),
                    exam.period or None,
                    exam.start,
                    code,
                    room["building"] or None,
                    _floor(room["floor"]),
                    _w(room["gender"], lang),
                    capacity,
                    room["seated_at_save"],
                    seated_now,
                    capacity - seated_now if capacity is not None else None,
                    round(seated_now / capacity, 4) if capacity else None,
                    len(set(room["exams"])),
                    _yes(any(model.exams[e].is_online for e in room["exams"]), lang),
                    count,
                    _JOIN.join(f"{e} " + " + ".join(labels) for e, labels in by_exam.items()),
                ),
            )
        )
    for (slot, gender), count in unassigned_in_file.items():
        exam = exams_by_slot[slot]
        saved = sum(
            part.student_count
            for group in model.groups
            if model.exams[group.exam].slot_index == slot and group.gender == gender
            for part in group.parts
            if not part.assigned
        )
        records.append(
            (
                (exam.day_no or 0, exam.start or time.max, slot, 1, (), GENDERS.index(gender), ()),
                (
                    exam.day_no,
                    exam.day,
                    context.date(exam),
                    exam.period or None,
                    exam.start,
                    _w("room_not_assigned", lang),
                    None,
                    None,
                    _w(gender, lang),
                    None,
                    saved,
                    context.unassigned_now[(slot, gender)],
                    None,
                    None,
                    None,
                    None,
                    count,
                    None,
                ),
            )
        )
    records.sort(key=lambda item: item[0])
    return [row for _key, row in records]


def _group_label(model: RosterModel, key: tuple[str, str, str], lang: str) -> str:
    group = model.group_index.get(key)
    return _section_label(group, lang) if group else ""


def _seats_for(context: _Context, rooms: set[tuple[int, str]]) -> tuple[int, int]:
    capacity = sum(context.model.saved.rooms[key]["capacity"] or 0 for key in rooms)
    seated = sum(context.room_now[key] for key in rooms)
    return capacity, seated


def _day_rows(context: _Context, sittings: list[Sitting]) -> list[tuple]:
    model, lang = context.model, context.lang
    by_day: dict[str, list[Sitting]] = defaultdict(list)
    unscheduled: list[Sitting] = []
    for sitting in sittings:
        if sitting.exam.scheduled:
            by_day[sitting.exam.day].append(sitting)
        else:
            unscheduled.append(sitting)
    rows = []
    for number, day in enumerate(model.saved.days, 1):
        own = by_day.get(day)
        if not own:
            continue
        rooms = {
            (s.exam.slot_index, s.part.room_code)  # type: ignore[union-attr]
            for s in own
            if s.basis in {BASIS_WHOLE, BASIS_SPLIT} and s.part is not None
        }
        capacity, _seated = _seats_for(context, rooms)
        students = {s.student_id for s in own}
        when = context.options.dates.get(day)
        rows.append(
            (
                number,
                day,
                _weekday(day, when, lang),
                when,
                len({s.exam.code for s in own}),
                len({(s.group.exam, s.group.section_key, s.group.gender) for s in own}),
                len(rooms),
                capacity,
                len(own),
                len(students),
                len({s.student_id for s in own if s.group.gender == "M"}),
                len({s.student_id for s in own if s.group.gender == "F"}),
                sum(1 for sid in students if len(model.flags.days.get(sid, {}).get(day, ())) >= 2),
                sum(1 for sid in students if model.flags.day_has_clash(sid, day)),
            )
        )
    if unscheduled:
        students = {s.student_id for s in unscheduled}
        rows.append(
            (
                None,
                _w(BASIS_NOT_SCHEDULED, lang),
                None,
                None,
                len({s.exam.code for s in unscheduled}),
                len({(s.group.exam, s.group.section_key, s.group.gender) for s in unscheduled}),
                None,
                None,
                len(unscheduled),
                len(students),
                len({s.student_id for s in unscheduled if s.group.gender == "M"}),
                len({s.student_id for s in unscheduled if s.group.gender == "F"}),
                None,
                None,
            )
        )
    return rows


def _period_rows(context: _Context, sittings: list[Sitting]) -> list[tuple]:
    by_slot: dict[int, list[Sitting]] = defaultdict(list)
    for sitting in sittings:
        if sitting.exam.scheduled:
            by_slot[sitting.exam.slot_index].append(sitting)
    rows = []
    for slot in sorted(
        by_slot,
        key=lambda s: (by_slot[s][0].exam.day_no or 0, by_slot[s][0].exam.start or time.max, s),
    ):
        own = by_slot[slot]
        exam = own[0].exam
        rooms = {
            (slot, s.part.room_code)  # type: ignore[union-attr]
            for s in own
            if s.basis in {BASIS_WHOLE, BASIS_SPLIT} and s.part is not None
        }
        capacity, seated = _seats_for(context, rooms)
        students = {s.student_id for s in own}
        rows.append(
            (
                exam.day_no,
                exam.day,
                context.date(exam),
                exam.period or None,
                exam.start,
                len({s.exam.code for s in own}),
                len(rooms),
                capacity,
                len(own),
                len(students),
                len({s.student_id for s in own if s.flags and s.flags.clash}),
                round(seated / capacity, 4) if capacity else None,
            )
        )
    return rows


def _program_day_rows(
    context: _Context, sittings: list[Sitting]
) -> tuple[tuple[str, ...], list[tuple]]:
    model = context.model
    students: dict[str, set[int]] = defaultdict(set)
    per_day: dict[tuple[str, str], set[int]] = defaultdict(set)
    unscheduled: dict[str, set[int]] = defaultdict(set)
    for sitting in sittings:
        program = model.students[sitting.student_id].program
        students[program].add(sitting.student_id)
        if sitting.exam.scheduled:
            per_day[(program, sitting.exam.day)].add(sitting.student_id)
        else:
            unscheduled[program].add(sitting.student_id)
    rows = []
    for program in sorted(students, key=lambda p: (_department_rank(p), p)):
        rows.append(
            (
                program or None,
                _department_label(program, context.lang),
                len(students[program]),
                *(len(per_day.get((program, day), ())) for day in model.saved.days),
                len(unscheduled.get(program, ())),
            )
        )
    return tuple(model.saved.days), rows


def _check_rows(
    context: _Context, sittings: list[Sitting], file_groups: tuple[str, ...]
) -> list[tuple]:
    model, options, lang = context.model, context.options, context.lang
    qa = model.saved.qa
    saved_groups = [g for g in model.groups if g.membership != NEW]
    whole = (
        options.scope_kind == "all" and not options.programs and set(file_groups) == set(GENDERS)
    )
    suffix = "" if whole else _pick(lang, " — whole timetable", " — الجدول كاملاً")
    per_day_counts = [len(exams) for days in model.flags.days.values() for exams in days.values()]
    max_per_day = qa.get("max_per_day") if isinstance(qa.get("max_per_day"), int) else 2
    over_limit = sum(
        1
        for days in model.flags.days.values()
        if any(len(exams) > max_per_day for exams in days.values())
    )
    two_plus = sum(
        1 for days in model.flags.days.values() if any(len(e) >= 2 for e in days.values())
    )
    # The build counts a split section per (course, section_key) across both
    # cohorts (``derive_multi_sitting_details``); so does this recount.
    parts_per_section: Counter = Counter()
    for group in model.groups:
        parts_per_section[(group.exam, group.section_key)] += len(group.parts)
    split_now = sum(1 for parts in parts_per_section.values() if parts > 1)
    unassigned_saved = len((qa.get("rooms") or {}).get("unassigned_room_sections") or [])
    unassigned_now = sum(
        1
        for group in model.groups
        if any(model.basis(group, sid) == BASIS_UNASSIGNED for sid in group.members)
        and model.exams[group.exam].scheduled
        and model.saved.assign_rooms
    )

    def num(value: object) -> int | None:
        return value if isinstance(value, int) and not isinstance(value, bool) else None

    def result(saved: int | None, now: int | None) -> str:
        if saved is None or now is None:
            return _w("info", lang)
        return _w("ok" if saved == now else "differs", lang)

    def screen(en: str, ar: str) -> str:
        return _pick(lang, f"Full summary › {en}", f"الملخص الكامل › {ar}")

    not_shown = _pick(lang, "Not on the screen", "غير معروضة في الشاشة")
    checks: list[tuple[str, str, int | None, int | None, str, str, str, str]] = [
        (
            "Courses",
            "المقررات",
            len(model.exams),
            sum(1 for e in model.exams.values() if e.in_lists),
            screen("Courses", "المقررات"),
            "Exams whose course is still in the lists.",
            "الاختبارات التي ما زال مقررها في القوائم.",
        ),
        (
            "Students",
            "الطلاب",
            num(model.saved.students_count),
            len(model.students),
            screen("Students", "الطلاب"),
            "Distinct students in the timetable.",
            "الطلاب المختلفون في الجدول.",
        ),
        (
            "Student exam sittings",
            "جلسات الطلاب في الاختبارات",
            sum(g.saved_count for g in saved_groups),
            sum(len(g.members) for g in model.groups),
            _pick(
                lang, "Sum of exam-card student counts", "مجموع أعداد الطلاب في بطاقات الاختبارات"
            ),
            "One sitting is one student in one exam.",
            "الجلسة طالب واحد في اختبار واحد.",
        ),
        (
            "Section groups identical to the saved timetable",
            "مجموعات الشعب المطابقة للجدول المحفوظ",
            len(saved_groups),
            sum(1 for g in model.groups if g.membership == MATCHES),
            _pick(lang, "Not on the screen (checksums)", "غير معروضة في الشاشة (البصمات)"),
            "Each section's students checked against the saved checksum.",
            "طلاب كل شعبة مطابَقون مع البصمة المحفوظة.",
        ),
        (
            "Sections whose program mix matches",
            "الشعب التي يطابق فيها توزيع البرامج",
            len(saved_groups),
            sum(1 for g in model.groups if g.program_mix == MATCHES and g.membership != NEW),
            not_shown,
            "Students per programme against the saved timetable.",
            "الطلاب لكل برنامج مقارنة بالجدول المحفوظ.",
        ),
        (
            "Conflicts (student × period)",
            "التعارضات (طالب × فترة)",
            num(qa.get("conflict_count")),
            sum(model.flags.clash_periods(sid) for sid in model.flags.days),
            screen("Conflicts", "التعارضات"),
            "Periods with two exams for one student.",
            "فترات فيها اختباران لطالب واحد.",
        ),
        (
            "Max exams/day/student",
            "أقصى اختبارات/يوم/طالب",
            num(qa.get("max_exams_per_day_per_student")),
            max(per_day_counts, default=0),
            screen("Max exams/day/student", "أقصى اختبارات/يوم/طالب"),
            "The most exams any student has on one day.",
            "أكبر عدد اختبارات لطالب في يوم واحد.",
        ),
        (
            "Students over limit/day",
            "طلاب تجاوزوا الحد",
            num(qa.get("students_over_limit_per_day")),
            over_limit,
            screen("Students over limit/day", "طلاب تجاوزوا الحد"),
            f"Students with more than {max_per_day} exams on a day.",
            f"طلاب لهم أكثر من {max_per_day} اختبارات في يوم.",
        ),
        (
            "Unassigned room groups",
            "مجموعات دون قاعة",
            unassigned_saved,
            unassigned_now,
            screen("Unassigned room groups", "مجموعات دون قاعة"),
            "Sections with students the timetable could not room.",
            "شعب فيها طلاب لم يجد لهم الجدول قاعة.",
        ),
        (
            "Sections split across rooms",
            "شعب موزعة على قاعات",
            num(qa.get("multi_sitting_sections")),
            split_now,
            screen("Sections split across rooms", "شعب موزعة على قاعات"),
            "Sections the timetable split across rooms.",
            "شعب وزعها الجدول على أكثر من قاعة.",
        ),
    ]
    rows = [
        (
            _pick(lang, en, ar) + suffix,
            saved,
            now,
            result(saved, now),
            where,
            _pick(lang, note_en, note_ar),
        )
        for en, ar, saved, now, where, note_en, note_ar in checks
    ]
    # 11: rows against the saved sittings of this file's exams, programmes and groups.
    expected: int | None = None
    if options.rows == "all":
        if options.scope_kind == "room":
            if not options.programs:
                expected = sum(
                    part.student_count
                    for group in model.groups
                    if group.gender in file_groups
                    for part in group.parts
                    if part.room_code == options.room_code
                    and model.exams[group.exam].slot_index == options.slot_index
                )
        else:
            expected = sum(
                count
                for group in model.groups
                if group.membership != NEW
                and group.gender in file_groups
                and _group_in_scope(group, options, model)
                for program, count in group.saved_program_counts.items()
                if not options.programs or program in options.programs
            )
    rows.append(
        (
            _pick(
                lang,
                "Rows in this file = saved sittings for this file's exams, programs and groups",
                "أسطر هذا الملف = جلسات الجدول المحفوظ لاختبارات هذا الملف وبرامجه وفئاته",
            ),
            expected,
            len(sittings),
            result(expected, len(sittings)),
            _pick(lang, "Not on the screen (recomputed)", "غير معروضة في الشاشة (أعيد حسابها)"),
            _pick(
                lang,
                "Info when rows are limited to flagged students, or a room file is limited to programmes.",
                "معلومة إذا اقتصرت الأسطر على ذوي التنبيهات أو اقتصر ملف القاعة على برامج.",
            ),
        )
    )
    bases = Counter(s.basis for s in sittings)
    seated = bases[BASIS_WHOLE] + bases[BASIS_SPLIT]
    accounted = seated + bases[BASIS_NO_SEAT] + bases[BASIS_UNASSIGNED] + bases[BASIS_NOT_SCHEDULED]
    rows.append(
        (
            _pick(
                lang,
                "Seated + No seat + Room not assigned + Not scheduled = rows",
                "بمقعد + بلا مقعد + دون قاعة + غير مجدول = الأسطر",
            ),
            len(sittings),
            accounted,
            result(len(sittings), accounted),
            _pick(
                lang,
                "Not on the screen (internal consistency)",
                "غير معروضة في الشاشة (اتساق داخلي)",
            ),
            _pick(
                lang,
                f"Seated {seated} · No seat {bases[BASIS_NO_SEAT]} · Room not assigned {bases[BASIS_UNASSIGNED]} · Not scheduled {bases[BASIS_NOT_SCHEDULED]}",
                f"بمقعد {seated} · بلا مقعد {bases[BASIS_NO_SEAT]} · دون قاعة {bases[BASIS_UNASSIGNED]} · غير مجدول {bases[BASIS_NOT_SCHEDULED]}",
            ),
        )
    )
    rows.append(
        (
            _pick(
                lang, "Students with 2+ exams on one day", "طلاب لهم اختباران أو أكثر في يوم واحد"
            )
            + suffix,
            None,
            two_plus,
            _w("info", lang),
            not_shown,
            _pick(
                lang,
                "Every such student is flagged Same day or Clash.",
                "كل طالب منهم عليه تنبيه اليوم نفسه أو تعارض.",
            ),
        )
    )
    return rows


def _change_rows(context: _Context, groups: list[SectionGroup]) -> list[tuple]:
    model, lang, options = context.model, context.lang, context.options
    rows = []
    for group in groups:
        exam = model.exams[group.exam]
        membership = group.membership
        if membership == MATCHES and group.program_mix == MATCHES:
            continue
        now = len(group.members)
        extra = now - group.saved_count
        no_seat = sum(1 for sid in group.members if model.basis(group, sid) == BASIS_NO_SEAT)
        split = len(group.parts) > 1
        # Seats exist only for a scheduled exam in a timetable that assigned
        # rooms; anything else is said as what it is, never as "no seat".
        roomed = exam.scheduled and model.saved.assign_rooms
        unroomed = (
            _pick(lang, "the exam is not scheduled", "الاختبار غير مجدول")
            if not exam.scheduled
            else _pick(
                lang,
                "rooms were not assigned in this timetable",
                "لم تُوزَّع القاعات في هذا الجدول",
            )
        )
        if membership == MATCHES:
            what = _w("mix_changed", lang)
            effect = _pick(
                lang,
                "The same students, but their programmes differ from the saved timetable; department counts on the screen may differ.",
                "الطلاب أنفسهم لكن برامجهم تختلف عن الجدول المحفوظ؛ قد تختلف أعداد الأقسام في الشاشة.",
            )
        elif membership == GONE:
            what = _w(GONE, lang)
            effect = _pick(
                lang,
                "No rows: the lists have no students in this section now.",
                "لا أسطر: لا يوجد طلاب في هذه الشعبة في القوائم الآن.",
            )
        elif membership == NEW:
            what = _w(NEW, lang)
            effect = (
                _pick(
                    lang,
                    f"{now} students have no seat: this section did not exist when the timetable was saved.",
                    f"{now} طالباً بلا مقعد: لم تكن هذه الشعبة موجودة عند حفظ الجدول.",
                )
                if roomed
                else _pick(
                    lang,
                    f"{now} students: this section did not exist when the timetable was saved, and {unroomed}.",
                    f"{now} طالباً: لم تكن هذه الشعبة موجودة عند حفظ الجدول، و{unroomed}.",
                )
            )
        else:
            what = _w(CHANGED, lang)
            if extra > 0 and roomed:
                effect = _pick(
                    lang,
                    f"{no_seat} students have no seat; rooms were sized for {group.saved_count}.",
                    f"{no_seat} طالباً بلا مقعد؛ حُددت القاعات لعدد {group.saved_count}.",
                )
            elif extra > 0:
                effect = _pick(
                    lang,
                    f"{extra} more students than saved; {unroomed}.",
                    f"عدد الطلاب أكثر بـ {extra} مما حُفظ؛ {unroomed}.",
                )
            elif extra < 0 and roomed:
                effect = _pick(
                    lang,
                    f"{-extra} fewer students than saved; rooms were sized for {group.saved_count}.",
                    f"عدد الطلاب أقل بـ {-extra} مما حُفظ؛ حُددت القاعات لعدد {group.saved_count}.",
                )
            elif extra < 0:
                effect = _pick(
                    lang,
                    f"{-extra} fewer students than saved; {unroomed}.",
                    f"عدد الطلاب أقل بـ {-extra} مما حُفظ؛ {unroomed}.",
                )
            else:
                effect = _pick(
                    lang,
                    "The same number of students, but not the same students.",
                    "عدد الطلاب نفسه لكنهم ليسوا الطلاب أنفسهم.",
                )
            if split and roomed:
                effect += _pick(lang, " Split boundary recomputed.", " أُعيد حساب حدود التقسيم.")
        if membership == MATCHES:
            todo = _pick(
                lang,
                "Timetable › Check changes › Save to refresh the programme counts.",
                "الجدول › التحقق من التغييرات › حفظ لتحديث أعداد البرامج.",
            )
        elif roomed:
            todo = _pick(
                lang,
                "Timetable › Check changes › Save to resize rooms.",
                "الجدول › التحقق من التغييرات › حفظ لإعادة تحديد القاعات.",
            )
        else:
            todo = _pick(
                lang,
                "Timetable › Check changes › Save to update the saved counts.",
                "الجدول › التحقق من التغييرات › حفظ لتحديث الأعداد المحفوظة.",
            )
        rows.append(
            (
                exam.code,
                _section_label(group, lang),
                _w(group.gender, lang),
                group.saved_count,
                now,
                extra,
                what,
                effect,
                todo,
            )
        )
    for code in model.missing_exams:
        if not any(
            _group_in_scope(group, options, model) for group in model.groups if group.exam == code
        ):
            continue
        saved = sum(group.saved_count for group in model.groups if group.exam == code)
        rows.append(
            (
                code,
                None,
                None,
                saved,
                0,
                -saved,
                _w("exam_missing", lang),
                _pick(
                    lang,
                    "No rows for this exam: its course is no longer in the lists.",
                    "لا أسطر لهذا الاختبار: لم يعد مقرره في القوائم.",
                ),
                _pick(
                    lang,
                    "Rebuild the timetable if the course was withdrawn.",
                    "أعد بناء الجدول إذا سُحب المقرر.",
                ),
            )
        )
    return rows


# ── Rendering ──────────────────────────────────────────────────


#: Everything outside XML 1.0's ``Char`` production. openpyxl's own
#: ILLEGAL_CHARACTERS_RE covers only C0 controls: U+FFFE, U+FFFF and lone
#: surrogates pass it, and the writer used in production (no lxml) then saves
#: a workbook that neither Excel nor openpyxl will open.
XML_ILLEGAL_RE = re.compile(r"[^\t\n\r\x20-\ud7ff\ue000-\ufffd\U00010000-\U0010ffff]")


def _clean(value: str) -> str:
    """Drop what an XML part cannot hold, from every string a cell gets."""
    return XML_ILLEGAL_RE.sub("", value)


class _Styles:
    """Shared cell styles: made once per workbook, copied onto cells by reference."""

    def __init__(self, ws) -> None:
        self.header = self._style(
            ws,
            font=Font(name="Arial", size=11, bold=True, color="FFFFFF"),
            fill=PatternFill("solid", fgColor=NAVY),
            alignment=Alignment(vertical="center"),
        )
        self.number = {
            kind: self._style(ws, number_format=fmt) for kind, fmt in _NUMBER_FORMATS.items()
        }

    @staticmethod
    def _style(ws, **attributes):
        cell = WriteOnlyCell(ws, "x")
        for name, value in attributes.items():
            setattr(cell, name, value)
        return cell._style


def _write_row(ws, values: Sequence, plan: list[tuple[int, Any]], texts: list[int]) -> None:
    cells: list[Any] = list(values)
    for index in texts:
        value = cells[index]
        if isinstance(value, str):
            value = _clean(value)
            if value.startswith("="):
                # write-only mode turns any "=..." string into a formula
                cell = WriteOnlyCell(ws, value)
                cell.data_type = "s"
                cells[index] = cell
            else:
                cells[index] = value
    for index, style in plan:
        value = cells[index]
        if value is not None and not isinstance(value, Cell):
            cell = WriteOnlyCell(ws, value)
            cell._style = style
            cells[index] = cell
    ws.append(cells)


def _header_cells(ws, headers: Sequence[str], styles: _Styles) -> list:
    cells = []
    for text in headers:
        cell = WriteOnlyCell(ws, text)
        cell.data_type = "s"
        cell._style = styles.header
        cells.append(cell)
    return cells


@dataclass
class _SheetWriter:
    ws: Any
    lang: str
    styles: _Styles
    row: int = 0  # explicit counter: never ws.max_row
    tables: list[tuple[TableSpec, str, int, int, list[str]]] = field(default_factory=list)

    def blank(self, count: int = 1) -> None:
        for _ in range(count):
            self.ws.append([])
            self.row += 1

    def table(
        self,
        spec: TableSpec,
        rows: Sequence[tuple],
        *,
        headers: list[str] | None = None,
        kinds: list[str] | None = None,
        totals: tuple | None = None,
        extra: dict[int, tuple[int, Any]] | None = None,
    ) -> tuple[int, int]:
        headers = headers or spec.headers(self.lang)
        kinds = kinds or [column.kind for column in spec.columns]
        headers = [" ".join(str(h).split()) for h in headers]
        self.ws.append(_header_cells(self.ws, headers, self.styles))
        self.row += 1
        start = self.row
        plan = [(i, self.styles.number[k]) for i, k in enumerate(kinds) if k in _NUMBER_FORMATS]
        texts = [i for i, k in enumerate(kinds) if k in {"text", "enum"}]
        if not rows:
            self.ws.append([])
            self.row += 1
        for index, values in enumerate(rows):
            if extra and index in extra:
                column, cell = extra[index]
                values = (*values, *([None] * (column - len(values))), cell)
            _write_row(self.ws, values, plan, texts)
            self.row += 1
        if totals is not None:
            _write_row(self.ws, totals, plan, texts)
            self.row += 1
        self.tables.append((spec, headers[0], start, self.row, headers))
        return start, self.row

    def finish_tables(self, program_days_totals: bool = False) -> None:
        for spec, _first, header_row, end_row, headers in self.tables:
            last = get_column_letter(len(headers))
            totals = program_days_totals and spec.name == "ProgramDays"
            table = Table(
                displayName=spec.name,
                ref=f"A{header_row}:{last}{end_row}",
                totalsRowCount=1 if totals else None,
                # The filter covers the data rows only, never a totals row.
                autoFilter=AutoFilter(ref=f"A{header_row}:{last}{end_row - (1 if totals else 0)}"),
            )
            table._initialise_columns()
            for column, name in zip(table.tableColumns, headers, strict=True):
                column.name = name
            if totals:
                table.tableColumns[0].totalsRowLabel = _w("total", self.lang)
                for column in table.tableColumns[2:]:
                    column.totalsRowFunction = "sum"
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            with warnings.catch_warnings():
                # write-only add_table warns unconditionally; the header cells
                # were written above, and each column name was set to match them.
                warnings.simplefilter("ignore", UserWarning)
                self.ws.add_table(table)


def _structured_ref(table: str, header: str) -> str:
    escaped = re.sub(r"(['\[\]#@])", r"'\1", header)
    return f"{table}[[{escaped}]]"


def _setup_sheet(
    ws,
    spec: SheetSpec,
    lang: str,
    widths: list[float],
    *,
    header_left: str,
    header_right: str,
    selected: bool,
) -> None:
    """Everything written before the first row: views, widths, sheet properties."""
    ws.sheet_view.rightToLeft = lang == "ar"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_view.tabSelected = selected
    if spec.freeze:
        ws.freeze_panes = spec.freeze
    ws.sheet_properties.tabColor = spec.tab
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True, autoPageBreaks=False)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 24 if spec.key != "about" else 30
    ws.page_setup.orientation = "landscape" if spec.landscape else "portrait"
    ws.page_setup.paperSize = spec.paper
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    header = ws.HeaderFooter.oddHeader
    header.left.text = header_left.replace("&", "&&")
    header.left.size = 9
    header.right.text = header_right.replace("&", "&&")
    header.right.size = 9
    footer = ws.HeaderFooter.oddFooter
    footer.center.text = "&P / &N"
    footer.center.size = 9


def _widths(tables: Sequence[TableSpec], lang: str, extra: int = 0) -> list[float]:
    widths: list[float] = []
    for table in tables:
        for index, column in enumerate(table.columns):
            width = (column.width_ar or column.width) if lang == "ar" else column.width
            if index < len(widths):
                widths[index] = max(widths[index], width)
            else:
                widths.append(width)
    widths.extend([10.0] * extra)
    return widths


def _cf(ws, ref: str, formula: str, *, color: str, fill: str | None = None, bold=False) -> None:
    ws.conditional_formatting.add(
        ref,
        FormulaRule(
            formula=[formula],
            font=Font(color=color, bold=bold),
            fill=PatternFill("solid", fgColor=fill, bgColor=fill) if fill else None,
        ),
    )


def _quote(text: str) -> str:
    return '"' + text.replace('"', '""') + '"'


def _apply_formatting(ws, spec: TableSpec, start: int, end: int, lang: str) -> None:
    first, last = start + 1, max(start + 1, end)
    full = f"A{first}:{get_column_letter(len(spec.columns))}{last}"

    def col(key: str) -> str:
        return f"${spec.letter(key)}{first}"

    yes = _quote(_w("yes", lang))
    if spec.name == "StudentExams":
        _cf(ws, full, f"{col('clash')}={yes}", color=DANGER, fill=DANGER_SOFT, bold=True)
        _cf(ws, full, f"{col('same_day')}={yes}", color=WARN, fill=WARN_SOFT)
        unseated = ",".join(
            f"{col('room_basis')}={_quote(_w(b, lang))}"
            for b in (BASIS_NO_SEAT, BASIS_UNASSIGNED, BASIS_NOT_SCHEDULED)
        )
        _cf(ws, full, f"OR({unseated})", color=DANGER, bold=True)
        _cf(ws, full, f'{col("change")}<>""', color=WARN)
    elif spec.name == "ExamSections":
        matches = _quote(_w(MATCHES, lang))
        _cf(
            ws,
            full,
            f"OR({col('membership')}<>{matches},{col('program_mix')}<>{matches})",
            color=WARN,
            fill=WARN_SOFT,
        )
        _cf(ws, full, f"{col('no_seat')}>0", color=DANGER)
    elif spec.name == "ExamRooms":
        _cf(ws, full, f'AND({col("free")}<>"",{col("free")}<0)', color=DANGER, bold=True)
        _cf(ws, full, f'AND({col("use")}<>"",{col("use")}>=0.95)', color=WARN)
    elif spec.name == "ExamFlags":
        clash = _quote(_w("clash", lang))
        _cf(ws, full, f"{col('flag')}={clash}", color=DANGER, fill=DANGER_SOFT, bold=True)
    elif spec.name == "ExportChecks":
        _cf(ws, full, f"{col('result')}={_quote(_w('differs', lang))}", color=WARN, fill=WARN_SOFT)
        _cf(ws, full, f"{col('result')}={_quote(_w('ok', lang))}", color=OK_TEXT)


def _discard_write_only(wb: Workbook) -> None:
    """Close and delete every write-only spool file; they hold student rows."""
    for ws in list(getattr(wb, "_sheets", [])):
        writer = getattr(ws, "_writer", None)
        if writer is None:
            continue
        rows = getattr(ws, "_rows", None)
        for closer in (getattr(rows, "close", None), writer.close):
            if closer is None:
                continue
            try:
                closer()
            except Exception:  # noqa: BLE001 - best effort; the file is removed below
                pass
        if writer.out in ALL_TEMP_FILES:
            try:
                writer.cleanup()
            except OSError:
                pass


def _new_workbook() -> Workbook:
    wb = Workbook(write_only=True)
    # Arial 11 is the workbook default, set before any sheet or style exists:
    # font 0 and the Normal style's font. Body cells never carry a font.
    arial = Font(name="Arial", size=11, family=2, color=BODY)
    wb._fonts = IndexedList([arial])
    wb._named_styles["Normal"].font = arial
    wb.calculation.fullCalcOnLoad = True
    return wb


@dataclass
class _Render:
    prepared: PreparedExport
    item: PreparedFile
    reference: str
    audit_hash: str

    @property
    def lang(self) -> str:
        return self.prepared.options.language

    @property
    def model(self) -> RosterModel:
        return self.prepared.model

    @property
    def options(self) -> ExportOptions:
        return self.prepared.options


def render_file(
    prepared: PreparedExport, item: PreparedFile, *, reference: str, audit_hash: str
) -> bytes:
    """Write one workbook. Spool files are removed whether or not this succeeds."""
    wb = _new_workbook()
    try:
        _render_workbook(wb, _Render(prepared, item, reference, audit_hash))
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    finally:
        _discard_write_only(wb)


def render_export(
    prepared: PreparedExport, *, reference: str, audit_hash: str
) -> tuple[bytes, str, str]:
    """The download: one workbook, or several in an uncompressed zip."""
    files = [
        (
            item.name(reference),
            render_file(prepared, item, reference=reference, audit_hash=audit_hash),
        )
        for item in prepared.files
    ]
    if len(files) == 1:
        return files[0][1], files[0][0], XLSX_TYPE
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_STORED) as archive:
        for name, content in files:
            archive.writestr(name, content)
    return buffer.getvalue(), prepared.filename(reference), ZIP_TYPE


def _layout(options: ExportOptions) -> list[SheetSpec]:
    return [sheet for sheet in SHEETS if options.contents == "full" or sheet.in_summary]


def _render_workbook(wb: Workbook, render: _Render) -> None:
    lang, model, item = render.lang, render.model, render.item
    run_id = model.saved.run_id
    wb.properties.title = _pick(
        lang, f"Exam student data — #{run_id}", f"بيانات الطلاب للاختبارات — {run_id}"
    )
    wb.properties.subject = render.reference
    wb.properties.creator = "Exam Timetable Builder"
    wb.properties.lastModifiedBy = "Exam Timetable Builder"
    wb.properties.keywords = f"run-{run_id}"
    wb.properties.description = _confidential(lang)
    wb.custom_doc_props.append(IntProperty(name="run_id", value=run_id))
    wb.custom_doc_props.append(StringProperty(name="reference", value=render.reference))
    wb.custom_doc_props.append(StringProperty(name="lists_code", value=model.lists_code_now))
    wb.custom_doc_props.append(StringProperty(name="data_sha256", value=item.data_sha256))

    header_left = _pick(
        lang,
        f"Exam student data · #{run_id} · {render.reference}",
        f"بيانات الطلاب للاختبارات · {run_id} · {render.reference}",
    )
    header_right = _pick(
        lang, "Confidential — exam administration only", "سري — لاستخدام لجنة الاختبارات فقط"
    )
    layout = _layout(render.options)
    sheet_rows = _sheet_row_counts(item)
    for position, spec in enumerate(layout):
        ws = wb.create_sheet(spec.title(lang))
        writers: dict[str, Callable] = {
            "about": _write_about,
            "student_exams": _write_simple,
            "students": _write_simple,
            "flags": _write_simple,
            "sections": _write_simple,
            "rooms": _write_simple,
            "by_day": _write_by_day,
            "program_days": _write_program_days,
            "checks": _write_checks,
            "file_info": _write_file_info,
        }
        writers[spec.key](
            ws,
            spec,
            render,
            layout=layout,
            sheet_rows=sheet_rows,
            header=(header_left, header_right),
            selected=position == 0,
        )
    wb.active = 0


def _sheet_row_counts(item: PreparedFile) -> dict[str, int]:
    tables = item.tables
    return {
        "student_exams": len(tables["StudentExams"]),
        "students": len(tables["ExamStudents"]),
        "flags": len(tables["ExamFlags"]),
        "sections": len(tables["ExamSections"]),
        "rooms": len(tables["ExamRooms"]),
        "by_day": len(tables["DaySummary"]),
        "program_days": len(tables["ProgramDays"]),
        "checks": len(tables["ExportChecks"]),
        "file_info": len(FILE_INFO_KEYS),
    }


def _write_simple(ws, spec: SheetSpec, render: _Render, *, header, selected, **_kw) -> None:
    table = spec.tables[0]
    _setup_sheet(
        ws,
        spec,
        render.lang,
        _widths(spec.tables, render.lang),
        header_left=header[0],
        header_right=header[1],
        selected=selected,
    )
    writer = _SheetWriter(ws, render.lang, _Styles(ws))
    start, end = writer.table(table, render.item.tables[table.name])
    _apply_formatting(ws, table, start, end, render.lang)
    writer.finish_tables()
    ws.print_title_rows = "1:1"


def _write_by_day(ws, spec: SheetSpec, render: _Render, *, header, selected, **_kw) -> None:
    _setup_sheet(
        ws,
        spec,
        render.lang,
        _widths(spec.tables, render.lang),
        header_left=header[0],
        header_right=header[1],
        selected=selected,
    )
    writer = _SheetWriter(ws, render.lang, _Styles(ws))
    writer.table(DAY_SUMMARY, render.item.tables["DaySummary"])
    writer.blank(2)
    writer.table(PERIOD_SUMMARY, render.item.tables["PeriodSummary"])
    writer.finish_tables()
    ws.print_title_rows = "1:1"


def _program_day_columns(item: PreparedFile, lang: str) -> tuple[list[str], list[str]]:
    fixed = [column.header(lang) for column in _PROGRAM_DAYS_FIXED]
    taken = {header.casefold() for header in fixed} | {_PROGRAM_DAYS_NS.header(lang).casefold()}
    days = []
    for day in item.day_headers:
        # The only header not fixed by this module: a timetable day label.
        header = _clean(" ".join(day.split()))
        while header.casefold() in taken:
            header += _pick(lang, " (day)", " (يوم)")
        taken.add(header.casefold())
        days.append(header)
    headers = [*fixed, *days, _PROGRAM_DAYS_NS.header(lang)]
    kinds = [c.kind for c in _PROGRAM_DAYS_FIXED] + ["int"] * (len(days) + 1)
    return headers, kinds


_PROGRAM_DAYS_TABLE = TableSpec(
    "ProgramDays",
    (*_PROGRAM_DAYS_FIXED, _PROGRAM_DAYS_DAY, _PROGRAM_DAYS_NS),
    "one programme",
    "برنامج واحد",
)


def _write_program_days(ws, spec: SheetSpec, render: _Render, *, header, selected, **_kw) -> None:
    lang = render.lang
    headers, kinds = _program_day_columns(render.item, lang)
    widths = [c.width_ar or c.width if lang == "ar" else c.width for c in _PROGRAM_DAYS_FIXED]
    widths += [max(9.0, len(h) + 2.0) for h in headers[3:]]
    _setup_sheet(
        ws, spec, lang, widths, header_left=header[0], header_right=header[1], selected=selected
    )
    writer = _SheetWriter(ws, lang, _Styles(ws))
    totals = (
        _w("total", lang),
        None,
        *(f"=SUBTOTAL(109,{_structured_ref('ProgramDays', name)})" for name in headers[2:]),
    )
    writer.table(
        _PROGRAM_DAYS_TABLE,
        render.item.tables["ProgramDays"],
        headers=headers,
        kinds=kinds,
        totals=totals,
    )
    writer.finish_tables(program_days_totals=True)
    ws.print_title_rows = "1:1"


def _write_checks(ws, spec: SheetSpec, render: _Render, *, header, selected, layout, **_kw) -> None:
    lang = render.lang
    _setup_sheet(
        ws,
        spec,
        lang,
        _widths(spec.tables, lang, extra=1),
        header_left=header[0],
        header_right=header[1],
        selected=selected,
    )
    writer = _SheetWriter(ws, lang, _Styles(ws))
    rows = list(render.item.tables["ExportChecks"])
    extra = None
    if any(sheet.key == "student_exams" for sheet in layout):
        # Check 11 also recounts the Student exams rows live in column G, so a
        # row deleted after download shows up as a difference in Excel itself.
        exams_sheet = SHEET_BY_KEY["student_exams"].title(lang).replace("'", "''")
        extra = {10: (6, f"=COUNTA('{exams_sheet}'!A:A)-1")}
        *head, note = rows[10]
        rows[10] = (
            *head,
            note
            + _pick(
                lang,
                " Column G recounts the Student exams rows in Excel.",
                " العمود G يعيد عدّ أسطر اختبارات الطلاب في Excel.",
            ),
        )
    start, end = writer.table(EXPORT_CHECKS, rows, extra=extra)
    _apply_formatting(ws, EXPORT_CHECKS, start, end, lang)
    writer.blank(2)
    changes = render.item.tables["ChangeLog"]
    if not changes:
        changes = [(_no_changes_text(render), *([None] * (len(CHANGE_LOG.columns) - 1)))]
    writer.table(CHANGE_LOG, changes)
    writer.finish_tables()


def _no_changes_text(render: _Render) -> str:
    """The one ChangeLog row of a file none of whose sections changed.

    "Every section matches" is said only when it is true of the whole
    timetable; a scoped file can be untouched while sections elsewhere differ,
    and About (whole-timetable) says so on the same workbook.
    """
    lang, model = render.lang, render.model
    run_id = model.saved.run_id
    if model.unchanged:
        return _pick(
            lang,
            f"No changes: every section matches saved timetable #{run_id}.",
            f"لا تغييرات: كل الشعب مطابقة للجدول المحفوظ رقم {_iso(run_id, lang)}",
        )
    elsewhere = len(model.differing_groups)
    return _pick(
        lang,
        f"No changes in this file's sections; {elsewhere:,} sections elsewhere in the timetable "
        f"differ from saved timetable #{run_id}.",
        f"لا تغييرات في شعب هذا الملف. الشعب المختلفة في بقية الجدول عن الجدول المحفوظ رقم "
        f"{_iso(run_id, lang)}: {elsewhere:,}.",
    )


def _write_file_info(
    ws, spec: SheetSpec, render: _Render, *, header, selected, layout, **_kw
) -> None:
    lang = render.lang
    _setup_sheet(
        ws,
        spec,
        lang,
        _widths(spec.tables, lang),
        header_left=header[0],
        header_right=header[1],
        selected=selected,
    )
    writer = _SheetWriter(ws, lang, _Styles(ws))
    writer.table(FILE_INFO, _file_info_rows(render))
    writer.blank(2)
    writer.table(COLUMN_GUIDE, _column_guide_rows(render, layout))
    writer.finish_tables()


def _confidential(lang: str) -> str:
    return _pick(
        lang,
        "Confidential — contains student names and IDs. For exam administration only. "
        "Delete when the examination period ends.",
        "سري — يحتوي على أسماء الطلاب وأرقامهم الجامعية. لاستخدام لجنة الاختبارات فقط، "
        "ويُحذف بعد انتهاء الاختبارات.",
    )


def _generated_line(render: _Render) -> str:
    lang = render.lang
    when = timezone.localtime(render.prepared.generated_at).strftime("%Y-%m-%d %H:%M")
    role = _ROLE_NAMES.get(render.prepared.generated_role, (render.prepared.generated_role,) * 2)
    who = render.prepared.generated_by
    return _pick(
        lang,
        f"{when} by {who} ({role[0]})",
        f"{_iso(when, lang)} بواسطة {_iso(who, lang)} ({role[1]})",
    )


def _scope_line(render: _Render) -> str:
    lang, options, model = render.lang, render.options, render.model
    kind = options.scope_kind
    if kind == "all":
        return _pick(lang, "Whole timetable", "الجدول كاملاً")
    if kind in {"course", "section"}:
        exam = model.exams[options.exam]
        if kind == "course":
            return _pick(
                lang,
                f"One course: {exam.code} · {exam.name}",
                f"مقرر واحد: {_iso(exam.code + ' · ' + exam.name, lang)}",
            )
        group = next(
            g
            for g in model.groups
            if (g.exam, g.section_key, g.gender)
            == (options.exam, options.section_key, options.gender)
        )
        label = f"{exam.code} {_section_label(group, lang)}"
        return _pick(
            lang,
            f"This section: {label} ({_w(group.gender, lang)})",
            f"هذه الشعبة: {_iso(label, lang)} ({_w(group.gender, lang)})",
        )
    if kind == "day":
        when = options.dates.get(options.day)
        weekday = _weekday(options.day, when, lang)
        detail = " ".join(
            part for part in (weekday or "", when.isoformat() if when else "") if part
        )
        text = options.day + (f" ({detail})" if detail else "")
        return _pick(lang, f"One day: {text}", f"يوم واحد: {_iso(text, lang)}")
    exam = scheduled_slots(model)[options.slot_index]  # type: ignore[index]
    slot = f"{exam.day} {exam.period}"
    if kind == "period":
        return _pick(lang, f"One period: {slot}", f"فترة واحدة: {_iso(slot, lang)}")
    return _pick(
        lang,
        f"One room: {options.room_code} · {slot}",
        f"قاعة واحدة: {_iso(options.room_code + ' · ' + slot, lang)}",
    )


def _students_line(render: _Render) -> str:
    lang, options, item = render.lang, render.options, render.item
    programs = (
        _pick(lang, "Programs ", "البرامج ") + _iso(", ".join(options.programs), lang)
        if options.programs
        else _pick(lang, "All programs", "كل البرامج")
    )
    groups = (item.gender,) if item.gender else options.groups
    rows = (
        _pick(lang, "every student", "جميع الطلاب")
        if options.rows == "all"
        else _pick(
            lang,
            "only students with a clash or same-day exam",
            "الطلاب الذين لديهم تعارض أو اختبار آخر في اليوم نفسه فقط",
        )
    )
    return _JOIN.join([programs, " / ".join(_w(g, lang) for g in groups), rows])


def _status_text(render: _Render) -> tuple[str, bool]:
    lang, model = render.lang, render.model
    run_id = model.saved.run_id
    total = len(model.groups)
    if model.unchanged:
        return (
            _pick(
                lang,
                f"✓ Student lists match saved timetable #{run_id}: {total:,} of {total:,} sections identical.",
                f"✓ قوائم الطلاب مطابقة للجدول المحفوظ رقم {_iso(run_id, lang)}: {total:,} من {total:,} شعبة متطابقة.",
            ),
            True,
        )
    # Students or programme counts: either one is a difference the file marks.
    differ = len(model.differing_groups)
    no_seat = whole_run_no_seat(model)
    missing = len(model.missing_exams)
    extra_en = f"; {missing} exams are no longer in the lists" if missing else ""
    extra_ar = f"؛ {missing} اختبارات لم تعد في القوائم" if missing else ""
    return (
        _pick(
            lang,
            f"≠ Changed since #{run_id} was saved: {differ:,} of {total:,} sections differ; "
            f"{no_seat:,} students have no seat{extra_en}. See Checks and changes ›",
            f"≠ تغيّرت القوائم بعد حفظ الجدول رقم {_iso(run_id, lang)}: {differ:,} من {total:,} شعبة مختلفة؛ "
            f"{no_seat:,} طالباً بلا مقعد{extra_ar}. انظر المطابقة والتغييرات ›",
        ),
        False,
    )


def _about_cell(
    ws, value: str, *, font: Font | None = None, fill: str | None = None, link: str | None = None
) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, _clean(value))
    cell.data_type = "s"
    if font is not None:
        cell.font = font
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if link:
        cell.hyperlink = Hyperlink(ref="A1", location=link, display=value)
    return cell


def _sheet_link(title: str) -> str:
    return "'" + title.replace("'", "''") + "'!A1"


def _write_about(
    ws, spec: SheetSpec, render: _Render, *, header, selected, layout, sheet_rows, **_kw
) -> None:
    lang, model, options, item = render.lang, render.model, render.options, render.item
    saved = model.saved
    _setup_sheet(
        ws,
        spec,
        lang,
        [34, 14, 90],
        header_left=header[0],
        header_right=header[1],
        selected=selected,
    )
    writer = _SheetWriter(ws, lang, _Styles(ws))
    title = _pick(
        lang,
        f"Exam student data — timetable #{saved.run_id}",
        f"بيانات الطلاب للاختبارات — الجدول رقم {saved.run_id}",
    )
    ws.append([_about_cell(ws, title, font=Font(name="Arial", size=18, bold=True, color=NAVY))])
    ws.append(
        [
            _about_cell(
                ws, _confidential(lang), font=Font(name="Arial", size=11, bold=True, color=DANGER)
            )
        ]
    )
    status, matches = _status_text(render)
    checks_title = SHEET_BY_KEY["checks"].title(lang)
    status_cell = _about_cell(
        ws,
        status,
        font=Font(name="Arial", size=11, bold=not matches, color=OK_TEXT if matches else WARN),
        fill=OK_SOFT if matches else WARN_SOFT,
        link=None if matches else _sheet_link(checks_title),
    )
    ws.append([_about_cell(ws, _pick(lang, "Status", "الحالة")), status_cell])
    ws.append([])
    writer.row = 4
    saved_local = timezone.localtime(saved.saved_at).strftime("%Y-%m-%d %H:%M")
    label = saved.label
    lines: list[tuple[str, str]] = []
    if options.prepared_for:
        lines.append((_pick(lang, "Prepared for", "أُعد لـ"), options.prepared_for))
    lines += [
        (
            _pick(lang, "Exam timetable", "جدول الاختبارات"),
            _pick(
                lang,
                f"#{saved.run_id} · {label} · saved {saved_local} (Riyadh, +03:00)",
                f"رقم {_iso(saved.run_id, lang)} · {_iso(label, lang)} · حُفظ {_iso(saved_local, lang)} (بتوقيت الرياض)",
            ),
        ),
        (
            _pick(lang, "Term", "الفصل"),
            _pick(
                lang,
                f"{saved.academic_year}, term {saved.term}",
                f"{_iso(saved.academic_year, lang)}، الفصل {_iso(saved.term, lang)}",
            ),
        ),
        (
            _pick(lang, "Student lists", "قوائم الطلاب"),
            _pick(
                lang,
                "Registrar class timetables imported from the university portal; read "
                + timezone.localtime(model.checked_at).strftime("%Y-%m-%d %H:%M"),
                "جداول الطلاب المستوردة من بوابة الجامعة؛ قُرئت "
                + _iso(timezone.localtime(model.checked_at).strftime("%Y-%m-%d %H:%M"), lang),
            ),
        ),
        (_pick(lang, "Exams in this file", "الاختبارات في هذا الملف"), _scope_line(render)),
        (_pick(lang, "Students in this file", "الطلاب في هذا الملف"), _students_line(render)),
        (
            _pick(lang, "Contents", "المحتوى"),
            _pick(
                lang,
                f"{item.rows:,} student-exam rows · {item.students:,} students · {item.exams:,} exams · {item.sections:,} sections",
                f"{item.rows:,} سطراً للطلاب والاختبارات · {item.students:,} طالباً · {item.exams:,} اختباراً · {item.sections:,} شعبة",
            )
            + (
                _pick(
                    lang,
                    " · summaries only — no names or IDs",
                    " · ملخصات فقط — دون أسماء أو أرقام",
                )
                if options.contents == "summary"
                else ""
            ),
        ),
        (
            _pick(lang, "Exam dates", "تواريخ الاختبارات"),
            _pick(lang, "Entered for this export", "أُدخلت لهذا التصدير")
            if options.dates
            else _pick(lang, "Not set — Date columns are empty", "غير محددة — أعمدة التاريخ فارغة"),
        ),
        (_pick(lang, "Generated", "أُنشئ"), _generated_line(render)),
        (
            _pick(lang, "Reference", "المرجع"),
            _pick(
                lang,
                f"{render.reference} — quote it with any question about this file",
                f"{_iso(render.reference, lang)} — اذكره مع أي استفسار عن هذا الملف",
            ),
        ),
    ]
    for key, value in lines:
        ws.append([_about_cell(ws, key), _about_cell(ws, value)])
    writer.row += len(lines)
    writer.blank(1)
    guide_rows = []
    for sheet in layout[1:]:
        count = sheet_rows[sheet.key]
        first = sheet.tables[0] if sheet.tables else _PROGRAM_DAYS_TABLE
        one_row = first.one_row_ar if lang == "ar" else first.one_row_en
        if sheet.key == "flags" and count == 0:
            one_row = _pick(
                lang,
                "0 rows — no student in this file has two exams on one day",
                "0 أسطر — لا يوجد طالب في هذا الملف له اختباران في يوم واحد",
            )
        link = WriteOnlyCell(ws, sheet.title(lang))
        link.hyperlink = Hyperlink(
            ref="A1", location=_sheet_link(sheet.title(lang)), display=sheet.title(lang)
        )
        guide_rows.append((link, count, one_row))
    writer.table(SHEET_GUIDE, guide_rows)
    writer.blank(2)
    ws.append(
        [
            _about_cell(
                ws,
                _pick(lang, "How to read this file", "كيف تقرأ هذا الملف"),
                font=Font(name="Arial", size=11, bold=True, color=NAVY),
            )
        ]
    )
    writer.row += 1
    for term, sentence in _how_to_read(lang):
        ws.append(
            [
                _about_cell(ws, term, font=Font(name="Arial", size=11, bold=True, color=BODY)),
                _about_cell(ws, sentence),
            ]
        )
        writer.row += 1
    writer.finish_tables()
    ws.protection.sheet = True


def _how_to_read(lang: str) -> list[tuple[str, str]]:
    en = [
        ("Clash", "Another exam in the same period."),
        ("Same day", "Another exam the same day in a different period."),
        (
            "Flags scope",
            "Flags and 'in timetable' counts cover the whole saved timetable, not only this file.",
        ),
        (
            "Split by student ID",
            "The section was split across rooms by the timetable. Seats follow ascending student ID in room order, filling each room up to its saved count. The timetable itself does not seat individuals.",
        ),
        (
            "No seat",
            "The section has more students now than when the timetable was saved; rooms were sized for the saved count.",
        ),
        ("Room not assigned", "The timetable found no room."),
        ("Not scheduled", "An exam the timetable could not place."),
        (
            "Section Not recorded / Unclear",
            "The imported timetables name no section for these students, or name more than one.",
        ),
    ]
    ar = [
        ("تعارض", "اختبار آخر في الفترة نفسها."),
        ("اليوم نفسه", "اختبار آخر في اليوم نفسه في فترة أخرى."),
        (
            "نطاق التنبيهات",
            "التنبيهات وأعداد «الجدول» تغطي الجدول المحفوظ كاملاً، لا هذا الملف وحده.",
        ),
        (
            "مقسّمة حسب الرقم الجامعي",
            "وزع الجدول الشعبة على أكثر من قاعة. تُوزع المقاعد بترتيب الأرقام الجامعية تصاعدياً وبترتيب القاعات، وتُملأ كل قاعة حتى عددها المحفوظ. الجدول نفسه لا يحدد مقعد كل طالب.",
        ),
        (
            "بلا مقعد",
            "عدد طلاب الشعبة الآن أكبر مما كان عند حفظ الجدول؛ حُددت القاعات للعدد المحفوظ.",
        ),
        ("لم تُحدَّد قاعة", "لم يجد الجدول قاعة."),
        ("غير مجدول", "اختبار تعذر على الجدول وضعه."),
        (
            "الشعبة غير مسجلة / غير محددة",
            "لا تذكر الجداول المستوردة شعبة لهؤلاء الطلاب، أو تذكر أكثر من شعبة.",
        ),
    ]
    return ar if lang == "ar" else en


def _file_info_rows(render: _Render) -> list[tuple]:
    lang, model, options, item = render.lang, render.model, render.options, render.item
    saved = model.saved
    counts = Counter(group.membership for group in model.groups)
    groups = (item.gender,) if item.gender else options.groups
    values: dict[str, object] = {
        "format": FORMAT,
        "format_version": str(FORMAT_VERSION),
        "generator_commit": os.environ.get("RENDER_GIT_COMMIT") or "local",
        "run_id": str(saved.run_id),
        "run_label": saved.label,
        "run_saved_at_utc": saved.saved_at.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "run_status": saved.primary_status,
        "run_input_fingerprint16": saved.input_fingerprint[:16],
        "enrollment_source": "scraper_timetable",
        "academic_year": saved.academic_year,
        "term": saved.term,
        "lists_code_saved": model.lists_code_saved,
        "lists_code_now": model.lists_code_now,
        "lists_match": "true" if model.lists_match else "false",
        "sections_total": str(len(model.groups)),
        "sections_matching": str(counts[MATCHES]),
        "sections_changed": str(counts[CHANGED]),
        "sections_new": str(counts[NEW]),
        "sections_gone": str(counts[GONE]),
        "lists_checked_at_utc": model.checked_at.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "scope_kind": options.scope_kind,
        "scope_value": scope_value(options, model),
        "programs": ",".join(options.programs),
        "groups": ",".join(groups),
        "file_group": item.gender or "",
        "file_part": f"{item.part_no}/{item.part_total}",
        "rows_option": options.rows,
        "contents": options.contents,
        "language": options.language,
        "exam_dates_source": "entered" if options.dates else "not_set",
        "prepared_for": options.prepared_for,
        "rows": str(item.rows),
        "students": str(item.students),
        "exams": str(item.exams),
        "sections_in_file": str(item.sections),
        "data_sha256": item.data_sha256,
        "seat_rule": SEAT_RULE,
        "flag_rule": FLAG_RULE,
        "generated_at_utc": render.prepared.generated_at.astimezone(UTC).strftime(
            "%Y-%m-%dT%H:%M:%SZ"
        ),
        "generated_by": render.prepared.generated_by,
        "reference": render.reference,
        "audit_entry_hash16": render.audit_hash[:16],
    }
    return [
        (key, str(values[key]) or None, _FILE_INFO_MEANING[key][1 if lang == "ar" else 0])
        for key in FILE_INFO_KEYS
    ]


def _column_guide_rows(render: _Render, layout: list[SheetSpec]) -> list[tuple]:
    lang = render.lang
    rows = []
    for sheet in layout[1:]:
        tables: Sequence[TableSpec] = sheet.tables or (_PROGRAM_DAYS_TABLE,)
        for table in tables:
            for column in table.columns:
                header = column.header(lang)
                if table.name == "ProgramDays" and column.key == "day_students":
                    header = _pick(lang, "Each exam day (W1-Sun …)", "كل يوم اختبار (W1-Sun …)")
                rows.append(
                    (
                        sheet.title(lang),
                        header,
                        column.meaning_ar if lang == "ar" else column.meaning_en,
                        _TYPE_NAMES[column.kind][1 if lang == "ar" else 0],
                        column.example or None,
                    )
                )
    return rows


def sittings_by_group(sittings: Iterable[Sitting]) -> dict[str, int]:
    counts = Counter(s.group.gender for s in sittings)
    return {gender: counts.get(gender, 0) for gender in GENDERS}


# ── Preflight: counts only, never a student ────────────────────


def _check_summary(model: RosterModel) -> dict:
    counts = Counter(group.membership for group in model.groups)
    return {
        "status": "matches" if model.unchanged else "changed",
        "sections_total": len(model.groups),
        "sections_matching": counts[MATCHES],
        "sections_changed": counts[CHANGED],
        "sections_new": counts[NEW],
        "sections_gone": counts[GONE],
        "program_mix_changed": _mix_only_changes(model),
        "exams_missing": list(model.missing_exams),
        "changed": [
            {
                "exam": group.exam,
                "section": group.section,
                "mapping_status": group.mapping_status,
                "gender": group.gender,
                "saved": group.saved_count,
                "now": len(group.members),
                "membership": group.membership,
                "program_mix": group.program_mix,
            }
            for group in model.differing_groups
        ],
        "no_seat": whole_run_no_seat(model),
        "lists_code_saved": model.lists_code_saved,
        "lists_code_now": model.lists_code_now,
    }


def export_choices(model: RosterModel, sittings: list[Sitting] | None = None) -> dict:
    """What the dialog's pickers offer: timetable facts with row counts."""
    sittings = all_sittings(model) if sittings is None else sittings
    by_exam = Counter(s.exam.code for s in sittings)
    by_group = Counter((s.group.exam, s.group.section_key, s.group.gender) for s in sittings)
    by_slot = Counter(s.exam.slot_index for s in sittings if s.exam.scheduled)
    by_day = Counter(s.exam.day for s in sittings if s.exam.scheduled)
    by_room = Counter(
        (s.exam.slot_index, s.part.room_code)  # type: ignore[union-attr]
        for s in sittings
        if s.basis in {BASIS_WHOLE, BASIS_SPLIT} and s.part is not None
    )
    by_program = Counter(model.students[s.student_id].program for s in sittings)
    exams = []
    for code, exam in sorted(
        model.exams.items(),
        key=lambda item: (
            item[1].day_no is None,
            item[1].day_no or 0,
            item[1].start or time.max,
            natural_key(item[0]),
        ),
    ):
        exams.append(
            {
                "code": code,
                "name": exam.name,
                "course_code": exam.source_code,
                "day": exam.day if exam.scheduled else "",
                "period": exam.period,
                "slot_index": exam.slot_index if exam.scheduled else None,
                "scheduled": exam.scheduled,
                "in_lists": exam.in_lists,
                "rows": by_exam[code],
                "sections": [
                    {
                        "section_key": group.section_key,
                        "gender": group.gender,
                        "section": group.section,
                        "mapping_status": group.mapping_status,
                        "membership": group.membership,
                        "rows": by_group[(group.exam, group.section_key, group.gender)],
                    }
                    for group in model.groups
                    if group.exam == code
                ],
            }
        )
    available = available_programs(model)
    return {
        "exams": exams,
        "periods": [
            {
                "slot_index": slot,
                "day": exam.day,
                "day_no": exam.day_no,
                "period": exam.period,
                "rows": by_slot[slot],
            }
            for slot, exam in scheduled_slots(model).items()
        ],
        "days": [
            {"day": day, "day_no": number, "rows": by_day[day]}
            for number, day in enumerate(model.saved.days, 1)
        ],
        "rooms": [
            {
                "slot_index": slot,
                "room_code": code,
                "gender": room["gender"],
                "building": room["building"],
                "rows": by_room[(slot, code)],
            }
            for (slot, code), room in sorted(
                model.saved.rooms.items(), key=lambda item: (item[0][0], natural_key(item[0][1]))
            )
        ],
        "programs": [
            {"program": program, "department": _department(program)[0], "rows": by_program[program]}
            for program in available
        ],
        "departments": [
            {
                "id": key,
                "name_en": english,
                "name_ar": arabic,
                "programs": [program for program in members if program in available],
            }
            for key, english, arabic, members in _DEPARTMENTS
            if any(program in available for program in members)
        ],
        "groups": sittings_by_group(sittings),
    }


def choices_digest(choices: dict) -> str:
    """A short checksum of the picker choices, so a dialog that holds them is not re-sent them."""
    canonical = json.dumps(choices, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:16]


def preflight_summary(
    model: RosterModel,
    options: ExportOptions | None,
    *,
    pickers: object = None,
    known_choices: object = None,
) -> dict:
    """The dialog's check panel, pickers and counts. No student ID or name.

    ``pickers`` may carry every picker's current value (exam, section_key,
    gender, slot_index, room_code, day) so one call prices every scope choice.
    ``choices`` (~200 KB on a whole run) is sent only when it differs from the
    ``known_choices`` digest the dialog already holds; ``choices_digest`` is
    always sent. The choices do not depend on the options, only on the run and
    the lists, so a dialog re-pricing its options gets counts alone.
    """
    saved = model.saved
    sittings = all_sittings(model)
    choices = export_choices(model, sittings)
    digest = choices_digest(choices)
    body: dict[str, Any] = {
        "ok": True,
        "run": {
            "id": saved.run_id,
            "label": saved.label,
            "saved_at": saved.saved_at.isoformat(),
            "academic_year": saved.academic_year,
            "term": saved.term,
        },
        "check": _check_summary(model),
        "choices_digest": digest,
        "mode": "sync",
    }
    if known_choices != digest:
        body["choices"] = choices
    if options is None:
        return body
    chosen = [s for s in sittings if _selected(s, options, model)]
    values = pickers if isinstance(pickers, dict) else {}
    scope_rows: dict[str, int] = {}
    for kind in SCOPE_KINDS:
        if kind == options.scope_kind:
            fields = {name: getattr(options, name) for name in _SCOPE_FIELDS[kind]}
        else:
            fields = {name: values.get(name) for name in _SCOPE_FIELDS[kind]}
        try:
            _scope_values(fields, kind, model)
        except ExportOptionsError:
            continue
        probe = ExportOptions(
            scope_kind=kind,
            programs=options.programs,
            groups=options.groups,
            rows=options.rows,
            **fields,
        )
        scope_rows[kind] = sum(1 for s in sittings if _selected(s, probe, model))
    any_group = [s for s in sittings if _selected(s, options, model, check_groups=False)]
    files = []
    if options.one_file_per_group:
        for gender in GENDERS:
            rows = sum(1 for s in chosen if s.group.gender == gender)
            if rows:
                files.append(
                    {
                        "name": f"{file_stem(options, model, gender)}_<REF>.xlsx",
                        "gender": gender,
                        "rows": rows,
                    }
                )
    elif chosen:
        files.append(
            {"name": f"{file_stem(options, model)}_<REF>.xlsx", "gender": None, "rows": len(chosen)}
        )
    body["counts"] = {
        "rows": len(chosen),
        "students": len({s.student_id for s in chosen}),
        "scope_rows": scope_rows,
        "groups": sittings_by_group(any_group),
    }
    body["files"] = files
    if len(files) == 1:
        body["download_name"] = files[0]["name"]
    else:
        body["download_name"] = f"{file_stem(options, model)}_<REF>.zip" if files else ""
    return body
