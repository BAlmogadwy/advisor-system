"""Department operations workbooks from saved, scraped exam enrolments.

Exports do not schedule exams, assign students to rooms, or consult live
registrations. The row grain is one teaching section within a canonical exam;
department counts remain exact even when the section shares or spans rooms.
"""

from __future__ import annotations

import hashlib
import math
import re
from collections import Counter
from datetime import date, datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.services.exam_operations_snapshot import EXAM_OPERATIONS_SNAPSHOT_VERSION
from core.services.exam_run_schema import load_normalised_run
from core.services.exam_sections import EXAM_ENROLLMENT_SOURCE

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class DepartmentExportUnavailable(ValueError):
    code = "operations_snapshot_required"


# These are programme memberships, never course-code or room ownership rules.
_DEPARTMENTS = (
    ("ai-ds", "AI & DS", "الذكاء الاصطناعي وعلوم البيانات", ("AI", "AI2", "DS", "DS2")),
    ("cs", "Computer Science", "علوم الحاسب", ("CS", "CS2")),
    ("is", "Information Systems", "نظم المعلومات", ("IS", "IS2")),
    ("coe", "Computer Engineering", "هندسة الحاسب", ("COE", "COE2")),
    ("cyp", "Cybersecurity", "الأمن السيبراني", ("CYP", "CYP2", "CYB", "CYB2")),
)

_HEADERS = {
    "en": [
        "Course",
        "Teaching section",
        "Day",
        "Date",
        "Period",
        "Gender",
        "Course name",
        "Department students",
        "Whole section students",
        "Building",
        "Room allocation (whole section)",
        "Invigilators",
        "Supervisor",
        "Course instructor",
        "Notes",
        "Record ID",
    ],
    "ar": [
        "رمز المقرر",
        "الشعبة",
        "اليوم",
        "التاريخ",
        "الفترة",
        "الشطر",
        "اسم المقرر",
        "طلاب القسم",
        "إجمالي الشعبة",
        "المبنى",
        "توزيع القاعات للشعبة كاملة",
        "المراقبون",
        "المشرف",
        "أستاذ المقرر",
        "ملاحظات",
        "معرف السجل",
    ],
}

_WEEKDAYS = {
    "mon": 0,
    "monday": 0,
    "الاثنين": 0,
    "الإثنين": 0,
    "tue": 1,
    "tuesday": 1,
    "الثلاثاء": 1,
    "wed": 2,
    "wednesday": 2,
    "الأربعاء": 2,
    "الاربعاء": 2,
    "thu": 3,
    "thursday": 3,
    "الخميس": 3,
    "fri": 4,
    "friday": 4,
    "الجمعة": 4,
    "sat": 5,
    "saturday": 5,
    "السبت": 5,
    "sun": 6,
    "sunday": 6,
    "الأحد": 6,
    "الاحد": 6,
}


def _positive_count(value: object) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise DepartmentExportUnavailable(
            "Saved department enrolment counts are invalid. Check changes and save again."
        )
    return value


def _validated_courses(data: dict) -> list[tuple[dict, dict]]:
    def text_value(value: object, *, blank: bool = True) -> bool:
        return (
            isinstance(value, str)
            and (blank or bool(value.strip()))
            and not re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]", value)
        )

    def gender_value(value: object) -> bool:
        return isinstance(value, str) and value in {"M", "F", "U"}

    def program_value(value: object) -> bool:
        return text_value(value) and len(value) <= 32 and value == value.strip().upper()

    if not isinstance(data, dict):
        raise DepartmentExportUnavailable("The saved timetable is invalid.")
    if data.get("status") != "ok":
        raise DepartmentExportUnavailable(
            "This saved timetable cannot be rendered by this application version. Rebuild it or use a compatible version before exporting."
        )
    if data.get("enrollment_source") != EXAM_ENROLLMENT_SOURCE:
        raise DepartmentExportUnavailable(
            "Load Courses from actual scraped timetables and rebuild before exporting."
        )
    snapshot = data.get("operations_snapshot")
    if (
        not isinstance(snapshot, dict)
        or type(snapshot.get("version")) is not int
        or snapshot.get("version") != EXAM_OPERATIONS_SNAPSHOT_VERSION
    ):
        raise DepartmentExportUnavailable(
            "Check changes and save this timetable to prepare department exports."
        )
    if snapshot.get("enrollment_source") != EXAM_ENROLLMENT_SOURCE:
        raise DepartmentExportUnavailable(
            "The saved department source is unavailable. Check changes and save again."
        )
    courses = snapshot.get("courses")
    schedule = data.get("schedule")
    if not isinstance(courses, dict) or not isinstance(schedule, list) or not schedule:
        raise DepartmentExportUnavailable(
            "The saved timetable has no department enrolment details."
        )
    if any(not isinstance(entry, dict) for entry in schedule):
        raise DepartmentExportUnavailable("The saved timetable is invalid.")
    codes = [entry.get("course_code") for entry in schedule]
    if (
        any(not text_value(code, blank=False) for code in codes)
        or len(set(codes)) != len(codes)
        or set(codes) != set(courses)
    ):
        raise DepartmentExportUnavailable(
            "Department enrolments do not match the saved timetable. Check changes and save again."
        )
    original_sections = data.get("section_enrollment")
    if not isinstance(original_sections, dict):
        raise DepartmentExportUnavailable("Saved room-allocation section details are invalid.")
    slots = data.get("slots", [])
    if not isinstance(slots, list) or any(
        not isinstance(slot, dict)
        or not text_value(slot.get("day"), blank=False)
        or not text_value(slot.get("period"))
        for slot in slots
    ):
        raise DepartmentExportUnavailable("Saved timetable day and period details are invalid.")
    result = []
    identities = set()
    for entry in schedule:
        if (
            not text_value(entry.get("day"), blank=False)
            or not text_value(entry.get("period"))
            or type(entry.get("slot_index")) is not int
            or entry["slot_index"] < 0
        ):
            raise DepartmentExportUnavailable("Saved exam placements are invalid.")
        course = courses[entry["course_code"]]
        if (
            not isinstance(course, dict)
            or any(
                not text_value(course.get(key), blank=key == "course_name")
                or course.get(key) != entry.get(key)
                for key in ("course_identity", "source_course_code", "course_name")
            )
            or course["course_identity"] in identities
        ):
            raise DepartmentExportUnavailable(
                "Saved course identities do not match department enrolments."
            )
        identities.add(course["course_identity"])
        sections = course.get("sections")
        if not isinstance(sections, list) or not sections:
            raise DepartmentExportUnavailable("Saved teaching-section details are missing.")
        aggregate = Counter()
        keys = set()
        section_totals = {}
        section_by_key = {}
        for section in sections:
            if not isinstance(section, dict) or not gender_value(section.get("gender")):
                raise DepartmentExportUnavailable("Saved teaching-section details are invalid.")
            key = (section.get("section_key"), section["gender"])
            if not text_value(key[0], blank=False) or key in keys:
                raise DepartmentExportUnavailable(
                    "Saved teaching sections are duplicated or unidentified."
                )
            keys.add(key)
            count = _positive_count(section.get("student_count"))
            counts = section.get("program_counts")
            if not isinstance(counts, dict) or not counts:
                raise DepartmentExportUnavailable("Saved section programme counts are missing.")
            if any(not program_value(program) for program in counts):
                raise DepartmentExportUnavailable("Saved programme identifiers are invalid.")
            if sum(_positive_count(value) for value in counts.values()) != count:
                raise DepartmentExportUnavailable(
                    "Saved section programme counts do not reconcile."
                )
            status = section.get("mapping_status")
            if not isinstance(status, str) or status not in {"mapped", "missing", "ambiguous"}:
                raise DepartmentExportUnavailable("Saved section attribution is invalid.")
            if not text_value(section.get("section")) or not text_value(
                section.get("mapping_source")
            ):
                raise DepartmentExportUnavailable(
                    "Saved section labels or attribution sources are invalid."
                )
            section_id = section.get("term_section_id")
            if status == "mapped":
                if (
                    type(section_id) is not int
                    or section_id <= 0
                    or key[0] != f"term-section:{section_id}"
                    or not section["section"].strip()
                    or section["mapping_source"] != EXAM_ENROLLMENT_SOURCE
                ):
                    raise DepartmentExportUnavailable(
                        "Saved recorded-section attribution is inconsistent."
                    )
            elif (
                section_id is not None
                or section["section"]
                or section["mapping_source"]
                or key[0] != f"unmapped:{section['gender']}:{status}"
            ):
                raise DepartmentExportUnavailable(
                    "Saved unresolved-section attribution is inconsistent."
                )
            instructors = section.get("instructors", [])
            if (
                not isinstance(instructors, list)
                or any(not text_value(name, blank=False) for name in instructors)
                or not text_value(section.get("instructor_source", ""))
                or not text_value(section.get("instructor_status", ""))
            ):
                raise DepartmentExportUnavailable("Saved section instructor details are invalid.")
            if status != "mapped" and instructors:
                raise DepartmentExportUnavailable(
                    "An unresolved section cannot have recorded instructors."
                )
            for program, value in counts.items():
                aggregate[(program, section["gender"])] += value
            section_totals[key] = count
            section_by_key[key] = section
        declared = Counter()
        program_counts = course.get("program_counts")
        if not isinstance(program_counts, list):
            raise DepartmentExportUnavailable("Saved course programme counts are missing.")
        for item in program_counts:
            if (
                not isinstance(item, dict)
                or not program_value(item.get("program"))
                or not gender_value(item.get("gender"))
            ):
                raise DepartmentExportUnavailable("Saved course programme counts are invalid.")
            key = (item["program"], item["gender"])
            if key in declared:
                raise DepartmentExportUnavailable("Saved programme counts are duplicated.")
            declared[key] = _positive_count(item.get("student_count"))
        if declared != aggregate or sum(aggregate.values()) != _positive_count(
            entry.get("enrolled_count")
        ):
            raise DepartmentExportUnavailable(
                "Saved department totals do not match the exam population."
            )
        original = original_sections.get(entry["course_code"])
        if not isinstance(original, list):
            raise DepartmentExportUnavailable("Saved room-allocation section details are invalid.")
        original_totals = {}
        for row in original:
            if (
                not isinstance(row, dict)
                or not text_value(row.get("section_key"), blank=False)
                or not gender_value(row.get("gender"))
            ):
                raise DepartmentExportUnavailable(
                    "Saved room-allocation section details are invalid."
                )
            key = (row["section_key"], row["gender"])
            captured = section_by_key.get(key)
            if (
                key in original_totals
                or captured is None
                or any(
                    row.get(field) != captured.get(field)
                    for field in (
                        "section",
                        "term_section_id",
                        "mapping_status",
                        "mapping_source",
                        "academic_year",
                        "term",
                    )
                )
            ):
                raise DepartmentExportUnavailable(
                    "Saved department sections do not match the room-allocation source."
                )
            original_totals[key] = _positive_count(row.get("student_count"))
        if original_totals != section_totals or len(original) != len(section_totals):
            raise DepartmentExportUnavailable(
                "Saved department sections do not match the room-allocation source."
            )
        result.append((entry, course))
    return result


def department_export_options(data: dict, language: str = "en") -> dict:
    courses = _validated_courses(data)
    available = {item["program"] for _, course in courses for item in course["program_counts"]}
    profiles = []
    handled = set()
    for key, english, arabic, members in _DEPARTMENTS:
        present = available.intersection(members)
        handled.update(members)
        if present:
            profiles.append(
                {
                    "id": key,
                    "name": arabic if language == "ar" else english,
                    "programs": sorted(present),
                }
            )
    for program in sorted(available - handled):
        profiles.append(
            {
                "id": "program-" + hashlib.sha256(program.encode()).hexdigest()[:12],
                "name": program
                or ("البرنامج غير مسجل" if language == "ar" else "Unspecified programme"),
                "programs": [program],
            }
        )
    for profile in profiles:
        population = set(profile["programs"])
        profile["course_count"] = sum(
            any(item["program"] in population for item in course["program_counts"])
            for _, course in courses
        )
        profile["student_sittings"] = sum(
            item["student_count"]
            for _, course in courses
            for item in course["program_counts"]
            if item["program"] in population
        )
    days = list(
        dict.fromkeys(
            slot["day"]
            for slot in data.get("slots", [])
            if isinstance(slot, dict)
            and isinstance(slot.get("day"), str)
            and slot["day"] != "OVERFLOW"
        )
    )
    genders = [
        gender
        for gender in ("M", "F", "U")
        if any(
            item["gender"] == gender for _, course in courses for item in course["program_counts"]
        )
    ]
    return {"departments": profiles, "days": days, "genders": genders}


def _request_options(
    data: dict, payload: dict
) -> tuple[list[dict], set[str], str, dict[str, date]]:
    if not isinstance(payload, dict):
        raise ValueError("Export options must be an object.")
    if set(payload) - {"departments", "genders", "language", "dates"}:
        raise ValueError("Unknown department export option.")
    language = payload.get("language", "ar")
    if not isinstance(language, str) or language not in {"ar", "en"}:
        raise ValueError("Choose Arabic or English for the workbook.")
    options = department_export_options(data, language)
    selected = payload.get("departments")
    profiles = {item["id"]: item for item in options["departments"]}
    if (
        not isinstance(selected, list)
        or not selected
        or any(not isinstance(value, str) or value not in profiles for value in selected)
        or len(set(selected)) != len(selected)
    ):
        raise ValueError("Select at least one available department without duplicates.")
    genders = payload.get("genders", options["genders"])
    if (
        not isinstance(genders, list)
        or not genders
        or any(
            not isinstance(gender, str) or gender not in options["genders"] for gender in genders
        )
        or len(set(genders)) != len(genders)
    ):
        raise ValueError("Select at least one available student group without duplicates.")
    dates = parse_exam_dates(payload.get("dates", {}), options["days"])
    return [profiles[key] for key in selected], set(genders), language, dates


def day_weekday(day: str) -> int | None:
    """Monday-based weekday a grid day label names ("W1-Sun" -> 6), if any."""
    return _WEEKDAYS.get(re.sub(r"^w\d+[-\s]+", "", day.strip().casefold()))


def parse_exam_dates(raw_dates: object, days: list[str]) -> dict[str, date]:
    """Validate per-day exam dates entered for an export; they are never stored.

    Dates must be ISO, Excel-representable, match the weekday a label names and
    follow the timetable's day order. Shared by every exam export.
    """
    if not isinstance(raw_dates, dict) or set(raw_dates) - set(days):
        raise ValueError("Enter dates only for days in this timetable.")
    dates = {}
    for day, value in raw_dates.items():
        if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("Exam dates must use YYYY-MM-DD.")
        try:
            dates[day] = date.fromisoformat(value)
        except ValueError as exc:
            raise ValueError("Enter a valid exam date.") from exc
        if dates[day].year < 1900:
            raise ValueError("Exam dates must be supported by Excel (1900 or later).")
        weekday = day_weekday(day)
        if weekday is not None and dates[day].weekday() != weekday:
            raise ValueError(f"The date for {day} must match its weekday.")
    ordered_dates = [dates[day] for day in days if day in dates]
    if any(right <= left for left, right in zip(ordered_dates, ordered_dates[1:], strict=False)):
        raise ValueError("Exam dates must be distinct and follow timetable day order.")
    return dates


def _room_distribution(entry: dict, sections: list[dict], language: str) -> dict:
    """Use only saved section fragments; never apportion shared-room students."""
    ar = language == "ar"
    rows = {(section["section_key"], section["gender"]): [] for section in sections}
    rooms = entry.get("rooms")
    if rooms is None:
        rooms = []
    if not isinstance(rooms, list):
        raise DepartmentExportUnavailable("Saved room allocations are invalid.")
    section_by_key = {(section["section_key"], section["gender"]): section for section in sections}
    seen_rooms = set()
    for room in rooms:
        if not isinstance(room, dict):
            raise DepartmentExportUnavailable("Saved room allocations are invalid.")
        room_code = room.get("room_code")
        gender = room.get("gender")
        if (
            not isinstance(room_code, str)
            or not room_code.strip()
            or not isinstance(gender, str)
            or gender not in {"M", "F", "U"}
            or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]", room_code)
        ):
            raise DepartmentExportUnavailable(
                "Saved room identifiers or student groups are invalid."
            )
        room_key = (gender, room_code)
        if room_code != "UNASSIGNED" and room_key in seen_rooms:
            raise DepartmentExportUnavailable("A saved exam room is duplicated.")
        seen_rooms.add(room_key)
        parts = room.get("section_parts")
        if not isinstance(parts, list) or not parts:
            raise DepartmentExportUnavailable(
                "Saved room allocations lack teaching-section details. Check changes and save again."
            )
        fragments = Counter()
        for part in parts:
            if not isinstance(part, dict):
                raise DepartmentExportUnavailable("Saved section fragments are invalid.")
            part_gender = part.get("gender", gender)
            if (
                not isinstance(part.get("section_key"), str)
                or not isinstance(part_gender, str)
                or part_gender != gender
            ):
                raise DepartmentExportUnavailable(
                    "Saved room and teaching-section student groups do not match."
                )
            key = (part["section_key"], part_gender)
            if key not in rows:
                raise DepartmentExportUnavailable(
                    "A saved room refers to an unknown teaching section."
                )
            if any(
                part.get(field) != section_by_key[key].get(field)
                for field in ("section", "term_section_id", "mapping_status")
            ):
                raise DepartmentExportUnavailable(
                    "Saved room fragments do not match their teaching section."
                )
            fragments[key] += _positive_count(part.get("student_count"))
        total = _positive_count(room.get("student_count"))
        if sum(fragments.values()) != total:
            raise DepartmentExportUnavailable("Saved room totals do not match section fragments.")
        for key, count in fragments.items():
            label = (
                ("قاعة غير مخصصة" if ar else "UNASSIGNED")
                if room_code == "UNASSIGNED"
                else room_code
            )
            building = room.get("building")
            if building is not None and (
                not isinstance(building, str)
                or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]", building)
            ):
                raise DepartmentExportUnavailable("Saved room building details are invalid.")
            building = (building or "").strip()
            text = f"{label}: {count}"
            if total != count:
                text += f" ({'إجمالي القاعة' if ar else 'room total'}: {total})"
            rows[key].append(
                {
                    "text": text,
                    "count": count,
                    "assigned": room_code != "UNASSIGNED",
                    "room_label": label,
                    "building": building if room_code != "UNASSIGNED" else "",
                }
            )
    for section in sections:
        allocation = rows[(section["section_key"], section["gender"])]
        if allocation and sum(part["count"] for part in allocation) != section["student_count"]:
            raise DepartmentExportUnavailable(
                "Saved room distribution does not cover the whole teaching section."
            )
    return rows


def _building_display(allocations: list[dict], language: str) -> str:
    buildings = {part["building"] for part in allocations}
    if not any(buildings):
        return ""
    if len(buildings) == 1:
        return next(iter(buildings))
    # Keep each location tied to its room when a section spans buildings or
    # only some rooms have a recorded building. Never infer it from the code.
    missing = "غير مسجل" if language == "ar" else "Not recorded"
    return "\n".join(f"{part['room_label']}: {part['building'] or missing}" for part in allocations)


def _section_order(label: str) -> tuple:
    """Keep recorded labels, but list M2 before M10 in operational handouts."""
    return tuple(
        (1, int(part)) if part.isdecimal() else (0, part.casefold())
        for part in re.split(r"(\d+)", label)
    )


def _department_rows(
    data: dict, profile: dict, genders: set[str], language: str, dates: dict
) -> list[dict]:
    ar = language == "ar"
    records = []
    programs = set(profile["programs"])
    for entry, course in _validated_courses(data):
        sections = course["sections"]
        if not any(
            section["gender"] in genders and programs.intersection(section["program_counts"])
            for section in sections
        ):
            continue
        distributions = _room_distribution(entry, sections, language)
        for section in sections:
            count = sum(
                value for program, value in section["program_counts"].items() if program in programs
            )
            if not count or section["gender"] not in genders:
                continue
            allocated = distributions[(section["section_key"], section["gender"])]
            notes = []
            shared = count != section["student_count"]
            if shared:
                notes.append(
                    "شعبة مشتركة مع برامج أخرى" if ar else "Section shared with other programmes"
                )
            if shared and len(allocated) > 1:
                notes.append(
                    "أعداد طلاب القسم في كل قاعة غير محددة"
                    if ar
                    else "Department students per room are not determined"
                )
            if len(allocated) > 1:
                notes.append("الشعبة موزعة على عدة قاعات" if ar else "Section split across rooms")
            if section["mapping_status"] != "mapped":
                notes.append(
                    ("الشعبة غير مسجلة" if ar else "Section not recorded")
                    if section["mapping_status"] == "missing"
                    else ("الشعبة تحتاج مراجعة" if ar else "Ambiguous teaching section")
                )
            if entry.get("is_online"):
                notes.append("اختبار إلكتروني" if ar else "Online exam")
            if entry.get("day") == "OVERFLOW":
                notes.append("موعد الاختبار غير محدد" if ar else "Exam time not assigned")
            if (not allocated and not entry.get("is_online")) or any(
                not part["assigned"] for part in allocated
            ):
                notes.append("راجع تخصيص القاعات" if ar else "Room allocation needs review")
            section_label = section.get("section") or ("غير مسجلة" if ar else "Not recorded")
            if section["mapping_status"] == "ambiguous":
                section_label = "تحتاج مراجعة" if ar else "Ambiguous section"
            instructors = section.get("instructors") or []
            recorded = (
                section.get("instructor_status") == "recorded"
                and section.get("instructor_source") == "recorded_section_meetings"
            )
            room_text = "\n".join(part["text"] for part in allocated) or (
                ("إلكتروني — لا توجد قاعة مخصصة" if ar else "Online — no room assigned")
                if entry.get("is_online")
                else ("غير مخصصة" if ar else "Not allocated")
            )
            records.append(
                {
                    "day": entry.get("day", ""),
                    "date": dates.get(entry.get("day")),
                    "period": entry.get("period", ""),
                    "gender": section["gender"],
                    "code": entry["course_code"],
                    "identity": course["course_identity"],
                    "name": course["course_name"],
                    "section": str(section_label),
                    "department_students": count,
                    "section_students": section["student_count"],
                    "instructor": "\n".join(instructors) if recorded else "",
                    "rooms": room_text,
                    "building": _building_display(allocated, language),
                    "notes": "; ".join(notes),
                    "slot_index": entry.get("slot_index", 999999),
                    "section_key": section["section_key"],
                    "record_id": hashlib.sha256(
                        "\x1f".join(
                            (course["course_identity"], section["gender"], section["section_key"])
                        ).encode()
                    ).hexdigest(),
                    # Reserve writing space for two invigilator names per room.
                    # Longer manually entered names can still require AutoFit.
                    "duty_height": 12
                    + 14 * max(2, 2 * sum(part["assigned"] for part in allocated)),
                    "print_notes": "; ".join(
                        item
                        for item in notes
                        if item
                        not in {
                            "شعبة مشتركة مع برامج أخرى",
                            "Section shared with other programmes",
                            "الشعبة موزعة على عدة قاعات",
                            "Section split across rooms",
                            "اختبار إلكتروني",
                            "Online exam",
                        }
                    ),
                }
            )
    records.sort(
        key=lambda row: (
            row["slot_index"],
            row["gender"],
            row["code"],
            row["identity"],
            _section_order(row["section"]),
            row["section"],
            row["section_key"],
        )
    )
    if not records:
        raise ValueError("The selected department has no students in the selected groups.")
    return records


def _write(
    ws, row: int, values: list, *, fill: str | None = None, bold: bool = False, ar: bool = False
) -> None:
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        # All imported labels/names are literal text, including leading '='.
        if isinstance(value, str):
            cell.data_type = "s"
        cell.font = Font(
            name="Arial", size=11, bold=bold, color="FFFFFF" if fill == "213A59" else "20334D"
        )
        cell.alignment = Alignment(
            horizontal="right" if ar else "left",
            vertical="center",
            wrap_text=True,
            readingOrder=2 if ar else 1,
        )
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if isinstance(value, date | datetime):
            cell.number_format = "yyyy-mm-dd"
        elif isinstance(value, int):
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.number_format = "#,##0"


def _sheet_base(ws, ar: bool, widths: list[int]) -> None:
    ws.sheet_view.rightToLeft = ar
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.center.size = 9
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _wrapped_height(values: list, widths: list[int], minimum: int = 37) -> float:
    """Include wrapping within lines, not only explicit line breaks."""
    lines = max(
        sum(
            max(1, math.ceil(len(line) / max(1, width - 3)))
            for line in str(value or "").split("\n")
        )
        for value, width in zip(values, widths, strict=True)
    )
    return min(409.5, max(minimum, 18 * lines + 8))


def _banner(
    ws,
    row: int,
    text: str,
    last_column: int,
    *,
    ar: bool,
    size: int = 10,
    fill: str | None = None,
    bold: bool = False,
) -> None:
    _write(ws, row, [text], ar=ar, fill=fill, bold=bold)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    ws.cell(row, 1).font = Font(
        name="Arial", size=size, bold=bold, color="20334D" if bold else "53657A"
    )


def _duty_formula(sheet: str, record_id: str, column: str, end: int, language: str) -> str:
    # Row numbers change when staff sort the Excel table. The ID travels with
    # the complete row, so duties remain tied to the canonical exam section.
    sheet = sheet.replace("'", "''")
    value = f"INDEX('{sheet}'!${column}$5:${column}${end},MATCH(\"{record_id}\",'{sheet}'!$P$5:$P${end},0))"
    missing = "راجع التكليف" if language == "ar" else "Review assignment"
    return f'=IFERROR(IF({value}="","",{value}),"{missing}")'


def _course_cell(record: dict, ar: bool) -> CellRichText:
    blocks = [
        TextBlock(InlineFont(rFont="Arial", sz=11, b=True, color="20334D"), record["code"]),
        TextBlock(InlineFont(rFont="Arial", sz=10, color="33465D"), "\n" + record["name"]),
    ]
    if record["instructor"]:
        label = "أستاذ المقرر: " if ar else "Instructor: "
        blocks.append(
            TextBlock(
                InlineFont(rFont="Arial", sz=9, color="53657A"), "\n" + label + record["instructor"]
            )
        )
    return CellRichText(blocks)


def _exam_week_by_day(data: dict) -> dict[str, int]:
    """Use saved exam weeks, including short unprefixed Sunday rollovers.

    Read all timetable days so a department with no exams in an earlier week
    keeps the same week numbers as the main timetable. Multiple periods on
    one day must not advance the week. Unrecognised days stay in full views.
    """
    weeks: dict[str, int] = {}
    seen: set[str] = set()
    week: int | None = 1
    previous_weekday: int | None = None
    for slot in data.get("slots", []):
        day = slot["day"]
        if day in seen:
            continue
        seen.add(day)
        label = day.strip().casefold()
        prefixed = re.fullmatch(r"w([1-9][0-9]*)[-\s]+(.+)", label)
        weekday = _WEEKDAYS.get(prefixed[2] if prefixed else label)
        if weekday is None:
            continue
        # Match the builder's Sunday-first week boundary, not ISO Monday.
        weekday = (weekday + 1) % 7
        if prefixed:
            # Saved/custom labels are not bounded by the builder's 60-day UI.
            # Reject unrepresentable week tabs before converting long digits.
            if len(f"الأسبوع {prefixed[1]} - الطالبات") > 31:
                week, previous_weekday = None, None
                continue
            week = int(prefixed[1])
        else:
            if week is None:
                continue
            if previous_weekday is not None and weekday < previous_weekday:
                week += 1
        if len(f"الأسبوع {week} - الطالبات") > 31:
            week, previous_weekday = None, None
            continue
        weeks[day] = week
        previous_weekday = weekday
    return weeks


def _write_print_sheet(
    wb,
    details,
    selected: list[dict],
    title: str,
    context: str,
    gender_name: str,
    gender: str,
    language: str,
    end: int,
    *,
    week: int | None = None,
) -> None:
    ar = language == "ar"
    week_name = (f"الأسبوع {week}" if ar else f"Week {week}") if week is not None else ""
    sheet_name = (
        (f"{week_name} - {gender_name}" if ar else f"{week_name} ({gender})")
        if week_name
        else (gender_name if ar else f"Print ({gender})")
    )
    printed = wb.create_sheet(sheet_name)
    widths = [43, 13, 14, 15, 22, 47, 32, 27]
    _sheet_base(printed, ar, widths)
    printed.sheet_properties.tabColor = "227F86"
    print_title = f"{title} — {gender_name}" + (f" — {week_name}" if week_name else "")
    _banner(printed, 1, print_title, 8, ar=ar, size=18, bold=True)
    _banner(printed, 2, context, 8, ar=ar)
    note = (
        "أعداد القاعات تخص الشعبة كاملة. — غير مسجل. أدخل التكليفات في التفاصيل، ثم اضبط ارتفاع الصفوف تلقائياً عند الحاجة قبل الطباعة."
        if ar
        else "Room counts cover the whole section. — Not recorded. Enter duties on Details; AutoFit row heights as needed before printing."
    )
    _banner(printed, 3, note, 8, ar=ar, size=9)
    headers = (
        [
            "المقرر",
            "الشعبة",
            "طلاب القسم",
            "إجمالي الشعبة",
            "المبنى",
            "القاعات وأعداد طلاب الشعبة",
            "المراقبون",
            "المشرف",
        ]
        if ar
        else [
            "Course",
            "Section",
            "Department students",
            "Whole section students",
            "Building",
            "Rooms and section students",
            "Invigilators",
            "Supervisor",
        ]
    )
    _write(printed, 4, headers, fill="213A59", bold=True, ar=ar)
    for row, height in ((1, 30), (2, 20), (3, 23), (4, 30)):
        printed.row_dimensions[row].height = height
    printed.print_title_rows = "1:4"
    printed.oddFooter.left.text = (
        gender_name + " · " + title + (" · " + week_name if week_name else "")
    ).replace("&", "&&")
    printed.oddFooter.left.size = 8
    row_number, previous, page_height = 5, None, 103
    for index, record in enumerate(selected):
        group = (record["day"], record["period"])
        course = _course_cell(record, ar)
        values = [
            str(course),
            record["section"],
            record["department_students"],
            record["section_students"],
            record["building"] or "—",
            record["rooms"],
            "",
            "",
        ]
        height = max(_wrapped_height(values, widths, minimum=42), min(409.5, record["duty_height"]))
        warning = record["print_notes"]
        note_height = max(21, 14 * math.ceil(len(warning) / 160) + 7) if warning else 0
        new_group = group != previous
        if row_number > 5 and (
            page_height + height + note_height + (25 if new_group else 0) > 680
            or group[0] != previous[0]
        ):
            printed.row_breaks.append(Break(id=row_number - 1))
            page_height, new_group = 103, True
        if new_group:
            date_label = (
                record["date"].isoformat()
                if record["date"]
                else ("التاريخ غير محدد" if ar else "Date not entered")
            )
            # Isolate the time range: otherwise an Arabic date label can make
            # a right-to-left print engine display the end time before start.
            period_label = f"\u2066{record['period']}\u2069" if ar else record["period"]
            label = "  |  ".join((record["day"], date_label, period_label))
            _banner(printed, row_number, label, 8, ar=ar, size=11, fill="DDECEF", bold=True)
            printed.row_dimensions[row_number].height = 25
            row_number += 1
            page_height += 25
        _write(printed, row_number, values, fill="F4F7FA" if index % 2 else "FFFFFF", ar=ar)
        printed.cell(row_number, 1, course).data_type = "s"
        for col in (2, 3, 4, 5):
            printed.cell(row_number, col).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for print_column, detail_column in ((7, "L"), (8, "M")):
            printed.cell(
                row_number,
                print_column,
                _duty_formula(details.title, record["record_id"], detail_column, end, language),
            )
            printed.cell(row_number, print_column).fill = PatternFill("solid", fgColor="FBF8EF")
            printed.cell(row_number, print_column).font = Font(
                name="Arial", size=10, color="20334D"
            )
        printed.row_dimensions[row_number].height = height
        for cell in printed[row_number]:
            cell.border = Border(bottom=Side(style="hair", color="D9E2EC"))
        printed.cell(row_number, 7).border = Border(
            bottom=Side(style="hair", color="D9E2EC"),
            **{"left" if ar else "right": Side(style="thin", color="D9E2EC")},
        )
        row_number += 1
        if warning:
            _banner(printed, row_number, warning, 8, ar=ar, size=9, fill="FFF1DB")
            printed.row_dimensions[row_number].height = note_height
            row_number += 1
        page_height += height + note_height
        previous = group
    printed.print_area = f"A1:H{row_number - 1}"
    printed.freeze_panes = None


def _workbook_bytes(
    run, data: dict, profile: dict, genders: set[str], language: str, dates: dict
) -> bytes:
    ar = language == "ar"
    records = _department_rows(data, profile, genders, language, dates)
    wb = Workbook()
    wb.properties.title = ("تشغيل الاختبارات — " if ar else "Exam operations — ") + profile["name"]
    wb.properties.subject = f"{run.label} · #{run.pk}"
    wb.properties.creator = "Exam Timetable Builder"
    ws = wb.active
    ws.title = "التفاصيل" if ar else "Details"
    widths = [17, 16, 13, 14, 18, 12, 40, 16, 17, 22, 45, 28, 25, 30, 55, 12]
    _sheet_base(ws, ar, widths)
    ws.sheet_properties.tabColor = "213A59"
    title = wb.properties.title
    programs = ", ".join(
        program or ("غير محدد" if ar else "unspecified") for program in profile["programs"]
    )
    context = f"{run.label} · #{run.pk} · {programs}"
    _banner(ws, 1, title, 15, ar=ar, size=18, bold=True)
    _banner(ws, 2, context, 15, ar=ar)
    note = (
        "أدخل التكليفات في الأعمدة المظللة. أوراق الطباعة مرتبطة بها. اضبط ارتفاع الصفوف تلقائياً إذا امتدت الأسماء إلى أسطر إضافية."
        if ar
        else "Enter duties in the shaded columns; print sheets follow these assignments. AutoFit row heights if names wrap onto extra lines."
    )
    _banner(ws, 3, note, 15, ar=ar, size=10)
    for row, height in ((1, 30), (2, 20), (3, 23), (4, 32)):
        ws.row_dimensions[row].height = height
    _write(ws, 4, _HEADERS[language], fill="213A59", bold=True, ar=ar)
    gender_names = (
        {"M": "الطلاب", "F": "الطالبات", "U": "غير محدد"}
        if ar
        else {"M": "Male", "F": "Female", "U": "Unspecified"}
    )
    for index, record in enumerate(records, 5):
        values = [
            record["code"],
            record["section"],
            record["day"],
            record["date"],
            record["period"],
            gender_names[record["gender"]],
            record["name"],
            record["department_students"],
            record["section_students"],
            record["building"],
            record["rooms"],
            "",
            "",
            record["instructor"],
            record["notes"],
            record["record_id"],
        ]
        _write(ws, index, values, fill="F4F7FA" if index % 2 else "FFFFFF", ar=ar)
        ws.row_dimensions[index].height = max(
            _wrapped_height(values[:15], widths[:15], minimum=32), min(409.5, record["duty_height"])
        )
        ws.cell(index, 1).font = Font(name="Arial", size=11, bold=True, color="20334D")
        for col in (12, 13):
            ws.cell(index, col).fill = PatternFill("solid", fgColor="FFF3D6")
        ws.cell(index, 12).border = Border(
            **{"left" if ar else "right": Side(style="thin", color="D9E2EC")}
        )
        if record["print_notes"]:
            ws.cell(index, 15).fill = PatternFill("solid", fgColor="FFF1DB")
        ws.cell(index, 15).font = Font(name="Arial", size=10, color="53657A")
    end = 4 + len(records)
    table = Table(displayName="DepartmentAssignments", ref=f"A4:P{end}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    # The table owns its filter. A worksheet AutoFilter over the same range
    # conflicts with it and makes Excel repair/remove the table on opening.
    ws.column_dimensions["P"].hidden = True
    ws.freeze_panes = "C5"
    ws.print_title_rows = "1:4"
    _write(
        ws,
        end + 2,
        ["مجموع تسجيلات الاختبارات للقسم" if ar else "Department exam sittings"],
        bold=True,
        ar=ar,
    )
    ws.cell(end + 2, 8, f"=SUM(H5:H{end})").number_format = "#,##0"
    ws.print_area = f"A1:O{end + 2}"
    for gender in ("M", "F", "U"):
        selected = [record for record in records if record["gender"] == gender]
        if selected:
            _write_print_sheet(
                wb, ws, selected, title, context, gender_names[gender], gender, language, end
            )
    weekly: dict[tuple[int, str], list[dict]] = {}
    weeks = _exam_week_by_day(data)
    for record in records:
        week = weeks.get(record["day"])
        if week is not None and record["gender"] in {"M", "F"}:
            weekly.setdefault((week, record["gender"]), []).append(record)
    for week, gender in sorted(weekly, key=lambda key: (key[0], key[1] != "M")):
        _write_print_sheet(
            wb,
            ws,
            weekly[week, gender],
            title,
            context,
            gender_names[gender],
            gender,
            language,
            end,
            week=week,
        )
    wb.active = 0
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_department_workbooks(run, payload: dict) -> tuple[bytes, str, str]:
    """Return a complete XLSX or ZIP without shared paths or database writes."""
    data = load_normalised_run(run)
    profiles, genders, language, dates = _request_options(data, payload)
    files = []
    for profile in profiles:
        filename = f"exam_department_{profile['id']}_{run.pk}.xlsx"
        files.append((filename, _workbook_bytes(run, data, profile, genders, language, dates)))
    if len(files) == 1:
        return files[0][1], files[0][0], XLSX_TYPE
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        for filename, content in files:
            archive.writestr(filename, content)
    return buffer.getvalue(), f"exam_departments_{run.pk}.zip", "application/zip"
